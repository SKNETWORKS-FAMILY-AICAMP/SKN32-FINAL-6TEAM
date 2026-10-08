# -*- coding: utf-8 -*-
"""글자로 만들기 — 파일·글 하나를 **줄 번호가 붙을 글**로. `[2026-09-27]` 설계서 §3-2

    종류 판별     확장자가 아니라 **첫 바이트**(magic bytes)로
    글            UTF-8(안 되면 CP949)
    PDF           글자가 있는 쪽은 PyMuPDF 로, 글자가 거의 없는 쪽(스캔)은 이미지로 그려 받아쓰기
    docx          문단 + 표(칸을 「 | 」로 잇는다)
    xlsx          **계산된 값**(openpyxl data_only) — 행의 빈칸이 아닌 칸을 「 | 」로 잇는다
    사진          로컬 비전 모델에 「한 줄씩 그대로 받아 적어라」만 시킨다(구조화를 시키지 않는다)

★받아쓰기 누락 검사 — 사진은 **위·아래 반쪽**으로도 받아 적어, 전체 받아쓰기에 없는 **시각 줄·일차 머리줄**이
  반쪽에만 있으면 `missing` 으로 적는다. 합치지 않는다(순서를 지어내지 않는다) — 확인 화면이 보여 준다.
  실측(triPilot : RAG, 2026-09-26): 한 번에 구조화를 시키면 2일차가 통째로 빠졌다. 받아쓰기는 6/6.
★`[2026-10-06 사용자 지시 — 값 변조 줄이기]` **바뀐 글자도 잡는다** — 같은 줄을 전체 받아쓰기와 반쪽 받아쓰기가 **다르게** 읽었으면(`differs`) 그 줄을 「확인 필요」로 알린다.
  평가셋 60건의 값 변조 12/350 = 3.4% 가 전부 사진 · 스캔 받아쓰기의 글자 오독(「창경궁」→「청경궁」)이었는데 위 누락 검사는 빠진 줄만 보고 바뀐 글자는 못 봤다.
  두 번 다 같게 오독한 줄은 못 잡는다(이 검사의 한계) — 그 줄은 확인 화면에서 고객이 원본과 대조한다.
★암호 걸린 PDF·매크로 파일은 열지 않는다(MVP 밖).
"""
from __future__ import annotations

from dataclasses import dataclass, field
import io
import re
from typing import Any, Callable

TRANSCRIBE_PROMPT = (
    "이 이미지에 적힌 글을 **한 줄씩 그대로** 받아 적어라. 줄바꿈을 원문대로 지키고, 고치거나 요약하거나 번역하지 마라. "
    "표는 한 행을 한 줄로, 칸 사이는 「 | 」로 적어라. 글이 아닌 것(사진·아이콘)은 적지 마라. 받아 적은 글만 내라.")

#: 글자가 이만큼보다 적은 PDF 쪽은 스캔으로 보고 그려서 받아쓴다(우리가 고른 값)
SCANNED_PAGE_CHARS = 20
#: 이미지로 그릴 때의 해상도(우리가 고른 값) — 받아쓰기 품질과 시간의 사이
RENDER_DPI = 200

_CHECK_LINE = re.compile(r"(\d{1,2}:\d{2}|\d{1,2}\s*시|\d{1,2}\s*일\s*차|day\s*\d{1,2})", re.IGNORECASE)


class UnsupportedSource(ValueError):
    """읽지 않는 형식 — 부르는 쪽이 `fatal_unsupported_format` 으로 남긴다."""


@dataclass
class SourceText:
    kind: str                         # text · pdf · docx · xlsx · image
    text: str
    pages: int = 1
    transcribed_pages: list[int] = field(default_factory=list)   # 받아쓰기로 만든 쪽(1부터)
    missing: list[dict[str, Any]] = field(default_factory=list)  # 누락 검사가 찾은, 전체에 없는 줄
    #: 같은 줄을 전체 · 반쪽 받아쓰기가 **다르게** 읽은 곳 — [{line(전체 글 기준 줄 번호), text, other, half, ratio, page?}]
    differs: list[dict[str, Any]] = field(default_factory=list)
    seconds: float = 0.0

    def numbered(self) -> list[tuple[int, str]]:
        return list(enumerate(self.text.splitlines(), start=1))


def sniff(data: bytes, *, filename: str = "") -> str:
    """첫 바이트로 종류를 가린다. ★확장자는 믿지 않는다."""
    head = data[:16]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\x89PNG") or head.startswith(b"\xff\xd8\xff") or head[:6] in (b"GIF87a", b"GIF89a") \
            or (head.startswith(b"RIFF") and data[8:12] == b"WEBP"):
        return "image"
    if head.startswith(b"PK\x03\x04"):
        import zipfile

        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile as exc:
            raise UnsupportedSource("zip 파일이 깨져 있다") from exc
        if any(n.startswith("word/") for n in names):
            if any(n.lower().endswith("vbaproject.bin") for n in names):
                raise UnsupportedSource("매크로가 든 문서는 열지 않는다")
            return "docx"
        if any(n.startswith("xl/") for n in names):
            if any(n.lower().endswith("vbaproject.bin") for n in names):
                raise UnsupportedSource("매크로가 든 문서는 열지 않는다")
            return "xlsx"
        raise UnsupportedSource("모르는 압축 형식이다")
    if b"\x00" in data[:1024]:
        raise UnsupportedSource(f"모르는 이진 형식이다 ({filename or '이름 없음'})")
    return "text"


def to_text(data: bytes, *, filename: str = "", see: Callable[[str, bytes], str] | None = None,
            check_halves: bool = True) -> SourceText:
    """글자로. 사진·스캔 쪽은 `see(prompt, png)`(비전 받아쓰기)가 있어야 한다 — 없으면 `UnsupportedSource`."""
    import time

    started = time.perf_counter()
    kind = sniff(data, filename=filename)
    if kind == "text":
        result = SourceText("text", _decode(data))
    elif kind == "docx":
        result = SourceText("docx", _docx(data))
    elif kind == "xlsx":
        result = SourceText("xlsx", _xlsx(data))
    elif kind == "pdf":
        result = _pdf(data, see=see, check_halves=check_halves)
    else:
        if see is None:
            raise UnsupportedSource("사진을 읽을 받아쓰기 모델이 설정되지 않았다")
        text, missing, differs = _transcribe(data, see=see, check_halves=check_halves)
        result = SourceText("image", text, transcribed_pages=[1], missing=missing, differs=differs)
    result.seconds = round(time.perf_counter() - started, 2)
    return result


# ── 형식별 ───────────────────────────────────────────────────────
def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnsupportedSource("글자 인코딩을 알 수 없다")


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _xlsx(data: bytes) -> str:
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell(v) for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _cell(value: Any) -> str:
    from datetime import date, datetime, time

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M") if (value.hour or value.minute) else value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value).strip()


def _pdf(data: bytes, *, see, check_halves: bool) -> SourceText:
    import fitz

    document = fitz.open(stream=data, filetype="pdf")
    if document.needs_pass:
        raise UnsupportedSource("암호 걸린 PDF 는 열지 않는다")
    parts: list[str] = []
    transcribed: list[int] = []
    missing: list[dict[str, Any]] = []
    differs: list[dict[str, Any]] = []
    for number, page in enumerate(document, start=1):
        text = page.get_text("text")
        if len(text.strip()) >= SCANNED_PAGE_CHARS:
            parts.append(text.rstrip("\n"))
            continue
        if see is None:
            raise UnsupportedSource(f"{number}쪽이 스캔인데 받아쓰기 모델이 설정되지 않았다")
        png = page.get_pixmap(dpi=RENDER_DPI).tobytes("png")
        written, lost, unlike = _transcribe(png, see=see, check_halves=check_halves)
        # 이 쪽의 첫 줄이 전체 글에서 몇 번째 줄인가 — 앞 쪽들을 「\n」로 이은 뒤 이어 붙는다
        offset = len(("\n".join(parts) + "\n").splitlines()) if parts else 0
        parts.append(written)
        transcribed.append(number)
        missing += [{**m, "page": number} for m in lost]
        differs += [{**d, "line": d["line"] + offset, "page": number} for d in unlike]
    return SourceText("pdf", "\n".join(parts), pages=len(document), transcribed_pages=transcribed, missing=missing, differs=differs)


def _see_once_more(see):
    """받아쓰기 호출을 **한 번만** 다시 한다 — 모델 서버의 일시 오류(HTTP 5xx · 빈 받아쓰기)만.
    ★읽기 전용 호출이라 다시 불러도 부작용이 없다. 두 번째도 실패하면 그대로 올린다(성공으로 추정하지 않는다).
    ☆2026-09-28 평가셋 60건: 일시 오류 3건(HTTP 500 1 · 빈 받아쓰기 2)이 사례를 통째로 잃게 했다(항목 21개)."""
    def call(prompt, image):
        try:
            return see(prompt, image)
        except Exception as exc:                  # noqa: BLE001 — 아래에서 고른 것만 다시 부른다
            text = str(exc)
            if "빈 받아쓰기" not in text and "HTTP 5" not in text:
                raise
            return see(prompt, image)
    return call


def _transcribe(image: bytes, *, see, check_halves: bool) -> tuple[str, list[dict[str, Any]], list[dict[str, Any]]]:
    """전체 받아쓰기 + (반쪽 둘과 맞춰 본) 빠진 줄 · 다르게 읽은 줄. `check_halves` 가 꺼져 있으면 전체만."""
    see = _see_once_more(see)
    full = see(TRANSCRIBE_PROMPT, image)
    if not check_halves:
        return full, [], []
    missing: list[dict[str, Any]] = []
    halves: list[tuple[str, str]] = []
    have = {_key(line) for line in full.splitlines() if _CHECK_LINE.search(line)}
    for name, half in _halves(image):
        written = see(TRANSCRIBE_PROMPT, half)
        halves.append((name, written))
        for line in written.splitlines():
            if _CHECK_LINE.search(line) and _key(line) not in have:
                missing.append({"half": name, "text": line.strip()})
                have.add(_key(line))
    return full, missing, differing_lines(full, halves)


#: 전체 받아쓰기의 줄과 반쪽 받아쓰기의 줄이 **이만큼 이상 닮았는데 같지 않으면** 같은 줄을 다르게 읽은 것으로 본다(우리가 고른 값 — 평가셋 60건으로 조정, 아래 `DIFFER_MIN_RATIO` 근거 참고).
#: 그보다 덜 닮으면 다른 줄이다(반쪽에 없는 줄 · 줄 나눔이 다른 줄은 여기서 말하지 않는다 — 누락 검사가 본다).
DIFFER_MIN_RATIO = 0.5
#: 이 글자 수(공백 · 구두점 뺀)보다 짧은 줄은 비교하지 않는다 — 「1일차」 · 「점심」 같은 짧은 줄은 닮은 정도가 의미가 없다
DIFFER_MIN_CHARS = 4


def differing_lines(full: str, halves: list[tuple[str, str]]) -> list[dict[str, Any]]:
    """전체 받아쓰기의 줄마다, 반쪽 받아쓰기에 **같은 줄**이 있으면 통과, **닮았지만 다른 줄**만 있으면 「다르게 읽었다」로 적는다.

    ★반쪽 경계에서 잘린 줄(반쪽의 줄이 전체 줄의 앞 · 뒤 조각)은 같은 줄로 본다. 같은 줄을 둘 다 같게 오독했으면 못 잡는다(한계).
    돌려주는 값: [{line(1부터), text, other, half, ratio}] — 줄 순서."""
    import difflib

    candidates = [(name, line.strip(), _key(line)) for name, written in halves for line in written.splitlines() if _key(line)]
    out: list[dict[str, Any]] = []
    for number, line in enumerate(full.splitlines(), start=1):
        key = _key(line)
        if len(key) < DIFFER_MIN_CHARS:
            continue
        best: tuple[float, str, str] | None = None
        same = False
        for name, text, other in candidates:
            cut = (len(other) >= max(DIFFER_MIN_CHARS, len(key) // 2) and other in key) or (len(key) >= max(DIFFER_MIN_CHARS, len(other) // 2) and key in other)
            if other == key or cut:                                                           # 같은 줄 · 경계에서 잘린 줄(절반 넘게 남은 조각)
                same = True
                break
            ratio = difflib.SequenceMatcher(None, key, other).ratio()
            if best is None or ratio > best[0]:
                best = (ratio, name, text)
        if not same and best is not None and best[0] >= DIFFER_MIN_RATIO:
            out.append({"line": number, "text": line.strip(), "other": best[2], "half": best[1], "ratio": round(best[0], 2)})
    return out


def other_reading(value: str, text: str, other: str) -> str | None:
    """`text`(전체 받아쓰기의 줄)에서 읽은 값 `value` 가 **`other`(반쪽 받아쓰기의 같은 줄)에서는 무엇으로 읽혔나**. `[2026-10-07 uiux 요청]`

    두 줄을 글자 단위로 맞춰(`difflib`) `value` 가 놓인 자리를 `other` 쪽 자리로 옮겨 자른다 — 「11:00 청경궁 관람」의 「청경궁 관람」 → 「11:00 창경궁 관람」의 「창경궁 관람」.
    화면이 「지금 값 / 다른 읽기」 두 단추를 그리고 고른 값을 그대로 보낼 수 있게 한다(줄 전체를 제목으로 보낼 수는 없다).
    ★못 정하면 `None` — 값이 줄 안에서 글자 그대로 안 보이거나(모델이 다듬은 값) · 옮긴 결과가 비었거나 · 길이가 너무 달라 맞춤이 어긋난 것으로 보일 때 · 같은 값일 때.
      그때 화면은 두 읽기 문장(note)만 보인다. 지어내지 않는다."""
    import difflib

    value = (value or "").strip()
    if not value or not text or not other:
        return None
    start = text.find(value)
    if start < 0:
        return None
    end = start + len(value)
    ops = difflib.SequenceMatcher(None, text, other, autojunk=False).get_opcodes()

    def at(position: int, *, is_end: bool) -> int:
        for tag, i1, i2, j1, j2 in ops:
            if is_end:
                inside = i1 < position <= i2 or (i1 == i2 == position)
            else:
                inside = i1 <= position < i2 or (i1 == i2 == position)
            if not inside:
                continue
            if tag == "equal":
                return j1 + (position - i1)
            if position == i1:
                return j1
            if position == i2:
                return j2
            return j1 + round((position - i1) * (j2 - j1) / max(1, i2 - i1))
        return len(other)

    got = other[at(start, is_end=False):at(end, is_end=True)].strip()
    if not got or got == value:
        return None
    if len(got) > 3 * len(value) + 2 or len(got) * 3 + 2 < len(value):
        return None
    return got


def _key(line: str) -> str:
    """반쪽 받아쓰기와 맞춰 볼 때의 열쇠 — 공백·구두점을 빼고 비교한다(글자 자체는 바꾸지 않는다)."""
    return re.sub(r"[\s·|,.:;()\[\]-]+", "", line)


def _halves(image: bytes) -> list[tuple[str, bytes]]:
    """위·아래 반쪽(가운데를 10% 겹친다 — 경계에 걸린 줄을 잃지 않게)."""
    from PIL import Image

    picture = Image.open(io.BytesIO(image))
    width, height = picture.size
    overlap = int(height * 0.05)
    out = []
    for name, box in (("top", (0, 0, width, height // 2 + overlap)),
                      ("bottom", (0, height // 2 - overlap, width, height))):
        buffer = io.BytesIO()
        picture.crop(box).save(buffer, format="PNG")
        out.append((name, buffer.getvalue()))
    return out


__all__ = ["SourceText", "TRANSCRIBE_PROMPT", "UnsupportedSource", "sniff", "to_text"]
