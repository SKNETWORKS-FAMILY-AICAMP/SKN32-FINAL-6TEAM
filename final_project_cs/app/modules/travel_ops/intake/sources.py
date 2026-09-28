# -*- coding: utf-8 -*-
"""글자로 만들기 — 파일·글 하나를 **줄 번호가 붙을 글**로. `[2026-09-27]` 설계서 §3-2

    종류 판별     확장자가 아니라 **첫 바이트**(magic bytes)로
    글            UTF-8(안 되면 CP949)
    PDF           글자가 있는 쪽은 PyMuPDF 로, 글자가 거의 없는 쪽(스캔)은 이미지로 그려 받아쓰기
    docx          문단 + 표(칸을 「 | 」로 잇는다)
    xlsx          **계산된 값**(openpyxl data_only) — 행의 빈칸이 아닌 칸을 「 | 」로 잇는다
    사진          로컬 비전 모델에 「한 줄씩 그대로 받아 적어라」만 시킨다(구조화를 시키지 않는다)

★받아쓰기 누락 검사 — 사진은 **위·아래 반쪽**으로도 받아 적어, 전체 받아쓰기에 없는 **시각 줄·일차 머리줄**이
  반쪽에만 있으면 `missing` 으로 적는다. 합치지 않는다(순서를 지어내지 않는다) — 확인 화면이 보여 준다.
  실측(triPilot : RAG, 2026-09-26): 한 번에 구조화를 시키면 2일차가 통째로 빠졌다. 받아쓰기는 6/6.
★암호 걸린 PDF·매크로 파일은 열지 않는다(MVP 밖).
"""
from __future__ import annotations

from dataclasses import dataclass, field
import io
import re
from typing import Any, Callable

TRANSCRIBE_PROMPT = (
    "이 이미지에 적힌 글을 **한 줄씩 그대로** 받아 적어라. 줄바꿈을 원문대로 지키고, 고치거나 요약하거나 번역하지 마라. "
    "표는 한 행을 한 줄로, 칸 사이는 「 | 」로 적어라. 글이 아닌 것(사진·아이콘)은 적지 마라. 받아 적은 글만 내라.")

#: 글자가 이만큼보다 적은 PDF 쪽은 스캔으로 보고 그려서 받아쓴다(우리가 고른 값)
SCANNED_PAGE_CHARS = 20
#: 이미지로 그릴 때의 해상도(우리가 고른 값) — 받아쓰기 품질과 시간의 사이
RENDER_DPI = 200

_CHECK_LINE = re.compile(r"(\d{1,2}:\d{2}|\d{1,2}\s*시|\d{1,2}\s*일\s*차|day\s*\d{1,2})", re.IGNORECASE)


class UnsupportedSource(ValueError):
    """읽지 않는 형식 — 부르는 쪽이 `fatal_unsupported_format` 으로 남긴다."""


@dataclass
class SourceText:
    kind: str                         # text · pdf · docx · xlsx · image
    text: str
    pages: int = 1
    transcribed_pages: list[int] = field(default_factory=list)   # 받아쓰기로 만든 쪽(1부터)
    missing: list[dict[str, Any]] = field(default_factory=list)  # 누락 검사가 찾은, 전체에 없는 줄
    seconds: float = 0.0

    def numbered(self) -> list[tuple[int, str]]:
        return list(enumerate(self.text.splitlines(), start=1))


def sniff(data: bytes, *, filename: str = "") -> str:
    """첫 바이트로 종류를 가린다. ★확장자는 믿지 않는다."""
    head = data[:16]
    if head.startswith(b"%PDF"):
        return "pdf"
    if head.startswith(b"\x89PNG") or head.startswith(b"\xff\xd8\xff") or head[:6] in (b"GIF87a", b"GIF89a") \
            or (head.startswith(b"RIFF") and data[8:12] == b"WEBP"):
        return "image"
    if head.startswith(b"PK\x03\x04"):
        import zipfile

        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile as exc:
            raise UnsupportedSource("zip 파일이 깨져 있다") from exc
        if any(n.startswith("word/") for n in names):
            if any(n.lower().endswith("vbaproject.bin") for n in names):
                raise UnsupportedSource("매크로가 든 문서는 열지 않는다")
            return "docx"
        if any(n.startswith("xl/") for n in names):
            if any(n.lower().endswith("vbaproject.bin") for n in names):
                raise UnsupportedSource("매크로가 든 문서는 열지 않는다")
            return "xlsx"
        raise UnsupportedSource("모르는 압축 형식이다")
    if b"\x00" in data[:1024]:
        raise UnsupportedSource(f"모르는 이진 형식이다 ({filename or '이름 없음'})")
    return "text"


def to_text(data: bytes, *, filename: str = "", see: Callable[[str, bytes], str] | None = None,
            check_halves: bool = True) -> SourceText:
    """글자로. 사진·스캔 쪽은 `see(prompt, png)`(비전 받아쓰기)가 있어야 한다 — 없으면 `UnsupportedSource`."""
    import time

    started = time.perf_counter()
    kind = sniff(data, filename=filename)
    if kind == "text":
        result = SourceText("text", _decode(data))
    elif kind == "docx":
        result = SourceText("docx", _docx(data))
    elif kind == "xlsx":
        result = SourceText("xlsx", _xlsx(data))
    elif kind == "pdf":
        result = _pdf(data, see=see, check_halves=check_halves)
    else:
        if see is None:
            raise UnsupportedSource("사진을 읽을 받아쓰기 모델이 설정되지 않았다")
        text, missing = _transcribe(data, see=see, check_halves=check_halves)
        result = SourceText("image", text, transcribed_pages=[1], missing=missing)
    result.seconds = round(time.perf_counter() - started, 2)
    return result


# ── 형식별 ───────────────────────────────────────────────────────
def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnsupportedSource("글자 인코딩을 알 수 없다")


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _xlsx(data: bytes) -> str:
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    lines: list[str] = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell(v) for v in row if v is not None and str(v).strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _cell(value: Any) -> str:
    from datetime import date, datetime, time

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M") if (value.hour or value.minute) else value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return str(value).strip()


def _pdf(data: bytes, *, see, check_halves: bool) -> SourceText:
    import fitz

    document = fitz.open(stream=data, filetype="pdf")
    if document.needs_pass:
        raise UnsupportedSource("암호 걸린 PDF 는 열지 않는다")
    parts: list[str] = []
    transcribed: list[int] = []
    missing: list[dict[str, Any]] = []
    for number, page in enumerate(document, start=1):
        text = page.get_text("text")
        if len(text.strip()) >= SCANNED_PAGE_CHARS:
            parts.append(text.rstrip("\n"))
            continue
        if see is None:
            raise UnsupportedSource(f"{number}쪽이 스캔인데 받아쓰기 모델이 설정되지 않았다")
        png = page.get_pixmap(dpi=RENDER_DPI).tobytes("png")
        written, lost = _transcribe(png, see=see, check_halves=check_halves)
        parts.append(written)
        transcribed.append(number)
        missing += [{**m, "page": number} for m in lost]
    return SourceText("pdf", "\n".join(parts), pages=len(document), transcribed_pages=transcribed, missing=missing)


def _see_once_more(see):
    """받아쓰기 호출을 **한 번만** 다시 한다 — 모델 서버의 일시 오류(HTTP 5xx · 빈 받아쓰기)만.
    ★읽기 전용 호출이라 다시 불러도 부작용이 없다. 두 번째도 실패하면 그대로 올린다(성공으로 추정하지 않는다).
    ☆2026-09-28 평가셋 60건: 일시 오류 3건(HTTP 500 1 · 빈 받아쓰기 2)이 사례를 통째로 잃게 했다(항목 21개)."""
    def call(prompt, image):
        try:
            return see(prompt, image)
        except Exception as exc:                  # noqa: BLE001 — 아래에서 고른 것만 다시 부른다
            text = str(exc)
            if "빈 받아쓰기" not in text and "HTTP 5" not in text:
                raise
            return see(prompt, image)
    return call


def _transcribe(image: bytes, *, see, check_halves: bool) -> tuple[str, list[dict[str, Any]]]:
    see = _see_once_more(see)
    full = see(TRANSCRIBE_PROMPT, image)
    if not check_halves:
        return full, []
    missing: list[dict[str, Any]] = []
    have = {_key(line) for line in full.splitlines() if _CHECK_LINE.search(line)}
    for name, half in _halves(image):
        written = see(TRANSCRIBE_PROMPT, half)
        for line in written.splitlines():
            if _CHECK_LINE.search(line) and _key(line) not in have:
                missing.append({"half": name, "text": line.strip()})
                have.add(_key(line))
    return full, missing


def _key(line: str) -> str:
    """반쪽 받아쓰기와 맞춰 볼 때의 열쇠 — 공백·구두점을 빼고 비교한다(글자 자체는 바꾸지 않는다)."""
    return re.sub(r"[\s·|,.:;()\[\]-]+", "", line)


def _halves(image: bytes) -> list[tuple[str, bytes]]:
    """위·아래 반쪽(가운데를 10% 겹친다 — 경계에 걸린 줄을 잃지 않게)."""
    from PIL import Image

    picture = Image.open(io.BytesIO(image))
    width, height = picture.size
    overlap = int(height * 0.05)
    out = []
    for name, box in (("top", (0, 0, width, height // 2 + overlap)),
                      ("bottom", (0, height // 2 - overlap, width, height))):
        buffer = io.BytesIO()
        picture.crop(box).save(buffer, format="PNG")
        out.append((name, buffer.getvalue()))
    return out


__all__ = ["SourceText", "TRANSCRIBE_PROMPT", "UnsupportedSource", "sniff", "to_text"]
