"""make_review_sheet — 분류 · 결 검수 시트(드롭다운 엑셀). `[2026-10-12]`

DB · 데이터 폴더를 읽지 않는다 — 가게 · 인허가 · 이전 시트를 시험이 만든다.
"""
from __future__ import annotations

import importlib.util
import os
import sys

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
MODULE = os.path.join(HERE, "..", "..", "..", "scripts", "dining", "make_review_sheet.py")


@pytest.fixture(scope="module")
def rs():
    spec = importlib.util.spec_from_file_location("dining_make_review_sheet", os.path.abspath(MODULE))
    mod = importlib.util.module_from_spec(spec)
    sys.modules["dining_make_review_sheet"] = mod
    spec.loader.exec_module(mod)
    return mod


def shop(name, *, category="한식", style=None, method=None, michelin=False, firstmenu=None, sources=("tourapi_kor_food",)):
    return {"name": name, "address": "서울특별시 종로구 종로 1", "road": "서울특별시 종로구 종로 1", "jibun": None,
            "category": category, "style": style, "style_method": method, "sources": list(sources),
            "michelin": michelin, "firstmenu": firstmenu}


def licence(rs, name, uptae):
    return {"관리번호": "x", "상호": name, "상태": "영업/정상", "도로명": "서울특별시 종로구 종로 1", "지번": "", "업태": uptae}


def test_결_초안은_업태가_결_하나로_모일_때만(rs):
    assert rs.style_draft(["냉면집"]) == "면"
    assert rs.style_draft(["횟집", "횟집"]) == "해산물"
    assert rs.style_draft(["한식"]) == ""                        # 넓은 업태는 비운다
    assert rs.style_draft(["냉면집", "분식"]) == ""              # 둘로 갈리면 비운다
    assert rs.style_draft([]) == ""


def test_줄은_미상_결없음_이전시트를_모두_담고_사람이_적은_칸을_옮긴다(rs):
    shops = {"u-cat": shop("미상집", category="미상", style="면"),
             "u-style": shop("냉면옥", firstmenu="물냉면"),
             "u-none": shop("정한집", method="manual"),             # 사람이 「없음」으로 정함 — 다시 묻지 않는다
             "u-done": shop("다된집", style="고기구이"),
             "u-prev": shop("이전집", style="면")}
    previous = [{"place_uid": "u-prev", "[확인] 분류": "한식", "확인일": "2026-10-13", "메모": "사람이 적음"},
                {"place_uid": "u-gone", "상호": "닫힌집", "확인일": ""}]
    rows, _ = rs.build_rows(shops, previous, rs.lc.index([licence(rs, "냉면옥", "냉면집")]))
    by = {r["place_uid"]: r for r in rows}
    assert set(by) == {"u-cat", "u-style", "u-prev", "u-gone"}
    assert by["u-style"]["결 초안"] == "면" and by["u-style"]["할 일"] == "[확인] 결에 적고 확인일"
    assert by["u-cat"]["할 일"].startswith("[확인] 분류에 적고 확인일")
    assert by["u-prev"]["[확인] 분류"] == "한식" and by["u-prev"]["확인일"] == "2026-10-13" and "사람이 적음" in by["u-prev"]["메모"]
    assert by["u-prev"]["할 일"].startswith("이미 채워짐")
    assert by["u-gone"]["할 일"] == "닫힌 가게 — 적지 않는다"
    assert [r["번호"] for r in rows] == [str(i) for i in range(1, len(rows) + 1)]


def test_엑셀은_드롭다운이_있고_적재기가_그대로_읽는다(rs, tmp_path):
    shops = {"u1": shop("미쉐린집", category="미상", michelin=True)}
    rows, _ = rs.build_rows(shops, [], rs.lc.index([]))
    path = str(tmp_path / "분류_검수_20261013_결.xlsx")
    rs.write_xlsx(rows, path)

    from openpyxl import load_workbook
    wb = load_workbook(path)
    assert wb.sheetnames == ["결 목록", rs.mc.XLSX_SHEET]
    lists = {str(d.sqref).split(":")[0][0]: d.formula1 for d in wb[rs.mc.XLSX_SHEET].data_validations.dataValidation
             if d.type == "list"}
    assert all(not f.startswith("=") for f in lists.values())    # 엑셀 XML 은 「=」 없이 적는다
    listed = [c.value for c in wb["결 목록"]["A"][1:1 + len(rs.mc.STYLES) + 1]]
    assert listed == [*rs.mc.STYLES, rs.mc.NO_STYLE]
    read = rs.mc.read_rows(path)
    assert read[0]["place_uid"] == "u1" and read[0]["묶음"] == "미쉐린"
    assert rs.mc.build(path)[1] == {"분류": 0, "결": 0, "확인 전": 1}


def test_결_설명은_결_목록과_빠짐없이_같다(rs):
    assert set(rs.STYLE_HELP) == {*rs.mc.STYLES, rs.mc.NO_STYLE}
