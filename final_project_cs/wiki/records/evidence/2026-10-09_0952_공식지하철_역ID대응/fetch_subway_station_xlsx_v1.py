from pathlib import Path
from datetime import datetime
import hashlib
import json
import httpx
import openpyxl

root = Path(__file__).resolve().parents[2]
out = root / 'datasets/mobility/raw/subway_realtime_station_20261009_v1'
out.mkdir(exist_ok=False)
url = 'https://datafile.seoul.go.kr/bigfile/iot/inf/nio_download.do?&useCache=false'
params = {'infId': 'OA-12764', 'seqNo': '', 'seq': '57', 'infSeq': '2'}
r = httpx.post(url, data=params, timeout=25, follow_redirects=True)
r.raise_for_status()
assert r.content.startswith(b'PK'), r.headers.get('content-type')
p = out / 'realtime_station_20260902.xlsx'
p.write_bytes(r.content)
meta = {'created_at': datetime.now().astimezone().isoformat(), 'source_page': 'https://data.seoul.go.kr/dataList/OA-12764/F/1/datasetView.do',
        'source_title': '실시간도착_역정보(20260902).xlsx', 'download_url': url, 'form': params,
        'sha256': hashlib.sha256(r.content).hexdigest(), 'bytes': len(r.content), 'status': r.status_code,
        'license': '공공누리 제1유형, 출처: 서울특별시', 'map_developer_api_calls': 0}
(out / 'acquisition-v1.json').write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding='utf8')
w = openpyxl.load_workbook(p, read_only=True, data_only=True)
print(json.dumps(meta, ensure_ascii=False))
for s in w:
    print(s.title, s.max_row, s.max_column)
    print(list(s.iter_rows(values_only=True))[:5])
