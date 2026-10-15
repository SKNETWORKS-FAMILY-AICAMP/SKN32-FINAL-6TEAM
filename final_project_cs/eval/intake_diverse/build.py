# -*- coding: utf-8 -*-
"""시험지 만들기 · 검사. `cases.py` 의 정답 정의 → `cases.jsonl` + 엑셀 원본(`inputs/*.xlsx`)

    python -m eval.intake_diverse.build                       # 만들기 + 검사
    python -m eval.intake_diverse.build --real-export <json>  # 실제 접수 원문이 내보낸 원문과 글자까지 같은지도 본다

★검사(어기면 멈춘다):
  - 정답 장소(`place`)는 입력 글(엑셀은 서버와 같은 `sources.to_text` 결과)의 **부분 문자열**이어야 한다 — 정답이 원문을 바꿔 적지 않게.
  - 이동의 from/to 도 원문의 부분 문자열이어야 한다.
  - 날짜 · 시각 모양(YYYY-MM-DD · HH:MM), 요청 · 빠진 값 · 이동 수단은 정해진 말만.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from pathlib import Path
import re
import sys

from .cases import ALL

HERE = Path(__file__).resolve().parent
INPUTS = HERE / "inputs"
OUT = HERE / "cases.jsonl"

REQUESTS = {"register", "plan_ask", "question", "correction"}
MISSING = {"first_day_date", "lodging_name", "booked_place_name"}
MODES = {"taxi", "subway", "bus", "transit", "walk", "car", "bike", "cablecar"}
KINDS = {"activity", "dining", "lodging"}
DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
HHMM = re.compile(r"^\d{2}:\d{2}$")


def xlsx_bytes(rows: list[list], merges: list[str]) -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "일정"
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            if value is not None and value != "":
                sheet.cell(row=r, column=c, value=value)
    for span in merges:
        sheet.merge_cells(span)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def source_text(case: dict) -> str:
    """서버가 읽는 글 — 글 · 메시지는 줄바꿈으로 잇고(★메시지 여러 개를 한 접수로 묶는 것은 이 실행기의 가정), 엑셀은 서버의 `to_text`."""
    parts: list[str] = []
    for one in case["inputs"]:
        if one["kind"] in ("text", "message"):
            parts.append(one["text"])
        elif one["kind"] == "xlsx":
            from app.domains.travel_ops.components.intake.sources import to_text

            path = INPUTS / f"{case['case_id']}.xlsx"
            parts.append(to_text(path.read_bytes(), filename=path.name).text)
        else:
            raise ValueError(f"{case['case_id']}: 모르는 입력 {one['kind']}")
    return "\n".join(parts)


def _plain(value):
    from datetime import date, datetime, time

    if isinstance(value, (datetime, date, time)):
        return {"__type__": type(value).__name__, "value": value.isoformat()}
    return value


def check(case: dict, text: str) -> list[str]:
    problems: list[str] = []
    g = case["gold"]
    for i, item in enumerate(g["items"]):
        if item["place"] not in text:
            problems.append(f"items[{i}].place 「{item['place']}」가 원문에 없다")
        if item["date"] and not DATE.match(item["date"]):
            problems.append(f"items[{i}].date 모양 {item['date']}")
        for key in ("start", "end"):
            if item[key] and not HHMM.match(item[key]):
                problems.append(f"items[{i}].{key} 모양 {item[key]}")
        if item["kind"] not in KINDS:
            problems.append(f"items[{i}].kind {item['kind']}")
    for m in g["moves"]:
        if m["mode"] not in MODES:
            problems.append(f"이동 수단 {m['mode']}")
        for key in ("from", "to"):
            if m[key] and m[key] not in text:
                problems.append(f"이동 {key} 「{m[key]}」가 원문에 없다")
    for r in g["requests"] + g["requests_optional"]:
        if r not in REQUESTS:
            problems.append(f"요청 {r}")
    for m in g["missing"] + g["missing_optional"]:
        if m not in MISSING:
            problems.append(f"빠진 값 {m}")
    return problems


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--real-export", help="x600 서버에서 내보낸 접수 원문 JSON(git 밖 datasets/) — 실제 원문이 글자까지 같은지 본다")
    args = parser.parse_args(argv)
    INPUTS.mkdir(exist_ok=True)
    exported: set[str] = set()
    if args.real_export:
        exported = {hashlib.md5((r.get("transcript") or "").encode("utf-8")).hexdigest()
                    for r in json.loads(Path(args.real_export).read_text(encoding="utf-8"))}
    lines, bad = [], 0
    for case in ALL:
        for one in case["inputs"]:
            if one["kind"] == "xlsx":
                (INPUTS / f"{case['case_id']}.xlsx").write_bytes(xlsx_bytes(one["rows"], one["merges"]))
        text = source_text(case)
        if case["origin"] == "real":
            digest = hashlib.md5(text.encode("utf-8")).hexdigest()
            case["provenance"]["transcript_md5"] = digest
            if exported and digest not in exported:
                print(f"{case['case_id']}: 실제 원문과 글자가 다르다(md5 {digest})")
                bad += 1
        problems = check(case, text)
        if problems:
            bad += 1
            print(case["case_id"], *problems, sep="\n  ")
        row = {**case, "source_text": text,
               "inputs": [{**one, "rows": [[_plain(v) for v in r] for r in one["rows"]]} if one["kind"] == "xlsx" else one
                          for one in case["inputs"]]}
        lines.append(json.dumps(row, ensure_ascii=False))
    OUT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    synthetic = sum(1 for c in ALL if c["origin"] == "synthetic")
    print(f"{len(ALL)}건(합성 {synthetic} · 실제 {len(ALL) - synthetic}) → {OUT.name}, 검사 실패 {bad}건")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
