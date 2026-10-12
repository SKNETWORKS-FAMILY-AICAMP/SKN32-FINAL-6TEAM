"""분류 · 음식 결 검수 시트(드롭다운 엑셀)를 만든다. `[2026-10-12]`

    python scripts/dining/make_review_sheet.py --db dining_rebuild           → category/분류_검수_<오늘>_결.xlsx
    python scripts/dining/make_review_sheet.py --db dining_rebuild --date 20261020

무엇을 담나.
    대표 분류가 미상인 가게 + 음식 결이 없는 가게(사람이 「없음」으로 정한 곳은 뺀다) + **이전 검수 시트의 모든 줄**.
    `make_category_sql.py` 는 가장 최근 검수 시트 하나만 읽는다 — 그래서 이전 시트의 줄을 하나도 빼지 않고 옮기고,
    사람이 적은 칸([확인] 분류 · [확인] 결 · 확인일 · 확인자 · 메모)과 분류 초안을 그대로 둔다. 새 가게가 들어온 뒤 다시 돌리면 된다.

초안.
    결 초안은 인허가 업태에서만 낸다(같은 주소 · 같은 상호 — `license_closure` 의 대조 함수). 업태가 결 하나로 정해질 때만
    (냉면집 → 면 · 횟집 → 해산물 · 호프 → 술집 …). 한식 · 일식 · 경양식처럼 넓은 업태는 초안을 비운다.
    분류 초안이 비었고 분류가 미상이면 결에서 낸다 — 결이 나라 하나로 정해질 때만(분식 → 한식 · 초밥 → 일식 …, `STYLE_CATEGORY`).
    고기구이 · 술집 · 닭요리처럼 나라가 갈리는 결은 비운다.
    미쉐린은 가이드 요리 종류를 메모에 적는다(결은 231 `style_of_guide` 가 분명한 것만 채운다 — 프렌치 · 컨템퍼러리 같은 것은 사람 몫).
    구글 · 카카오 값은 넣지 않는다(저장 금지) — 지도 링크만 둔다. 사람이 보고 판단한다.

엑셀.
    「검수」 시트의 노란 칸만 적는다. [확인] 분류 · [확인] 결은 「결 목록」 시트를 원본으로 한 드롭다운이고(목록 밖 값은 엑셀이 막는다),
    확인일은 2026-10-13 꼴만 받는다. 목록은 `make_category_sql.ALLOWED` · `STYLES` 그대로다(231 제약과 시험으로 맞춘다).

인허가 자료가 없으면 멈춘다 — 초안 없이 조용히 만들지 않는다.
"""
from __future__ import annotations

import argparse
import csv
import glob
import os
import sys
import urllib.parse
from collections import Counter
from datetime import date
from typing import Any

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import license_closure as lc  # noqa: E402
import make_category_sql as mc  # noqa: E402

CATEGORY_DIR = os.path.join(mc.DINING_DATA, "category")

#: 업태가 결 하나로 정해지는 것만. 없는 업태(한식 · 일식 · 경양식 · 중국식 …)는 초안을 비운다
UPTAE_STYLE = {"식육(숯불구이)": "고기구이", "냉면집": "면", "탕류(보신용)": "국물요리", "분식": "분식", "김밥(도시락)": "분식",
               "횟집": "해산물", "호프/통닭": "술집", "정종/대포집/소주방": "술집", "커피숍": "커피·차", "다방": "커피·차",
               "까페": "커피·차", "전통찻집": "커피·차", "제과점영업": "베이커리·디저트", "뷔페식": "뷔페",
               "패스트푸드": "버거·샌드위치", "외국음식전문점(인도,태국등)": "아시아요리"}
#: 결 → 대표 분류. 나라가 하나로 정해지는 결만. 없는 결(고기구이 · 면 · 해산물 · 닭요리 · 덮밥 · 아시아요리 · 술집 · 뷔페 ·
#: 코스·다이닝)은 분류 초안을 비운다 `[2026-10-12 사용자]`
STYLE_CATEGORY = {"국물요리": "한식", "한상차림": "한식", "분식": "한식", "족발·보쌈": "한식", "초밥": "일식",
                  "돈가스·튀김": "일식", "중식요리": "중식", "이탈리안": "양식", "버거·샌드위치": "양식", "브런치": "양식",
                  "커피·차": "카페디저트", "베이커리·디저트": "카페디저트"}
SOURCE_LABEL = {"tourapi_kor_food": "관광공사", "michelin_guide": "미쉐린", "nopo_kakao_curation": "노포", "vegan_curated": "비건",
                "halal_curated": "할랄", "localdata_food": "인허가", "localdata_rest": "인허가"}
#: 사람이 보기 좋게 묶는 순서 — 메뉴 원문이 있는 곳부터, 미쉐린은 따로
GROUPS = ("관광공사(메뉴 있음)", "비건 · 할랄", "노포", "기타", "미쉐린")
LIVE = ("영업/정상", "영업", "정상")
#: 사람이 적는 칸 — 이전 시트에서 그대로 옮긴다
CARRY = ("분류 초안", "초안 근거", "[확인] 분류", "[확인] 결", "확인일", "확인자", "메모")
INPUT = ("[확인] 분류", "[확인] 결", "확인일", "확인자", "메모")
COLUMNS = ["번호", "묶음", "할 일", "place_uid", "상호", "주소", "들어온 곳", "대표메뉴(관광공사)", "현재 분류", "분류 초안", "초안 근거",
           "[확인] 분류", "현재 결", "결 초안", "결 초안 근거", "[확인] 결", "확인일", "확인자", "메모", "카카오맵", "네이버 지도"]
STYLE_HELP = {
    "고기구이": "삼겹살 · 갈비 · 한우 · 곱창 · 숯불구이, 스테이크 · 바비큐", "국물요리": "국밥 · 탕 · 찌개 · 전골 · 설렁탕 · 해장국 · 샤브샤브",
    "면": "냉면 · 칼국수 · 막국수 · 국수 · 라멘 · 우동 · 짜장 · 짬뽕", "해산물": "횟집 · 생선구이 · 조개 · 낙지 · 게장 · 아구찜 · 복집",
    "한상차림": "한정식 · 백반 · 쌈밥 · 비빔밥 · 솥밥 · 정찬", "분식": "떡볶이 · 김밥 · 순대 · 튀김",
    "닭요리": "삼계탕 · 닭갈비 · 닭한마리 · 찜닭 · 치킨", "족발·보쌈": "족발 · 보쌈 · 수육", "초밥": "스시 · 초밥 · 오마카세",
    "돈가스·튀김": "돈가스 · 카츠 · 텐동", "덮밥": "규동 · 덮밥 · 돈부리",
    "중식요리": "탕수육 · 딤섬 · 마라 · 코스 중식 (짜장 · 짬뽕이 중심이면 면)", "이탈리안": "파스타 · 피자 · 리조또",
    "버거·샌드위치": "버거 · 샌드위치 · 핫도그", "브런치": "브런치 · 샐러드 · 오믈렛 · 팬케이크",
    "아시아요리": "쌀국수 · 커리 · 팟타이 · 케밥 · 중동 음식", "베이커리·디저트": "빵집 · 케이크 · 빙수 · 디저트",
    "커피·차": "카페 · 찻집", "술집": "호프 · 포차 · 막걸리집 · 이자카야 · 와인바",
    "뷔페": "뷔페 · 샐러드바 (무한리필 고깃집은 고기구이)", "코스·다이닝": "무엇의 코스인지 한 갈래로 말하기 어려운 코스 · 파인다이닝",
    mc.NO_STYLE: "보고도 어느 결인지 정할 수 없을 때 — 규칙이 다시 채우지 않는다",
}
GUIDE = [
    "적는 법",
    "· 노란 칸만 적는다: [확인] 분류 · [확인] 결 · 확인일 · 확인자 · 메모. 다른 칸은 고치지 않는다(place_uid 로 원장 가게를 찾는다).",
    "· [확인] 분류 · [확인] 결은 칸을 누르면 나오는 목록에서 고른다. 목록 밖 값이 있으면 원장 적재가 멈춘다.",
    "· 결은 대표메뉴(가게가 내세우는 것) 기준 한 가지. 나라는 말하지 않는다 — 일식 · 한식은 분류 칸이 가른다.",
    "· 코스 식당이라도 무엇의 코스인지 분명하면 그 결(오마카세 → 초밥 · 한정식 코스 → 한상차림).",
    "· 확인일은 2026-10-13 처럼 적는다. 확인일이 없는 줄은 원장에 들어가지 않는다.",
    "· 줄의 「할 일」이 적을 칸을 알려 준다. 「현재 결」이 맞으면 [확인] 결은 비워 둔다.",
    "",
    "예시 (시트에는 넣지 않았다)",
    "  상호 「노루정」 · 대표메뉴 「한우모듬」 → [확인] 분류 한식 · [확인] 결 고기구이 · 확인일 2026-10-13 · 확인자 이름",
    "  상호 「아리차이」 · 대표메뉴 「아리차이 코스요리」 → [확인] 분류 중식 · [확인] 결 코스·다이닝",
    "",
    "다 채우면 파일 이름을 바꾸지 말고 그대로 둔다 — make_category_sql.py 가 가장 최근 검수 시트(이 파일)를 읽는다.",
]


def load_license_any() -> list[dict[str, str]]:
    """인허가 자료 — 새 형식(`license_closure.load_license`)이 없으면 예전 형식(한글 칸 · CP949) 파일. 둘 다 없으면 멈춘다."""
    records = lc.load_license()
    if records:
        return records
    for name in ("서울시 일반음식점 인허가 정보.csv", "서울시 휴게음식점 인허가 정보.csv"):
        path = os.path.join(lc.RAW, name)
        if not os.path.exists(path):
            continue
        with open(path, encoding="cp949", errors="replace") as fh:
            records += [{"관리번호": r["관리번호"].strip(), "상호": r["사업장명"].strip(), "상태": r["영업상태명"].strip(),
                         "도로명": r["도로명주소"].strip(), "지번": r["지번주소"].strip(), "업태": r["업태구분명"].strip()}
                        for r in csv.DictReader(fh)]
    if not records:
        raise SystemExit(f"인허가 자료가 없다 — {lc.RAW} 에 「서울시 일반음식점 인허가 정보」를 받아 둔다(fetch_license.py)")
    return records


def style_draft(uptae: list[str]) -> str:
    """업태들이 결 하나로 모일 때만 그 결. 아니면 ""."""
    styles = {UPTAE_STYLE.get(u) for u in uptae}
    return next(iter(styles)) if len(styles) == 1 and None not in styles else ""


def task(needs_category: bool, needs_style: bool) -> str:
    fields = [name for name, need in (("[확인] 분류", needs_category), ("[확인] 결", needs_style)) if need]
    if not fields:
        return "이미 채워짐 — 현재 분류 · 결이 틀렸을 때만 [확인] 칸에 적고 확인일"
    text = " · ".join(fields) + "에 적고 확인일"
    return text + (" — 현재 결이 틀렸으면 [확인] 결도" if needs_category and not needs_style else "")


def group_of(sources: str, michelin: bool, firstmenu: str | None) -> str:
    if michelin:
        return "미쉐린"
    if "노포" in sources:
        return "노포"
    if "비건" in sources or "할랄" in sources:
        return "비건 · 할랄"
    return "관광공사(메뉴 있음)" if firstmenu else "기타"


def build_rows(shops: dict[str, dict[str, Any]], previous: list[dict[str, str]],
               license_by: Any) -> tuple[list[dict[str, str]], Counter]:
    """시트 줄. `shops` 는 영업 중 가게 {place_uid: {...}}, `previous` 는 이전 시트 줄(빼지 않는다)."""
    prev = {row["place_uid"]: row for row in previous if row.get("place_uid")}
    wanted = set(prev) | {uid for uid, s in shops.items()
                          if s["category"] in (None, "미상") or (s["style"] is None and s["style_method"] != "manual")}
    rows, n = [], Counter()
    for uid in wanted:
        shop, old = shops.get(uid), prev.get(uid, {})
        if shop is None:                              # 이전 시트에 있었는데 지금은 닫힌 가게 — 줄은 남기고 할 일은 없다
            n["닫힘(이전 줄)"] += 1
            rows.append({**{c: old.get(c, "") for c in COLUMNS}, "묶음": "기타", "할 일": "닫힌 가게 — 적지 않는다"})
            continue
        found = [r for r in lc.candidates(license_by, shop["road"], shop["jibun"]) if lc.same_name(shop["name"], r["상호"])]
        use = [r for r in found if r["상태"] in LIVE] or found
        uptae = sorted({r["업태"] for r in use if r["업태"]})
        needs_category, needs_style = shop["category"] in (None, "미상"), shop["style"] is None
        draft = style_draft(uptae) if needs_style else ""
        sources = " · ".join(dict.fromkeys(SOURCE_LABEL[s] for s in shop["sources"] if s in SOURCE_LABEL)) or "-"
        memo = [old["메모"]] if old.get("메모") else []
        if shop["michelin"]:
            cuisine = shop.get("cuisine")
            memo.append(f"미쉐린 요리 「{cuisine}」" if cuisine else "미쉐린 — 가이드 페이지의 요리 종류를 참고한다")
        category_draft = {}
        if needs_category and not old.get("분류 초안") and STYLE_CATEGORY.get(shop["style"]):
            category_draft = {"분류 초안": STYLE_CATEGORY[shop["style"]], "초안 근거": f"결 「{shop['style']}」"}
        if needs_style:
            memo.append(f"결이 없다고 보면 [확인] 결에 「{mc.NO_STYLE}」")
        query = urllib.parse.quote(f"{shop['name']} {shop['address'] or ''}".strip())
        rows.append({
            "묶음": group_of(sources, shop["michelin"], shop["firstmenu"]), "할 일": task(needs_category, needs_style),
            "place_uid": uid, "상호": shop["name"], "주소": shop["address"] or "", "들어온 곳": sources,
            "대표메뉴(관광공사)": shop["firstmenu"] or "", "현재 분류": shop["category"] or "",
            **{c: old.get(c, "") for c in CARRY}, **category_draft,
            "현재 결": shop["style"] or "", "결 초안": draft,
            "결 초안 근거": ("규칙(관광공사 메뉴 · 상호)" if not needs_style
                         else f"인허가 업태 「{' · '.join(uptae)}」" if uptae else "인허가 자료에서 같은 주소 · 같은 상호를 못 찾음"),
            "메모": " · ".join(dict.fromkeys(memo)),
            "카카오맵": f"https://map.kakao.com/?q={query}", "네이버 지도": f"https://map.naver.com/p/search/{query}"})
        n["결 초안 있음" if draft else "결 필요 · 초안 없음" if needs_style else "결 있음"] += 1
    rows.sort(key=lambda r: (GROUPS.index(r["묶음"]), not r["결 초안"], r["상호"]))
    for i, row in enumerate(rows, 1):
        row["번호"] = str(i)
    return rows, n


def write_xlsx(rows: list[dict[str, str]], path: str) -> None:
    """「결 목록」(드롭다운 원본 · 적는 법) + 「검수」(노란 칸 · 드롭다운 · 지도 링크)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    font, thin = "Arial", Side(style="thin", color="D9D9D9")
    wb = Workbook()
    ls = wb.active
    ls.title = "결 목록"
    ls.append(["결", "이런 가게"])
    styles = [*mc.STYLES, mc.NO_STYLE]
    for style in styles:
        ls.append([style, STYLE_HELP[style]])
    ls.append([])
    ls.append(["분류"])
    cat_start = ls.max_row + 1
    for cat in mc.ALLOWED:
        ls.append([cat])
    ls.append([])
    guide_start = ls.max_row + 1
    for line in GUIDE:
        ls.append([line])
    for row in ls.iter_rows():
        for cell in row:
            cell.font = Font(name=font, size=10)
    for ref in ("A1", "B1", f"A{cat_start - 1}", f"A{guide_start}"):
        ls[ref].font = Font(name=font, bold=True)
    ls.column_dimensions["A"].width, ls.column_dimensions["B"].width = 18, 60

    ws = wb.create_sheet(mc.XLSX_SHEET)
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row.get(c, "") for c in COLUMNS])
    last = len(rows) + 1
    for j, name in enumerate(COLUMNS, 1):
        head = ws.cell(1, j)
        head.font = Font(name=font, bold=True, color="FFFFFF")
        head.fill = PatternFill("solid", fgColor="BF8F00" if name in INPUT else "44546A")
        head.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = {"할 일": 34, "상호": 22, "주소": 30, "초안 근거": 30, "결 초안 근거": 26,
                                                            "메모": 34, "대표메뉴(관광공사)": 22}.get(name, 12)
        for i in range(2, last + 1):
            cell = ws.cell(i, j)
            cell.font, cell.border = Font(name=font, size=10), Border(bottom=thin)
            if name in INPUT:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
            if name in ("place_uid", "번호", "확인일"):
                cell.number_format = "@"                    # 확인일을 날짜로 바꾸지 않게
            if name in ("카카오맵", "네이버 지도") and cell.value:
                cell.hyperlink, cell.value = cell.value, "열기"
                cell.font = Font(name=font, size=10, color="0563C1", underline="single")
    ws.freeze_panes, ws.row_dimensions[1].height = "F2", 30
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"

    def column(name: str) -> str:
        return get_column_letter(COLUMNS.index(name) + 1)

    lists = (("[확인] 분류", f"'결 목록'!$A${cat_start}:$A${cat_start + len(mc.ALLOWED) - 1}", "분류"),
             ("[확인] 결", f"'결 목록'!$A$2:$A${1 + len(styles)}", "음식 결"))
    for name, source, title in lists:                    # formula1 은 「=」 없이 — 엑셀 XML 이 그렇게 적는다
        dv = DataValidation(type="list", formula1=source, allow_blank=True, showErrorMessage=True, showInputMessage=True,
                            errorTitle=f"{title} 값", error="목록에 있는 값만 쓸 수 있어요 — 「결 목록」 시트를 보세요.",
                            promptTitle=title, prompt="목록에서 고르세요")
        ws.add_data_validation(dv)
        dv.add(f"{column(name)}2:{column(name)}{last}")
    d = column("확인일")
    dv = DataValidation(type="custom", allow_blank=True, showErrorMessage=True, errorTitle="확인일",
                        error="확인일은 2026-10-13 처럼 적어 주세요.",
                        formula1=f'OR({d}2="",AND(LEN({d}2)=10,MID({d}2,5,1)="-",MID({d}2,8,1)="-"))')
    ws.add_data_validation(dv)
    dv.add(f"{d}2:{d}{last}")
    wb.active = 1
    wb.save(path)


def load_shops(dsn: str) -> dict[str, dict[str, Any]]:
    import psycopg

    with psycopg.connect(dsn) as conn:
        rows = conn.execute("""
            SELECT p.place_uid::text, p.name_ko, coalesce(p.road_address, p.jibun_address), p.road_address, p.jibun_address,
                   p.category, p.style, p.style_method,
                   (SELECT array_agg(DISTINCT r.source_code) FROM dining.dn_source_record r WHERE r.place_uid = p.place_uid),
                   EXISTS (SELECT 1 FROM dining.dn_attribute a WHERE a.place_uid = p.place_uid AND a.attr_code = 'michelin'
                            AND a.retired_at IS NULL),
                   (SELECT s.raw_json ->> 'firstmenu' FROM dining.dn_source_record s
                     WHERE s.place_uid = p.place_uid AND s.source_code = 'tourapi_kor_food' LIMIT 1),
                   (SELECT s.raw_json ->> '요리' FROM dining.dn_source_record s
                     WHERE s.place_uid = p.place_uid AND s.source_code = 'michelin_guide' LIMIT 1)
              FROM dining.dn_place p WHERE p.record_status <> 'closed' AND NOT p.is_synthetic""").fetchall()
    keys = ("name", "address", "road", "jibun", "category", "style", "style_method", "sources", "michelin", "firstmenu",
            "cuisine")
    return {row[0]: {**dict(zip(keys, row[1:])), "sources": row[8] or []} for row in rows}


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--db", help="원장 DB 이름(예: dining_rebuild). 없으면 core_db 규칙")
    ap.add_argument("--date", default=date.today().strftime("%Y%m%d"), help="파일 이름의 날짜(YYYYMMDD)")
    args = ap.parse_args()
    import core_db  # noqa: E402 — 같은 폴더(license_closure 와 같은 DB 이름 규칙)

    dsn = (core_db.dsn() if not args.db
           else f"postgresql://postgres@localhost:{os.environ.get('DINING_PG_PORT', '5433')}/{args.db}")
    sheets = sorted(path for pattern in mc.SHEETS for path in glob.glob(pattern))
    previous = mc.read_rows(sheets[-1]) if sheets else []
    out = os.path.join(CATEGORY_DIR, f"분류_검수_{args.date}_결.xlsx")
    if sheets and os.path.abspath(sheets[-1]) == os.path.abspath(out):
        raise SystemExit(f"{out} 가 이미 가장 최근 시트다 — 다른 --date 로 만든다(덮어쓰지 않는다)")
    rows, n = build_rows(load_shops(dsn), previous, lc.index(load_license_any()))
    os.makedirs(CATEGORY_DIR, exist_ok=True)
    write_xlsx(rows, out)
    print(f"이전 시트 {os.path.basename(sheets[-1]) if sheets else '-'} ({len(previous)}줄) → {out} ({len(rows)}줄) {dict(n)}")


if __name__ == "__main__":
    main()
