from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path(r"C:\Users\playdata2\.codex\visualizations\2026\09\11\01a08f9b-1114-71a0-ace8-2a38da4df27e\triPilot_경쟁서비스_기능비교_근거링크.docx")
SRC_DIR = Path(r"C:\Users\playdata2\Documents\final_workspace\team_branch\자료조사")


def file_uri(filename):
    return (SRC_DIR / filename).resolve().as_uri()


INTERNAL = {
    "김지혜": file_uri("김지혜.md"),
    "서유현": file_uri("서유현.md"),
    "송채영": file_uri("송채영.docx"),
    "최상욱": file_uri("최상욱.md"),
    "최연우": file_uri("최연우.txt"),
    "정세환": file_uri("정세환.txt"),
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def font_run(run, size=9, bold=False, color="000000"):
    run.font.name = "Malgun Gothic"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Malgun Gothic")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url, color="0563C1", underline=True):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Malgun Gothic")
    r_fonts.set(qn("w:hAnsi"), "Malgun Gothic")
    r_fonts.set(qn("w:eastAsia"), "맑은 고딕")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "17")
    r_pr.append(sz)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def cell_text(cell, text, bold=False, size=8.5, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    font_run(r, size=size, bold=bold, color=color)


def add_lines(cell, lines, size=8.3):
    cell.text = ""
    for idx, line in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run("• " + line)
        font_run(r, size=size)


def add_sources(cell, sources):
    cell.text = ""
    for idx, (label, url) in enumerate(sources):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        add_hyperlink(p, label, url)


def style_table(table, widths, header=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for row_idx, row in enumerate(table.rows):
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_margins(cell)
            if row_idx == 0 and header:
                set_cell_shading(cell, "17365D")
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F2F6FA")
    if header:
        repeat_header(table.rows[0])


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    font_run(r, size=15 if level == 1 else 11.5, bold=True)
    return p


services = [
    ("TripIt Pro", ["예약 이메일을 일정으로 자동 정리", "실시간 항공편 지연·취소·게이트 알림", "Fare Tracker, Seat Tracker, 대체 항공편 검색", "체크인·출발 알림, 공항 지도·환승·수하물 안내", "국가별 정보, 일정 자동 공유"], "항공편 운영 지원이 강점이다. 다른 예약에 미치는 연쇄 영향 검증은 내부 조사에서 확인되지 않았다.", [("공식 Pro 기능", "https://www.tripit.com/web/pro"), ("공식 기능 비교", "https://help.tripit.com/en/support/solutions/articles/103000063396-tripit-or-tripit-pro-"), ("내부 조사", INTERNAL["김지혜"])]),
    ("GOKO", ["AI 일정 생성", "여행 커뮤니티와 일정 연결", "계획 인식 및 이동수단 안내", "돌발상황 입력 시 우회 이동수단 안내"], "직접 테스트에서 일부 장소·시각·이동수단이 누락됐다.", [("공식 서비스", "https://go-ko.app/"), ("직접 테스트", INTERNAL["서유현"])]),
    ("Layla", ["취향·예산 기반 AI 일정", "호텔·항공·액티비티 고려", "상세 이동 안내", "예약 전문가·사업자 연결", "대화형 일정 수정"], "직접 테스트에서 역 개수와 명절 휴무일 안내 오류가 있었다.", [("공식 서비스", "https://layla.ai/"), ("직접 테스트", INTERNAL["서유현"])]),
    ("Wanderlog", ["AI 일정 초안과 추천", "일정·지도 통합", "예약 이메일 가져오기", "실시간 공동 편집", "예산·비용 분담, 체크리스트", "오프라인 이용과 경로 최적화"], "내부 테스트에서는 초안 이후 세부 수정이 사용자 작업 중심이었다.", [("공식 기능", "https://wanderlog.com/pages/help-center"), ("도움말", "https://help.wanderlog.com/hc/en-us"), ("직접 테스트", INTERNAL["서유현"])]),
    ("Booking.com", ["대화형 목적지·숙소 탐색", "개인 조건 기반 추천", "도시·국가·지역 일정 생성", "숙소 가격·상세 페이지·예약 흐름 연동", "생성형 AI 고객지원과 예약 문의"], "지역별 제공 범위가 다르며 내부 조사 당시 한국에서는 AI Planner를 확인하지 못했다.", [("AI Planner 발표", "https://news.booking.com/bookingcom-launches-new-ai-trip-planner-to-enhance-travel-planning-experience/"), ("Agentic AI 발표", "https://news.booking.com/bookingcom-debuts-agentic-ai-innovations-adding-to-its-robust-suite-of-genai-tools-for-customers/"), ("직접 테스트", INTERNAL["서유현"])]),
    ("Plzgo", ["도시·여행 스타일 기반 일별 AI 일정", "식당·카페·관광지 시간순 구성", "대중교통 이동시간·환승 안내", "장소 영업 여부와 영업시간 검증 배지", "카메라 기반 한국어 메뉴 번역"], "내부 테스트 당시 사용자 초안 입력과 생성 대안 수정이 어려웠다.", [("Google Play", "https://play.google.com/store/apps/details?id=app.plzgo"), ("직접 테스트", INTERNAL["서유현"])]),
    ("AskBami", ["한국 여행 AI 채팅", "교통 경로와 저장 장소", "카메라 번역과 문화·상황 설명", "스마트 여행 캘린더", "자연어 하루 일정과 돌발상황 대안"], "내부 테스트에서 세션 기억과 영업시간 반영에 한계가 있었다.", [("공식 서비스", "https://www.askbami.com/fr"), ("직접 테스트", INTERNAL["서유현"])]),
    ("GuideGeek", ["개인화 일정 생성", "실시간 여행 조언과 현지 추천", "다국어 대화", "항공·숙박·식당·경험 가격 비교와 예약 연결", "WhatsApp·Instagram·Messenger 제공"], "내부 테스트에서 실시간 교통 조회, 사전 알림, 변경 전 일정 회상은 확인되지 않았다.", [("공식 기능", "https://guidegeek.com/destinations"), ("공식 서비스", "https://guidegeek.com/"), ("직접 테스트", INTERNAL["서유현"])]),
    ("노닐다 Nonilda", ["13개 언어 여행계획", "AI 경로 최적화", "입장권·교통 일괄 예매", "여행사·가이드 매칭", "한국관광공사 TourAPI 기반 관광정보", "해외 카드·보안 결제"], "내부 테스트에서 원하지 않는 장소 삭제와 일정 수정 흐름이 원활하지 않았다.", [("공식 서비스", "https://www.vin-co.kr/"), ("Google Play", "https://play.google.com/store/apps/details?id=kr.vinco.nonilda"), ("직접 테스트", INTERNAL["서유현"])]),
    ("KAYAK Trips", ["예약 확인 메일·Gmail 동기화", "통합 일정 생성·관리", "항공편 지연·취소·게이트 실시간 업데이트", "가격 추적", "공유 일정과 공동 편집", "대화형 여행 검색"], "내부 테스트에서는 서울 실시간 교통·지하철 시간표와 명절 휴무 정보가 부정확했다.", [("공식 Trips", "https://www.kayak.com/trips"), ("공식 도움말", "https://www.kayak.com/c/help/account-trips/"), ("직접 테스트", INTERNAL["서유현"])]),
    ("Tripomatic", ["AI 일정 생성·수정", "장소·시각·메모·숙소 직접 편집", "장소 사진·설명·영업시간", "도보·차량·자전거·대중교통 경로", "비용 추정·수정", "공동 편집·예약 연결·오프라인 지도·내보내기"], "서울 테스트에서 대중교통 경로가 제한됐고 장소 영업시간이 다른 지도 서비스와 달랐다.", [("공식 기능", "https://tripomatic.com/en"), ("AI 도움말", "https://support.tripomatic.com/tripomatic-26/faq-ai-assistant.html"), ("직접 테스트", INTERNAL["최상욱"])]),
    ("IVisitKorea", ["무료 1~14일 한국 AI 일정", "날짜·예산·스타일 기반 개인화", "장소 추가·삭제·교체·순서 변경", "일정 저장·공유·공개 일정 복사", "예상 비용과 예약 파트너 연결"], "내부 테스트에서 확정되지 않은 경기와 조건이 불명확한 가격을 일정에 포함했다.", [("공식 FAQ", "https://www.ivisitkorea.com/faq/"), ("직접 테스트", INTERNAL["최상욱"])]),
    ("Triple", ["AI 일정 생성과 맞춤 정보·상품 추천", "선택 장소 기반 동선 구성", "동행자 공동 계획", "여행 가계부", "현지 이용자 정보 공유", "항공·호텔 예약"], "지속적인 일정 감시와 프롬프트 기반 변경 루프는 내부 조사에서 확인되지 않았다.", [("Google Play", "https://play.google.com/store/apps/details?id=com.titicacacorp.triple"), ("내부 조사", INTERNAL["최연우"])]),
    ("Mindtrip", ["대화형 맞춤 일정 생성", "사진·링크·스크린샷·PDF에서 일정 생성", "지도·사진·리뷰와 장소 추천", "실시간 공동 계획·그룹 채팅", "예약·영수증 업로드와 이메일 전달", "항공·호텔·식당·활동 예약 연결"], "내부 조사에서는 변화를 자동 추적하고 반영하는 구체적인 사후관리 기능과 외부 에이전트 지원을 확인하지 못했다.", [("공식 서비스", "https://mindtrip.ai/"), ("공식 FAQ", "https://resources.mindtrip.ai/travelers/help/traveler-faqs"), ("내부 조사", INTERNAL["최연우"])]),
    ("NOL", ["숙소·항공·레저·공연 예약", "지역·일정·동선·취향 기반 추천 추진", "예약상품과 일정 연결 추진", "상품 조회·주문·결제·취소 API", "가격·재고 동기화와 결과 콜백"], "공개 API는 상거래 기능 중심이며 지속적인 일정 운영 API는 내부 조사에서 확인되지 않았다.", [("통합 발표", "https://nol-universe.com/ko/newsroom/pressRelease/detail?prNo=2507"), ("공개 API", "https://open-api.swallow.goglobal.travel/"), ("내부 조사", INTERNAL["최연우"])]),
    ("마이리얼트립", ["항공·숙소·렌터카·투어 예약", "eSIM 판매", "여행 플래너 상담", "Custom Trip 맞춤 일정"], "외부 개인 에이전트 접근과 여행 중 지속 일정 관리는 내부 조사에서 확인되지 않았다.", [("공식 서비스", "https://www.myrealtrip.com/"), ("맞춤여행", "https://customtrip.myrealtrip.com/"), ("내부 조사", INTERNAL["최연우"])]),
    ("네이버 브랜드·쇼핑 커넥트", ["크리에이터와 브랜드 제휴", "상품별 전용 링크 발급", "콘텐츠·커머스 제휴", "성과 집계와 수수료 정산", "캠페인 관리와 메시지"], "여행 일정 운영 서비스가 아니라 상품 홍보·예약 전환 목적의 제휴 플랫폼으로 분류된다.", [("공식 소개", "https://brandconnect.naver.com/about/creator"), ("운영정책", "https://brandconnect.naver.com/service/policy/creator"), ("내부 조사", INTERNAL["최연우"])]),
    ("Gemini Spark", ["Maps·Flights·Hotels 실시간 데이터로 일정 생성", "Gmail 예약 정보 반영", "Google Docs 일정표 생성·수정", "Calendar 일정 추가·변경", "백그라운드 반복 작업", "Chrome에서 예약 입력 단계까지 수행"], "여러 Google 도구를 연결하는 방식이며 내부 조사에서는 독립된 여행 운영 플랫폼과 지속 감지·검증 루프를 확인하지 못했다.", [("Google 공식 설명", "https://blog.google/products-and-platforms/products/gemini/how-gemini-plans-trips/"), ("내부 조사", INTERNAL["최연우"])]),
    ("Expedia Romie", ["그룹 채팅 기반 여행계획", "이메일 여행정보 수집", "AI 맞춤 일정과 추천", "날씨·돌발상황 모니터링 및 대안 제시", "일정 실시간 업데이트", "예약 변경·취소 셀프서비스"], "공식 발표 당시 Romie는 실험용 알파였으며 기능별 제공 지역·시점이 달랐다.", [("공식 발표", "https://www.expedia.com/newsroom/spring-product-release-2024/"), ("내부 조사", INTERNAL["최연우"])]),
    ("Tripsy", ["AI 에이전트·어시스턴트 MCP 연결", "대화로 일정 생성·활동 추가·순서 변경", "날짜·활동·시간 수정", "여행·숙소·교통·비용 데이터 접근", "공개 API와 OAuth"], "에이전트 접근과 상태 CRUD는 강점이다. 자동 감지 후 전체 여행을 통제하는 지속 루프는 내부 조사에서 확인되지 않았다.", [("공식 AI·MCP", "https://tripsy.app/ai"), ("공식 API", "https://docs.api.tripsy.app/"), ("내부 조사", INTERNAL["최연우"])]),
    ("대한민국 구석구석", ["국내 관광지·음식점·숙박 정보", "테마·지역별 여행기사", "AI 콕콕플래너 지역 코스", "지도 경로와 연관 기사"], "공식 앱 설명은 관광정보 전반을 확인해 주며, AI 콕콕플래너 세부 기능은 내부 조사 기준이다.", [("공식 앱", "https://play.google.com/store/apps/details?id=com.visitkorea.kr"), ("내부 조사", INTERNAL["정세환"])]),
    ("스투비플래너", ["국가·테마별 여행 코스와 템플릿", "지도 기반 여행 루트와 교통정보", "항공권·호텔 검색·비교", "추천 숙소", "여행 경비 계산", "전문가 여행 의뢰"], "여행 전 계획·비교 기능이 중심이다.", [("공식 약관·기능 범위", "https://trip.stubbyplanner.com/terms"), ("항공권 비교", "https://flights.stubbyplanner.com/"), ("경비 계산", "https://www.stubbyplanner.com/kb/europe-budget-calculator.html"), ("내부 조사", INTERNAL["정세환"])]),
    ("레바캉스", ["도시·날짜 기반 AI 일정", "최적 동선과 명소 설계", "전 세계 도시 여행 가이드", "박물관·미술관 정보", "회원 일정·가이드북 공유", "외부 예약 링크"], "날씨·환율·교통 등 세부 정보 범위는 내부 조사 문서에 기록돼 있다.", [("공식 서비스", "https://www.lesvacances.co.kr/"), ("내부 조사", INTERNAL["정세환"])]),
    ("Stippl", ["AI 일별 일정과 경로", "편집·공유 일정", "예산·비용 관리", "짐 목록과 체크리스트", "예약 이메일 전달", "지도·여행 저널·PDF 내보내기"], "일정·예산·준비물을 한 앱에서 관리하는 기능이 강점이다.", [("AI Planner", "https://jecbmcms.stippl.io/ai-travel-planner"), ("짐 목록", "https://go.stippl.io/packing-list"), ("내부 조사", INTERNAL["정세환"])]),
    ("TripMazer", ["AI 여행계획·예약·재계획", "실시간 가격 비교와 여행 중 지원", "예약·문서 이해와 선호 기억", "날씨·지연 등 변경 대응", "예산·장소 추천", "실시간 AI 가격·일정 제안"], "실시간 재계획이 직접 경쟁 기능이다. 가격·일정 제안은 추정치이며 최종 가격·재고는 예약 파트너 기준이라고 명시한다.", [("공식 제품 설명", "https://mazerlabs.in/products/trip-mazer/"), ("가격·AI 면책", "https://tripmazer.com/legal/cancellation-refunds"), ("내부 조사", INTERNAL["정세환"])]),
    ("Trip.com", ["AI 맞춤 일정 생성", "기존 예약 가져오기", "지도 기반 경로 시각화", "일정 이름·순서·장소·메모 편집", "플로팅 AI의 실시간 추천", "예약·여행정보·고객상담 연결"], "내부 테스트에서는 일정 연동 자동 변경, 자동 검증, 대안 자동 반영이 확인되지 않았다.", [("Trip.Planner 발표", "https://www.trip.com/newsroom/trip-com-launches-trip-planner-smart-itineraries-tailored-to-your-travel-style-with-real-time-recommendations/"), ("TripGenie 발표", "https://www.trip.com/newsroom/introducing-tripgenie-groundbreaking-ai-travel-assistant/"), ("직접 테스트", INTERNAL["송채영"])]),
]


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = doc.styles
styles["Normal"].font.name = "Malgun Gothic"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
styles["Normal"].font.size = Pt(9.5)
for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 11.5)):
    st = styles[name]
    st.font.name = "Malgun Gothic"
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)

title_ppr = styles["Title"].element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(8)
r = title.add_run("triPilot 경쟁 서비스 기능 비교")
font_run(r, size=24, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)
r = p.add_run("서비스별 기능과 공식 근거 링크를 정리하고 triPilot triPilot의 차별점을 실무 관점에서 압축한 자료입니다. 조사 기준일은 2026년 9월 11일입니다.")
font_run(r, size=10)

add_heading(doc, "핵심 차별점 세 가지", 1)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(7)
r = p.add_run("개별 기능보다 아래 세 요소가 하나의 운영 흐름으로 결합된다는 점이 핵심입니다.")
font_run(r, size=9.5)

summary = doc.add_table(rows=1, cols=3)
headers = ["핵심", "triPilot이 하는 일", "기존 서비스와의 차이"]
for i, h in enumerate(headers):
    cell_text(summary.rows[0].cells[i], h, bold=True, size=9.2, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
summary_rows = [
    ("1  여행 전체 지속 감시", "항공·숙박·활동·이동·예산을 하나의 상태로 유지하고 변경이 남은 일정에 미치는 영향을 다시 계산한다.", "TripIt은 항공편, TripMazer는 날씨·지연 대응이 강하지만 여러 예약의 연쇄 충돌을 결정론적으로 검증하는 구조는 공개 자료에서 확인되지 않았다."),
    ("2  검증된 대안만 제시", "LLM이 만든 후보를 영업시간·이동시간·예산·예약조건과 대조한다. 위반 후보는 폐기하고 검증 불가 시 결과를 내보내지 않는다.", "여러 서비스의 직접 테스트에서 휴무일·운영시간·경기 일정 오류가 확인됐다. TripMazer도 AI 가격·일정 제안을 추정치로 규정한다."),
    ("3  개인 에이전트용 상태형 API", "개인 에이전트가 예약·선호·피드백을 전달하면 triPilot이 이를 지속 상태에 반영하고 수락·거절 이후에도 재계획한다.", "Tripsy는 MCP·API와 CRUD를 제공하지만 자동 감지·영향 분석·검증·재제안을 묶은 운영 루프는 공개 자료에서 확인되지 않았다."),
]
for a, b, c in summary_rows:
    cells = summary.add_row().cells
    cell_text(cells[0], a, bold=True, size=9.3, color="17365D")
    cell_text(cells[1], b, size=9)
    cell_text(cells[2], c, size=9)
style_table(summary, [Inches(1.35), Inches(2.45), Inches(3.15)])

add_heading(doc, "실무용 한 문장", 1)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(7)
r = p.add_run("triPilot은 여행 중 변화를 전체 예약과 대조하고 실제 데이터 검증을 통과한 대안만 실행 후보로 제시하는 여행 운영 서비스다.")
font_run(r, size=11, bold=True, color="17365D")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("표현 주의  ")
font_run(r, size=9, bold=True)
r = p.add_run("triPilot 기능은 제공 문서에서 설계 또는 제공 예정으로 표현돼 있다. 현재 구현 범위와 검증 결과를 별도로 표시해야 한다.")
font_run(r, size=9)

land = doc.add_section(WD_SECTION.NEW_PAGE)
land.orientation = WD_ORIENT.LANDSCAPE
land.page_width, land.page_height = land.page_height, land.page_width
land.top_margin = Cm(1.35)
land.bottom_margin = Cm(1.35)
land.left_margin = Cm(1.4)
land.right_margin = Cm(1.4)
land.footer.is_linked_to_previous = False


def configure_landscape(sec):
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21.0)
    sec.top_margin = Cm(1.35)
    sec.bottom_margin = Cm(1.35)
    sec.left_margin = Cm(1.4)
    sec.right_margin = Cm(1.4)

add_heading(doc, "서비스별 기능과 근거", 1)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
r = p.add_run("공식 자료는 서비스가 공개적으로 설명한 기능을 뜻합니다. 직접 테스트의 오류·제한은 팀 조사 문서를 근거로 별도 표시했습니다.")
font_run(r, size=9)

for group_index, service_group in enumerate([services[:13], services[13:-3], services[-3:]]):
    if group_index:
        doc.add_page_break()
        continued = add_heading(doc, "서비스별 기능과 근거 계속", 2)
        continued.paragraph_format.space_before = Pt(0)
    table = doc.add_table(rows=1, cols=4)
    for i, h in enumerate(["서비스", "확인된 기능", "직접 테스트 및 한계", "근거 링크"]):
        cell_text(table.rows[0].cells[i], h, bold=True, size=9, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for service, features, limitation, sources in service_group:
        row = table.add_row()
        cells = row.cells
        cell_text(cells[0], service, bold=True, size=8.7, color="17365D")
        add_lines(cells[1], features, size=8.1)
        cell_text(cells[2], limitation, size=8.1)
        add_sources(cells[3], sources)
    style_table(table, [Inches(1.35), Inches(4.25), Inches(2.7), Inches(1.55)])

doc.add_page_break()
diff_heading = add_heading(doc, "세 가지 차별점의 근거 비교", 1)
diff = doc.add_table(rows=1, cols=4)
for i, h in enumerate(["차별점", "경쟁 서비스에서 확인된 수준", "triPilot 설계", "핵심 근거"]):
    cell_text(diff.rows[0].cells[i], h, bold=True, size=9, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
diff_rows = [
    ("여행 전체 지속 감시", "TripIt은 항공 상태, KAYAK은 예약·항공 업데이트, Expedia와 TripMazer는 날씨·돌발상황 대응을 제공한다. 공개 근거에서는 변경 이후 모든 예약·동선·예산을 함께 재검증하는 과정이 명확하지 않다.", "변경 이벤트가 발생하면 항공·숙박·활동·이동시간·예산을 다시 대조해 남은 여행 전체의 실행 가능성을 판정한다.", [("TripIt", "https://www.tripit.com/web/pro"), ("KAYAK", "https://www.kayak.com/trips"), ("Expedia", "https://www.expedia.com/newsroom/spring-product-release-2024/"), ("TripMazer", "https://mazerlabs.in/products/trip-mazer/"), ("triPilot 내부 설계", INTERNAL["김지혜"])]),
    ("검증된 대안만 제시", "일부 서비스는 최신 정보 또는 실시간 추천을 표방하지만 직접 테스트에서 휴무일·영업시간·경기 일정 오류가 관찰됐다. TripMazer는 AI 가격과 일정이 추정치임을 명시한다.", "LLM은 후보를 만들고, 코드는 DB의 영업시간·이동시간·예산·원 예약값과 대조한다. 위반 후보를 폐기하고 검증되지 않은 값은 출력하지 않는다.", [("TripMazer 면책", "https://tripmazer.com/legal/cancellation-refunds"), ("테스트 보고서", INTERNAL["서유현"]), ("Tripomatic·IVisitKorea 테스트", INTERNAL["최상욱"]), ("triPilot 검증 설계", INTERNAL["정세환"])]),
    ("개인 에이전트용 상태형 API", "Tripsy는 외부 AI와 연결되는 MCP·API를 제공하며 일정 CRUD가 가능하다. NOL API는 상품·주문·결제 중심이다. 다수 서비스는 자체 앱·웹·채팅 인터페이스 중심이다.", "개인 에이전트가 희망 장소·예약·선호·피드백을 전달하면 서버가 여행 상태에 반영한다. 이후 변경 감지, 영향 분석, 대안 검증, 수락·거절 반영을 반복한다.", [("Tripsy MCP", "https://tripsy.app/ai"), ("Tripsy API", "https://docs.api.tripsy.app/"), ("NOL API", "https://open-api.swallow.goglobal.travel/"), ("triPilot 내부 설계", INTERNAL["최연우"])]),
]
for a, b, c, links in diff_rows:
    cells = diff.add_row().cells
    cell_text(cells[0], a, bold=True, size=9, color="17365D")
    cell_text(cells[1], b, size=8.5)
    cell_text(cells[2], c, size=8.5)
    add_sources(cells[3], links)
style_table(diff, [Inches(1.45), Inches(3.35), Inches(3.35), Inches(1.7)])

add_heading(doc, "판단 기준", 1)
for text in [
    "기능 존재 여부는 공식 서비스 페이지, 공식 발표, 도움말, API 문서를 우선 사용했다.",
    "공식 설명과 직접 사용 결과가 다르면 두 근거를 함께 표시했다.",
    "미확인은 기능이 없다는 뜻이 아니라 제공 자료와 공개 근거에서 확인하지 못했다는 뜻이다.",
    "예정·베타·알파 기능은 현재 상용 기능과 구분해야 한다.",
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("• " + text)
    font_run(r, size=9)

for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    for run in list(p.runs):
        p._p.remove(run._r)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("triPilot 경쟁 서비스 기능 비교  |  2026-09-11")
    font_run(r, size=8, color="666666")

doc.save(OUT)
print(OUT)
