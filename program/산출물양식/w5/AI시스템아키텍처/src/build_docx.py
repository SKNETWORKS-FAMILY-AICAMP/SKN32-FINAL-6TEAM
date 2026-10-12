# -*- coding: utf-8 -*-
"""[모델링 및 평가] AI 시스템 아키텍처(멀티 에이전트 아키텍처) — 부트캠프 양식에 채워 docx 로 만든다.

    python build_docx.py

- 양식 파일을 그대로 열어 머리 표 · 정보 표 · 절 제목 띠 표는 두고, 절마다 본문을 새로 채운다.
- 문단 · 글머리표(● / ○) · 표 · 그림 칸은 양식에 있던 것을 복제해 서식을 그대로 쓴다.
- 그림은 ../그림/*_v1.png (make_figs.py 가 만든다).
- 사실 근거는 그림 스크립트 머리 주석과 ../근거조사_*.md 에 있다.
"""
import copy
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
FIG = ROOT / "그림"
TEMPLATE = ROOT.parent.parent / "[모델링 및 평가] AI 시스템 아키텍처 (멀티 에이전트 아키텍처)_양식.docx"
OUT = ROOT / "[모델링 및 평가] AI 시스템 아키텍처 (멀티 에이전트 아키텍처)_6팀.docx"

from doc_content import S1, S2, S3, S4, S5_AFTER, S5_P, S5_T, S6  # noqa: E402

sys.path.insert(0, str(ROOT.parent / "_공용"))
import term_pass as TP  # noqa: E402

_PARTICLE = re.compile(r"(?<=[A-Za-z0-9)])\s+(?=(?:은|는|이|가|을|를|의|에|에서|에게|로|으로|와|과|도|만|까지|부터|이며|이고|입니다)(?=[\s,.)」]|$))")


def tidy(text):
    """영어 낱말 뒤 조사를 붙이고, 가운뎃점 앞뒤 빈칸을 없앤다."""
    return _PARTICLE.sub("", text.replace(" · ", "·"))


# ───────────────────────── 조립 ─────────────────────────
def set_runs(p_el, parts):
    """p_el 의 run 들을 지우고 (rPr 원본, 글) 목록으로 다시 채운다."""
    for r in p_el.findall(qn("w:r")):
        p_el.remove(r)
    for proto, text in parts:
        r = copy.deepcopy(proto)
        for t in r.findall(qn("w:t")):
            r.remove(t)
        t = r.makeelement(qn("w:t"), {})
        t.text = tidy(text)
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
        p_el.append(r)
    return p_el


def main():
    d = docx.Document(str(TEMPLATE))
    body = d.element.body
    paras = d.paragraphs
    T = d.tables
    # 원형
    P_H3 = paras[[p.text for p in paras].index("1.1 서비스 개요 및 목적")]._p
    P_TXT = P_H3.getnext()
    p_users = next(p for p in paras if p.text.startswith("주요 사용자"))
    p_sub = next(p for p in paras if p.text.startswith("복합 질의 분기"))
    R_LABEL, R_TEXT = p_users.runs[0]._r, p_users.runs[2]._r
    P_L0, P_L1 = p_users._p, p_sub._p
    R_PLAIN = paras[[p.text for p in paras].index(P_TXT.xpath("string(.)"))].runs[0]._r if P_TXT.xpath("string(.)") else R_TEXT
    R_H3 = P_H3.findall(qn("w:r"))[0]
    TBL = T[5]._tbl
    BOX = T[4]._tbl
    heads = {1: T[2]._tbl, 2: T[3]._tbl, 3: T[6]._tbl, 4: T[7]._tbl, 5: T[8]._tbl, 6: T[10]._tbl}
    protos = [copy.deepcopy(x) for x in (P_H3, P_TXT, P_L0, P_L1, TBL, BOX)]
    P_H3, P_TXT, P_L0, P_L1, TBL, BOX = protos
    R_LABEL, R_TEXT, R_H3 = copy.deepcopy(R_LABEL), copy.deepcopy(R_TEXT), copy.deepcopy(R_H3)
    R_PLAIN = copy.deepcopy(R_PLAIN)

    # 머리 표 · 정보 표
    for p in T[0].rows[0].cells[0].paragraphs:
        for r in p.runs:
            r.text = r.text.replace("00기", "32기").replace("_____팀", "6팀(A-COPilot)")
    meta = {"산출물 단계": "모델링 및 평가", "제출 일자": "2026. 10. 6.",
            "깃허브 경로": "https://github.com/SKNETWORKS-FAMILY-AICAMP/SKN32-FINAL-6TEAM", "작성 팀원": "6팀(A-COPilot)"}
    for row in T[1].rows:
        k = row.cells[0].text.strip()
        if k in meta:
            c = row.cells[1]
            ps = c.paragraphs
            if ps[0].runs:
                ps[0].runs[0].text = meta[k]
                for r in ps[0].runs[1:]:
                    r.text = ""
            else:
                ps[0].add_run(meta[k])

    # 절 제목 띠 사이의 옛 본문을 지운다(머리 표 · 정보 표 · 띠 표 · sectPr 는 남김)
    keep = set(heads.values()) | {T[0]._tbl, T[1]._tbl}
    started = False
    for el in list(body.iterchildren()):
        if el is heads[1]:
            started = True
            continue
        if not started or el in keep or el.tag == qn("w:sectPr"):
            continue
        body.remove(el)

    usable = Cm(17.0)

    def para(text, center=False, size=None, italic=False):
        el = copy.deepcopy(P_TXT)
        set_runs(el, [(R_PLAIN, text)])
        p = docx.text.paragraph.Paragraph(el, d._body)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            if size:
                r.font.size = Pt(size)
            if italic:
                r.italic = True
        return el

    def h3(text):
        el = set_runs(copy.deepcopy(P_H3), [(R_H3, text)])
        docx.text.paragraph.Paragraph(el, d._body).paragraph_format.keep_with_next = True
        return el

    def bullet(label, text, lvl):
        el = copy.deepcopy(P_L1 if lvl else P_L0)
        parts = ([(R_LABEL, label)] if label else []) + ([(R_TEXT, (" " if label else "") + text)] if text else [])
        return set_runs(el, parts)

    def picture(fname, caption, max_h_cm=22.0):
        from PIL import Image
        path = FIG / fname
        w, h = Image.open(path).size
        width = usable
        if (h / w) * 17.0 > max_h_cm:
            width = Cm(max_h_cm * w / h)
        el = para("", center=True)
        p = docx.text.paragraph.Paragraph(el, d._body)
        p.runs[0].add_picture(str(path), width=width) if p.runs else p.add_run().add_picture(str(path), width=width)
        p.paragraph_format.keep_with_next = True
        return [el, para(caption, center=True, size=9)]

    def box(fname, caption):
        el = copy.deepcopy(BOX)
        tbl = docx.table.Table(el, d._body)
        cell = tbl.rows[0].cells[0]
        cp = cell.paragraphs[0]
        for r in cp.runs:
            r.text = ""
        from PIL import Image
        path = FIG / fname
        w, h = Image.open(path).size
        width = Cm(16.4) if (h / w) * 16.4 <= 20 else Cm(20 * w / h)
        (cp.runs[0] if cp.runs else cp.add_run()).add_picture(str(path), width=width)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return [el, para(caption, center=True, size=9)]

    def table(head, rows, ratio):
        el = copy.deepcopy(TBL)
        trs = el.findall(qn("w:tr"))
        h_tr, b_tr = trs[0], trs[1]
        h_tc = h_tr.findall(qn("w:tc"))[0]
        b_tcs = b_tr.findall(qn("w:tc"))  # 첫 칸(굵게) · 나머지 칸(보통)
        for tr in trs:
            el.remove(tr)
        total = 9510
        widths = [int(total * r / sum(ratio)) for r in ratio]
        grid = el.find(qn("w:tblGrid"))
        for g in list(grid):
            grid.remove(g)
        for wv in widths:
            g = grid.makeelement(qn("w:gridCol"), {qn("w:w"): str(wv)})
            grid.append(g)

        long_cols = {i for i in range(len(head)) if max(len(r[i]) for r in rows) > 40}

        def make_tr(proto_tr, proto_tc, cells, body=False):
            tr = copy.deepcopy(proto_tr)
            for tc in tr.findall(qn("w:tc")):
                tr.remove(tc)
            for ci, (wv, text) in enumerate(zip(widths, cells)):
                tc = copy.deepcopy((b_tcs[0] if ci == 0 else b_tcs[-1]) if body else proto_tc)
                tcpr = tc.find(qn("w:tcPr"))
                if tcpr is not None:
                    tw = tcpr.find(qn("w:tcW"))
                    if tw is not None:
                        tw.set(qn("w:w"), str(wv)); tw.set(qn("w:type"), "dxa")
                    gs = tcpr.find(qn("w:gridSpan"))
                    if gs is not None:
                        tcpr.remove(gs)
                ps = tc.findall(qn("w:p"))
                for extra in ps[1:]:
                    tc.remove(extra)
                p = ps[0]
                rs = p.findall(qn("w:r"))
                proto_r = rs[0] if rs else R_TEXT
                set_runs(p, [(proto_r, text)])
                if body and ci in long_cols:
                    docx.text.paragraph.Paragraph(p, d._body).alignment = WD_ALIGN_PARAGRAPH.LEFT
                tr.append(tc)
            return tr

        el.append(make_tr(h_tr, h_tc, head))
        for row in rows:
            el.append(make_tr(b_tr, None, row, body=True))
        return el

    def render(blocks):
        out = []
        for b in blocks:
            k = b[0]
            if k == "h":
                out.append(h3(b[1]))
            elif k == "p":
                out.append(para(b[1]))
            elif k == "b":
                out.append(bullet(b[1], b[2], b[3]))
            elif k == "t":
                if out:
                    docx.text.paragraph.Paragraph(out[-1], d._body).paragraph_format.keep_with_next = True
                out += [table(b[1], b[2], b[3]), para("")]
            elif k == "img":
                out += picture(b[1], b[2])
            elif k == "box":
                out += box(b[1], b[2])
        return out

    S5 = [("p", S5_P), ("t",) + S5_T, ("p", S5_AFTER)]
    # 용어 정비(2026-10-06 사용자 지적 「너만 아는 단어 다 수정」): 실무 용어 · 1절 맨 앞 「용어 정의」 표 · 줄표 없앰
    def pre(x):  # 문서에 실린 모양(tidy 뒤)과 같게 맞춘 뒤 다시 쓰기를 건다
        if isinstance(x, str):
            return tidy(x)
        if isinstance(x, (list, tuple)):
            return [pre(v) for v in x]
        return x
    secs = [[[b[0], b[1]] + pre(list(b[2:])) if b[0] in ("img", "box") else pre(b) for b in s] for s in (S1, S2, S3, S4, S5, S6)]
    secs, miss = TP.run_blocks(secs, "arch")
    print("용어 정비: 못 찾은 문장", len(miss))

    def nodash(x):
        if isinstance(x, str):
            x = x.replace(" — ", ": ").replace("—", "없음")
            return x
        if isinstance(x, list):
            return [nodash(v) for v in x]
        return x
    secs = [[[b[0], b[1]] + nodash(list(b[2:])) if b[0] in ("img", "box") else nodash(b) for b in s] for s in secs]
    for n, blocks in zip(range(1, 7), secs):
        anchor = heads[n]
        for el in reversed(render(blocks) + [para("")]):
            anchor.addnext(el)
    # 양식의 파란 글자(예시 표시)를 검정으로
    for c in body.iter(qn("w:color")):
        if (c.get(qn("w:val")) or "").upper() == "0000FF":
            c.set(qn("w:val"), "000000")
    d.save(str(OUT))
    print("saved", OUT.name)


if __name__ == "__main__":
    main()
