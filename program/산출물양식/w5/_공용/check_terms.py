"""정비 뒤 남은 내부 말 · 금지 표현을 센다 (2026-10-06). 사용: python check_terms.py <docx|xlsx> ..."""
from __future__ import annotations

import re
import sys

WATCH = ["짚", "거름", "가르친", "가르치", "독립 ", "색인 글", "정답 줄", "문턱", "새어", "요청 둘", "한 줄", "섞은", "판별기", "정본", "§",
         "이쪽만", "수정 판", "사전 등록 판", "보낼 곳", "끼움자리", "되잡기", "바깥함", "일꾼", "국적", "4인", "목 서버", "가짜 서버"]
BARE_JEV = re.compile(r"(?<![A-Za-z가-힣])Jev(?!\s*계열|\s*API|Bench|K5|-Omni|ify|bench)")


def text_of(path):
    if path.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        return "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str))
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for r in tb.rows:
            for c in r.cells:
                parts.append(c.text)
    return "\n".join(parts)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    for path in sys.argv[1:]:
        t = text_of(path)
        hits = {w: t.count(w) for w in WATCH if w in t}
        print(path.split("/")[-1][:40], "남은 말:", hits, "| Jev 단독:", len(BARE_JEV.findall(t)), "| 줄표:", t.count("—"))
        if "-v" in sys.argv[0:1]:
            pass


if __name__ == "__main__":
    main()
