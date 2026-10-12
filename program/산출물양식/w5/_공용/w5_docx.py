# -*- coding: utf-8 -*-
"""5주차 산출물 docx 공용 빌더 — 부트캠프 양식을 열어 표지 · 정보 표 · 절 제목 띠는 두고 본문만 채운다.

블록:
  ("h", 제목)                       소제목(양식의 Heading 3 서식 복제)
  ("p", 글)                         본문 문단
  ("b", 라벨, 글, 단계)             글머리표(0 = ●, 1 = ○). 라벨은 굵게
  ("t", 머리줄, 행들, 비율)         표(양식의 첫 데이터 표 서식 복제)
  ("img", 그림경로, 캡션)           그림 + 캡션
  ("code", 머리글, 코드글)          소스코드 발췌(고정폭 작은 글씨, 회색 바탕)
  ("note", 글)                      작은 회색 글(출처 · 주의)

같은 날 다른 세션이 만든 w6/AI시스템아키텍처/src/build_docx.py 의 방식(양식 서식 복제)을 일반화했다.
"""
import copy
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

GITHUB = "https://github.com/SKNETWORKS-FAMILY-AICAMP/SKN32-FINAL-6TEAM"
TEAM = "6팀(A-COPilot)"
DATE = "2026. 10. 6."
FONT = "맑은 고딕"
BANNED = {"—": "줄표", "–": "줄표", "★": "별", "☆": "별", "✓": "체크", "✔": "체크", "…": "말줄임표"}


def check(text):
    for ch, name in BANNED.items():
        if ch in text:
            raise ValueError(f"금지 문자 {ch}({name}): {text[:60]}")
    return text


def _font(run, size=None, bold=None, mono=False, color=None):
    name = "Consolas" if mono else FONT
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(k), name)
    fonts.set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_runs(p_el, parts):
    for r in p_el.findall(qn("w:r")):
        p_el.remove(r)
    for proto, text in parts:
        r = copy.deepcopy(proto)
        for t in r.findall(qn("w:t")):
            r.remove(t)
        t = r.makeelement(qn("w:t"), {})
        t.text = check(text)
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(t)
        p_el.append(r)
    return p_el


def build(template, out, *, cover_replace=None, stage="모델링 및 평가", banners=None, sections, max_img_h_cm=21.0):
    """sections: 절 제목 띠 순서대로 블록 목록의 목록. banners: {띠 순번: 새 제목}(양식 번호 오류 바로잡기용)."""
    d = docx.Document(str(template))
    body = d.element.body
    T = d.tables
    # 표지 · 정보 표
    for p in T[0].rows[0].cells[0].paragraphs:
        for r in p.runs:
            r.text = r.text.replace("00기", "32기").replace("_____팀", TEAM).replace("___팀", TEAM)
            for old, new in (cover_replace or []):
                r.text = r.text.replace(old, new)
    meta = {"산출물 단계": stage, "제출 일자": DATE, "깃허브 경로": GITHUB, "작성 팀원": TEAM}
    for row in T[1].rows:
        k = row.cells[0].text.strip()
        if k in meta:
            ps = row.cells[1].paragraphs
            if ps[0].runs:
                ps[0].runs[0].text = meta[k]
                for r in ps[0].runs[1:]:
                    r.text = ""
            else:
                _font(ps[0].add_run(meta[k]), size=10)
    # 절 제목 띠(1행 · 셀 하나가 가로로 합쳐진 표, 표지 · 정보 표 다음)
    heads = []
    for t in T[2:]:
        if len(t.rows) == 1 and len({id(c._tc) for c in t.rows[0].cells}) == 1 and t.rows[0].cells[0].text.strip()[:1].isdigit():
            heads.append(t._tbl)
    if len(heads) != len(sections):
        raise ValueError(f"절 띠 {len(heads)}개인데 내용은 {len(sections)}절")
    for h in heads:  # 절 제목 띠가 쪽 끝에 홀로 남지 않게 다음 내용과 붙인다
        for tc in docx.table.Table(h, d._body).rows[0].cells[:1]:
            for p in tc.paragraphs:
                p.paragraph_format.keep_with_next = True
    for i, new in (banners or {}).items():
        tbl = docx.table.Table(heads[i], d._body)
        p = tbl.rows[0].cells[0].paragraphs[0]
        p.runs[0].text = new
        for r in p.runs[1:]:
            r.text = ""
    # 원형: 소제목 · 데이터 표
    h3 = next((p for p in d.paragraphs if p.style.name.lower().startswith("heading 3")), None) or \
        next(p for p in d.paragraphs if p.style.name.lower().startswith("heading"))
    P_H3 = copy.deepcopy(h3._p)
    R_H3 = copy.deepcopy(h3.runs[0]._r)
    data_tbl = next(t for t in T[2:] if t._tbl not in heads and len(t.rows) >= 2 and len(t.columns) >= 3)
    TBL = copy.deepcopy(data_tbl._tbl)
    # 띠 사이 옛 본문 지우기
    keep = set(heads) | {T[0]._tbl, T[1]._tbl}
    started = False
    for el in list(body.iterchildren()):
        if el is heads[0]:
            started = True
            continue
        if not started or el in keep or el.tag == qn("w:sectPr"):
            continue
        body.remove(el)

    def para(text="", *, center=False, size=10, bold=False, color=None, indent_cm=0.0, after=4):
        p = d.add_paragraph()
        p.style = d.styles["Normal"]
        pf = p.paragraph_format
        pf.space_after = Pt(after)
        pf.line_spacing = 1.25
        if indent_cm:
            pf.left_indent = Cm(indent_cm)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if text:
            _font(p.add_run(check(text)), size=size, bold=bold, color=color)
        el = p._p
        el.getparent().remove(el)
        return el

    def heading(text):
        el = set_runs(copy.deepcopy(P_H3), [(R_H3, text)])
        docx.text.paragraph.Paragraph(el, d._body).paragraph_format.keep_with_next = True
        return el

    def bullet(label, text, lvl):
        p = d.add_paragraph()
        p.style = d.styles["Normal"]
        pf = p.paragraph_format
        pf.left_indent = Cm(0.6 + 0.7 * lvl)
        pf.first_line_indent = Cm(-0.45)
        pf.space_after = Pt(2)
        pf.line_spacing = 1.2
        _font(p.add_run("● " if lvl == 0 else "○ "), size=10)
        if label:
            _font(p.add_run(check(label)), size=10, bold=True)
        if text:
            _font(p.add_run((" " if label else "") + check(text)), size=10, bold=False)
        el = p._p
        el.getparent().remove(el)
        return el

    def note(text):
        return para(text, size=8.5, color="595959", after=6)

    def picture(path, caption, max_h=None):
        from PIL import Image
        path = Path(path)
        w, h = Image.open(path).size
        width_cm = 16.5
        limit = max_h or max_img_h_cm
        if (h / w) * width_cm > limit:
            width_cm = limit * w / h
        el = para(center=True, after=2)
        p = docx.text.paragraph.Paragraph(el, d._body)
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        p.paragraph_format.keep_with_next = True
        return [el, para(caption, center=True, size=9, color="404040", after=8)]

    def code(title, text):
        out = [para(title, size=9, bold=True, after=1)]
        tbl = d._body.add_table(1, 1, Cm(16.5))   # 양식의 여백 값이 소수라 d.add_table 은 쓰지 못한다
        if "Table Grid" in [s.name for s in d.styles]:
            tbl.style = d.styles["Table Grid"]
        cell = tbl.rows[0].cells[0]
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F3F4F1")
        tcpr.append(shd)
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(0)
        cp.paragraph_format.line_spacing = 1.0
        lines = text.rstrip("\n").split("\n")
        for i, line in enumerate(lines):
            r = cp.add_run(line.replace("\t", "    "))
            _font(r, size=7.5, mono=True)
            if i < len(lines) - 1:
                r.add_break()
        el = tbl._tbl
        el.getparent().remove(el)
        out += [el, para(after=4)]
        return out

    def table(head, rows, ratio):
        el = copy.deepcopy(TBL)
        trs = el.findall(qn("w:tr"))
        h_tr, b_tr = trs[0], trs[1]
        h_tc = h_tr.findall(qn("w:tc"))[0]
        b_tcs = b_tr.findall(qn("w:tc"))
        for tr in trs:
            el.remove(tr)
        total = 9510
        widths = [int(total * r / sum(ratio)) for r in ratio]
        grid = el.find(qn("w:tblGrid"))
        for g in list(grid):
            grid.remove(g)
        for wv in widths:
            grid.append(grid.makeelement(qn("w:gridCol"), {qn("w:w"): str(wv)}))

        def make_tr(proto_tr, proto_tc, cells, is_body=False):
            tr = copy.deepcopy(proto_tr)
            for tc in tr.findall(qn("w:tc")):
                tr.remove(tc)
            for ci, (wv, text) in enumerate(zip(widths, cells)):
                tc = copy.deepcopy((b_tcs[0] if ci == 0 else b_tcs[-1]) if is_body else proto_tc)
                tcpr = tc.find(qn("w:tcPr"))
                if tcpr is not None:
                    tw = tcpr.find(qn("w:tcW"))
                    if tw is not None:
                        tw.set(qn("w:w"), str(wv)); tw.set(qn("w:type"), "dxa")
                    for tag in ("w:gridSpan", "w:vMerge"):
                        x = tcpr.find(qn(tag))
                        if x is not None:
                            tcpr.remove(x)
                ps = tc.findall(qn("w:p"))
                for extra in ps[1:]:
                    tc.remove(extra)
                p = ps[0]
                rs = p.findall(qn("w:r"))
                proto_r = rs[0] if rs else OxmlElement("w:r")
                set_runs(p, [(proto_r, str(text))])
                pp = docx.text.paragraph.Paragraph(p, d._body)
                for r in pp.runs:
                    r.font.color.rgb = RGBColor(0, 0, 0)
                    r.font.size = Pt(9)
                if is_body and len(str(text)) > 25:
                    pp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                tr.append(tc)
            return tr

        el.append(make_tr(h_tr, h_tc, head))
        for row in rows:
            el.append(make_tr(b_tr, None, row, is_body=True))
        # 표가 쪽 사이에서 갈라지지 않게: 줄은 쪼개지 않고, 마지막 줄 말고는 다음 줄과 붙인다(줄 수가 적은 표만)
        trs = el.findall(qn("w:tr"))
        for i, tr in enumerate(trs):
            trpr = tr.find(qn("w:trPr"))
            if trpr is None:
                trpr = OxmlElement("w:trPr")
                tr.insert(0, trpr)
            if trpr.find(qn("w:cantSplit")) is None:
                trpr.append(OxmlElement("w:cantSplit"))
            if len(trs) <= 14 and i < len(trs) - 1:
                for pe in tr.iter(qn("w:p")):
                    docx.text.paragraph.Paragraph(pe, d._body).paragraph_format.keep_with_next = True
        return el

    def render(blocks):
        out = []
        for b in blocks:
            k = b[0]
            if k == "h":
                out.append(heading(b[1]))
            elif k == "p":
                out.append(para(b[1]))
            elif k == "b":
                if len(b) == 3 and isinstance(b[2], int):  # ["b", "라벨: 글", 단계] 모양
                    lab, _, txt = b[1].partition(": ")
                    b = ["b", lab + ":", txt, b[2]] if txt else ["b", "", b[1], b[2]]
                out.append(bullet(b[1], b[2], b[3] if len(b) > 3 else 0))
            elif k == "t":
                out += [table(b[1], b[2], b[3]), para(after=2)]
            elif k == "img":
                out += picture(b[1], b[2], b[3] if len(b) > 3 else None)
            elif k == "code":
                out += code(b[1], b[2])
            elif k == "note":
                out.append(note(b[1]))
            else:
                raise ValueError(f"모르는 블록 {k}")
        return out

    for anchor, blocks in zip(heads, sections):
        for el in reversed(render(blocks) + [para(after=2)]):
            anchor.addnext(el)
    for c in body.iter(qn("w:color")):
        if (c.get(qn("w:val")) or "").upper() == "0000FF":
            c.set(qn("w:val"), "000000")
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists():
        import datetime, shutil
        bk = out.parent / "_backup" / (datetime.datetime.now().strftime("%Y-%m-%d_%H%M") + "_" + out.stem)
        bk.mkdir(parents=True, exist_ok=True)
        shutil.copy2(out, bk / out.name)
    d.save(str(out))
    return out


def to_pdf(docx_path, pdf_path=None):
    """워드로 열어 PDF 로 내보낸다(모양 확인용)."""
    import win32com.client as wc
    docx_path = str(Path(docx_path).resolve())
    pdf_path = str(Path(pdf_path).resolve()) if pdf_path else docx_path[:-5] + ".pdf"
    app = wc.DispatchEx("Word.Application")
    app.Visible = False
    try:
        doc = app.Documents.Open(docx_path, ReadOnly=True)
        doc.ExportAsFixedFormat(pdf_path, 17)
        doc.Close(False)
    finally:
        app.Quit()
    return pdf_path


def pages_png(pdf_path, out_dir, zoom=0.7, pages=None):
    import fitz
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    d = fitz.open(pdf_path)
    names = []
    for i in range(d.page_count):
        if pages and i not in pages:
            continue
        p = out_dir / f"p{i + 1:02d}.png"
        d[i].get_pixmap(matrix=fitz.Matrix(zoom, zoom)).save(str(p))
        names.append(p)
    return d.page_count, names
