"""5주차(2026-09-29 ~ 10-06) 스탠드업 발표 v2 — Claude 판 빌더.

v1 에 대한 사용자 지적(2026-10-06): 「설명 스킬 잘 써서 작성해. 문서 내에 너만 아는 내용 있으면 안 됨. 시각화 자료 최대한」.
바꾼 것: 장마다 그림 하나가 주인공(흐름도 · 네이티브 차트 · 화면 캡처) · 처음 나오는 말은 그 장의 「용어」 칸에서 풀이 ·
내부 말(한 줄로 이은 시험 · 요청 둘 · 짚기 · 가르친 · 독립 B …)을 실무 용어로 · 실제 시험 문제 예를 그대로 보여 줌 · 맨 끝에 용어 풀이 부록.
재료: ../_재료/주간_사실.md(숫자) · ../_재료/쉬운말_재료_v2.md(풀이 · 실제 예). 출력: ../v2/ (v1 은 그대로 둔다)
"""
from __future__ import annotations

import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
W5 = ROOT.parent
OUTDIR = ROOT / "v2"
OUTDIR.mkdir(exist_ok=True)
OUT = OUTDIR / "[발표] triPilot_진행경과_2026-10-06_Claude_v2.pptx"
CAP = W5 / "화면설계서" / "assets"
THUMB = OUTDIR / "_thumbs"
PAGES = {"vec": "24쪽", "multi": "55쪽", "arch": "28쪽", "scr": "25쪽", "w6": "23쪽"}
FIG_FAIR = W5 / "멀티에이전트_테스트보고서" / "그림" / "상용_출제자제외_BC.png"

FONT = "맑은 고딕"
GREEN = RGBColor(0x2D, 0x5F, 0x48)
DEEP = RGBColor(0x1E, 0x3F, 0x31)
MINT = RGBColor(0xEE, 0xF4, 0xEF)
LEAF = RGBColor(0x8C, 0xB8, 0x9F)
GOLD = RGBColor(0xC8, 0x91, 0x3A)
SAND = RGBColor(0xF6, 0xEA, 0xD4)
RED = RGBColor(0xB0, 0x3A, 0x2E)
ROSE = RGBColor(0xFB, 0xEE, 0xE9)
GREY = RGBColor(0xA9, 0xB1, 0xAC)
INK = RGBColor(0x22, 0x2A, 0x26)
MUTED = RGBColor(0x5F, 0x6B, 0x64)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xBF, 0xD8, 0xCB)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
SCRIPT = []


# ---------------------------------------------------------------- 도구
def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    lines = runs if isinstance(runs, list) and runs and isinstance(runs[0], list) else [runs]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.space_after = Pt(space)
        items = [(line, {})] if isinstance(line, str) else line
        for t, o in items:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=MINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, lw=1.0):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def label_box(slide, x, y, w, h, title, sub=None, fill=MINT, line=None, tcolor=DEEP, tsize=15, ssize=11, scolor=MUTED, align=PP_ALIGN.CENTER):
    box(slide, x, y, w, h, fill=fill, line=line)
    lines = [[(title, {"bold": True, "size": tsize, "color": tcolor})]]
    if sub:
        lines.append([(sub, {"size": ssize, "color": scolor})])
    text(slide, x + 0.12, y + 0.08, w - 0.24, h - 0.16, lines, align=align, anchor=MSO_ANCHOR.MIDDLE, space=2)


def arrow(slide, x1, y1, x2, y2, color=MUTED, w=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    te = etree.SubElement(ln, qn("a:tailEnd"))
    te.set("type", "triangle")
    te.set("w", "med")
    te.set("len", "med")
    return c


def line(slide, x1, y1, x2, y2, color=MUTED, w=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(w)
    return c


def circle_num(slide, x, y, n, d=0.42, fill=GREEN, color=WHITE, size=14):
    box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    text(slide, x, y, d, d, str(n), size=size, bold=True, color=color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def picture(slide, path, x, y, w, h, frame=None):
    iw, ih = Image.open(path).size
    sc = min(w / iw, h / ih)
    pw, ph = iw * sc, ih * sc
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    if frame is not None:
        pic.line.color.rgb = frame
        pic.line.width = Pt(0.75)
    return px, py, pw, ph


def title(slide, t, sub=None):
    size = 26 if len(t) <= 31 else 22
    text(slide, 0.6, 0.38, 12.1, 0.9, t, size=size, bold=True, color=DEEP)
    if sub:
        text(slide, 0.6, 1.12, 12.1, 0.4, sub, size=13, color=MUTED)


def footer(slide, t):
    text(slide, 0.6, 7.02, 12.1, 0.35, t, size=9.5, color=MUTED)


def terms_box(slide, x, y, w, h, rows, size=10.5):
    """처음 나오는 말을 그 자리에서 푼다 — 「이 장의 말」 칸."""
    box(slide, x, y, w, h, fill=WHITE, line=LEAF)
    lines = [[("이 장의 말", {"bold": True, "size": size + 0.5, "color": GREEN})]]
    for k, v in rows:
        lines.append([(k + "  ", {"bold": True, "size": size, "color": INK}), (v, {"size": size, "color": MUTED})])
    text(slide, x + 0.18, y + 0.12, w - 0.36, h - 0.2, lines, space=3)


def stat(slide, x, y, w, h, big, label, how, color=GREEN, big_size=34, fill=MINT):
    box(slide, x, y, w, h, fill=fill)
    text(slide, x + 0.22, y + 0.16, w - 0.44, 0.7, big, size=big_size, bold=True, color=color)
    text(slide, x + 0.22, y + 0.2 + big_size / 72 * 1.22, w - 0.44, 0.5, label, size=13, bold=True, color=INK)
    text(slide, x + 0.22, y + h - 0.58, w - 0.44, 0.5, how, size=9.5, color=MUTED)


def bubble(slide, x, y, w, h, t, who, right=False):
    fill = SAND if not right else MINT
    box(slide, x, y, w, h, fill=fill, radius=0.25)
    text(slide, x + 0.15, y + 0.06, w - 0.3, 0.25, who, size=9, bold=True, color=MUTED, align=PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT)
    text(slide, x + 0.15, y + 0.3, w - 0.3, h - 0.34, t, size=12.5, color=INK, align=PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT)


def bar_chart(slide, x, y, w, h, cats, vals, colors, *, horizontal=True, vmin=None, vmax=None, fmt='0', label_size=11, cat_size=11, gap=60):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("값", vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    ch.font.name = FONT
    ch.font.size = Pt(cat_size)
    ch.font.color.rgb = INK
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(label_size)
    dl.font.bold = True
    dl.font.name = FONT
    ser = plot.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c
    va = ch.value_axis
    va.has_major_gridlines = False
    va.visible = False
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    ca = ch.category_axis
    ca.format.line.color.rgb = GREY
    ca.tick_labels.font.size = Pt(cat_size)
    ca.tick_labels.font.name = FONT
    if horizontal:
        ca.reverse_order = True
    return ch


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t
    SCRIPT.append(t)


def dark(slide):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = DEEP


def new(t=None, sub=None):
    s = prs.slides.add_slide(BLANK)
    if t:
        title(s, t, sub)
    return s


# ================================================================ 1 표지
s = new()
dark(s)
text(s, 0.9, 1.5, 11.5, 0.5, "SKN32 Final 6팀 · A-COPilot", size=18, color=PALE)
text(s, 0.9, 2.15, 11.5, 1.1, "triPilot 5주차 진행 경과", size=46, bold=True, color=WHITE)
text(s, 0.9, 3.4, 11.5, 0.9, "자유여행자가 올린 여행 계획을 여행이 끝날 때까지 지켜 주는 서비스", size=22, color=WHITE)
text(s, 0.9, 5.3, 11.5, 0.5, "2026.09.29 ~ 10.06 스탠드업", size=16, color=PALE)
notes(s, "triPilot 5주차 진행 경과를 말씀드리겠습니다. triPilot 은 자유여행자가 이미 짠 여행 계획을 받아서, 여행이 끝날 때까지 문제가 생기지 않는지 지켜보고 바꿔 주는 서비스입니다.")

# ================================================================ 2 서비스 흐름 + 이번 주 손댄 곳
s = new("올린 계획을 여행 내내 지키는 흐름에서, 이번 주엔 세 곳을 손봤습니다",
        "초록 칸은 서비스가 하는 일, 금색 꼬리표는 이번 주에 손댄 곳입니다")
steps = [("계획 올리기", "글 · 사진 · 파일"), ("검사 · 고객 확인", "시간 겹침 · 이동 가능 여부"),
         ("여행 중 위험 점검", "3분마다 휴관 · 마감 · 날씨"), ("바꿀 안 + 알림", "식당 · 활동 담당 AI 가 안을 만듦")]
X0, Y0, BW, BH, GAP = 0.6, 2.0, 2.65, 1.35, 0.55
for i, (t, sub) in enumerate(steps):
    x = X0 + i * (BW + GAP)
    label_box(s, x, Y0, BW, BH, t, sub, fill=MINT, line=GREEN, tsize=16)
    circle_num(s, x - 0.12, Y0 - 0.18, i + 1)
    if i < 3:
        arrow(s, x + BW + 0.05, Y0 + BH / 2, x + BW + GAP - 0.05, Y0 + BH / 2, color=GREEN, w=2.5)
# 되돌아가는 고리: ④ → ③
x3 = X0 + 2 * (BW + GAP) + BW / 2
x4 = X0 + 3 * (BW + GAP) + BW / 2
yb = Y0 + BH + 0.45
line(s, x4, Y0 + BH, x4, yb, color=GREEN, w=2.5)
line(s, x4, yb, x3, yb, color=GREEN, w=2.5)
arrow(s, x3, yb, x3, Y0 + BH + 0.03, color=GREEN, w=2.5)
text(s, x3 + 0.3, yb + 0.05, 3.2, 0.35, "여행이 끝날 때까지 되풀이", size=12, bold=True, color=GREEN)
# 고객 채팅 입력
label_box(s, 10.55, 4.35, 2.55, 0.85, "고객 채팅 · 개인 AI 비서", "「거기 웨이팅 심해?」", fill=WHITE, line=GREY, tsize=12.5, ssize=10.5)
arrow(s, 11.8, 4.35, 11.8, Y0 + BH + 0.03, color=GREY, w=1.75)
# 이번 주 꼬리표
tags = [(0.6, 5.55, 5.4, "화면 12개를 cs 프로젝트 서버에 연결", "계획 올리기부터 여행 화면 · 재난 시 일정 정지까지 (4장)"),
        (6.25, 5.55, 3.1, "실시간 진행 · 개인 AI 입구", "오래 걸리는 일의 진행 표시, AI 비서 연결 (5장)"),
        (9.6, 5.55, 3.1, "고객 말 → 일정 항목 찾기 모델", "파인튜닝 시험 중, 제품 미연결 (6~9장)")]
for x, y, w, t, sub in tags:
    box(s, x, y, w, 1.15, fill=SAND)
    text(s, x + 0.18, y + 0.12, w - 0.36, 0.95, [[(t, {"bold": True, "size": 13, "color": INK})], [(sub, {"size": 10.5, "color": MUTED})]], space=3)
line(s, 3.3, Y0 + BH, 3.3, 5.55, color=GOLD, w=1.25)
line(s, 7.8, Y0 + BH + 0.45, 7.8, 5.55, color=GOLD, w=1.25)
line(s, 11.15, 5.2, 11.15, 5.55, color=GOLD, w=1.25)
footer(s, "흐름은 「아키텍처 한눈에 v26」 그림 1(계획 확인) · 그림 4(에이전트 구성)를 줄인 것입니다. 알림은 디스코드 등으로, 바뀐 내용은 로그인 없이 열리는 여행계획서 링크에도 반영됩니다.")
notes(s, "서비스 흐름부터 짚겠습니다. 고객이 계획을 올리면 코드가 시간 겹침과 이동 가능 여부를 검사하고 고객이 확인합니다. 등록하면 여행 중에 3분마다 휴관, 마감, 날씨 같은 위험을 점검하고, 문제가 생기면 식당이나 활동 담당 AI 가 바꿀 안을 만들어 알립니다. 이 점검과 수정은 여행이 끝날 때까지 되풀이됩니다. 이번 주에는 화면 연결, 실시간 진행 표시와 개인 AI 입구, 그리고 고객 말이 어느 일정 항목을 가리키는지 찾는 모델을 다뤘습니다.")

# ================================================================ 3 한 주 규모
s = new("한 주 동안 커밋 531건, 연구 문서와 웹 기능이 가장 많았습니다",
        "커밋 = 코드나 문서를 저장소에 한 번 저장한 단위. 팀원 브랜치까지 모든 브랜치를 셌습니다")
days = ["9/29", "9/30", "10/1", "10/2", "10/3", "10/4", "10/5", "10/6"]
cnt = [4, 91, 85, 18, 112, 74, 76, 71]
text(s, 0.6, 1.75, 6.4, 0.4, "날짜별 커밋 수 (모든 브랜치, 합계 531)", size=13, bold=True, color=DEEP)
bar_chart(s, 0.5, 2.15, 6.6, 4.5, days, cnt, [GREEN] * 8, horizontal=False, vmin=0, vmax=130, cat_size=12, gap=45)
areas = ["연구 문서", "웹 기능", "모델 비교 도구", "계획 문서", "여행 기능", "액티비티", "이동"]
acnt = [149, 51, 23, 19, 12, 12, 11]
text(s, 7.5, 1.75, 5.3, 0.4, "무엇을 했나 (커밋 제목 머리로 나눈 277건)", size=13, bold=True, color=DEEP)
bar_chart(s, 7.4, 2.15, 5.4, 4.5, areas, acnt, [GOLD] + [LEAF] * 6, vmin=0, vmax=175, cat_size=12, gap=45)
footer(s, "git log, 2026-09-29 ~ 10-06. 작업 브랜치(role-manager)만 세면 402건. 오른쪽은 제목에 영역 머리(docs(research) 등)가 붙은 277건만 나눴고, 나머지 254건은 머리가 없어 나누지 않았습니다.")
notes(s, "한 주 규모입니다. 모든 브랜치를 합쳐 커밋이 531건이었고, 10월 3일이 112건으로 가장 많았습니다. 영역별로는 모델 실험을 기록한 연구 문서가 149건으로 가장 많고, 그다음이 웹 기능 51건입니다.")

# ================================================================ 4 웹 화면
s = new("고객이 계획을 올리고 여행이 끝날 때까지 쓰는 화면을 이어 붙였습니다",
        "휴대폰 화면 기준, 왼쪽에서 오른쪽으로 고객이 지나가는 순서입니다")
shots = [(CAP / "capture_docs" / "04_trips_new_390.png", "계획 올리기", "글을 붙여 넣거나 사진 · 파일"),
         (CAP / "capture_2026-10-06" / "q_read_390.png", "계획 읽는 중", "기다리는 동안 취향 질문"),
         (CAP / "capture_2026-10-06" / "card_390.png", "항로 지킴이", "문제 생기면 먼저 바꿀지 고름"),
         (CAP / "capture_2026-10-06" / "day_overview.png", "계획 확인", "날짜별 일정 · 지도 · 수정"),
         (CAP / "capture_docs" / "09_trip_home_390.png", "여행 화면", "일정 · 지도 · 채팅 · 알림"),
         (CAP / "capture_2026-10-06" / "safety_top.png", "재난 시 일정 정지", "안전 안내 · 119 · 대피 장소")]
SW, SG = 1.86, 0.2
for i, (p, t, sub) in enumerate(shots):
    x = 0.6 + i * (SW + SG)
    picture(s, p, x, 1.7, SW, 3.9, frame=GREY)
    circle_num(s, x + SW / 2 - 0.19, 5.68, i + 1, d=0.38, size=12)
    text(s, x - 0.05, 6.1, SW + 0.1, 0.3, t, size=12.5, bold=True, align=PP_ALIGN.CENTER)
    text(s, x - 0.05, 6.4, SW + 0.1, 0.3, sub, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 5:
        arrow(s, x + SW + 0.02, 5.87, x + SW + SG - 0.02, 5.87, color=GREEN, w=1.75)
footer(s, "mock 서버(같은 모양의 답만 흉내 내는 시험용 서버)로 화면 반응을 확인. 실서버 확인은 계획 확인 화면 일부만(10-03), 바뀐 화면은 아직입니다.")
notes(s, "웹 화면입니다. 고객은 계획을 붙여 넣고, 읽는 동안 취향 질문에 답하고, 항로 지킴이를 켤지 고릅니다. 항로 지킴이는 여행 중 문제가 생기면 서비스가 먼저 바꿔 두고 알릴지, 먼저 물어볼지를 정하는 스위치입니다. 그다음 계획을 확인하고, 여행 중에는 일정, 지도, 채팅을 씁니다. 재난 문자가 오면 남은 일정을 멈추고 안전 안내를 보여 줍니다. 이 화면들은 mock 서버로 화면 반응을 확인했고, 진짜 서버 응답까지 확인한 것은 계획 확인 화면 일부입니다.")

# ================================================================ 5 서버 기능
s = new("고객의 AI 비서가 들어올 입구와, 기다리는 동안의 진행 표시가 생겼습니다")
# 왼쪽: 개인 AI 입구
text(s, 0.6, 1.45, 6.0, 0.4, "개인 AI 입구", size=17, bold=True, color=DEEP)
label_box(s, 0.6, 2.05, 1.85, 1.2, "고객의 AI 비서", "고객이 평소 쓰는 AI", fill=WHITE, line=GREY, tsize=13)
label_box(s, 3.0, 2.05, 1.85, 1.2, "triPilot 입구", "본인 키 · 본인 여행만", fill=MINT, line=GREEN, tsize=13)
label_box(s, 5.4, 2.05, 1.2, 1.2, "내 여행", None, fill=SAND, tsize=13)
arrow(s, 2.48, 2.65, 2.97, 2.65, color=GREEN)
arrow(s, 4.88, 2.65, 5.37, 2.65, color=GREEN)
text(s, 0.6, 3.45, 6.0, 1.0, [[("읽기", {"bold": True}), ("  일정 · 위험 점검 결과 · 이동 판정을 물어봄  ", {"color": MUTED}), ("켜짐", {"bold": True, "color": GREEN})],
                              [("고치기", {"bold": True}), ("  일정 수정 · 신고  ", {"color": MUTED}), ("기본 꺼짐(설정으로 켬)", {"bold": True, "color": RED})]], size=12, space=4)
stat(s, 0.6, 4.45, 2.9, 1.6, "15/15", "자동 시험 통과", "2026-10-02, 본인 여행만 열리는지 · 꺼진 쓰기가 막히는지", big_size=30)
# 오른쪽: 실시간 진행
text(s, 7.1, 1.45, 5.7, 0.4, "실시간 진행 표시", size=17, bold=True, color=DEEP)
prog = ["받았어요", "일정 읽기", "장소 · 운영시간", "정리 완료"]
for i, t in enumerate(prog):
    cx = 7.35 + i * 1.5
    box(s, cx, 2.2, 0.36, 0.36, fill=GREEN if i < 2 else WHITE, line=GREEN, shape=MSO_SHAPE.OVAL)
    text(s, cx - 0.5, 2.68, 1.36, 0.3, t, size=10.5, align=PP_ALIGN.CENTER)
    if i < 3:
        line(s, cx + 0.38, 2.38, cx + 1.48, 2.38, color=GREEN if i < 1 else GREY, w=2.5)
text(s, 7.1, 3.15, 5.7, 1.2, [[("계획 읽기 · 일정 짜기 · 채팅처럼 수십 초 걸리는 일을 하는 동안, 지금 몇 단계인지와 ", {}), ("3초마다 「살아 있음」 신호", {"bold": True}), ("를 화면에 보냅니다. 멈췄는지 기다리는 중인지 고객이 알 수 있습니다.", {})]], size=12)
stat(s, 7.1, 4.45, 2.75, 1.6, "21/21", "자동 시험 통과", "2026-10-02, 채팅 · 일정 짜기 · 계획 읽기", big_size=30)
stat(s, 10.05, 4.45, 2.75, 1.6, "517/517", "식당 담당 AI 시험", "팀원 정세환, 2026-10-06", big_size=30)
stat(s, 3.7, 4.45, 2.9, 1.6, "10/10", "담당 다시 고르기", "맡을 담당 AI 가 없을 때 한 번 더 고름, mock 서버 시험", big_size=30)
footer(s, "개인 AI 입구는 MCP 서버로 만들었습니다. 디스코드 · 텔레그램 연결과 여행 삭제는 코드에 들어갔지만 x600 서버 반영과 실제 연결은 아직 확인하지 않았습니다.")
notes(s, "서버 쪽은 두 가지가 컸습니다. 왼쪽은 개인 AI 입구입니다. 고객이 평소 쓰는 AI 비서가 triPilot 에 들어와 본인 여행만 읽을 수 있고, 고치는 기능은 기본으로 꺼 두었습니다. 오른쪽은 실시간 진행 표시입니다. 오래 걸리는 일을 하는 동안 몇 단계인지와 살아 있다는 신호를 3초마다 보내서, 고객이 멈췄는지 기다리는 중인지 알 수 있게 했습니다.")

# ================================================================ 6 모델이 푸는 문제
s = new("모델이 푸는 문제: 고객의 말이 일정표의 몇 번 항목을 가리키는지 찾기",
        "어느 항목인지 알아야 그 장소의 영업시간 · 예약을 확인하고 바꿀 수 있습니다")
box(s, 0.6, 1.75, 4.1, 2.35, fill=WHITE, line=GREY)
text(s, 0.8, 1.88, 3.8, 0.3, "고객 일정표 (일부)", size=12, bold=True, color=DEEP)
rows = [("23번", "금 13:30", "대림미술관 전시 관람"), ("24번", "금 15:30", "호텔 체크인"), ("25번", "금 16:30", "피크닉 전시 보기")]
for i, (n, tm, t) in enumerate(rows):
    y = 2.3 + i * 0.55
    hi = n == "25번"
    box(s, 0.8, y, 3.7, 0.45, fill=SAND if hi else MINT)
    text(s, 0.95, y + 0.08, 3.5, 0.3, [[(n + "  ", {"bold": True, "color": GOLD if hi else GREEN}), (tm + "  ", {"color": MUTED}), (t, {"bold": hi})]], size=12)
# 대화
bubble(s, 5.0, 1.75, 3.6, 0.72, "금요일에 전시 뭐 봐?", "고객")
bubble(s, 5.6, 2.55, 3.9, 0.95, "금요일엔 13:30 대림미술관 전시 관람, 16:30 피크닉 전시 보기가 있어요.", "triPilot", right=True)
bubble(s, 5.0, 3.58, 3.6, 0.72, "뒤에 거 예매해야 돼?", "고객")
arrow(s, 8.7, 3.94, 9.75, 3.94, color=GOLD, w=2.5)
label_box(s, 9.85, 3.5, 2.95, 0.9, "정답: 25번", "「뒤에 거」 = 피크닉 전시", fill=SAND, tsize=17, tcolor=GOLD)
# 없음 예
bubble(s, 5.0, 4.6, 3.6, 0.72, "남산타워 몇 시에 올라가?", "고객")
arrow(s, 8.7, 4.96, 9.75, 4.96, color=GREY, w=2.5)
label_box(s, 9.85, 4.55, 2.95, 0.85, "정답: 없음", "일정에 없는 곳은 지어내지 않음", fill=MINT, tsize=16)
# 시험지 설명
box(s, 0.6, 4.35, 4.1, 2.45, fill=MINT)
text(s, 0.8, 4.47, 3.75, 2.25, [
    [("시험지 769문제", {"bold": True, "size": 15, "color": DEEP})],
    [("헷갈리게 만든 11가지 유형: 「거기」 같은 가리킴 말 · 비슷한 이름(대림미술관 / 대림창고) · 「~말고」 · 오타 · 일정과 무관한 잡담 · 그럴듯하지만 없는 일정 등", {"size": 11})],
    [("AI 모델(Codex · Claude)이 쓴 합성 문제. 처음 문제와 체크리스트는 사람이 썼고 일부를 사람이 검수", {"size": 10, "color": MUTED})]], space=4)
text(s, 5.0, 5.6, 7.8, 1.2, [[("769 = 835 − 66", {"bold": True, "color": DEEP}),
                             ("  한 문장에 항목 두 개를 묻는 66문제(「점심이랑 저녁 다 알려 줘」 유형)는 채점 기준을 따로 두어 주 점수에서 뺐습니다.", {"color": MUTED})]], size=11.5)
footer(s, "예시는 시험 세트 mix12(Claude 작성)의 실제 문제입니다. 일정은 가상 여행자 일정이고 날짜는 2027-11-05(금)입니다.")
notes(s, "여기서부터 모델 이야기입니다. 우리 모델이 푸는 문제는 하나입니다. 고객이 한 말이 일정표의 몇 번 항목을 가리키는지 찾는 것입니다. 예를 들어 금요일 전시를 물은 뒤 뒤에 거 예매해야 돼라고 하면, 정답은 25번 피크닉 전시입니다. 일정에 없는 남산타워를 물으면 정답은 없음입니다. 이걸 맞혀야 그 장소의 영업시간이나 예약을 확인하고 바꿀 수 있습니다. 시험지는 헷갈리게 만든 11가지 유형, 769문제입니다.")

# ================================================================ 7 파인튜닝 결과
s = new("약 4B 모델을 파인튜닝하니 26B · 27B 로컬 모델만큼 맞혔습니다",
        "같은 769문제, 같은 채점. 막대 끝 숫자가 맞힌 문제 수입니다")
cats = ["Gemma 4 E4B (약 4B) 파인튜닝 전", "Gemma 4 E4B 파인튜닝 1회차", "Gemma 4 E4B 파인튜닝 2회차 (시드만 바꿈)",
        "Winnow-E4B (Jev 계열, 약 4B) 파인튜닝", "Gemma 4 26B-A4B (파인튜닝 없음)", "Qwen3.8 27B (파인튜닝 없음)"]
vals = [688, 750, 760, 756, 732, 749]
bar_chart(s, 0.5, 1.65, 8.0, 4.9, cats, vals, [GREY, GREEN, GREEN, LEAF, GOLD, GOLD], vmin=600, vmax=780, cat_size=11.5, label_size=13, gap=45)
terms_box(s, 8.75, 1.7, 4.0, 3.55, [
    ("학습 자료", "연습문제 1,760개. 처음 문제와 체크리스트는 사람이 쓰고, 이를 바탕으로 Claude 가 늘림(일부 사람 검수)"),
    ("방법", "QLoRA, 같은 설정으로 시드만 바꿔 두 번 → 750 · 760"),
    ("Jev 계열 모델", "선택지 중 하나를 고르는 결정 전용 공개 모델 묶음. Winnow-E4B 는 그중 하나")], size=10)
stat(s, 8.75, 5.4, 4.0, 1.45, "평균 4.55초", "GPU 없이 CPU 4코어로 이어지는 질문", "파인튜닝 E4B 를 8비트로 줄인 7.5GB 파일. 새 일정 첫 질문은 평균 67.5초", big_size=24)
footer(s, "파인튜닝 전 대비 파인튜닝 1회차만 맞힌 문제 69 : 반대 7 (p=6.4e-14, 우연으로 보기 어려움). 4B 모델끼리 차이(750 · 756 · 760)는 시드 흔들림 안이라 순위를 매기지 않습니다. 상용 모델은 다음 장에서 출제자를 빼고 비교합니다.")
notes(s, "결과입니다. 약 4B 크기의 Gemma 4 E4B 를 우리 연습문제로 파인튜닝하니 688개에서 750개, 다시 학습하면 760개를 맞혔습니다. 파인튜닝하지 않은 26B, 27B 로컬 모델이 732개, 749개라서, 작은 모델로 큰 모델만큼 맞힌 셈입니다. 다만 같은 방법으로 두 번 학습해도 10문제쯤 흔들리기 때문에, Jev 계열 Winnow-E4B 의 756 과 우리 E4B 의 우열은 가리지 않습니다. GPU 없이 CPU 만으로도 이어지는 질문은 평균 4.5초에 답합니다.")

# ================================================================ 7-2 출제자를 뺀 상용 비교
s = new("문제를 낸 AI 는 그 문제에서 빼고 다시 재도, 상용 모델이 가장 높았습니다",
        "Codex 가 쓴 문제에서는 Codex 를, Claude 가 쓴 문제에서는 Claude 를 빼고 B · C 두 채점으로 비교했습니다")
import json as _json
_fc = _json.loads((W5 / "멀티에이전트_테스트보고서" / "_재료" / "fair_commercial.json").read_text(encoding="utf-8"))
for j, (sub, lab) in enumerate((("codex_made", "Codex 가 쓴 659문제 · Codex 제외"), ("claude_made", "Claude 가 쓴 176문제 · Claude 제외"))):
    x = 0.5 + j * 4.25
    text(s, x + 0.1, 1.6, 4.1, 0.3, lab, size=12.5, bold=True, color=DEEP)
    rows = _fc[sub]
    cd = CategoryChartData()
    cd.categories = [r["model"].replace("(상용)", "").replace("(파인튜닝 없음, 로컬)", " 로컬").replace("(약 4B, 로컬)", " 로컬").strip() for r in rows]
    cd.add_series("B 채점", [round(r["B"] / r["n"] * 100, 1) for r in rows])
    cd.add_series("C 채점", [round(r["C"] / r["n"] * 100, 1) for r in rows])
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(1.95), Inches(4.2), Inches(4.75), cd).chart
    ch.has_title = False
    ch.has_legend = j == 0
    if j == 0:
        ch.legend.font.size = Pt(10)
        ch.legend.include_in_layout = False
    ch.font.name = FONT
    ch.font.size = Pt(9.5)
    pl = ch.plots[0]
    pl.gap_width = 40
    pl.has_data_labels = True
    pl.data_labels.number_format = '0.0'
    pl.data_labels.number_format_is_linked = False
    pl.data_labels.font.size = Pt(8.5)
    for sr, c in zip(pl.series, (GREEN, LEAF)):
        sr.format.fill.solid()
        sr.format.fill.fore_color.rgb = c
    ch.value_axis.minimum_scale = 80
    ch.value_axis.maximum_scale = 102
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.reverse_order = True
    ch.category_axis.tick_labels.font.size = Pt(9.5)
terms_box(s, 9.05, 1.65, 3.75, 2.85, [
    ("왜 빼나", "문제를 쓴 AI 는 자기 말투 · 정답 기준에 익숙해 점수가 부풀 수 있음"),
    ("B 채점", "한 문장에 항목 둘을 묻는 질문은 하나라도 맞으면 정답"),
    ("C 채점", "낸 항목 묶음이 정답 묶음과 완전히 같아야 정답")], size=10)
box(s, 9.05, 4.65, 3.75, 1.95, fill=SAND)
text(s, 9.22, 4.75, 3.45, 1.8, [
    [("읽는 법", {"bold": True, "color": GOLD, "size": 12.5})],
    [("B 에서 파인튜닝 E4B 는 27B 와 같은 642/659. C 에서 로컬이 낮은 건 항목을 하나만 내는 구조라 「둘 다」 질문을 놓쳐서입니다.", {"size": 10.5})]], space=3)
footer(s, "저장된 문제별 풀이를 다시 셈(새 호출 없음). 남는 치우침: 「둘 다」 질문의 정답 묶음과 파인튜닝 E4B 의 학습 자료는 Claude 가 만들었습니다. mix12 는 176문제라 1~2문제 차이는 구분 안 됨.")
notes(s, "앞 장의 상용 모델 점수는 자기가 낸 문제까지 푼 것이라 공정하지 않았습니다. 그래서 Codex 가 쓴 문제에서는 Codex 를, Claude 가 쓴 문제에서는 Claude 를 빼고 다시 셌습니다. 그래도 상용 모델이 가장 높았습니다. 파인튜닝한 4B 모델은 하나라도 맞으면 정답으로 보는 B 채점에서 27B 와 같은 점수였고, 둘 다 맞혀야 하는 C 채점에서는 항목을 하나만 내는 구조 때문에 낮았습니다.")

# ================================================================ 8 E2E 파이프라인 테스트
s = new("앞단에 규칙 · 작은 분류 모델을 두니 점수는 같고 호출은 19% 줄었습니다",
        "E2E 파이프라인 테스트: 단계를 실제로 이어 붙이고 문제마다 처음부터 끝까지 돌렸습니다 (GPU 서버 C, 2026-10-06)")
fun = [(835, "문제 835개 입력", "복수 항목 질의 66개 포함", GREY),
       (710, "규칙이 125개를 바로 처리", "「첫 번째」 · 「~말고」 · 같은 장소 같은 말 패턴", LEAF),
       (674, "ML 사전 필터가 36개를 처리", "일정과 무관한 잡담이면 「없음」으로 끝냄", LEAF),
       (674, "파인튜닝 E4B 호출 674번", "혼자 다 풀면 835번", GREEN)]
for i, (n, t, sub, c) in enumerate(fun):
    y = 1.75 + i * 1.18
    wbar = 6.2 * n / 835
    box(s, 0.6, y, wbar, 0.95, fill=c, radius=0.12)
    text(s, 0.8, y + 0.08, 6.0, 0.4, t, size=14, bold=True, color=WHITE if c in (GREEN, GREY) else INK)
    text(s, 0.8, y + 0.5, 6.0, 0.35, sub, size=10.5, color=WHITE if c in (GREEN, GREY) else INK)
    if i < 3:
        arrow(s, 0.9, y + 0.96, 0.9, y + 1.17, color=MUTED, w=1.5)
cats = ["파인튜닝 E4B 혼자", "규칙 + ML 사전 필터 + 파인튜닝 E4B", "라우터에 베이스 Winnow-E4B (Jev 계열, 파인튜닝 전) 추가"]
bar_chart(s, 7.2, 1.7, 5.6, 2.6, cats, [748, 745, 697], [GREEN, LEAF, RED], vmin=600, vmax=790, cat_size=10.5, label_size=13, gap=40)
text(s, 7.3, 4.35, 5.4, 0.3, "맞힌 문제 수 / 769", size=10.5, color=MUTED)
box(s, 7.25, 4.75, 5.5, 2.1, fill=ROSE)
text(s, 7.45, 4.85, 5.1, 1.95, [
    [("읽는 법", {"bold": True, "color": RED, "size": 13})],
    [("748 대 745 는 구분되지 않습니다(갈린 문제 3 : 6, p=0.51).", {"size": 11})],
    [("697 은 파인튜닝 안 한 모델을 파인튜닝한 모델과 이은 결과라 공정한 비교가 아닙니다. 같은 방법으로 파인튜닝한 두 모델을 섞은 구조는 747~752 로 혼자와 구분되지 않았습니다.", {"size": 11})]], space=4)
footer(s, "라우터 = 질문을 어느 모델로 보낼지 고르는 단계. 697 은 시험 문제 일부를 본 뒤 고친 사후 수정 버전이고, 미리 고정한 버전은 472 였습니다. 표의 점수는 모두 AI 가 쓴 합성 문제 기준입니다.")
notes(s, "모델을 여러 단계로 이어 붙이면 좋아지는지 실제로 돌려 봤습니다. 835문제를 넣으면 정해 둔 말 패턴의 규칙이 125개를 바로 처리하고, 작은 분류 모델이 잡담 36개를 걸러서, 큰 모델은 674번만 부르면 됩니다. 점수는 혼자 풀 때와 같고 호출은 19% 줄었습니다. 반면 파인튜닝하지 않은 Jev 계열 모델을 라우터로 넣으면 697로 떨어졌는데, 학습 조건이 달라 공정한 비교는 아닙니다.")

# ================================================================ 9 혼합 트래픽 테스트
s = new("일정과 무관한 말이 섞이면, 앞단이 엉뚱한 일정을 고르는 실수를 줄였습니다",
        "혼합 트래픽 테스트: 실제 채팅처럼 일정 질문 사이에 잡담 · 관광 질문을 섞은 2,735문제")
mix = [("일정 질문", 835, GREEN), ("관광지 질문", 600, LEAF), ("일상 대화", 1000, GOLD), ("구어체 문장", 300, GREY)]
x = 0.6
text(s, 0.6, 1.7, 12.0, 0.35, "섞은 구성 (2,735문제)", size=13, bold=True, color=DEEP)
for t, n, c in mix:
    w = 12.1 * n / 2735
    box(s, x, 2.1, w - 0.04, 0.85, fill=c, radius=0.1)
    text(s, x + 0.1, 2.18, w - 0.2, 0.7, [[(t, {"bold": True, "size": 12.5, "color": WHITE if c != LEAF else INK})], [(f"{n:,}", {"size": 12, "color": WHITE if c != LEAF else INK})]], align=PP_ALIGN.CENTER)
    x += w
text(s, 0.6, 3.2, 6.0, 0.35, "일정과 무관한 말에 일정 항목을 잘못 고른 비율 (낮을수록 좋음)", size=12.5, bold=True, color=DEEP)
bar_chart(s, 0.5, 3.55, 6.0, 2.6, ["파인튜닝 E4B 혼자", "규칙 + ML 사전 필터 + 파인튜닝 E4B"], [4.5, 3.2], [GREY, GREEN], vmin=0, vmax=6, fmt='0.0"%"', cat_size=11, label_size=14, gap=50)
text(s, 6.9, 3.2, 6.0, 0.35, "전체 정답률 (높을수록 좋음)", size=12.5, bold=True, color=DEEP)
bar_chart(s, 6.8, 3.55, 6.0, 2.6, ["파인튜닝 E4B 혼자", "규칙 + ML 사전 필터 + 파인튜닝 E4B"], [90.8, 91.9], [GREY, GREEN], vmin=85, vmax=94, fmt='0.0"%"', cat_size=11, label_size=14, gap=50)
footer(s, "일정과 무관한 말 1,786개 중 80 → 58개. 정답률은 복수 항목 질의를 뺀 2,669문제 기준 2,424 → 2,453개(갈린 문제 37 : 8, p=1.5e-5). 큰 모델 호출 2,735 → 2,117번. 빈 답 85건(3.1%)은 원인 미확인입니다.")
notes(s, "실제 채팅에는 일정 질문만 오지 않습니다. 그래서 잡담, 관광 질문, 구어체 문장을 섞은 2,735문제로 다시 돌렸습니다. 이렇게 섞으면 앞단의 규칙과 분류 모델이 값을 했습니다. 일정과 무관한 말에 엉뚱한 일정을 고르는 비율이 4.5%에서 3.2%로 줄고, 전체 정답률은 90.8%에서 91.9%로 올랐습니다.")

# ================================================================ 10 고객 말 분류
s = new("고객 말 꼬리표 네 개 중 「얼마나 급한가」가 학습으로 가장 많이 올랐습니다",
        "고객 말 분류: 한 문장에 의도 · 대상 · 감정 · 긴급도 꼬리표를 붙여 처리 순서와 담당을 정합니다")
bubble(s, 0.6, 1.8, 5.2, 0.9, "식당 왔는데 우리 네 명 자리 없대요. 어떡하죠?", "고객")
chips = [("의도", "사건 신고"), ("대상", "식당"), ("감정", "부정"), ("긴급도", "높음")]
for i, (k, v) in enumerate(chips):
    x = 0.6 + i * 1.32
    label_box(s, x, 3.0, 1.22, 0.95, v, k, fill=SAND if k == "긴급도" else MINT, tsize=14, ssize=10)
    arrow(s, 3.2, 2.72, x + 0.61, 2.98, color=GREY, w=1.25)
terms_box(s, 0.6, 4.3, 5.2, 2.5, [
    ("파인튜닝 전 E4B", "학습 없이 지시문만 주고 꼬리표를 고르게 한 결과"),
    ("학습한 분류 모델", "문장을 숫자로 바꾸는 공개 모델(multilingual-e5-base)에 분류 층을 얹어 우리 문장 720개로 학습"),
    ("시험 400문장", "Codex 가 쓰고 Claude 가 따로 꼬리표를 달아 맞춰 본 합성 문장")], size=10)
bar_chart(s, 6.3, 1.7, 6.5, 2.4, ["파인튜닝 전 E4B", "학습한 분류 모델"], [60.9, 88.4], [GREY, GREEN], vmin=0, vmax=100, fmt='0.0"%"', label_size=15, cat_size=12, gap=45)
text(s, 6.4, 1.45, 6.0, 0.3, "긴급도 정확도 (400문장)", size=12.5, bold=True, color=DEEP)
text(s, 6.4, 4.3, 6.0, 0.3, "학습한 분류 모델의 나머지 꼬리표 정확도", size=12.5, bold=True, color=DEEP)
bar_chart(s, 6.3, 4.6, 6.5, 2.3, ["의도", "감정", "대상"], [97.5, 93.8, 93.2], [GREEN, GREEN, GREEN], vmin=0, vmax=110, fmt='0.0"%"', label_size=13, cat_size=12, gap=55)
footer(s, "분류 시험 CX 400(합성). 학습한 분류 모델은 작은 인코더라 지시문만 받은 E4B 와 조건이 달라 통제된 우열 비교는 아닙니다.")
notes(s, "고객 말에는 꼬리표 네 개를 붙입니다. 무엇을 원하는지, 식당인지 활동인지, 감정은 어떤지, 얼마나 급한지입니다. 이 중 긴급도가 가장 어려웠는데, 학습 없이 지시문만 준 모델은 60.9%였고 우리 문장으로 학습한 작은 분류 모델은 88.4%였습니다.")

# ================================================================ 11 규정 검색
s = new("규정 검색은 고객용 문장 인덱싱으로 정답이 늘었지만 오응답도 늘었습니다",
        "규정 검색(RAG): 고객 질문에 맞는 회사 규정 조각을 찾아 저장된 문장을 그대로 보여 줍니다")
bubble(s, 0.6, 1.75, 4.6, 0.85, "창덕궁 임시 휴관이면 일정 어쩌죠?", "고객")
arrow(s, 2.9, 2.62, 2.9, 2.95, color=GREEN)
label_box(s, 0.6, 3.0, 4.6, 1.05, "정답 규정 조각", "「운영 중단과 통제」 문서 · 「임시 휴관과 정기 휴관을 나눈다」 절", fill=MINT, line=GREEN, tsize=13, ssize=10.5)
terms_box(s, 0.6, 4.35, 4.6, 2.5, [
    ("인덱싱 텍스트", "규정을 찾을 때 질문과 비교하는 글. 직원용 본문 → 「제목 + 절 제목 + 고객용 문장」으로 바꿔 봄"),
    ("블라인드 평가셋 48문제", "규정 문장을 보지 않고 따로 쓴 질문(AI 합성). 쉬운 개발용 세트와 나눠 이 세트를 기준으로 봄"),
    ("오응답", "규정 질문이 아닌 말에 규정을 붙여 보여 줌")], size=10)
text(s, 5.6, 1.5, 3.6, 0.3, "정답 규정을 보인 질문 /48", size=12.5, bold=True, color=DEEP)
bar_chart(s, 5.5, 1.8, 3.6, 3.0, ["지금", "고객용 문장"], [21, 31], [GREY, GREEN], vmin=0, vmax=40, label_size=15, cat_size=12, horizontal=False, gap=60)
text(s, 9.4, 1.5, 3.6, 0.3, "여행 대화에 규정을 잘못 붙인 비율", size=12.5, bold=True, color=RED)
bar_chart(s, 9.3, 1.8, 3.6, 3.0, ["지금", "고객용 문장"], [22.3, 36.7], [GREY, RED], vmin=0, vmax=45, fmt='0.0"%"', label_size=15, cat_size=12, horizontal=False, gap=60)
box(s, 5.55, 5.05, 7.25, 1.8, fill=ROSE)
text(s, 5.75, 5.15, 6.9, 1.6, [
    [("그래서 바로 넣지 않습니다", {"bold": True, "color": RED, "size": 13})],
    [("「일요일 08:44에 2명 예약해 주세요」 같은 예약 요청에도 규정이 붙었습니다. 제품에서는 앞의 분류 단계가 규정 질문만 이쪽으로 보내므로, 분류 뒤에서 다시 재고 정합니다.", {"size": 11})],
    [("지금 성능: 정답 규정이 상위 3개 안에 든 비율 30/48 = 62.5% (목표 90%)", {"size": 11, "bold": True})]], space=4)
footer(s, "2026-10-06, 제품의 규정 답 함수를 그대로 호출. 오응답 비율은 라벨 없는 여행 대화 300건 기준. 규정 문장은 모델이 쓰지 않고 저장된 문장을 그대로 실어 인용 정확도 334/334 입니다.")
notes(s, "규정 검색입니다. 고객이 창덕궁 임시 휴관이면 어쩌죠라고 물으면, 운영 중단 규정의 임시 휴관 절을 찾아 그대로 보여 줘야 합니다. 규정을 찾을 때 비교하는 글을 직원용 본문에서 고객용 문장으로 바꾸니 정답을 보인 질문이 21개에서 31개로 늘었습니다. 하지만 예약 요청 같은 말에도 규정을 붙이는 일이 22%에서 37%로 늘어서, 바로 넣지 않고 분류 단계 뒤에서 다시 재기로 했습니다.")

# ================================================================ 12 산출물
s = new("5주차 산출물 5종 · 6주차 보고서 1종, 모든 문서 첫머리에 용어 정의")
docs = [("벡터DB · GraphDB\n구축 결과서", "vec", "검색 성능 · 저장소별 쓰임"),
        ("멀티 에이전트 테스트\n계획 및 결과 보고서", "multi", "파인튜닝 · E2E 테스트"),
        ("AI 시스템\n아키텍처", "arch", "구성도 · 흐름 · 장애 대응"),
        ("화면설계서", "scr", "화면 12개 · 흐름"),
        ("서비스 테스트 계획\n및 결과 보고서(6주차)", "w6", "Pass/Fail · 결함 대장")]
for i, (n, key, how) in enumerate(docs):
    x = 0.6 + i * 2.5
    box(s, x, 1.55, 2.3, 4.75, fill=SAND if key == "w6" else MINT)
    th = THUMB / f"{key}.png"
    if th.exists():
        picture(s, th, x + 0.15, 1.7, 2.0, 2.75, frame=GREY)
    text(s, x + 0.15, 4.55, 2.0, 0.8, n, size=12.5, bold=True, color=DEEP)
    text(s, x + 0.15, 5.4, 2.0, 0.4, how, size=10, color=MUTED)
    pg = PAGES.get(key)
    if pg:
        text(s, x + 0.15, 5.8, 2.0, 0.4, pg, size=13, bold=True, color=GREEN if key != "w6" else GOLD)
# 요구사항 정의서: 진행 막대
text(s, 0.6, 6.45, 3.0, 0.35, "요구사항 정의서 85개", size=12.5, bold=True, color=DEEP)
tot = 85
x = 3.3
for lab, n, c in (("구현 끝 24", 24, GREEN), ("일부 구현 58", 58, LEAF), ("결정 대기 3", 3, GOLD)):
    w = 9.4 * n / tot
    box(s, x, 6.42, w - 0.03, 0.42, fill=c, radius=0.15)
    if n > 5:
        text(s, x, 6.48, w, 0.3, lab, size=10.5, bold=True, color=WHITE if c == GREEN else INK, align=PP_ALIGN.CENTER)
    x += w
text(s, 12.15, 6.88, 0.7, 0.25, "결정 대기 3", size=8.5, color=MUTED, align=PP_ALIGN.RIGHT)
footer(s, "모든 문서를 코덱스와 교차 검증했고, 우리끼리 쓰던 말(한 줄 시험 · 요청 둘 · 독립 세트 등)을 실무 용어로 바꾼 뒤 1절 맨 앞에 「읽기 전에: 용어 정의」 표를 두었습니다.")
notes(s, "산출물입니다. 5주차 네 종과 6주차 서비스 테스트 보고서를 만들었고, 모두 코덱스와 교차 검증했습니다. 요구사항 85개 중 구현이 끝난 것은 24개, 일부 구현이 58개, 결정을 기다리는 것이 3개입니다.")

# ================================================================ 13 남은 문제
s = new("남은 문제도 숫자로 드러나 있습니다", "왼쪽 숫자가 지금 값, 오른쪽 회색이 목표입니다")
issues = [("0건", "목표 35행", "사람 평가", "실제 사람이 답을 읽고 매긴 점수가 아직 없음. 평가 시트 35행만 준비"),
          ("62.5%", "목표 90%", "규정 검색 상위 3개 안 정답", "블라인드 평가셋 30/48"),
          ("3.1%", "목표 0%", "일정 항목 찾기 빈 답", "모델이 아무 답도 안 낸 경우 85/2,735. 원인 미확인"),
          ("1/172", "목표 0", "엉뚱한 일정을 바꾸려 함", "채팅 한마디를 할 일로 바꿀 때 다른 항목을 고른 1건"),
          ("1건", "목표 0", "알림 번역 금액 단위 바뀜", "일본어로 옮기며 「35,000원」 → 「35,000円」, 숫자만 대조하는 검사를 통과"),
          ("3.4%", "목표 0%", "계획 읽기 값 지어냄", "원문에 없는 제목 · 예약번호를 뽑은 경우 12/350")]
for i, (big, goal, lab, how) in enumerate(issues):
    x = 0.6 + (i % 3) * 4.1
    y = 1.65 + (i // 3) * 2.65
    box(s, x, y, 3.85, 2.4, fill=ROSE)
    text(s, x + 0.22, y + 0.15, 2.3, 0.7, big, size=30, bold=True, color=RED)
    text(s, x + 2.2, y + 0.35, 1.45, 0.4, goal, size=11, bold=True, color=MUTED, align=PP_ALIGN.RIGHT)
    text(s, x + 0.22, y + 0.95, 3.4, 0.4, lab, size=13.5, bold=True)
    text(s, x + 0.22, y + 1.4, 3.4, 0.95, how, size=10.5, color=MUTED)
footer(s, "모든 값은 AI 가 쓴 합성 시험 기준입니다(사람이 쓴 실제 고객 문장 시험은 아직 없음).")
notes(s, "남은 문제입니다. 가장 큰 것은 사람 평가가 아직 없다는 점입니다. 규정 검색은 목표 90%에 못 미치고, 일정 항목 찾기에서 빈 답이 3.1% 나오며, 알림 번역에서 원이 엔으로 바뀐 사례가 한 건 검사를 통과했습니다.")

# ================================================================ 14 다음 주
s = new()
dark(s)
text(s, 0.9, 0.6, 11.5, 1.0, "다음 주에는 실서버 확인과 사람 평가로 합성 점수의 빈칸을 메웁니다", size=24, bold=True, color=WHITE)
pairs = [("바뀐 화면 대부분 실서버 확인 안 함", "x600 서버의 cs 프로젝트 서버에서 바뀐 화면을 실제로 확인"),
         ("사람 평가 0건", "평가 시트 35행으로 사람 평가 진행"),
         ("4B 모델끼리 순위 미정 · 제품 미연결", "시드 3개 이상으로 다시 학습해 순위 확인, 제품 연결 검토"),
         ("규정 검색 62.5% · 오응답 증가", "고객용 문장 인덱싱을 분류 단계 뒤에서 다시 재고 결정"),
         ("다른 Jev 계열 후보 미시험", "Plumb-4B 등 3종을 같은 방법으로 파인튜닝해 비교 (GPU 서버 B)")]
text(s, 0.9, 1.85, 5.2, 0.4, "지금 막힌 것", size=14, bold=True, color=PALE)
text(s, 7.0, 1.85, 5.5, 0.4, "다음 주에 할 것", size=14, bold=True, color=PALE)
for i, (a, b) in enumerate(pairs):
    y = 2.35 + i * 0.95
    box(s, 0.9, y, 5.2, 0.75, fill=RGBColor(0x2F, 0x55, 0x45))
    text(s, 1.1, y + 0.2, 4.9, 0.4, a, size=13.5, color=WHITE)
    arrow(s, 6.15, y + 0.375, 6.9, y + 0.375, color=PALE, w=2.0)
    box(s, 6.95, y, 5.6, 0.75, fill=GREEN)
    text(s, 7.15, y + 0.2, 5.3, 0.4, b, size=13.5, bold=True, color=WHITE)
notes(s, "다음 주에는 왼쪽의 막힌 것을 하나씩 오른쪽 일로 풉니다. 바뀐 화면을 실서버에서 확인하고, 사람 평가를 진행하고, 시드를 늘려 모델 순위를 확인하고, 규정 검색 개선을 분류 단계 뒤에서 다시 재고, 다른 Jev 계열 후보를 같은 방법으로 비교하겠습니다. 감사합니다.")

# ================================================================ 15 부록: 용어 풀이
s = new("부록. 이 발표의 시험 · 자료 · 구성 이름", "팀이 붙인 이름만 모았습니다(표준 기술 용어는 풀지 않음)")
gl = [("mock 서버", "서버 응답 모양만 흉내 내는 시험용 대역 서버"),
      ("cs 프로젝트 서버", "고객 요청을 처리하는 triPilot 앱(api · worker · web · db)"),
      ("x600 서버", "cs 프로젝트 서버가 올라가 있는 팀 기계"),
      ("GPU 서버 A · B · C", "모델 학습 · 시험용으로 빌린 원격 GPU 기계"),
      ("항로 지킴이", "문제가 생기면 서비스가 먼저 바꿔 두고 알릴지 고르는 스위치"),
      ("Jev 계열 모델", "선택지 중 하나를 고르는 결정 전용 공개 모델 묶음(Winnow-E4B 등)"),
      ("Jev API(유료)", "같은 계열의 종량 과금 외부 결정 API"),
      ("시험 세트 mix8~12", "일정 항목 매칭 5묶음 835문제(AI 합성). mix8~11 Codex · mix12 Claude 작성"),
      ("769문제", "835문제에서 복수 항목 질의 66문제를 뺀 주 지표 분모"),
      ("복수 항목 질의", "한 문장에 일정 항목 두 개를 묻는 질문"),
      ("B · C 채점", "B: 복수 항목 질의는 하나라도 맞으면 정답 · C: 낸 항목 묶음이 정답 묶음과 완전히 같아야 정답"),
      ("E2E 파이프라인 테스트", "규칙 → ML 사전 필터 → 라우터 → 파인튜닝 E4B 를 문제마다 실제로 이어 부른 시험"),
      ("ML 사전 필터", "일정과 무관한 잡담을 먼저 거르는 작은 분류 모델(multilingual-e5-base 기반)"),
      ("혼합 트래픽 테스트", "일정 질문 835 + 관광지 질문 600 + 일상 대화 1,000 + 구어체 300 = 2,735문제"),
      ("블라인드 평가셋 B", "규정 문장을 보지 않고 따로 쓴 규정 질문 160문제(답 있는 48)"),
      ("CX 400", "고객 문장 분류 평가셋 400문장(Codex 작성 · Claude 교차 라벨)"),
      ("인덱싱 텍스트", "규정을 찾을 때 질문과 비교하는 글(직원용 본문 → 고객용 문장으로 바꿔 봄)"),
      ("합성 문제", "사람이 아니라 AI 모델(Codex · Claude)이 쓴 시험 문장")]
half = (len(gl) + 1) // 2
for col in range(2):
    for j, (k, v) in enumerate(gl[col * half:(col + 1) * half]):
        x = 0.6 + col * 6.2
        y = 1.65 + j * 0.58
        box(s, x, y, 5.95, 0.5, fill=MINT if j % 2 == 0 else WHITE)
        text(s, x + 0.15, y + 0.12, 2.1, 0.3, k, size=11, bold=True, color=DEEP)
        text(s, x + 2.3, y + 0.12, 3.55, 0.3, v, size=10, color=INK)
notes(s, "부록은 발표에서 쓴 말을 모은 것입니다. 질문이 있으시면 이 장을 보며 답하겠습니다.")

prs.save(OUT)
(OUTDIR / "[발표] triPilot_진행경과_2026-10-06_Claude_v2_대본.md").write_text(
    "# triPilot 5주차 스탠드업 v2 대본 (Claude 판)\n\n" + "\n\n".join(f"## {i + 1}장\n\n{t}" for i, t in enumerate(SCRIPT)) + "\n", encoding="utf-8")
print("saved", OUT)
