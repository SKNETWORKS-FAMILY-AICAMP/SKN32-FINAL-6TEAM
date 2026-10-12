"""[모델배포] 요구사항 정의서(업데이트 2026-10-06) 빌더.

원본: program/산출물양식/[기획] 요구사항 정의서_2차_여행.xlsx (2026-09-10 판, REQ-001~062) 는 건드리지 않고 복사해 쓴다.
재료: ../_재료/코덱스_변경목록_v2.json (코덱스가 코드 · 문서와 대조해 만든 변경 목록, Claude 검토 · 수정 요청 반영본)
더한 것: H열 「구현 상태(2026-10-06)」 · I열 「상태 근거」 · REQ-063 이후 새 요구 · 시트 「변경 요약」 · 「확인 못 한 것」.
"""
from __future__ import annotations

import copy
import datetime
import json
import math
import re
import shutil
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "_공용"))
import term_pass as TP  # noqa: E402

HERE = Path(__file__).resolve().parent
W5 = HERE.parents[1]
SRC = W5.parent / "[기획] 요구사항 정의서_2차_여행.xlsx"
MAT = HERE.parent / "_재료"
OUT = HERE.parent / "[모델배포] 요구사항 정의서_6팀.xlsx"
BAD = re.compile(r"[—–★☆✓✔…]|[぀-ヿ一-鿿]")
FONT = Font(name="맑은 고딕", size=10)
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")
STATUS_FILL = {"구현 끝": "E2F0D9", "일부": "FFF2CC", "결정 대기": "FCE4D6", "미구현": "F4CCCC"}
NONFUNC = ("보안", "개인정보", "접근성", "검증", "품질", "데이터 사용", "데이터·검증", "가용성", "성능", "확장", "권한")


def load(path: Path) -> dict:
    s = path.read_text(encoding="utf-8")
    m = re.search(r"\{.*\}", s, re.S)
    return json.loads(m.group(0))


def clean(t):
    if t is None:
        return None
    t = str(t)
    bad = BAD.findall(t)
    if bad:
        raise ValueError(f"금지 문자 {bad[:5]}: {t[:80]}")
    return t


def height(ws, row, cols="ABCDEFGHI"):
    lines = 1
    for c in cols:
        v = ws[f"{c}{row}"].value
        if not v:
            continue
        per = max(4, int((ws.column_dimensions[c].width or 10) / 1.85))
        n = sum(max(1, math.ceil(len(part) / per)) for part in str(v).split("\n"))
        lines = max(lines, n)
    ws.row_dimensions[row].height = min(409, 14.5 * lines + 4)


def style(cell, fill=None, bold=False):
    cell.font = Font(name="맑은 고딕", size=10, bold=bold)
    cell.alignment = WRAP
    cell.border = BOX
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def main(src_json: str = "코덱스_변경목록_v2.json"):
    data = load(MAT / src_json)
    v1 = {a["id"]: a for a in load(MAT / "코덱스_변경목록.md").get("adds", [])}  # 1차: B=대분류 · C=중분류
    if OUT.exists():
        bk = OUT.parent / "_backup" / (datetime.datetime.now().strftime("%Y-%m-%d_%H%M") + "_" + OUT.stem)
        bk.mkdir(parents=True, exist_ok=True)
        shutil.copy2(OUT, bk / OUT.name)
    wb = openpyxl.load_workbook(SRC)
    ws = wb["시트1"]
    ws["A1"].value = "요구사항 정의서 (업데이트 2026-10-06)"
    rows = {str(ws.cell(r, 4).value).strip(): r for r in range(5, ws.max_row + 1) if ws.cell(r, 4).value}
    last = max(rows.values())
    # 중분류 → (분류, 대분류): 합친 칸은 맨 위 칸에만 값이 있어 아래로 채워 읽는다
    group = {}
    cur_a = cur_b = None
    for rr in range(5, last + 1):
        cur_a = ws.cell(rr, 1).value or cur_a
        cur_b = ws.cell(rr, 2).value or cur_b
        cc = ws.cell(rr, 3).value
        if cc:
            group.setdefault(str(cc).strip(), (cur_a, cur_b))

    # 머리: H · I 열
    for col, w in (("E", 60), ("F", 95), ("G", 30)):
        ws.column_dimensions[col].width = w
    for col, title, w in (("H", "구현 상태\n(2026-10-06)", 13), ("I", "상태 근거", 70)):
        ws.column_dimensions[col].width = w
        ws[f"{col}3"].value = title
        ws.merge_cells(f"{col}3:{col}4")
        for r in (3, 4):
            ws[f"{col}{r}"]._style = copy.copy(ws[f"G{r}"]._style)

    summary = []
    for ch in data["changes"]:
        rid = ch["id"]
        r = rows[rid]
        for k, v in (ch.get("fields") or {}).items():
            ws[f"{k}{r}"].value = clean(v)
        st = ch.get("status") or "모름"
        ws[f"H{r}"].value = clean(st)
        ws[f"I{r}"].value = clean(ch.get("status_basis"))
        style(ws[f"H{r}"], STATUS_FILL.get(st))
        style(ws[f"I{r}"])
        for c in "EFG":
            ws[f"{c}{r}"].alignment = WRAP
        summary.append((rid, "수정" if ch["action"] == "modify" else "유지", st, clean(ch.get("reason"))))

    # 새 요구: 마지막 행 다음부터
    r = last
    for add in data.get("adds", []):
        r += 1
        a_kind = add.get("A")
        b, c = add.get("B"), add.get("C")
        if b in ("기능", "비기능"):          # 2차: B 에 기능/비기능, C 에 기존 중분류 이름
            if c in group:                    # 기존 중분류면 그 대분류를 시트에서 찾는다
                a_kind, b = group[c]
            else:                             # 새 묶음이면 1차의 대분류 · 중분류를 쓴다
                old = v1.get(add["id"], {})
                a_kind, b, c = b, old.get("B") or c, old.get("C") or c
        if not a_kind:
            a_kind = "비기능" if any(k in f"{b} {c}" for k in NONFUNC) else "기능"
        vals = {"A": a_kind, "B": b, "C": c, "D": add["id"], "E": add.get("E"), "F": add.get("F"),
                "G": add.get("G"), "H": add.get("status"), "I": add.get("status_basis")}
        for col, v in vals.items():
            cell = ws[f"{col}{r}"]
            cell.value = clean(v)
            style(cell, STATUS_FILL.get(v) if col == "H" else ("EAF1FB" if col in "ABCD" else None))
        summary.append((add["id"], "추가", add.get("status"), clean(add.get("reason") or add.get("E"))))

    for row in range(5, r + 1):
        if ws.cell(row, 4).value:
            height(ws, row)
    ws.freeze_panes = "E5"

    # 변경 요약 시트
    s2 = wb.create_sheet("변경 요약")
    s2.append(["요구사항 ID", "이번 판에서", "구현 상태(2026-10-06)", "바꾼 이유 · 새 요구 내용"])
    for row in summary:
        s2.append(list(row))
    for col, w in zip("ABCD", (13, 11, 16, 90)):
        s2.column_dimensions[col].width = w
    for i, rr in enumerate(s2.iter_rows(min_row=1, max_row=s2.max_row), start=1):
        for cell in rr:
            style(cell, "D9E1F2" if i == 1 else (STATUS_FILL.get(cell.value) if cell.column == 3 else None), bold=(i == 1))
        if i > 1:
            height(s2, i, "ABCD")
    cnt = {}
    for _, act, st, _ in summary:
        cnt[act] = cnt.get(act, 0) + 1
        cnt[st] = cnt.get(st, 0) + 1
    tot = len(summary)
    s2.append([])
    s2.append(["합계", f"유지 {cnt.get('유지', 0)} · 수정 {cnt.get('수정', 0)} · 추가 {cnt.get('추가', 0)}",
               f"구현 끝 {cnt.get('구현 끝', 0)}/{tot} · 일부 {cnt.get('일부', 0)}/{tot} · 결정 대기 {cnt.get('결정 대기', 0)}/{tot}",
               "구현 상태는 코드 · 문서 · 기존 시험 기록을 읽어 판정했다(2026-10-06). 새로 시험을 돌리거나 x600 서버에서 확인하지 않았다. "
               "「구현 끝」은 적힌 범위의 구현을 뜻하며 배포 완료가 아니다."])
    for cell in s2[s2.max_row]:
        style(cell, "F2F2F2", bold=True)
    height(s2, s2.max_row, "ABCD")

    # 확인 못 한 것
    s3 = wb.create_sheet("확인 못 한 것")
    s3.append(["번호", "확인하지 못했거나 서로 어긋나는 것 (2026-10-06)"])
    for i, t in enumerate(data.get("doubts", []), start=1):
        s3.append([i, clean(t)])
    s3.column_dimensions["A"].width = 7
    s3.column_dimensions["B"].width = 120
    for i, rr in enumerate(s3.iter_rows(min_row=1, max_row=s3.max_row), start=1):
        for cell in rr:
            style(cell, "D9E1F2" if i == 1 else None, bold=(i == 1))
        if i > 1:
            height(s3, i, "AB")
    # 용어 정비(2026-10-06 사용자 지적 「너만 아는 단어 다 수정」): 모든 칸을 실무 용어로 바꾸고 「용어 정의」 시트를 맨 앞에 둔다
    body = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and not cell.value.startswith("="):
                    cell.value = TP.run_text(cell.value, "req")
                    body.append(cell.value)
    rows = TP.glossary_rows("req", " ".join(body))
    sg = wb.create_sheet("용어 정의", 0)
    sg.append(["용어", "뜻"])
    for k, v in rows:
        sg.append([k, v])
    sg.column_dimensions["A"].width = 26
    sg.column_dimensions["B"].width = 110
    for i, rr in enumerate(sg.iter_rows(min_row=1, max_row=sg.max_row), start=1):
        for cell in rr:
            style(cell, "D9E1F2" if i == 1 else None, bold=(i == 1 or cell.column == 1))
        if i > 1:
            height(sg, i, "AB")
    wb.active = 1
    wb.save(OUT)
    print("saved", OUT, "| 행", r, "| 요약", cnt, "| 용어", len(rows))


if __name__ == "__main__":
    main(*sys.argv[1:])
