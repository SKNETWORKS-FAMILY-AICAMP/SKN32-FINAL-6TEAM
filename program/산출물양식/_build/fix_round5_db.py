# -*- coding: utf-8 -*-
"""DB 5차 — 기계 치환이 남긴 흔적 정리."""
import re
import shutil
import sys

from docx import Document

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
NAME = "[데이터 수집 및 저장] 데이터베이스저장소 설계 문서_2차_여행.docx"
P = "../w2/제출본/" + NAME
REP = [
    ("인계 근거를 남긴다..", "인계 근거를 남긴다."),
    ("본 뒤 채택한다..", "본 뒤 채택한다."),
    ("watch.py에는 관측·변화 저장만 있고", "감시 모듈에는 관측·변화 저장만 있고"),
    ("--는 DDL에 선언된 FK..는 제안·논리·다형 연결이다.", "실선(--)은 DDL에 선언된 FK, 점선(..)은 제안·논리·다형 연결이다."),
    ("--도 환경 적용 완료를 뜻하지 않는다.", "실선이라도 실제 환경에 적용됐다는 뜻은 아니다."),
    ("011·013에 연결 FK 없음.", "마이그레이션 011·013에 연결 FK 없음."),
    ("012의 target_kind·target_id는 논리 식별", "마이그레이션 012의 target_kind·target_id는 논리 식별"),
    ("012에 observation_id FK 없음.", "마이그레이션 012에 observation_id FK 없음."),
    ("004의 유일 인덱스 agent_runs_one_active_per_case가 강제한다.", "마이그레이션 004의 유일 인덱스 agent_runs_one_active_per_case가 강제한다."),
    ("013에서 추가.", "마이그레이션 013에서 추가."),
    ("012의 watch_changes_pending_idx는", "마이그레이션 012의 watch_changes_pending_idx는"),
    ("worker 상태와 무효화 제안은 C-03·11.", "배달 작업자의 상태 처리와 무효화 제안은 C-03·11."),
    ("worker의 버전 검사", "배달 작업자의 버전 검사"),
    ("한 스윕·한 Trip에 Case 하나", "한 번의 감시 회차·한 Trip에 Case 하나"),
    ("같은 스윕·Trip의", "같은 감시 회차·Trip의"),
    ("스윕", "감시 회차"),
    ("watch.py", "감시 모듈"),
]
DUP = " agent_runs의 활성 실행 유일 인덱스는 마이그레이션 004에 있다."
DUP2 = ". agent_runs의 활성 실행 유일 인덱스는 마이그레이션 004에 있다."


def main():
    d = Document(P)
    # 엔터티 표: agent_runs 가 아닌 행에서 복사된 문장 제거
    ent = [t for t in d.tables if t.rows[0].cells[1].text.strip() == "테이블"][0]
    dup_fixed = 0
    for row in ent.rows[1:]:
        if row.cells[1].text.strip() == "agent_runs":
            continue
        for c in row.cells:
            for pp in c.paragraphs:
                t = pp.text
                if DUP2 in t or DUP in t:
                    t = t.replace(DUP2, ".").replace(DUP, "")
                    pp.runs[0].text = t
                    for r in pp.runs[1:]:
                        r.text = ""
                    dup_fixed += 1
    paras = list(d.paragraphs) + [pp for t in d.tables for r in t.rows for c in r.cells for pp in c.paragraphs]
    seen, n = set(), 0
    for pp in paras:
        if id(pp._p) in seen or not pp.runs:
            continue
        seen.add(id(pp._p))
        t0 = pp.text
        t = t0
        for a, b in REP:
            t = t.replace(a, b)
        t = re.sub(r"(?<![.])\.\.$", ".", t)          # 칸 끝 「..」
        t = re.sub(r"(?<=[가-힣)])\.\.(?=\s)", ".", t)  # 문장 사이 「..」 (점선 기호 「(..)」 는 괄호 안이라 제외)
        if t != t0:
            pp.runs[0].text = t
            for r in pp.runs[1:]:
                r.text = ""
            n += 1
    d.save(P)
    shutil.copy(P, "../" + NAME)
    allt = [pp.text for pp in paras]
    print("DB 5차: 고친 문단 %d | 복사 문장 제거 %d" % (n, dup_fixed))
    print("남은 흔적: .py %d | 스윕 %d | worker %d | 마침표 중복 %d | 줄번호 %d" % (
        sum(len(re.findall(r"\w+\.py\b", x)) for x in allt), sum(x.count("스윕") for x in allt),
        sum(len(re.findall(r"\bworker\b", x)) for x in allt),
        sum(len(re.findall(r"[가-힣]\.\.(?!\))", x)) for x in allt),
        sum(len(re.findall(r"\b\d{3}:\d+", x)) for x in allt)))


if __name__ == "__main__":
    main()
