# -*- coding: utf-8 -*-
"""양식 docx 의 표지·머리 표는 그대로 두고, 본문을 통째로 새로 짜는 조립 도구 (4차 v4).

v3 와 다른 점(코덱스 명세 2026-09-28 + 반박 반영)
- 본문 10.5pt 고정 줄간격 16pt, 표 9.5pt 고정 14pt, 캡션 9pt. 제목 3단계(팀 13 · 주제 11.5 · 블록 10.5).
- 표는 폭 9,345 dxa 고정, 열 4개 이하(5열은 짧은 값만), 본문 표 8행·상세 근거 표 12행 이하, 머리행 반복, 행 쪼갬 금지.
- 그림은 삽입 폭 155~165mm, 300dpi 이상, 삽입 크기 기준 글자 9.5pt 이상(그림 쪽에서 크기를 맞춘다).
- 요약 카드는 2x2(지표 종류 / 값 / 조건).
- 글이 규격을 넘으면 `warnings` 에 쌓아 두었다가 빌드가 끝나기 전에 보여 준다(추정 줄 수. 실제 줄 수는 qa_docx.py 가 Word PDF 에서 잰다).
★양식의 절 제목·번호·표지·머리 표는 건드리지 않는다.
"""
from __future__ import annotations

import math
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from docxfill import check, esc, png_size  # noqa: E402

FONT = ('<w:rFonts w:ascii="Malgun Gothic" w:cs="Malgun Gothic" '
        'w:eastAsia="Malgun Gothic" w:hAnsi="Malgun Gothic"/>')
BLUE = "2A78D6"
ORANGE = "EB6834"
NAVY = "1F3864"
HEAD_FILL = "CFE2F3"
ZEBRA = "F4F7FB"
BORDER = "D9E2F3"
INK = "000000"
MUTED = "595959"
TABLE_W = 9345

BODY_SZ, TBL_SZ, CAP_SZ = 21, 19, 18         # half-point
BODY_LINE, TBL_LINE, CAP_LINE = 320, 280, 260  # twip, 고정(exact)
KOR_CM = {21: 0.37, 19: 0.335, 18: 0.32}      # 한글 한 글자 폭(cm) 추정


def _rpr(size=BODY_SZ, bold=False, color=None, italic=False):
    parts = [FONT]
    parts.append('<w:b w:val="1"/><w:bCs w:val="1"/>' if bold else '<w:b w:val="0"/><w:bCs w:val="0"/>')
    if italic:
        parts.append('<w:i w:val="1"/>')
    parts.append('<w:color w:val="%s"/>' % (color or "000000"))
    parts.append('<w:sz w:val="%d"/><w:szCs w:val="%d"/><w:rtl w:val="0"/>' % (size, size))
    return "<w:rPr>%s</w:rPr>" % "".join(parts)


def _runs(text, size=BODY_SZ, bold=False, color=None, italic=False):
    """`**굵게**` 표시를 굵은 글로 바꾼다."""
    check(text)
    out = []
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if not chunk:
            continue
        out.append('<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r>'
                   % (_rpr(size, bold or i % 2 == 1, color), esc(chunk)))
    return "".join(out)


def est_lines(text: str, width_cm: float, size=BODY_SZ) -> int:
    """글자 수로 줄 수를 어림한다(영문·숫자는 한글의 0.55배 폭)."""
    text = text.replace("**", "")
    total = 0
    for chunk in text.split("\n"):
        w = sum(0.55 if ord(ch) < 0x2E80 else 1.0 for ch in chunk) * KOR_CM.get(size, 0.37)
        total += max(1, math.ceil(w / width_cm))
    return total


class Composer:
    def __init__(self, template: Path):
        self.doc = Document(str(template))
        self._normalise_page_margins()
        body = self.doc.element.body
        self.sect = body.find(qn("w:sectPr"))
        head_table = self.doc.tables[2]._tbl
        removing = False
        for el in list(body.iterchildren()):
            if el is head_table:
                removing = True
                continue
            if removing and el is not self.sect:
                body.remove(el)
        self.figures = 0
        self.tables = 0
        self.warnings: list[str] = []
        self.registry: list[dict] = []      # 검사용 목록: 블록 종류·표·그림

    def warn(self, msg: str):
        self.warnings.append(msg)

    def _normalise_page_margins(self):
        for section in self.doc.element.body.iter(qn("w:pgMar")):
            for key, value in list(section.attrib.items()):
                if "." in value:
                    section.set(key, str(round(float(value))))
            if section.get(qn("w:gutter")) is None:
                section.set(qn("w:gutter"), "0")

    # -- 표지 ---------------------------------------------------------
    def cover(self, cover_title: str, stage: str, date: str, repo: str, authors: str):
        first = self.doc.tables[0].rows[0].cells[0].paragraphs[0]
        for r in list(first.runs)[1:]:
            r._element.getparent().remove(r._element)
        first.runs[0].text = cover_title
        meta = self.doc.tables[1]
        for row, value in zip(meta.rows, (stage, date, repo, authors)):
            cell = row.cells[1]
            p = cell.paragraphs[0]
            for extra in list(cell.paragraphs)[1:]:
                extra._element.getparent().remove(extra._element)
            for r in list(p.runs)[1:]:
                r._element.getparent().remove(r._element)
            if not p.runs:
                p.add_run("")
            check(value, "표지")
            p.runs[0].text = value

    def _add(self, xml: str):
        el = parse_xml(_with_ns(xml))
        self.sect.addprevious(el)
        return el

    def gap(self, twips=80, keep_next=False):
        self._add('<w:p><w:pPr>%s<w:spacing w:before="0" w:after="0" w:line="%d" w:lineRule="exact"/></w:pPr></w:p>'
                  % ('<w:keepNext w:val="1"/>' if keep_next else "", twips))

    # -- 제목 ---------------------------------------------------------
    def h1(self, text: str):
        """양식의 1단계 제목(번호 자동). 양식 그대로."""
        check(text)
        self._add('<w:p><w:pPr><w:keepNext w:val="1"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="3"/></w:numPr>'
                  '<w:spacing w:before="280" w:after="100" w:line="276" w:lineRule="auto"/>'
                  '<w:ind w:left="720" w:hanging="360"/></w:pPr>'
                  '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r></w:p>' % (_rpr(26, True, "1F3A5F"), esc(text)))
        self.registry.append(dict(kind="h1", text=text))

    def h2(self, text: str):
        """양식의 2단계 제목(2.1 등)."""
        check(text)
        self._add('<w:p><w:pPr><w:keepNext w:val="1"/><w:spacing w:before="220" w:after="80" w:line="276" w:lineRule="auto"/>'
                  '<w:ind w:left="240" w:firstLine="0"/></w:pPr>'
                  '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r></w:p>' % (_rpr(22, True, "1F2933"), esc(text)))
        self.registry.append(dict(kind="h2", text=text))

    def h3(self, text: str):
        """신규 제목 1단계: 팀. 13pt 굵게 1F3864, 앞 12pt 뒤 6pt."""
        check(text)
        self._add('<w:p><w:pPr><w:keepNext w:val="1"/><w:pBdr><w:left w:val="single" w:sz="24" w:space="6" w:color="%s"/></w:pBdr>'
                  '<w:spacing w:before="240" w:after="120"/><w:ind w:left="300" w:firstLine="0"/></w:pPr>'
                  '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r></w:p>' % (BLUE, _rpr(26, True, NAVY), esc(text)))
        self.registry.append(dict(kind="h3", text=text))

    def h4(self, text: str):
        """신규 제목 2단계: 주제. 11.5pt 굵게 1F3864, 앞 9pt 뒤 4pt."""
        check(text)
        self._add('<w:p><w:pPr><w:keepNext w:val="1"/><w:spacing w:before="180" w:after="80"/><w:ind w:left="240" w:firstLine="0"/></w:pPr>'
                  '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r></w:p>' % (_rpr(23, True, NAVY), esc(text)))
        self.registry.append(dict(kind="h4", text=text))

    def h5(self, text: str):
        """신규 제목 3단계: 블록 표제. 10.5pt 굵게 검정, 앞 6pt 뒤 3pt."""
        check(text)
        self._add('<w:p><w:pPr><w:keepNext w:val="1"/><w:spacing w:before="120" w:after="60"/><w:ind w:left="240" w:firstLine="0"/></w:pPr>'
                  '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r></w:p>' % (_rpr(BODY_SZ, True, "000000"), esc(text)))
        self.registry.append(dict(kind="h5", text=text))

    # -- 본문 ---------------------------------------------------------
    def _para_xml(self, text, keep=False, bold=False, color=None, indent=240, after=100):
        return ('<w:p><w:pPr>%s<w:widowControl w:val="1"/><w:spacing w:before="0" w:after="%d" w:line="%d" w:lineRule="exact"/>'
                '<w:ind w:left="%d"/><w:jc w:val="left"/></w:pPr>%s</w:p>'
                % ('<w:keepNext w:val="1"/>' if keep else "", after, BODY_LINE, indent, _runs(text, BODY_SZ, bold, color)))

    def para(self, text: str, keep=False, max_lines=4):
        n = est_lines(text, 16.0)
        if n > max_lines:
            self.warn(f"문단 {n}줄 추정(>{max_lines}): {text[:36]}")
        self._add(self._para_xml(text, keep=keep))
        self.registry.append(dict(kind="para", text=text))

    def lead(self, text: str, keep=True):
        """그 자리의 결론 한 문장(굵게). 표·그림 앞에서는 그것과 같은 쪽에 묶는다."""
        n = est_lines(text, 16.0)
        if n > 2:
            self.warn(f"결론 문장 {n}줄 추정(>2): {text[:36]}")
        self._add(self._para_xml(text, keep=keep, bold=True))
        self.registry.append(dict(kind="lead", text=text))

    def line(self, label: str, value: str):
        """양식의 「라벨 : 값」 글머리 줄. 3줄 이하."""
        check(label + value)
        n = est_lines(label + " : " + value, 15.0)
        if n > 3:
            self.warn(f"글머리 {n}줄 추정(>3): {label}")
        self._add('<w:p><w:pPr><w:widowControl w:val="1"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'
                  '<w:spacing w:before="0" w:after="80" w:line="%d" w:lineRule="exact"/><w:ind w:left="720" w:hanging="360"/>'
                  '<w:rPr><w:color w:val="000000"/></w:rPr></w:pPr>'
                  '<w:r>%s<w:t xml:space="preserve">%s : </w:t></w:r>%s</w:p>'
                  % (BODY_LINE, _rpr(BODY_SZ, True), esc(label), _runs(value, BODY_SZ)))
        self.registry.append(dict(kind="line", text=label + value))

    def bullet(self, text: str):
        check(text)
        n = est_lines(text, 15.0)
        if n > 3:
            self.warn(f"글머리 {n}줄 추정(>3): {text[:30]}")
        self._add('<w:p><w:pPr><w:widowControl w:val="1"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'
                  '<w:spacing w:before="0" w:after="80" w:line="%d" w:lineRule="exact"/><w:ind w:left="720" w:hanging="360"/>'
                  '<w:rPr><w:color w:val="000000"/></w:rPr></w:pPr>%s</w:p>' % (BODY_LINE, _runs(text, BODY_SZ)))
        self.registry.append(dict(kind="bullet", text=text))

    def note(self, text: str, keep=False):
        """표·그림 아래 출처·주의(9pt, 595959, 줄 간격 13pt)."""
        n = est_lines(text, 16.0, CAP_SZ)
        if n > 3:
            self.warn(f"출처글 {n}줄 추정(>3): {text[:30]}")
        self._add('<w:p><w:pPr>%s<w:widowControl w:val="1"/><w:spacing w:before="0" w:after="100" w:line="%d" w:lineRule="exact"/><w:ind w:left="240"/></w:pPr>%s</w:p>'
                  % ('<w:keepNext w:val="1"/>' if keep else "", CAP_LINE, _runs(text, CAP_SZ, False, MUTED)))
        self.registry.append(dict(kind="note", text=text))

    # -- 표 -----------------------------------------------------------
    def table(self, title, header, rows, widths, align=None, first_bold=True, detail=False, cont=False, five_ok=False):
        """폭 9,345 dxa. detail=True 는 상세 근거 표(12행), 아니면 본문 표(8행). widths 는 비율."""
        n_cols = len(header)
        limit = 12 if detail else 8
        if n_cols > 5:
            self.warn(f"표 열 {n_cols}개(>5): {title}")
        if len(rows) > limit:
            self.warn(f"표 행 {len(rows)}개(>{limit}): {title}")
        total = sum(widths)
        w = [round(TABLE_W * x / total) for x in widths]
        w[-1] += TABLE_W - sum(w)
        align = align or ["left"] * n_cols
        for i, h in enumerate(header):
            if len(h) > 16:
                self.warn(f"표 머리 {len(h)}자(>16): {title} / {h}")
        long5 = n_cols == 5
        for r in rows:
            if len(r) != n_cols:
                self.warn(f"열 수 불일치: {title}")
            for i, v in enumerate(r):
                cm = w[i] / TABLE_W * 16.48 - 0.28
                lines = est_lines(str(v), cm, TBL_SZ)
                if lines > 4:
                    self.warn(f"표 셀 {lines}줄 추정(>4): {title} / {str(v)[:26]}")
                if long5 and not five_ok and lines > 2:
                    self.warn(f"5열 표 셀 {lines}줄 추정(>2): {title} / {str(v)[:26]}")
                if i == 0 and len(str(v)) > 20 and not all(ord(ch) < 0x80 or ch == '·' for ch in str(v)):
                    self.warn(f"표 첫 열 {len(str(v))}자(>20): {title} / {str(v)[:20]}")
        self.tables += 1
        num = self.tables
        cap = "표 %d. %s%s" % (num, title, " (계속)" if cont else "")
        check(cap, "표 제목")
        grid = "".join('<w:gridCol w:w="%d"/>' % x for x in w)

        def cell(text, width, fill, bold, al, hdr=False, keep=False):
            check(str(text), "표")
            shd = '<w:shd w:val="clear" w:color="auto" w:fill="%s"/>' % fill if fill else ""
            paras = []
            for chunk in str(text).split("\n"):
                paras.append('<w:p><w:pPr>%s<w:keepLines w:val="1"/><w:spacing w:before="0" w:after="0" w:line="%d" w:lineRule="exact"/>'
                             '<w:jc w:val="%s"/></w:pPr>%s</w:p>'
                             % ('<w:keepNext w:val="1"/>' if keep else "", TBL_LINE, "center" if hdr else al,
                                _runs(chunk, TBL_SZ, bold or hdr, "000000" if hdr else "1F2933")))
            return ('<w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/><w:tcBorders>'
                    '<w:top w:val="single" w:sz="4" w:space="0" w:color="%s"/><w:left w:val="single" w:sz="4" w:space="0" w:color="%s"/>'
                    '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="%s"/><w:right w:val="single" w:sz="4" w:space="0" w:color="%s"/>'
                    '</w:tcBorders>%s<w:tcMar><w:top w:w="60" w:type="dxa"/><w:left w:w="80" w:type="dxa"/>'
                    '<w:bottom w:w="60" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tcMar><w:vAlign w:val="%s"/></w:tcPr>%s</w:tc>'
                    % (width, BORDER, BORDER, BORDER, BORDER, shd, "center" if hdr else "top", "".join(paras)))

        # 표 제목을 표의 첫 행(머리행)에 넣는다 — 쪽을 넘어가면 제목과 열 이름이 새 쪽에서 되풀이되어, 앞쪽 없이도 무슨 표인지 안다.
        title_cell = ('<w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/><w:gridSpan w:val="%d"/><w:tcBorders><w:top w:val="nil"/><w:left w:val="nil"/>'
                      '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="%s"/><w:right w:val="nil"/></w:tcBorders>'
                      '<w:tcMar><w:top w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/><w:bottom w:w="40" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tcMar></w:tcPr>'
                      '<w:p><w:pPr><w:keepNext w:val="1"/><w:keepLines w:val="1"/><w:spacing w:before="60" w:after="0" w:line="%d" w:lineRule="exact"/></w:pPr>%s</w:p></w:tc>'
                      % (TABLE_W, n_cols, BORDER, CAP_LINE, _runs(cap, CAP_SZ, True, MUTED)))
        trs = ['<w:tr><w:trPr><w:cantSplit w:val="1"/><w:tblHeader w:val="1"/></w:trPr>%s</w:tr>' % title_cell,
               '<w:tr><w:trPr><w:cantSplit w:val="1"/><w:tblHeader w:val="1"/></w:trPr>%s</w:tr>'
               % "".join(cell(h, w[i], HEAD_FILL, True, "center", hdr=True, keep=True) for i, h in enumerate(header))]
        for r, row in enumerate(rows):
            fill = ZEBRA if r % 2 == 1 else None
            keep = r == 0 and len(rows) > 1                       # 머리행 + 데이터 2행이 같은 쪽에 묶이도록
            trs.append('<w:tr><w:trPr><w:cantSplit w:val="1"/></w:trPr>%s</w:tr>'
                       % "".join(cell(v, w[i], fill, first_bold and i == 0 and len(str(v)) <= 20, align[i], keep=keep) for i, v in enumerate(row)))
        self._add('<w:tbl><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:jc w:val="center"/><w:tblLayout w:type="fixed"/>'
                  '<w:tblLook w:val="0420"/><w:tblCaption w:val="%s"/></w:tblPr><w:tblGrid>%s</w:tblGrid>%s</w:tbl>'
                  % (TABLE_W, esc("%s:%d" % ("detail" if detail else "body", num)), grid, "".join(trs)))
        self.registry.append(dict(kind="table", num=num, title=title, rows=len(rows), cols=n_cols, detail=detail))
        self.gap(100)
        return num

    def form_table(self, name, header, rows, widths, align=None, first_bold=True):
        """v5 양식 표 — 양식의 열 구성 그대로, 표 제목 없음(양식에 없다), 머리행 하나를 쪽마다 반복, 행 쪼갬 금지.
        name 은 검사·경고용 이름이다(문서에 찍히지 않는다)."""
        n_cols = len(header)
        if len(rows) > 12:
            self.warn(f"표 행 {len(rows)}개(>12): {name}")
        total = sum(widths)
        w = [round(TABLE_W * x / total) for x in widths]
        w[-1] += TABLE_W - sum(w)
        align = align or ["left"] * n_cols
        for r in rows:
            if len(r) != n_cols:
                self.warn(f"열 수 불일치: {name}")
            for i, v in enumerate(r):
                cm = w[i] / TABLE_W * 16.48 - 0.28
                lines = est_lines(str(v), cm, TBL_SZ)
                if lines > 4:
                    self.warn(f"표 셀 {lines}줄 추정(>4): {name} / {str(v)[:26]}")
        self.tables += 1
        grid = "".join('<w:gridCol w:w="%d"/>' % x for x in w)

        def cell(text, width, fill, bold, al, hdr=False, keep=False):
            check(str(text), "표")
            shd = '<w:shd w:val="clear" w:color="auto" w:fill="%s"/>' % fill if fill else ""
            paras = []
            for chunk in str(text).split("\n"):
                paras.append('<w:p><w:pPr>%s<w:keepLines w:val="1"/><w:spacing w:before="0" w:after="0" w:line="%d" w:lineRule="exact"/>'
                             '<w:jc w:val="%s"/></w:pPr>%s</w:p>'
                             % ('<w:keepNext w:val="1"/>' if keep else "", TBL_LINE, "center" if hdr else al,
                                _runs(chunk, TBL_SZ, bold or hdr, "000000" if hdr else "1F2933")))
            return ('<w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/><w:tcBorders>'
                    '<w:top w:val="single" w:sz="4" w:space="0" w:color="%s"/><w:left w:val="single" w:sz="4" w:space="0" w:color="%s"/>'
                    '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="%s"/><w:right w:val="single" w:sz="4" w:space="0" w:color="%s"/>'
                    '</w:tcBorders>%s<w:tcMar><w:top w:w="60" w:type="dxa"/><w:left w:w="80" w:type="dxa"/>'
                    '<w:bottom w:w="60" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tcMar><w:vAlign w:val="%s"/></w:tcPr>%s</w:tc>'
                    % (width, BORDER, BORDER, BORDER, BORDER, shd, "center" if hdr else "top", "".join(paras)))

        trs = ['<w:tr><w:trPr><w:cantSplit w:val="1"/><w:tblHeader w:val="1"/></w:trPr>%s</w:tr>'
               % "".join(cell(h, w[i], HEAD_FILL, True, "center", hdr=True, keep=True) for i, h in enumerate(header))]
        for r, row in enumerate(rows):
            fill = ZEBRA if r % 2 == 1 else None
            keep = r == 0 and len(rows) > 1
            trs.append('<w:tr><w:trPr><w:cantSplit w:val="1"/></w:trPr>%s</w:tr>'
                       % "".join(cell(v, w[i], fill, first_bold and i == 0 and len(str(v)) <= 20, align[i], keep=keep) for i, v in enumerate(row)))
        self._add('<w:tbl><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:jc w:val="center"/><w:tblLayout w:type="fixed"/>'
                  '<w:tblLook w:val="0420"/><w:tblCaption w:val="form:%d"/></w:tblPr><w:tblGrid>%s</w:tblGrid>%s</w:tbl>'
                  % (TABLE_W, self.tables, grid, "".join(trs)))
        self.registry.append(dict(kind="table", name=name, rows=len(rows), cols=n_cols))
        self.gap(100)

    def table_split(self, title, header, rows, widths, limit=12, **kw):
        """행이 limit 을 넘으면 물리적으로 나눈다. 조각마다 머리행이 반복되고 표 제목에 「(계속)」이 붙으며 마지막 조각은 2행 이상이다."""
        rows = list(rows)
        if len(rows) <= limit:
            return self.table(title, header, rows, widths, detail=True, **kw)
        n_parts = -(-len(rows) // limit)
        size = -(-len(rows) // n_parts)
        for i in range(n_parts):
            part = rows[i * size:(i + 1) * size]
            self.table(title, header, part, widths, detail=True, cont=i > 0, **kw)

    def kpi(self, cards):
        """2x2 카드. 카드 = (지표 종류, 값 또는 상태, 평가 대상·조건). 값 17pt · 표제 10pt · 조건 9pt(2줄 이하)."""
        assert len(cards) == 4, "카드는 2x2, 네 개"
        w = TABLE_W // 2
        w2 = TABLE_W - w

        def one(kind, value, cond, width):
            check(kind + value + cond, "카드")
            if est_lines(cond, 7.3, CAP_SZ) > 2:
                self.warn(f"카드 조건 {est_lines(cond, 7.3, CAP_SZ)}줄 추정(>2): {cond[:24]}")
            return ('<w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/><w:tcBorders><w:top w:val="single" w:sz="18" w:space="0" w:color="%s"/>'
                    '<w:left w:val="single" w:sz="8" w:space="0" w:color="FFFFFF"/><w:bottom w:val="single" w:sz="8" w:space="0" w:color="FFFFFF"/>'
                    '<w:right w:val="single" w:sz="8" w:space="0" w:color="FFFFFF"/></w:tcBorders>'
                    '<w:shd w:val="clear" w:color="auto" w:fill="F4F7FB"/><w:tcMar><w:top w:w="90" w:type="dxa"/><w:left w:w="140" w:type="dxa"/>'
                    '<w:bottom w:w="90" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar></w:tcPr>'
                    '<w:p><w:pPr><w:keepNext w:val="1"/><w:spacing w:before="0" w:after="0" w:line="260" w:lineRule="exact"/></w:pPr>%s</w:p>'
                    '<w:p><w:pPr><w:keepNext w:val="1"/><w:spacing w:before="0" w:after="20" w:line="400" w:lineRule="exact"/></w:pPr>%s</w:p>'
                    '<w:p><w:pPr><w:spacing w:before="0" w:after="0" w:line="%d" w:lineRule="exact"/></w:pPr>%s</w:p></w:tc>'
                    % (width, BLUE, _runs(kind, 20, True, NAVY), _runs(value, 34, True, BLUE), CAP_LINE, _runs(cond, CAP_SZ, False, MUTED)))
        rows = "".join('<w:tr><w:trPr><w:cantSplit w:val="1"/></w:trPr>%s</w:tr>' % "".join(one(*card, wd) for card, wd in zip(cards[i:i + 2], (w, w2))) for i in (0, 2))
        self._add('<w:tbl><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:jc w:val="center"/><w:tblLayout w:type="fixed"/>'
                  '<w:tblLook w:val="0000"/><w:tblCaption w:val="card"/></w:tblPr><w:tblGrid><w:gridCol w:w="%d"/><w:gridCol w:w="%d"/></w:tblGrid>%s</w:tbl>'
                  % (TABLE_W, w, w2, rows))
        self.registry.append(dict(kind="cards"))
        self.gap(100)

    def callout(self, title: str, lines, kind="info"):
        """공통 해석 주의 상자(문서에 하나). 줄은 4줄 이하."""
        fill, bar = {"info": ("EEF5FD", BLUE), "warn": ("FDF3E8", ORANGE)}[kind]
        n = sum(est_lines(x, 15.0, TBL_SZ) for x in lines)
        if n > 4:
            self.warn(f"상자 {n}줄 추정(>4): {title}")
        paras = ['<w:p><w:pPr><w:keepNext w:val="1"/><w:spacing w:before="0" w:after="40" w:line="%d" w:lineRule="exact"/></w:pPr>%s</w:p>' % (TBL_LINE, _runs(title, TBL_SZ, True, NAVY))]
        for line in lines:
            check(line, "상자")
            paras.append('<w:p><w:pPr><w:spacing w:before="0" w:after="20" w:line="%d" w:lineRule="exact"/><w:ind w:left="200" w:hanging="200"/></w:pPr>%s%s</w:p>'
                         % (TBL_LINE, _runs("· ", TBL_SZ, False, MUTED), _runs(line, TBL_SZ, False, "1F2933")))
        self._add('<w:tbl><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:jc w:val="center"/><w:tblLayout w:type="fixed"/><w:tblLook w:val="0000"/><w:tblCaption w:val="callout"/></w:tblPr>'
                  '<w:tblGrid><w:gridCol w:w="%d"/></w:tblGrid><w:tr><w:trPr><w:cantSplit w:val="1"/></w:trPr><w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/>'
                  '<w:tcBorders><w:top w:val="nil"/><w:left w:val="single" w:sz="30" w:space="0" w:color="%s"/><w:bottom w:val="nil"/><w:right w:val="nil"/></w:tcBorders>'
                  '<w:shd w:val="clear" w:color="auto" w:fill="%s"/><w:tcMar><w:top w:w="90" w:type="dxa"/><w:left w:w="160" w:type="dxa"/>'
                  '<w:bottom w:w="70" w:type="dxa"/><w:right w:w="140" w:type="dxa"/></w:tcMar></w:tcPr>%s</w:tc></w:tr></w:tbl>'
                  % (TABLE_W, TABLE_W, TABLE_W, bar, fill, "".join(paras)))
        self.registry.append(dict(kind="callout", title=title))
        self.gap(100)

    # -- 그림 ---------------------------------------------------------
    def figure(self, png: Path, caption: str):
        """그림은 원본 크기(300dpi 기준 inch) 그대로 넣는다. 155~165mm 를 벗어나면 경고."""
        w_px, h_px = png_size(str(png))
        dpi = 300
        width_cm = w_px / dpi * 2.54
        if not 15.5 <= width_cm <= 16.5:
            self.warn(f"그림 폭 {width_cm * 10:.0f}mm(155~165mm 밖): {png.name}")
        p = self.doc.add_paragraph()
        self.sect.addprevious(p._p)
        p.alignment = 1
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = 60
        p.paragraph_format.space_after = 0
        p.add_run().add_picture(str(png), width=Cm(width_cm))
        self.figures += 1
        caption = "그림 %d. %s" % (self.figures, caption)
        check(caption, "캡션")
        if est_lines(caption, 16.0, CAP_SZ) > 2:
            self.warn(f"캡션 {est_lines(caption, 16.0, CAP_SZ)}줄 추정(>2): {caption[:30]}")
        self._add('<w:p><w:pPr><w:widowControl w:val="1"/><w:spacing w:before="0" w:after="140" w:line="%d" w:lineRule="exact"/><w:jc w:val="center"/></w:pPr>%s</w:p>'
                  % (CAP_LINE, _runs(caption, CAP_SZ, False, MUTED)))
        self.registry.append(dict(kind="figure", num=self.figures, file=png.name, width_mm=round(width_cm * 10, 1), px=(w_px, h_px), dpi=round(w_px / (width_cm / 2.54))))

    # -- 저장 ---------------------------------------------------------
    def all_text(self):
        parts = [p.text for p in self.doc.paragraphs]
        for tb in self.doc.tables:
            for row in tb.rows:
                parts += [c.text for c in row.cells]
        return "\n".join(parts)

    def save(self, out: Path):
        check(self.all_text(), out.name)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))
        reopened = Document(str(out))
        check("\n".join(p.text for p in reopened.paragraphs), out.name)
        return out


def _with_ns(xml: str) -> str:
    return re.sub(r"^(<[\w:]+)", r"\1 " + nsdecls("w", "wp", "a", "pic", "r"), xml, count=1)
