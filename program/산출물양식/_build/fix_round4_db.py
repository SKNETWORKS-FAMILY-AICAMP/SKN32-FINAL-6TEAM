# -*- coding: utf-8 -*-
"""DB 4차 — 문장 가운데 남은 마이그레이션 줄번호, 코드 파일명, 엔터티 표 E-06~E-32 구분 칸 정리."""
import re
import shutil
import sys

from docx import Document

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
NAME = "[데이터 수집 및 저장] 데이터베이스저장소 설계 문서_2차_여행.docx"
P = "../w2/제출본/" + NAME

REP = [
    ("001:5·13. 12종 ENUM.", "12종 ENUM(001 마이그레이션)."),
    ("001:8·17. proposed/", "001 마이그레이션의 ENUM: proposed/"),
    ("001:19,003:14. UNIQUE(tenant_id,topic,dedupe_key).", "UNIQUE(tenant_id,topic,dedupe_key)(003 마이그레이션)."),
    ("001:19,003:14. 같은 테넌트·topic 내 유일.", "같은 테넌트·topic 내 유일(003 마이그레이션)."),
    ("현재 watch.py는 watched와 raw를 저장.", "현재 감시 코드는 watched와 raw를 저장."),
    ("현재 watch.py가 넣는 값은", "현재 감시 코드가 넣는 값은"),
    ("이번 확인한 watch.py에는 이 열을 채우는 경로가 없음.", "이번에 확인한 감시 코드에는 이 열을 채우는 경로가 없음."),
    ("transition.py는 Case 버전을", "Case 전이 코드는 Case 버전을"),
]
# 「001:14의 X」 → 「001 마이그레이션의 X」, 「010:41 X」 → 「X」
REF_OF = re.compile(r"\b(\d{3}):\d+(?:[~·]\d+)*의\s")
REF_LEAD = re.compile(r"\b\d{3}:\d+(?:[~·]\d+)*(?:,\s*\d{3}:\d+(?:[~·]\d+)*)*\.?\s+")
ENT_DDL = re.compile(r"^(E-\d{2}) / 범위=DDL; 구현 파일 존재 / 적용 여부 미확인\..*$", re.S)
ENT_DESIGN = re.compile(r"^(E-\d{2}) / \(제안\) 물리 설계 / 설계만 존재\..*$", re.S)


def main():
    d = Document(P)
    paras = list(d.paragraphs) + [pp for t in d.tables for r in t.rows for c in r.cells for pp in c.paragraphs]
    seen, n = set(), 0
    for pp in paras:
        if id(pp._p) in seen or not pp.runs:
            continue
        seen.add(id(pp._p))
        t0 = pp.text
        t = t0
        for a, b in REP:
            t = t.replace(a, b)
        t = ENT_DDL.sub(r"\1 / 구현 파일 존재 (마이그레이션 DDL)", t)
        t = ENT_DESIGN.sub(r"\1 / 설계만 존재 (제안)", t)
        t = REF_OF.sub(r"\1 마이그레이션의 ", t)
        t = REF_LEAD.sub("", t)
        if t != t0:
            pp.runs[0].text = t
            for r in pp.runs[1:]:
                r.text = ""
            n += 1
    d.save(P)
    shutil.copy(P, "../" + NAME)
    allt = [pp.text for pp in paras]
    left_ref = sum(len(re.findall(r"\b\d{3}:\d+", x)) for x in allt)
    left_py = sum(len(re.findall(r"\w+\.py\b", x)) for x in allt)
    left_long = sum(1 for x in allt if "범위=DDL; 구현 파일 존재 / 적용 여부 미확인" in x)
    print("DB 4차: 고친 문단 %d | 남은 줄번호 %d | 남은 .py %d | 남은 긴 상태 문장 %d" % (n, left_ref, left_py, left_long))


if __name__ == "__main__":
    main()
