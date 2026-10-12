# -*- coding: utf-8 -*-
"""DB 문서에 정의 없이 쓰인 내부 작업 코드(DB-01~08)를 문장에서 뺀다."""
import re, shutil, sys
from docx import Document
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
NAME = "[데이터 수집 및 저장] 데이터베이스저장소 설계 문서_2차_여행.docx"
P = "w2/제출본/" + NAME
A = re.compile("DB-0[0-9](?:[·,] ?0[0-9])*(?:에서 |의 |에 )")
B = re.compile(" ?\(?DB-0[0-9](?:[·,] ?(?:DB-)?0[0-9])*\)?(?=[,.;)]|$)")
C = re.compile("DB-0[0-9](?:[·,] ?(?:DB-)?0[0-9])*,? ?")
d = Document(P)
paras = list(d.paragraphs) + [pp for t in d.tables for r in t.rows for c in r.cells for pp in c.paragraphs]
seen, n = set(), 0
for pp in paras:
    if id(pp._p) in seen or not pp.runs:
        continue
    seen.add(id(pp._p))
    t0 = pp.text
    t = C.sub("", B.sub("", A.sub("", t0)))
    t = re.sub(" {2,}", " ", t).replace(" .", ".").replace(" ,", ",").replace(",.", ".").replace("; ;", ";").strip()
    if t != t0:
        pp.runs[0].text = t
        for r in pp.runs[1:]:
            r.text = ""
        n += 1
d.save(P); shutil.copy(P, NAME)
allt = [pp.text for pp in paras]
print("고친 문단", n, "| 남은 DB-0x", sum(len(re.findall("DB-0[0-9]", x)) for x in allt))
for x in allt:
    if "검증 환경 시험으로 확인한다" in x:
        print("  예:", x[:120]); break
