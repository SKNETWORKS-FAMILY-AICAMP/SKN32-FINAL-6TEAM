# -*- coding: utf-8 -*-
"""v2 부품과 09-11 양식을 바탕으로 만든 진행 공유 v3. 기존 출력은 덮어쓰지 않는다."""
from __future__ import annotations
import argparse
import io
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--forms', required=True, type=Path)
parser.add_argument('--assets', required=True, type=Path)
parser.add_argument('--out', required=True, type=Path)
args = parser.parse_args()
OUT, FORMS, ASSETS = args.out, args.forms, args.assets
if OUT.exists():
    raise SystemExit('이미 있는 출력 파일입니다. 덮어쓰지 않습니다: ' + str(OUT))
SRC = FORMS / '[발표] A-COP_진행경과_2026-09-11.pptx'
for required in [SRC, ASSETS / '1.png'] + [ASSETS / 'shots' / f for f in ['m_tripilot-intro-animated-nav-switch.png', 'm_tripilot-mobile-onboarding.png', 'm_tripilot-intake-wireframe.png']]:
    if not required.is_file():
        raise SystemExit('필수 입력 파일이 없습니다: ' + str(required))
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTE = RGBColor(0x47, 0x55, 0x69)
SOFT = RGBColor(0x64, 0x74, 0x8B)
LINE = RGBColor(0xE2, 0xE8, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C = {"blue": (RGBColor(0x25, 0x63, 0xEB), RGBColor(0xDB, 0xEA, 0xFE)),
     "green": (RGBColor(0x05, 0x96, 0x69), RGBColor(0xDC, 0xFC, 0xE7)),
     "amber": (RGBColor(0xB4, 0x53, 0x09), RGBColor(0xFE, 0xF3, 0xC7)),
     "violet": (RGBColor(0x6D, 0x28, 0xD9), RGBColor(0xED, 0xE9, 0xFE)),
     "slate": (RGBColor(0x47, 0x55, 0x69), RGBColor(0xF1, 0xF5, 0xF9)),
     "red": (RGBColor(0xB9, 0x1C, 0x1C), RGBColor(0xFE, 0xE2, 0xE2))}
FONT = "Pretendard"

src = Presentation(str(SRC))
BG_COVER = [s for s in src.slides[0].shapes if s.shape_type == 13][0].image.blob
BG = [s for s in src.slides[4].shapes if s.shape_type == 13][0].image.blob
p = Presentation()
p.slide_width, p.slide_height = src.slide_width, src.slide_height
BLANK = p.slide_layouts[6]


# ── 부품 ─────────────────────────────────────────────────────────
def text(s, x, y, w, h, lines, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate([lines] if isinstance(lines, str) else lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        spec = line if isinstance(line, tuple) else (line, size, bold, color)
        r = para.add_run()
        r.text = spec[0]
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(spec[1]), spec[2], spec[3], FONT
        para.space_after = Pt(3)
    return b


def card(s, x, y, w, h, title, body=None, tone="slate", title_size=14, body_size=12, align=PP_ALIGN.LEFT,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, bar=True):
    fg, bg = C[tone]
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid()
    b.fill.fore_color.rgb = bg
    b.line.color.rgb = LINE
    b.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        b.adjustments[0] = 0.08
    if bar:
        a = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.12), Inches(y + 0.14), Inches(0.06), Inches(0.3))
        a.fill.solid()
        a.fill.fore_color.rgb = fg
        a.line.fill.background()
    pad = 0.28 if bar and align == PP_ALIGN.LEFT else 0.12
    lines = [(title, title_size, True, fg if bar else INK)]
    for line in body or []:
        lines.append((line, body_size, False, MUTE))
    text(s, x + pad, y + 0.1, w - pad - 0.1, h - 0.15, lines, align=align,
         anchor=MSO_ANCHOR.TOP if body else MSO_ANCHOR.MIDDLE)
    return b


def pill_width(label):
    # 10pt 한글 0.16인치를 실제 딱지 글자 크기 12pt에 비례해 늘린다.
    return .32 + sum(.16 if ord(c) > 127 else .085 for c in label) * 1.2


def pill(s, x, y, label, tone="blue", w=None):
    fg, bg = C[tone]
    w = w or pill_width(label)
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.38))
    b.adjustments[0] = 0.5
    b.fill.solid()
    b.fill.fore_color.rgb = bg
    b.line.fill.background()
    b.shadow.inherit = False
    tf = b.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(.08)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(12), True, fg, FONT
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    return b


def arrow(s, x1, y1, x2, y2, color=RGBColor(0x94, 0xA3, 0xB8)):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(2)
    ln.line._get_or_add_ln().append(_arrow_head())
    return ln


def _arrow_head():
    from pptx.oxml.ns import qn
    from lxml import etree
    tail = etree.SubElement(etree.Element(qn("a:ln")), qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return tail


def page(title, badge, tone="blue", num=None):
    s = p.slides.add_slide(BLANK)
    s.shapes.add_picture(io.BytesIO(BG), 0, 0, p.slide_width, p.slide_height)
    text(s, 0.46, 0.3, 10.5, 0.5, title, 24, True)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.46), Inches(0.88), Inches(12.42), Inches(0.012))
    ln.fill.solid()
    ln.fill.fore_color.rgb = LINE
    ln.line.fill.background()
    pill(s, 12.88 - pill_width(badge), 0.4, badge, tone)
    if num:
        text(s, 12.2, 7.05, 0.7, 0.3, f"{num}", 10, False, SOFT, PP_ALIGN.RIGHT)
    return s


def kpi(s, x, y, w, value, label, tone):
    fg, bg = C[tone]
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.05))
    b.adjustments[0] = 0.1
    b.fill.solid()
    b.fill.fore_color.rgb = WHITE
    b.line.color.rgb = LINE
    b.shadow.inherit = False
    text(s, x + 0.15, y + 0.08, w - 0.3, 0.55, value, 24, True, fg)
    text(s, x + 0.15, y + 0.6, w - 0.3, 0.4, label, 12, False, MUTE)



def notes(s, value):
    assert 150 <= len(value) <= 300, (len(value), value)
    s.notes_slide.notes_text_frame.text = value


def table(s, rows, widths, y=1.7, height=4.5, headings=('항목', '상태', '근거와 남은 일')):
    tbl = s.shapes.add_table(len(rows)+1, len(widths), Inches(.5), Inches(y), Inches(sum(widths)), Inches(height)).table
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate([headings] + [r[:3] for r in rows]):
        for j, value in enumerate(row):
            c = tbl.cell(i, j)
            c.text = value
            c.margin_left = c.margin_right = Inches(.12)
            c.margin_top = c.margin_bottom = Inches(.08)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            tone = rows[i-1][3] if i else 'blue'
            c.fill.fore_color.rgb = INK if i == 0 else C[tone][1] if j == 1 else WHITE
            for para in c.text_frame.paragraphs:
                for r in para.runs:
                    r.font.name, r.font.size = FONT, Pt(12)
                    r.font.bold = i == 0 or j == 1
                    r.font.color.rgb = WHITE if i == 0 else C[tone][0] if j == 1 else INK
    return tbl


# 1. 표지
s = p.slides.add_slide(BLANK)
s.shapes.add_picture(io.BytesIO(BG_COVER), 0, 0, p.slide_width, p.slide_height)
text(s, .46, 2.05, 12, .85, 'triPilot 진행 공유', 40, True)
text(s, .46, 3.08, 12, .55, '팀 A-COPilot · 2026. 9. 28.', 18, color=MUTE)
text(s, .46, 3.72, 12, .45, '여행 운영 서비스의 구현 현황', 18, color=MUTE)
for i, label in enumerate(['자동 검사와 빌드', '유스케이스', '변경 흐름', '공통 플로우', '화면 목업']):
    pill(s, .46 + i*2.5, 4.62, label, w=2.3)

# 2. 목차
s = page('검사부터 화면까지 다섯 주제를 짚는다', '공유 순서', num=2)
agenda = [('01', '자동 검사와 빌드', '코드를 합칠 때 두 범위로 검사한다'),
          ('02', '유스케이스', '점검한 10개 항목 중 6개를 완료했다'),
          ('03', '변경 흐름', '감지와 요청은 같은 기준으로 변경을 정한다'),
          ('04', '공통 플로우', '공통 플로우는 변경 판단과 장소 탐색을 잇는다'),
          ('05', '화면 목업', '목업을 바탕으로 등록 화면을 서버에 연결했다')]
for i, (n, title, body) in enumerate(agenda):
    y = 1.35 + i * .91
    text(s, .6, y, .8, .5, n, 24, True, C['blue'][0])
    text(s, 1.6, y, 3.6, .5, title, 18, True)
    text(s, 5.2, y+.04, 7.6, .5, body, 16, color=MUTE)
notes(s, '오늘은 triPilot의 구현 현황을 다섯 주제로 말씀드리겠습니다. 먼저 코드를 합칠 때 어떤 자동 검사를 하는지 보고, 유스케이스의 범위와 준비 상태를 살펴보겠습니다. 이어서 시스템이 변화를 감지할 때와 여행자가 요청할 때의 변경 흐름을 비교하겠습니다. 액티비티와 요식업의 공통 플로우, 여행 등록 목업도 보겠습니다. 마지막에는 다음에 할 일과 팀이 결정할 내용을 정리하겠습니다.')

# 3. CI/CD
s = page('코드를 합칠 때 두 범위로 검사한다', '자동 검사', num=3)
text(s, .55, 1.12, 12.2, .45, '자동 검사(CI)는 변경마다 돌고, 자동 배포(CD)는 빌드까지다.', 16)
text(s, .55, 1.72, 2, .35, 'develop', 16, True)
card(s, .5, 2.2, 1.9, .72, '코드 올리기', ['push / PR'], 'blue')
card(s, 3.05, 1.72, 3.0, .65, 'ruff 린트', None, 'blue')
card(s, 3.05, 2.57, 3.0, .75, 'PostgreSQL 준비', ['스키마·프롬프트 등록'], 'blue')
card(s, 6.45, 2.57, 3.35, .75, '구조·계약·단위 시험', None, 'blue')
arrow(s, 2.42, 2.52, 3.02, 2.04)
arrow(s, 2.42, 2.57, 3.02, 2.94)
arrow(s, 6.08, 2.94, 6.42, 2.94)
text(s, 10.1, 2.68, 2.7, .45, '1,015건 모두 통과', 14, True, C['green'][0])
text(s, .55, 3.63, 2, .35, 'main', 16, True)
card(s, .5, 4.07, 1.9, .72, '코드 올리기', ['push / PR'], 'blue')
card(s, 3.05, 3.48, 3.0, .60, 'ruff 린트', None, 'blue')
card(s, 3.05, 4.23, 3.0, .75, 'PostgreSQL 준비', ['데모 데이터도 등록'], 'blue')
card(s, 6.45, 4.23, 1.85, .75, '구조 시험', None, 'blue')
card(s, 8.7, 4.23, 1.65, .75, '전체 시험', None, 'blue')
card(s, 3.05, 5.13, 3.0, .6, 'Docker 빌드', None, 'blue')
arrow(s, 2.42, 4.40, 3.02, 3.78)
arrow(s, 2.42, 4.44, 3.02, 4.60)
arrow(s, 2.42, 4.48, 3.02, 5.43)
arrow(s, 6.08, 4.60, 6.42, 4.60)
arrow(s, 8.33, 4.60, 8.67, 4.60)
text(s, 10.55, 4.30, 2.3, .62, '1,400건 모두 통과', 14, True, C['green'][0])
text(s, 6.45, 5.13, 6.3, .58, '전체 시험은 실제 모델·검색 통합 시험을 뺀다.\n린트·시험·빌드는 병렬이며 이미지는 올리지 않는다.', 12, color=MUTE)
text(s, .55, 5.85, 12.1, .35, '2026-09-24 깨끗한 환경 재현 기록, 각 범위 실패 0건. 제외한 시험은 로컬에서 돌린다.', 12, color=MUTE)
card(s, .5, 6.22, 6.05, .75, '도입하며 린트 오류 91건을 정리했다', ['안 쓰는 import와 중복 줄 등을 정리했다.'], 'blue')
card(s, 6.8, 6.22, 6.05, .75, '배포용 이미지는 빌드만 한다', ['실제 키와 서버 주소를 넣지 않는다.'], 'blue')
notes(s, '코드를 합칠 때 두 범위로 검사하며 배포는 빌드까지만 합니다. 09-24 재현에서 develop 1,015건과 main 1,400건이 각각 모두 통과했습니다. 두 yml에서 린트·시험·빌드는 병렬입니다. PostgreSQL 준비 뒤 시험하며 main은 구조 시험을 먼저 돌립니다. pgvector와 모의 모델을 쓰고 Dockerfile이 없으면 참고 구현체를 빌드합니다. 프롬프트 이력 시험은 이후 복구했습니다. 다음은 기능 범위입니다.\n근거: 커밋 53184c94, ce1f288c, e9095c6f.')

# 4. 유스케이스 전체 범위
s = page('여행자가 판단하고 개발팀이 운영을 맡는다', '유스케이스', num=4)
text(s, .55, 1.12, 12, .4, '2026-09-24 문서 기준: 여행자·개인 에이전트 27개, 개발·운영팀 18개', 14)
table(s, [
    ('여행자·개인 에이전트', 'A 여행 전 / D 요청', '설문과 초안을 받아 등록하고, 대안을 고르거나 되돌린다.', 'blue'),
    ('시계·외부 데이터', 'B 매일 안내 / C 감지', '하루 시작과 다음 일정, 이동을 안내하고 변화를 감지한다.', 'blue'),
    ('업체', 'E 예약', '모의 업체를 쓴다. 실제 예약은 최소 기능 제품 이후다.', 'slate'),
    ('개발·운영팀', 'F 결함 / G 실행 관리\nH 검증 / I 데이터·배포', '버그와 밀린 일, 알림 발송을 관리한다.\n시험과 시나리오를 재고 데이터와 배포를 관리한다.', 'blue')
], [2.8, 3.3, 6.25], 1.72, 2.95, ('행위자', '기능 묶음', '맡은 일'))
flow = [('03:00', '새벽 식당 확인'), ('08:00', '① 하루 시작'), ('일정 시작', '② 다음 일정'), ('이동 시작', '③ 이동 안내'), ('수시', '3분 감시')]
for i, (tm, desc) in enumerate(flow):
    x = .5+i*2.5
    card(s, x, 5.03, 2.25, .85, tm, [desc], 'blue', align=PP_ALIGN.CENTER, bar=False)
    if 0 < i < 4:
        arrow(s, x-.23, 5.45, x-.03, 5.45)
text(s, 10.53, 6.02, 2.25, .35, '감시는 별도로 돈다.', 12, color=MUTE)
text(s, .55, 6.5, 12, .45, '③ 출발 시각은 다음 일정 시작에서 이동 시간과 여유 10분을 뺀 때다.', 14)
notes(s, '여행자가 판단하고 개발·운영팀은 운영을 맡습니다. 09-24 문서에서 여행자·개인 에이전트는 27개, 개발·운영팀은 18개입니다. 서울 자유여행 외국인 최대 4인을 대상으로 개인 에이전트를 REST·A2A로 연결합니다. 외부 기상·교통·재난 정보 등을 받고 디스코드로 알립니다. D-020의 세 안내는 처리 건(Case)이나 언어 모델 없이 보냅니다. 치명 결함은 대체 소스로 대응하고 실패하면 서버를 끄는 사양이나 중단 구현은 미완입니다. 다음은 09-28 코드 대조 결과입니다.')

# 5. 점검표 완료 6개
s = page('점검한 10개 항목 중 6개를 완료했다', '준비 상태', num=5)
text(s, .55, 1.12, 12, .45, '2026-09-28 코드 대조: 완료 6개 / 부분 2개 / 없음 1개 / 팀 몫 1개', 14)
table(s, [
    ('여행 테마·구성·이동수단 등 설문', '완료', '내국인 여부와 우선순위 1·2도 받는다. 판정에는 아직 쓰지 않는다.', 'green'),
    ('13 실내·실외 / 14 세부 테마', '완료', '세부 선택지는 담당 팀이 정한다.', 'green'),
    ('15 일정이 꼬이면', '완료', '바로 적용하거나 먼저 묻는다. 보류 제안과 시험 18건이 있다.', 'green'),
    ('16 여유에 따른 밀도 목표', '완료', '0.40 / 0.55 / 0.70에 맞춰 하루 장소 수를 고른다.', 'green'),
    ('하루 기본 시간 08:00~22:00', '완료', '「절반 숙소 나감 10시」의 뜻은 정하지 않았다.', 'green'),
    ('식당 영업 확인·사용자 요청', '완료', '새벽 3시에 확인한다. 구글 무료 한도 내 호출 예산을 건다.', 'green')
], [4.2, 1.05, 7.1], 1.85, 4.85)
notes(s, '점검한 열 항목 중 완료는 여섯 개입니다. 기준은 2026-09-28 코드 대조이며, 전체 유스케이스 수와는 다른 분모입니다. 설문 계약에 칸이 있어도 모든 답을 판정에 쓰는 것은 아닙니다. D-020에서 15번은 바로 적용할지 먼저 물을지를, 16번은 밀도 목표를 정합니다. 시험 18건은 tests/e2e/test_ask_first.py의 기록입니다. 식당 확인에는 호출 예산이 걸려 있습니다. 이 완료 상태와 별개로 세부 선택지와 숙소 출발 문구는 결정이 남았습니다. 다음 장에서 나머지 네 항목을 보겠습니다.')

# 6. 점검표 남은 4개
s = page('남은 4개 항목은 구현과 팀 협의가 필요하다', '남은 범위', num=6)
text(s, .55, 1.12, 12, .45, '점검 항목 총 10개 중 부분 2개, 없음 1개, 팀 몫 1개', 14)
table(s, [
    ('0 언어 / 3 기피 음식 / 12 종교', '부분', '언어는 여행 설정으로 받는다. 기피 음식과 종교는 설문 계약에 없다.', 'amber'),
    ('밀도 초과 때 대안 1·2·3 선택', '부분', '선택 흐름은 있다. 밀도 초과로 그 흐름을 여는 연결이 없다.', 'amber'),
    ('방문 직전 식당 판정값 재사용', '없음', '09-28 코드 대조에서 새벽 판정값을 직전에 다시 쓰는 곳을 찾지 못했다.', 'red'),
    ('식당 대기시간 20분 주기 조회', '팀 몫', '요식업 팀이 수집 방법을 정한다. 우리 서버는 그 값을 받을 칸을 만든다.', 'slate')
], [4.2, 1.05, 7.1], 1.9, 3.9)
text(s, .55, 6.18, 12, .75, '부분 2개와 없음 1개는 우리가 만든다.\n대기시간 조회는 요식업 팀의 수집 방법이 나오면 붙인다.', 16)
notes(s, '나머지 네 항목은 부분 두 개, 없음 한 개, 팀 몫 한 개입니다. 언어는 locale로 받지만 기피 음식과 종교는 설문 계약에 칸이 없습니다. 대안 선택 자체는 있으나 밀도 초과를 계기로 연결하지 않았습니다. 방문 직전 판정값 재사용의 없음은 2026-09-28 코드 대조 범위의 결론입니다. 식당 대기시간 수집과 20분 주기 조회는 D-020에 적힌 요식업 팀의 기획입니다. 값을 얻는 방법이 정해지면 받는 쪽을 맞춰야 합니다. 다음은 이미 공유하는 변경 처리 흐름입니다.')

# 7. 변경 흐름
s = page('감지와 요청은 같은 기준으로 변경을 정한다', '변경 흐름', num=7)
card(s, .5, 1.35, 2.35, .95, '시스템이 감지한다', ['3분마다 살핀다.'], 'blue')
card(s, .5, 2.85, 2.35, .95, '여행자가 요청한다', ['신고를 분류한다.'], 'blue')
mid = [('대안을 만든다', '시각·종류를 맞춘다.'), ('적용 여부를 정한다', '바꿀지 물을지 정한다.'), ('변경을 기록한다', '새 일정을 저장한다.'), ('결과를 알린다', '알림이나 답을 보낸다.')]
for i, (t, d) in enumerate(mid):
    x = 3.18 + i*2.47
    card(s, x, 2.12, 2.25, .95, t, [d], 'blue')
    if i:
        arrow(s, x-.20, 2.60, x-.03, 2.60)
arrow(s, 2.87, 1.82, 3.15, 2.36)
arrow(s, 2.87, 3.32, 3.15, 2.80)
card(s, .5, 4.15, 6.0, 1.0, '먼저 묻는 경우', ['고정 일정이나 「먼저 물어봐줘」 여행이면', '대안 1·2·3을 보여 주고 기다린다.'], 'blue')
card(s, 6.8, 4.15, 6.05, 1.0, '바로 적용하는 경우', ['그 외에는 가장 좋은 대안을 적용하고 알린다.'], 'blue')
text(s, .55, 5.5, 12, .4, '다른 점 둘', 16, True)
text(s, .55, 6.0, 12, .42, '감지: 먼저 묻기 여행의 지진·재난·경보에는 안전 알림만 보낸다.', 14)
text(s, .55, 6.55, 12, .42, '직접 선택: 교체·선택·되돌리기는 적용 여부 판단 없이 바로 적용한다.', 14)
notes(s, '감지와 사용자 신고는 같은 적용 기준을 씁니다. D-020의 apply_or_ask와 pending.decide가 그 기준이고, 적용한 새 판은 itinerary_versions에 기록합니다. 먼저 물을 때는 대안 1·2·3을 보여 주며 새 판을 쓰지 않습니다. 사용자가 고정한 일정을 바꿔야 해도 묻습니다. 감지 쪽에서 먼저 묻기를 택한 여행의 지진·재난·경보는 안전 알림만 보냅니다. 직접 고른 교체·선택·되돌리기는 바로 적용합니다. 결과는 계획서 링크를 알리거나 처리 건(Case)을 닫고 답합니다. 다음은 공통 플로우입니다.')

# 8. 공통 플로우
s = page('공통 플로우는 변경 판단과 장소 탐색을 잇는다', '공통 플로우', num=8)
im = Image.open(ASSETS / '1.png').crop((128, 126, 1912, 842))
buf = io.BytesIO()
im.save(buf, 'PNG')
buf.seek(0)
text(s, .55, 1.08, 12, .4, 'CORE(여행 상태를 가진 우리 서버)가 변경을 판단한다.', 14)
s.shapes.add_picture(buf, Inches(1.65), Inches(1.55), width=Inches(10.0))
text(s, 1.65, 5.58, 10.0, .3, '팀 FigJam 보드, 2026-09-28 캡처. 공통 플로우의 앞부분이다.', 10, color=MUTE)
# 아래 도식은 원본의 작은 글을 읽을 수 있도록 다시 설명한다.
steps = [('CORE가 판단한다', '시작·종료와 변화를 본다.'), ('CORE가 조회한다', '외부 정보를 받는다.'), ('에이전트가 찾는다', '대체 장소 1~3순위를 정한다.'), ('CORE에 반영한다', '확인 필요 여부를 보낸다.')]
for i, (t, d) in enumerate(steps):
    x = .46 + i*3.15
    card(s, x, 5.95, 2.9, .95, t, [d], 'blue', 14, 12)
    if i:
        arrow(s, x-.23, 6.4, x-.03, 6.4)
notes(s, '공통 플로우는 CORE의 변경 판단과 에이전트의 장소 탐색을 연결합니다. CORE는 여행 상태를 가진 우리 서버입니다. 팀 FigJam 보드의 2026-09-28 캡처에서 앞부분을 가져왔습니다. CORE는 하루 시작과 종료, 사용자 요청과 외부 정보를 보고 변경 필요를 판단합니다. 에이전트는 대체 장소를 찾고 우선순위와 확인 필요 여부를 돌려줍니다. 캡처에 뒷부분이 없어 전체 연결과 코드 구현의 일치 여부는 확인하지 못했습니다. 다음은 목업을 바탕으로 서버에 연결한 화면입니다.')

# 9. 화면 목업
s = page('목업을 바탕으로 등록 화면을 서버에 연결했다', '화면 연결', num=9)
shots = [('m_tripilot-intro-animated-nav-switch.png', '인트로'), ('m_tripilot-mobile-onboarding.png', '취향 설정'), ('m_tripilot-intake-wireframe.png', '계획 등록')]
for i, (f, caption) in enumerate(shots):
    x = .55 + i*2.35
    s.shapes.add_picture(str(ASSETS/'shots'/f), Inches(x), Inches(1.4), height=Inches(4.95))
    text(s, x, 6.4, 2.27, .4, caption, 14, True, align=PP_ALIGN.CENTER)
    text(s, x, 6.88, 2.27, .3, '09-23 목업(기준선)', 10, color=MUTE, align=PP_ALIGN.CENTER)
text(s, 7.8, 1.08, 5.05, .35, '실제로 연결한 흐름', 16, True)
flow = [('온보딩', '언어·약관과 취향을 받는다.'), ('계획 등록', '글·파일을 접수한다.'), ('확인 화면', '고친 뒤 등록한다.'), ('여행 화면', '일정과 채팅을 연다.')]
for i, (t, d) in enumerate(flow):
    y = 1.62 + i*.91
    card(s, 7.8, y, 5.05, .72, t, [d], 'blue')
    if i:
        arrow(s, 10.33, y-.17, 10.33, y-.03)
card(s, 7.8, 5.24, 5.05, .72, '일정 짜 줘', ['확인 화면에서 일정을 생성해 등록한다.'], 'blue')
text(s, 7.83, 6.02, 5.0, .32, '실측 2026-09-28 · 브라우저 각 1회', 12, True, C['green'][0])
text(s, 7.83, 6.4, 5.0, .76, ['사진 받아쓰기 50.9초, 항목·장소 각 6/6 등록', '「2박3일·넷·실내 위주」: 실제 모델로 3일 등록'], 12)
notes(s, '09-23 목업을 바탕으로 웹 앱을 실제 서버에 연결했습니다. 09-27에는 글·파일 접수와 확인, 등록 뒤 여행 화면·채팅까지 연결했습니다. 09-28에는 확인 화면의 「일정 짜 줘」를 planner.plan_trip과 _create_trip에 연결했습니다. 같은 날 브라우저에서 각 1회 측정해 사진 받아쓰기 50.9초와 항목·장소 각 6/6 등록, 실제 모델의 3일 일정 등록을 기록했습니다. 반복 시험 결과는 아닙니다. 다음은 남은 작업입니다.\n근거: 커밋 257d239a, 9c0b80f7.')

# 10. 다음 단계
s = page('남은 연결을 구현하고 미정 규칙을 정한다', '다음 단계', num=10)
text(s, .55, 1.15, 6, .4, '할 일', 18, True, C['blue'][0])
items = [
    ('자동 검사와 빌드', ['서비스 빌드 파일을 준비한다.', '배포 대상은 최소 기능 제품 이후에 정한다.'], 'blue'),
    ('유스케이스', ['방문 직전 판정값 재사용과', '밀도 초과 때 대안 선택을 연결한다.'], 'blue'),
    ('공통 플로우', ['대체 장소 우선순위를 코드와 대조한다.', '요식업 팀과 대기시간 수집 방법을 맞춘다.'], 'blue'),
    ('화면과 평가', ['실제 연결 화면의 1회 확인을 반복 시험으로 늘린다.', '계획 읽기 평가셋 60건을 준비한다.'], 'blue')]
for i, (t, body, tone) in enumerate(items):
    card(s, .5+(i%2)*6.25, 1.72+(i//2)*1.45, 6.05, 1.15, t, body, tone, 16, 14)
text(s, .55, 4.83, 12, .45, '결정이 필요한 것', 18, True, C['amber'][0])
for i, (title, body) in enumerate([
    ('숙소 출발 시각', '「절반 숙소 나감 10시」의 뜻'),
    ('지연 때 처리', '밀도 70% 구간에서 할 행동'),
    ('세부 테마', '담당 팀이 정할 선택지')]):
    card(s, .5+i*4.2, 5.55, 3.95, 1.1, title, [body], 'amber', 16, 14)
notes(s, '다음에는 남은 연결을 구현하고 미정 규칙을 정하겠습니다. 서비스의 Dockerfile이 생기면 main이 해당 이미지를 빌드합니다. 밀도 초과의 대안 선택과 방문 직전 판정값 재사용을 연결하겠습니다. 공통 플로우의 우선순위와 대기시간 수집은 담당 팀과 맞추겠습니다. 실제 연결 화면은 1회 확인에서 반복 시험으로 늘리고 계획 읽기 평가셋 60건도 준비하겠습니다. D-020의 숙소 출발 문구와 지연 때 밀도 70%에서 할 행동, 담당 팀의 세부 테마 선택지는 결정이 필요합니다. 결정에 맞춰 진행하겠습니다.')

# 저장 직전에도 기존 출력이 없는지 확인한다. xb로 경쟁 상황의 덮어쓰기도 막는다.
assert len(p.slides) == 10
OUT.parent.mkdir(parents=True, exist_ok=True)
with OUT.open('xb') as stream:
    p.save(stream)
print(f'저장: {OUT.name} ({len(p.slides)}장)')
