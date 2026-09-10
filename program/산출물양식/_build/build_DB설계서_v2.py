# -*- coding: utf-8 -*-
"""데이터베이스·저장소 설계 문서 — A-COP 여행 지속관리 CS 플랫폼.

근거: program/plan/A-COP_구현계획서_v11.md §4-E·§6-A~§6-C·§7-A~§7-C,
      wiki/architecture/notifications.md,
      program/plan/A-COP_비용과_수익구조.md,
      final_project_cs/app/infrastructure/db/migrations/001~013.

2026-09-10 갱신: SQL로 확인한 저장소와 미구현 설계를 구분한다.
수신 확인을 통지 계약에서 제외하고, 조건 보존·대상별 멱등성·Trip 동시성을 반영한다.

★양식 처리 함수·경로·블록 복제·서식은 유지한다.
  엔터티 14행, 관계 11행, 정의서 9블록(6·6·7·7·5·8·10·7·5행),
  제약 13행, 변경 이력 5행을 유지한다.
★[실측]은 저장소의 SQL·코드 확인을 뜻한다. 운영 DB 적용·성능 검증을 뜻하지 않는다.
★outbox 제약은 UNIQUE(tenant_id, topic, dedupe_key)를 그대로 승계한다.
"""
import copy
import os
import sys

from docx import Document
from docx.shared import Pt

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
FORMS = os.path.dirname(HERE)
SRC = os.path.join(FORMS, "[데이터 수집 및 저장] 데이터베이스_저장소 설계 문서.docx")
OUT = os.path.join(FORMS, "[데이터 수집 및 저장] 데이터베이스저장소 설계 문서_2차_여행.docx")

d = Document(SRC)
T = d.tables


def setc(cell, text, bold=False):
    p = cell.paragraphs[0]
    # ★`[2026-09-10 정정]` 첫 문단만 바꾸면 양식 예시가 **둘째 문단 이후에 남는다.**
    #   실제로 수집 데이터 보고서에 「모델 응답 테스트 · 리파인먼트 · 내부 테스트」가
    #   새 본문 밑에 붙어 나갔다. 첫 문단은 서식을 살리려고 두고 나머지를 지운다.
    for extra in list(cell.paragraphs)[1:]:
        extra._element.getparent().remove(extra._element)
    for r in list(p.runs)[1:]:
        r._element.getparent().remove(r._element)
    if not p.runs:
        p.add_run("")
    r = p.runs[0]
    r.text = text
    r.bold = bold
    r.font.size = Pt(9)
    r.font.name = "맑은 고딕"


def grow(tbl, need):
    while len(tbl.rows) - 1 < need:
        tbl._tbl.append(copy.deepcopy(tbl.rows[-1]._tr))


def shrink(tbl, keep):
    while len(tbl.rows) - 1 > keep:
        tbl._tbl.remove(tbl.rows[-1]._tr)


def fill(tbl, rows, start=1):
    grow(tbl, start - 1 + len(rows))
    shrink(tbl, start - 1 + len(rows))
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if j < len(tbl.columns):
                setc(tbl.rows[start + i].cells[j], v)


def settext(para, text, bold=False):
    for r in list(para.runs)[1:]:
        r._element.getparent().remove(r._element)
    if not para.runs:
        para.add_run("")
    para.runs[0].text = text
    para.runs[0].bold = bold


# ── 1. 모델링 방법 ────────────────────────────────────────────────────
P = d.paragraphs
settext(P[7], "모델링 방법 : 개념 → 논리 → 물리. [설계·미구현] Trip을 집합체 루트로 두고 "
              "일정 버전을 append-only로 보존한다. MVP는 서울·한국어 하나이며, 접점은 "
              "개인 에이전트 API와 여행별 토큰의 여행계획서 링크(로그인 없음)다. "
              "링크가 최신 상태의 정본이며 읽기와 재요청을 제공한다. 번역은 개인 에이전트가 "
              "담당한다. [미확보] 고객에게 직접 가는 링크·알림에는 번역 주체가 없다. "
              "화면은 운영자 넷(일정·사건·승인·추적)과 고객 링크 하나다.")
settext(P[8], "정규화 수준 : 3NF 기준. Trip·항목 조건은 scope로 구분한 한 표에 보존한다. "
              "§7-B의 '19표 중 5표만 교체'는 전환 전 자산의 구분이다. 교체 대상은 "
              "orders·order_items·products·returns·shipments이며, 코어 14표 "
              "customer_cases·case_events·agent_runs·team_tasks·action_requests·"
              "action_approvals·outbox·llm_calls·prompts·customers·tenants·"
              "knowledge_documents·knowledge_chunks·feedback_analytics_reports는 승계한다. "
              "지식 표 구조는 유지하되 커머스 RAG 내용은 여행 규정으로 교체한다. "
              "[실측] 010은 기존 커머스 표를 삭제하지 않고 여행 표를 추가했다. "
              "19는 현재 전체 표 수나 여행 표 수가 아니다.")
settext(P[9], "도구·근거 : PostgreSQL 16 기준 설계·번호별 SQL 마이그레이션·pydantic 계약. "
              "[실측] 010은 places·bookings·supplier_bookings, 011은 place_catalog·"
              "source_sync_state, 012는 watch_observations·watch_changes를 정의한다. "
              "013은 places에 공급자 신원 열과 인덱스를 추가한다. "
              "[미구현] trips·itinerary_versions·itinerary_items·watch_schedule 및 조건 표의 "
              "마이그레이션은 없다. 아래 실제 표는 주요 열을 묶어 요약하며 전체 DDL은 SQL이 정본이다. "
              "설계 표의 열·키 표시는 제안이며 구현된 제약이 아니다.")

# ── 2.1 엔터티 목록 ──────────────────────────────────────────────────
# 구분 | 테이블 | 제약 대상 | 제약 유형 | 주요 속성
fill(T[3], [
    ("여행 [설계·미구현]", "trips", "trip_id", "PK 제안",
     "최신 일정 포인터 · 사건 입력 개정번호 · 재검토 필요 상태 · 종료 시각"),
    ("일정 [설계·미구현]", "itinerary_versions", "(trip_id, version_no)", "복합 PK 제안",
     "append-only · 변경 이유 · 되돌림 출처 · 제시한 다른 안 보존"),
    ("항목 [설계·미구현]", "itinerary_items", "(trip_id, version_no, item_id)", "복합 PK 제안",
     "시작·끝 · 장소 · 비용 · 잠김 · 선택적 booking_id; 무예약 활동도 대상"),
    ("조건 [설계·미구현]", "trip_constraints (가칭)", "constraint_id", "PK 제안",
     "scope(trip/item) · ref_id · kind · value; 두 수준의 조건을 모두 만족"),
    ("장소 [실측 010·013]", "places", "place_id / (tenant_id, name, kind)", "PK / UNIQUE",
     "예약이 참조하는 장소; 013 source_name·source_content_id·source_content_type_id·source_resolved_at"),
    ("예약 [실측 010]", "bookings", "booking_id / (tenant_id, booking_no)", "PK / UNIQUE",
     "customer_id · place_id · starts_at · party_size · capacity · amount_cents · locked"),
    ("공급자 원장 [실측 010]", "supplier_bookings",
     "supplier_booking_id / (tenant_id, supplier, supplier_ref)", "PK / UNIQUE",
     "booking_id FK · status · confirmed_at; 우리 예약과 별도 저장하여 대조"),
    ("카탈로그 [실측 011]", "place_catalog / source_sync_state",
     "(tenant_id, source, content_id) / (tenant_id, source, scope)", "각 UNIQUE",
     "공급자 장소 목록과 동기화 진행 상태; 수정 시각과 수집 시각 분리"),
    ("관측 [실측 012]", "watch_observations",
     "(tenant_id, target_kind, target_id, observed_at)", "UNIQUE",
     "source · fingerprint · payload; 판정 관련 값만 해시하여 직전 관측과 비교"),
    ("변화 [실측 012]", "watch_changes", "change_id", "PK",
     "before_json · after_json · summary · affected_bookings · notified_at · notify_error"),
    ("감시 일정 [설계·미구현]", "watch_schedule", "schedule_id", "PK 제안",
     "다음 확인·안내 시점; 최소 84초·기본 2분, 계획서에 걸린 대상만 감시"),
    ("케이스·이벤트 [실측 001]", "customer_cases / case_events",
     "case_id / (case_id, aggregate_version)", "PK / UNIQUE",
     "실제 Case 표명은 customer_cases. Trip 연결·trigger_source는 설계; 한 스윕·한 Trip당 Case 하나"),
    ("실행·발행 [실측 001·003]", "action_requests / outbox",
     "(tenant_id, idempotency_key) / (tenant_id, topic, dedupe_key)", "각 UNIQUE",
     "대상별 실행 키는 §4-E 설계; [실측] 현재 business_subject는 case_id 고정"),
    ("부가 저장소 [설계·미구현]", "evidences / suppliers / delegation_scopes / booking_actions / change_requests / notifications",
     "물리 키·관계", "DDL 없음",
     "기존 설계명이며 실제 표로 세지 않는다. 시연 자동 실행 게이트도 없다. 통지 기록은 outbox·watch_changes와 구분"),
])

# ── 2.2 엔터티 관계 ──────────────────────────────────────────────────
fill(T[4], [
    ("여행-일정버전 [설계]", "trips", "itinerary_versions",
     "1 : N. 기존 버전 보존; 되돌림도 최신 사실을 재검증한 새 버전"),
    ("일정버전-항목 [설계]", "itinerary_versions", "itinerary_items",
     "1 : N. 버전별 항목 저장; 대체 항목은 원래 항목과 대응 관계 유지"),
    ("여행·항목-조건 [설계]", "trips / itinerary_items", "trip_constraints",
     "각 1 : N. scope와 ref_id로 대상을 특정; Trip 조건이 항목 조건을 덮지 않음"),
    ("여행-케이스 [설계]", "trips", "customer_cases",
     "1 : N. 한 스윕·한 Trip은 묶음 Case 하나; 원인별 사건·근거는 각각 보존"),
    ("케이스-이벤트 [실측 001]", "customer_cases", "case_events",
     "1 : N. case_events.case_id FK; aggregate_version으로 이벤트 순서 보호"),
    ("고객·장소-예약 [실측 010]", "customers / places", "bookings",
     "각 1 : N. customer_id 필수 FK, place_id 선택 FK; Trip·항목 연결은 아직 설계"),
    ("예약-공급자 원장 [실측 010]", "bookings", "supplier_bookings",
     "1 : N. booking_id FK; 우리 예약 상태와 공급자 상태를 독립 보존"),
    ("장소-카탈로그 [실측 011·013]", "places", "place_catalog",
     "공급자 신원으로 논리 연결. DB FK는 없음; 이름 검색 결과를 확정 신원으로 간주하지 않음"),
    ("감시 대상-관측·변화 [실측 012]", "계획서에 걸린 대상", "watch_observations / watch_changes",
     "대상별 관측·변화 이력. target_kind·target_id 논리 식별; Trip·Case FK는 없음"),
    ("여행-감시 일정 [설계]", "trips", "watch_schedule",
     "1 : N. 감시와 여행·하루 시작 및 항목 출발 안내 시점을 구분; 종료 후 신규 작업 중지"),
    ("케이스-실행·발행 [실측·설계 구분]", "customer_cases", "action_requests / outbox",
     "[실측] action_requests.case_id FK. outbox는 payload로 연결하며 Case FK 없음. "
     "[설계] 일정 버전·포인터·Case·outbox를 함께 커밋; 안내는 Case 없이 발행"),
])

# ── 3.1 테이블 정의서 ────────────────────────────────────────────────
# 양식은 "테이블 : users" 문단 + 표[5] 한 벌이다. 나머지는 복제해 이어 붙인다.
TABLES = [
    ("trips  [설계·미구현: DDL 없음]", [
        ("trip_id", "uuid", "O", "", "O", "여행 한 건. PK 제안; Case 여럿을 가로지르는 집합체"),
        ("tenant_id", "text", "", "", "O", "테넌트 격리 기준. 토큰으로 특정한 여행과 모든 조회·변경의 소유를 대조"),
        ("status / affected_item_ids", "text / jsonb", "", "", "O",
         "설계안: 여행 상태와 영향 항목. 조정 실패 시 일정 내용은 유지하고 재검토 필요 상태를 표시"),
        ("incident_revision", "bigint", "", "", "O",
         "설계안: 사건 입력 개정번호. 새 사건·예약 잠금 변경도 반영; 일정 버전과 함께 동시성 검사"),
        ("current_version_no", "int", "", "", "O",
         "최신 일정 포인터. (trip_id, version_no) 참조 제안; 반영 직전 일정 버전·사건 개정번호를 SQL로 검사"),
        ("closed_at", "timestamptz", "", "", "",
         "종료 시각. 종료하면 새 점검·안내 작업을 만들지 않는다. 다음 확인 시점은 watch_schedule 설계로 분리"),
    ]),
    ("itinerary_versions  [설계·미구현: DDL 없음]", [
        ("trip_id", "uuid", "O", "O", "O", "trips.trip_id 참조 제안. version_no와 복합 PK"),
        ("version_no", "int", "O", "", "O", "Trip 안에서 증가. 기존 버전은 UPDATE·DELETE하지 않는다"),
        ("created_by", "text", "", "", "O", "customer / system / rollback. 생성 주체의 기록"),
        ("rolled_back_from", "int", "", "", "",
         "되돌림 원본 번호. 옛 일정의 휴업·취소 사실을 재확인하며, 재검증 없이 복원하지 않는다"),
        ("created_at", "timestamptz", "", "", "O", "새 버전 생성 시각. 묶음 조정 성공 시 버전은 하나만 증가"),
        ("reason / decision_json", "text / jsonb", "", "", "O",
         "설계안: 원인별 사건·근거, 선택 이유·추가 비용, 미적용 통과 대안 최대 둘, 업체 미확정 상태·변경 링크"),
    ]),
    ("itinerary_items  [설계·미구현: DDL 없음]", [
        ("item_id", "uuid", "O", "", "O",
         "버전 간 항목을 추적할 식별자. PK는 (trip_id, version_no, item_id); 대체 전후 대응을 보존"),
        ("trip_id / version_no", "uuid / int", "O", "O", "O",
         "itinerary_versions의 복합 키 참조 제안. 동일 항목의 과거 버전도 함께 보존"),
        ("kind / booking_id", "text / uuid", "", "", "",
         "kind 필수: activity / dining / mobility / lodging / flight. booking_id는 선택 FK 제안; 무예약이면 NULL"),
        ("is_locked", "bool", "", "", "O",
         "잠긴 예약의 시간·장소·내용 변경, 삭제 및 사실상 참석 불가능하게 만드는 변경 금지"),
        ("starts_at / ends_at", "timestamptz", "", "", "",
         "항목의 시작·끝. 예약 유무와 무관한 판정 입력; 필수 시각이 미확인이면 자동 적용 제외"),
        ("place_ref", "uuid", "", "O", "",
         "places.place_id 참조 제안. 공급자 신원은 places의 013 열로 확인; 외부 id와 내부 id를 혼동하지 않음"),
        ("cost", "numeric", "", "", "",
         "예산 판정 입력. 이동·새 예약·취소·기존 예약 잔여 부담을 포함; 미확정 환급은 확정 할인으로 차감하지 않음"),
    ]),
    ("trip_constraints  [설계·미구현: 가칭, DDL 없음]", [
        ("constraint_id", "uuid", "O", "", "O", "조건 한 건의 PK 제안. Trip·항목 조건을 같은 표에 저장"),
        ("tenant_id", "text", "", "", "O", "조건의 대상이 같은 테넌트·여행에 속하는지 서버에서 확인"),
        ("trip_id", "uuid", "", "O", "O", "소속 Trip 참조 제안. 항목 조건도 여행 단위로 함께 조회"),
        ("scope", "text", "", "", "O", "trip / item. 조건 수준을 나누는 칸 하나; 조건 표를 둘로 나누지 않음"),
        ("ref_id", "uuid", "", "", "O",
         "scope=trip이면 trip_id, scope=item이면 item_id. 다형 참조이므로 단일 FK로 강제했다고 적지 않음"),
        ("kind", "text", "", "", "O", "dietary / party_size / age_limit 등. 여행 식단과 항목 인원·정원 조건을 모두 보존"),
        ("value", "jsonb", "", "", "O",
         "조건값. 두 수준을 모두 만족해야 통과; 미확인과 조건 없음은 구분. 표·열명 및 버전별 보존 방식은 [미확보]"),
    ]),
    ("bookings  [실측: 010_domain_travel.sql, 열 묶음 요약]", [
        ("booking_id", "uuid", "O", "", "O", "실제 PK. 기본값 gen_random_uuid(); 서버가 확인할 예약 대상 id"),
        ("tenant_id / customer_id / booking_no", "text / uuid / text", "", "", "O",
         "모두 NOT NULL. customer_id는 customers FK; UNIQUE(tenant_id, booking_no)"),
        ("kind / status / starts_at", "text / text / timestamptz", "", "", "O",
         "모두 NOT NULL. 종류·예약 상태·시작 시각; SQL의 text 열에 값 CHECK 제약이 있다고 간주하지 않음"),
        ("place_id / locked", "uuid / boolean", "", "", "",
         "place_id는 선택 FK(places), locked는 NOT NULL DEFAULT false. Trip·item FK는 010에 없음"),
        ("party_size / capacity / amount_cents / created_at", "int / int / int / timestamptz", "", "", "",
         "앞 셋은 NULL 허용, created_at은 NOT NULL DEFAULT now(). "
         "[실측] read.booking은 booking_id가 없으면 시작 시각이 가장 이른 예약(과거·취소도 안 거른다)을 반환하고 case_id를 사용하지 않음"),
    ]),
    ("watch_observations  [실측: 012_watch_observations.sql]", [
        ("observation_id", "uuid", "O", "", "O", "실제 PK. 기본값 gen_random_uuid(); 관측 한 건"),
        ("tenant_id", "text", "", "", "O", "관측 대상의 테넌트"),
        ("target_kind", "text", "", "", "O", "place / weather / flight 등 대상 종류; SQL에는 text로 정의"),
        ("target_id", "text", "", "", "O", "종류 안의 식별자. place_id 또는 좌표·예보 시각 등; 물리 FK 없음"),
        ("source", "text", "", "", "O", "조회 출처. 출처를 구분하여 근거와 관측을 해석"),
        ("fingerprint", "text", "", "", "O",
         "판정에 쓰는 부분만 해시. 전체 응답의 무관한 필드 변경을 사건으로 만들지 않음"),
        ("payload", "jsonb", "", "", "O", "관측 내용. 직전 fingerprint와 비교하여 변화가 있을 때 watch_changes에 별도 저장"),
        ("observed_at", "timestamptz", "", "", "O",
         "DEFAULT now(). UNIQUE(tenant_id, target_kind, target_id, observed_at). "
         "같은 대상의 observed_at DESC 인덱스 존재; 수집 시각을 현장 확인 시각으로 바꾸어 표시하지 않음"),
    ]),
    ("places  [실측: 010·013, 열 묶음 요약]", [
        ("place_id", "uuid", "O", "", "O", "실제 PK. 예약이 참조하는 내부 장소 id"),
        ("tenant_id", "text", "", "", "O", "장소의 테넌트. UNIQUE(tenant_id, name, kind)의 일부"),
        ("name / kind", "text / text", "", "", "O",
         "장소 이름·종류. 이름의 일치만으로 공급자 장소를 특정하지 않음"),
        ("latitude / longitude", "double precision", "", "", "", "위도·경도. NULL은 미확인"),
        ("weather_sensitive", "boolean", "", "", "",
         "날씨 영향 여부. NULL은 모름; 실내라는 확정 사실과 구분하여 기록"),
        ("hours_confirmed_at / open_at_slot", "timestamptz / boolean", "", "", "",
         "영업 정보 확인 시각과 해당 시각 영업 여부. NULL은 미확인"),
        ("dietary / dietary_absent", "text[]", "", "", "O",
         "각 DEFAULT '{}'. 확인된 있음·확인된 없음을 구분; 양쪽 모두 없으면 모름"),
        ("source_name / source_content_id", "text / text", "", "", "",
         "013 추가. 해소한 공급자·공급자 id; (tenant_id, source_name, source_content_id) 인덱스, UNIQUE·FK 아님"),
        ("source_content_type_id", "text", "", "", "",
         "013 추가. 공급자의 장소 종류. 동명 장소를 종류와 함께 구분"),
        ("source_resolved_at", "timestamptz", "", "", "",
         "013 추가. 공급자 신원을 해소한 시각. 011 place_catalog의 fetched_at·source_modified_at과 별개"),
    ]),
    ("watch_changes  [실측: 012_watch_observations.sql, 열 묶음 요약]", [
        ("change_id", "uuid", "O", "", "O", "실제 PK. 변화 한 건; observation_id·trip_id·case_id FK는 없음"),
        ("tenant_id / target_kind / target_id / source", "text", "", "", "O",
         "모두 NOT NULL. 관측과 같은 대상·출처를 논리적으로 식별"),
        ("before_json / after_json", "jsonb", "", "", "",
         "before_json은 NULL 허용, after_json은 NOT NULL. 변경 전후 값 보존"),
        ("summary", "text", "", "", "O", "실제 필드 차이에서 만든 변화 요약. 관측되지 않은 원인을 생성하지 않음"),
        ("affected_bookings", "jsonb", "", "", "O",
         "DEFAULT '[]'::jsonb. 영향받은 고객·예약; 빈 목록도 보존. 무예약 항목·Trip 연결 확장은 [미구현]"),
        ("detected_at", "timestamptz", "", "", "O",
         "DEFAULT now(). [실측] 관측·변화 저장은 있으나 묶음 Case 생성 연결은 [미구현]"),
        ("notified_at / notify_error", "timestamptz / text", "", "", "",
         "통지 시각·오류. (tenant_id, notified_at, detected_at DESC) 인덱스. "
         "고객 수신·읽음 확인이 아니며 outbox의 발행 중복 제약을 대신하지 않음"),
    ]),
    ("watch_schedule  [설계·미구현: DDL 없음]", [
        ("schedule_id", "uuid", "O", "", "O", "감시·안내 일정 한 건의 PK 제안"),
        ("tenant_id / trip_id", "text / uuid", "", "", "O",
         "소속 테넌트·여행. trip_id FK 제안; 종료된 Trip에는 새 작업을 만들지 않음"),
        ("target_kind / target_id", "text / text", "", "", "O",
         "계획서에 걸린 대상만 등록. 현재 감시 대상은 예약에서 나오며 카탈로그 전수 폴링은 하지 않음"),
        ("next_check_at / notice_due_at", "timestamptz", "", "", "",
         "감시와 안내 시각을 별도 관리. 감시 최소 84초·기본 2분; 출발 안내는 시작−이동 시간−여유. 역산은 [미구현]"),
        ("notice_kind / target_version_no / state", "text / int / text", "", "", "",
         "설계안: 여행·하루 시작/항목 출발, 대상 버전, 대기·무효화 상태. "
         "일정 변경 시 미발행 출발 안내를 무효화하고 재계산; 구체 키·인덱스는 [미확보]"),
    ]),
]

anchor_para = None
for p in d.paragraphs:
    if p.text.strip().startswith("테이블 :"):
        anchor_para = p
        break

settext(anchor_para, "테이블 : trips  [설계·미구현: DDL 없음]", bold=True)
fill(T[5], TABLES[0][1])

after = T[5]._tbl
for name, cols in TABLES[1:]:
    new_p = copy.deepcopy(anchor_para._p)
    new_t = copy.deepcopy(T[5]._tbl)
    after.addnext(new_t)
    new_t.addprevious(new_p)
    after = new_t
    # 새로 붙인 문단·표를 python-docx 객체로 다시 감싼다
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    settext(Paragraph(new_p, anchor_para._parent), "테이블 : " + name, bold=True)
    fill(Table(new_t, d), cols)

# ── 3.2 제약 조건 및 무결성 ──────────────────────────────────────────
fill(T[6], [
    ("DB [실측 001]", "UNIQUE", "case_events(case_id, aggregate_version)",
     "실제 PK는 event_id이며 위 둘은 UNIQUE다. Case별 이벤트 순서를 보호한다. "
     "승계하는 낙관적 잠금은 customer_cases 한 행에만 적용되므로 서로 다른 Case의 같은 Trip 쓰기는 막지 못한다."),
    ("DB·애플리케이션", "UNIQUE 승계 / 대상 규칙 [미구현]", "action_requests(tenant_id, idempotency_key)",
     "키 재료는 tenant · 요청 id · 작업 종류 · 대상이다. 대상은 서버가 실재·여행 소유를 확인한 객체 id다. "
     "activity.*·dining.*는 예약이 있으면 booking_id, 없으면 item_id; booking.*는 booking_id; "
     "mobility.*는 item_id; trip.*는 trip_id다. 대상 불명이면 거부하며 Case id나 다른 예약으로 폴백하지 않는다. "
     "Team 값은 참고값이다. [실측] 현재 controller의 business_subject는 case_id 고정이며, "
     "같은 Case·같은 작업 종류·다른 예약의 둘째 제안이 사라지는 재현이 있다(DoD-24)."),
    ("DB [실측 003]", "UNIQUE", "outbox UNIQUE(tenant_id, topic, dedupe_key)",
     "tenant_id를 반드시 포함한다. [설계] trip.notice 변경 통지는 적용된 일정 버전당 하나다. "
     "안내 키는 여행·대상·안내 종류로 구성하며 하루·항목 등 해당 시점을 식별한다. "
     "같은 스윕·발행 재시도의 DB 중복을 막는 제약이며 고객 수신을 보증하는 제약이 아니다."),
    ("DB·애플리케이션 [설계·미구현]", "복합 키 / append-only", "itinerary_versions · itinerary_items",
     "일정 버전 PK는 (trip_id, version_no), 항목 PK는 (trip_id, version_no, item_id)를 제안한다. "
     "과거 일정·이벤트를 지우지 않는다. 재요청은 묶음 단위이며, 옛 버전과 다른 안 모두 최신 사실·필수 조건을 "
     "다시 검증한 뒤 새 버전으로 저장한다. 휴업 식당이 포함된 과거 안을 그대로 복원하지 않는다."),
    ("DB·애플리케이션 [설계·미구현]", "Trip 동시성 계약 다섯", "기준 일정 버전 + 사건 입력 개정번호",
     "① 기준 일정 버전과 사건 입력 개정번호를 함께 읽는다.\n"
     "② 후보 계산 중에는 긴 트랜잭션을 유지하지 않는다.\n"
     "③ 반영 직전 두 값이 그대로인지 확인하고 SQL 조건으로 다시 강제한다.\n"
     "④ 달라졌으면 후보를 버리고 최신 입력으로 한 번 재계산한다. 버전 숫자만 바꿔 같은 후보를 넣지 않는다.\n"
     "⑤ 성공할 때만 일정 버전 append·최신 포인터·Case 전이·outbox 삽입을 함께 커밋한다. 실패하면 전부 롤백한다.\n"
     "현재 Trip 보호 장치는 없다. 일정이 같아도 새 사건·예약 잠금 변경이 생길 수 있어 개정번호도 검사한다."),
    ("접근 제어 [승계·설계 구분]", "테넌트·여행 소유 확인", "tenant_id / trip_id / 여행별 토큰",
     "[실측] read.booking은 tenant_id·customer_id로 조회하지만 booking_id가 없으면 "
     "시작 시각이 가장 이른 예약(과거·취소도 안 거른다)을 반환하고 case_id는 사용하지 않는다. 이 동작으로 사건 대상을 특정했다고 간주하지 않는다. "
     "[설계] API·링크의 같은 재요청 경로에서 여행별 토큰과 대상 소유를 검증한다. "
     "전체 표의 RLS 적용 완료나 여행 토큰 저장소가 있다고 적지 않는다. 토큰 저장·열람 기록의 물리 구조는 [미확보]."),
    ("애플리케이션 [실측·설계 구분]", "상태 전이 / 라우팅 두 축", "customer_cases.status · issue_code · intent",
     "Case 상태 12종을 승계한다. [실측] case_type은 issue_code 접두에서 얻고 intent와 함께 라우팅한다. "
     "객체 종류는 activity·booking·mobility·dining·lodging·flight, 요청 종류는 "
     "itinerary_submit·incident_report·confirm_request·adjust_reject·other다. "
     "접두 없는 other는 사람 인계 경로다. 별도 case_type DB 열을 추가했다고 적지 않는다."),
    ("애플리케이션 [설계·미구현]", "묶음 조정 / 부분 반영 금지", "한 스윕·한 Trip = Case 하나",
     "탈락 조건 셋: ①필수 조건 위반·필수 판정값 미확인 ②잠긴 예약 훼손 "
     "③실행 권한·근거의 실재·소유 확인 부족. 통과 후보만 사전식으로 비교한다: "
     "①바뀐 항목 수 최소 ②추가 비용 최소 ③원래 일정과 시간 차이 최소(최대 이동량→합계) "
     "④되돌리기 쉬움(내부 복원→외부 보상, 동률이면 복원 가능 기간이 긴 쪽). "
     "끝까지 같으면 정규화 후보 id로 결정한다. 항목 대체 대응을 보존하여 변경 수 누락을 막는다. "
     "해 없음·근거 부족·탐색 상한 도달을 구분하고 실패 시 부분 적용하지 않는다. "
     "일정은 유지하되 재검토 필요·영향 항목·실패 근거를 링크에 표시하고 사람에게 넘긴다."),
    ("애플리케이션 [설계·미구현]", "실행 경계 / 위임 / 보상", "suppliers · delegation_scopes · booking_actions · change_requests",
     "Pre-CS는 우리 일정 버전에만 적용한다. 실제 업체 예약·결제는 승인 없이 실행하지 않는다. "
     "업체 건은 바꿀 항목·대안·차액·변경 링크로 인계한다. 계획서에는 기존 확정 예약과 "
     "'업체 변경 필요·미확정'을 함께 표시한다. 예약 확정 자체가 필수 조건이면 미확정 안은 자동 적용에서 제외한다. "
     "[실측] app/·tests/의 tier == 'simulated' 게이트는 0건이며 자동 실행 분기도 없다. "
     "시연 분기를 열기 전에 게이트·DoD-14·15 검증이 필요하다. 위임 상한·철회·실행 이유·금액·되돌림 기한·"
     "보상 실패 인계는 시연 설계이며 관련 표의 DDL은 없다."),
    ("데이터·애플리케이션 [설계·미구현]", "조건 보존 / 전체 재검증", "trip_constraints.scope · 일정 전체",
     "scope=trip과 scope=item 조건을 한 표에서 읽고 모두 만족해야 통과한다. "
     "예약 없는 Activity도 항목의 시각·장소·비용·잠김으로 판정하며 예약 규정은 선택 정보다. "
     "시간 충돌·이동 여유·예산·인원·식단은 최종 조합 전체에서 코드로 재검증한다. "
     "[실측] 현재 starts_at_proposed는 근거 대조에서 제외되어 있어 자동 적용 검증이 완료된 상태가 아니다. "
     "[추정] 탐색 상한은 4단계·단계당 3후보·81조합·외부 조회 24회·120초이며 더 이른 마감이 우선한다."),
    ("발행·고객 상태 [설계]", "알림 세 시점 / 링크 정본", "outbox · 최신 일정 · 재요청 이력",
     "①변경(생성·조정·재요청 반영 포함) ②여행·하루 시작 ③각 항목 출발에 알린다. "
     "① 최초 검증 통과 시 링크를 처음 제공한다. ②는 전날 저녁과 여행·매일 시작 안내를 포함한다. "
     "②·③은 Case·LLM 없이 문구틀로 발행한다. 변경과 안내는 각각 판정하며 둘 다 참이면 변경 통지가 먼저다. "
     "무료 채널 하나인 디스코드 웹훅을 사용한다. '닿았나'는 사양이 아니므로 sent/received/acted를 "
     "고객 확인 계약으로 두지 않는다. 링크는 최신 버전·변경 이유·다른 안·업체 미확정을 보여준다. "
     "열람 시각은 운영 정보이며 통과 조건이 아니다. 열람 기록 구현·장기 미열람 재알림은 [미확보]."),
    ("감시·안내 [설계·미구현]", "주기와 발행 시각 분리", "watch_schedule · read.route · read.transit",
     "감시는 최소 84초·기본 2분이며 계획서에 걸린 대상만 조회한다. 감시와 outbox 배달 루프를 분리한다. "
     "출발 시각은 항목 시작−이동 시간−여유이며 여유는 설정값이다. 사전 경로 조회로 발행 시각을 갱신하고 "
     "③ 출발 안내 때 경로를 다시 조회한다. 사전 조회 간격은 [미확보]. 일정 변경 시 미발행 출발 안내를 무효화·재계산한다. "
     "[실측] read.route·read.transit는 항상 None이므로 이동 시간 계산·저장·역산은 [미구현]. "
     "조회 실패는 모름으로 표시하고 링크를 제공한다. 종료 후 신규 감시·안내를 중지한다."),
    ("데이터·운영", "출처·시각·비용 근거", "API / RAG / 합성",
     "API는 장소·운영시간·이동·기상 조회(Mock 허용), RAG는 취소·환급·위약금·동행 규정으로 지식 표 내용 교체, "
     "합성은 일정·예약·고객 시드다. 공급자 수정 시각·수집 시각·현장 확인 시각을 구분하고 모르는 값은 채우지 않는다. "
     "Places·Routes 단가, 여행 LLM 단가, 인계 비율·시간은 [미확보]. "
     "[추정] 여행 LLM 3~37원은 커머스 단가를 차용한 값이며 전체 기계 원가가 아니다. "
     "안내는 LLM 없이도 경로 조회 비용이 발생한다. 출발 조회 20~40회/7일은 [추정]이며 사전 조회 비용도 별도다. "
     "월 기본료+초과 건당 종량은 수익 모델 제안이고 구매자·가격은 [미확보]."),
])

# ── 4. 변경 이력 ─────────────────────────────────────────────────────
fill(T[7], [
    ("2026-09-10", "전체", "v11·팀장 결정 열 반영. 서울·한국어·API와 여행계획서 링크·"
     "운영자 넷과 고객 하나를 기준으로 갱신. 착수 2026-09-08, 최종발표 2026-10-26, "
     "09-10 기준 남은 기간 6주 4일", "문서 기준선", "v11 §0-4·§1·§9"),
    ("2026-09-10", "전체", "[실측] 010의 places·bookings·supplier_bookings, "
     "011의 place_catalog·source_sync_state, 012의 watch_observations·watch_changes, "
     "013의 places 공급자 신원 열을 실제 SQL로 구분. Trip·일정·감시 일정은 설계로 정정",
     "실제 저장소 / 미구현 설계", "010~013 SQL · v11 §7-B"),
    ("2026-09-10", "전체", "19표 중 도메인 5표 교체·코어 14표 승계, scope 조건 한 표, "
     "서버가 확인한 대상별 멱등 키, outbox UNIQUE(tenant_id, topic, dedupe_key), "
     "Trip 동시성 계약 다섯과 묶음 원자 반영을 명시",
     "조건·정합성·동시성", "v11 §4-E·§6-C-7·§7-B"),
    ("2026-09-10", "전체", "수신 확인 계약 삭제. 변경(생성 포함)·여행/하루 시작·항목 출발의 "
     "세 시점, 경로 재조회, 오래된 안내 무효화를 반영. "
     "[실측] read.route·read.transit의 None 반환과 시연 게이트 0건을 미구현으로 명시",
     "감시·통지·실행 경계", "v11 §4-C·§6-A~§6-C · notifications.md"),
    ("[예정]", "최연우", "미구현 설계의 SQL 마이그레이션·빈 DB 적용 검증, "
     "조건 보존·다른 대상 제안 분리·Trip 충돌 거부와 1회 재계산·원자 커밋·"
     "최신 링크·안내 중복 방지를 검증한다. DoD는 26항목(자동 20·아키텍처 테스트 4·측정 2)이며 "
     "이 문서 갱신을 구현·검증 완료로 간주하지 않는다",
     "미구현 저장소·DoD-1~26", "v11 §12; 세부 일정 [미확보]"),
])

# ── 표지 ────────────────────────────────────────────────────────────
# ★`[2026-09-10]` 양식 자리표시자가 그대로 나가고 있었다 — "00기 : _____팀",
#   작성 팀원 "000", 제출 일자 "2026. 4. 18.". 심사 제출본에 남으면 안 된다.
#   기획서와 같은 값으로 맞춘다.
_COVER = {"00기 : _____팀": "32기 : 6팀", "2026. 4. 18.": "2026. 9. 10."}


def _cover_fix(doc):
    def swap(par):
        for k, v in _COVER.items():
            if k in par.text:
                for r in par.runs:
                    if k in r.text:
                        r.text = r.text.replace(k, v)
    for p in doc.paragraphs:
        swap(p)
    for tb in doc.tables[:2]:
        for row in tb.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    swap(p)
                if c.text.strip() == "000":
                    setc(c, "김지혜, 서유현, 송채영, 정세환, 최상욱, 최연우")
            if row.cells[0].text.strip() == "깃허브 경로" and not row.cells[1].text.strip():
                setc(row.cells[1],
                     "https://github.com/SKNETWORKS-FAMILY-AICAMP/SKN32-FINAL-6TEAM")


_cover_fix(d)

try:
    d.save(OUT)
except PermissionError:
    base, ext = os.path.splitext(OUT)
    OUT = base + "_새버전" + ext
    print("  원본이 열려 있어 새 이름으로 저장한다:", os.path.basename(OUT))
    d.save(OUT)
print("저장:", os.path.basename(OUT), "| 테이블 정의", len(TABLES), "개")
