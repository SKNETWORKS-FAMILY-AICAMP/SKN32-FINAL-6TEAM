# -*- coding: utf-8 -*-
"""양식 docx 의 표지·머리 표는 그대로 두고, 본문을 통째로 새로 짜는 조립 도구 (4차 v2).

왜 새로 만들었나
- 3차·4차 빌드(`build_데이터전처리_3종_여행*.py`)는 양식의 빈 표에 값을 채우는 방식이라, 팀별(코어·이동·액티비티·요식업)로
  절을 나누거나 요약 카드·그래프를 넣기가 어렵다.
- 이 도구는 양식의 **글꼴·표 색·글머리표 번호 체계**를 그대로 쓰되 본문 요소를 직접 만든다.
  ★양식의 절 제목(전처리 개요·원천 데이터 분석 등)과 번호는 바꾸지 않는다.

지키는 것
- 사람이 안 쓰는 기호(줄표·화살표·별표 등)는 `docxfill.check` 로 막는다.
- 표는 폭을 고정하고(합계 9,345 dxa) 제목 행을 쪽마다 되풀이하며 한 행을 쪽 경계에서 쪼개지 않는다.
- 색은 양식의 표 제목색(cfe2f3)과 표지 파랑에 맞춘다. 강조색은 두 가지(파랑·주황)만 쓴다.
"""
from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from docxfill import check, esc  # noqa: E402

FONT = ('<w:rFonts w:ascii="Malgun Gothic" w:cs="Malgun Gothic" '
        'w:eastAsia="Malgun Gothic" w:hAnsi="Malgun Gothic"/>')
BLUE = "2A78D6"
DARK = "1F3A5F"
HEAD_FILL = "CFE2F3"      # 양식 표 제목색
ZEBRA = "F5F8FC"
INK = "1F2933"
MUTED = "5B6B7C"
TABLE_W = 9345            # 양식 표 폭(dxa)


def _rpr(size=20, bold=False, color=None, italic=False):
    parts = [FONT]
    if bold:
        parts.append('<w:b w:val="1"/><w:bCs w:val="1"/>')
    else:
        parts.append('<w:b w:val="0"/><w:bCs w:val="0"/>')
    if italic:
        parts.append('<w:i w:val="1"/>')
    parts.append('<w:color w:val="%s"/>' % (color or "000000"))
    parts.append('<w:sz w:val="%d"/><w:szCs w:val="%d"/><w:rtl w:val="0"/>' % (size, size))
    return "<w:rPr>%s</w:rPr>" % "".join(parts)


def _runs(text, size=20, bold=False, color=None, italic=False):
    """`**굵게**` 표시를 굵은 글로 바꾼다(문장 안 일부 강조용)."""
    check(text)
    out = []
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if not chunk:
            continue
        out.append('<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r>'
                   % (_rpr(size, bold or i % 2 == 1, color, italic), esc(chunk)))
    return "".join(out)


class Composer:
    def __init__(self, template: Path):
        self.doc = Document(str(template))
        self._normalise_page_margins()
        body = self.doc.element.body
        self.sect = body.find(qn("w:sectPr"))
        head_table = self.doc.tables[2]._tbl            # 절 머리 표(「데이터 전처리 프로세스」 등)
        self.head_table = head_table
        removing = False
        for el in list(body.iterchildren()):
            if el is head_table:
                removing = True
                continue
            if removing and el is not self.sect:
                body.remove(el)
        self._img_id = 900
        self.figures = 0

    # -- 양식의 소수 twip 속성을 정수로 ------------------------------
    def _normalise_page_margins(self):
        for section in self.doc.element.body.iter(qn("w:pgMar")):
            for key, value in list(section.attrib.items()):
                if "." in value:
                    section.set(key, str(round(float(value))))
            if section.get(qn("w:gutter")) is None:                       # 양식(구글 문서 내보내기)에 빠진 필수 속성
                section.set(qn("w:gutter"), "0")

    # -- 표지 ---------------------------------------------------------
    def cover(self, cover_title: str, stage: str, date: str, repo: str, authors: str):
        first = self.doc.tables[0].rows[0].cells[0].paragraphs[0]
        for r in list(first.runs)[1:]:
            r._element.getparent().remove(r._element)
        first.runs[0].text = cover_title
        meta = self.doc.tables[1]
        for row, value in zip(meta.rows, (stage, date, repo, authors)):
            cell = row.cells[1]
            p = cell.paragraphs[0]
            for extra in list(cell.paragraphs)[1:]:
                extra._element.getparent().remove(extra._element)
            for r in list(p.runs)[1:]:
                r._element.getparent().remove(r._element)
            if not p.runs:
                p.add_run("")
            check(value, "표지")
            p.runs[0].text = value

    # -- 요소 추가 ----------------------------------------------------
    def _add(self, xml: str):
        el = parse_xml(_with_ns(xml))
        self.sect.addprevious(el)
        return el

    def gap(self, twips=120, keep_next=False):
        self._add('<w:p><w:pPr>%s<w:spacing w:before="0" w:after="%d"/></w:pPr></w:p>' % ('<w:keepNext w:val="1"/>' if keep_next else "", twips))

    def h1(self, text: str):
        """양식의 1단계 제목(번호 자동)."""
        check(text)
        self._add('<w:p><w:pPr><w:keepNext w:val="1"/><w:numPr><w:ilvl w:val="0"/><w:numId w:val="3"/></w:numPr>'
                  '<w:spacing w:before="280" w:after="100" w:line="276" w:lineRule="auto"/>'
                  '<w:ind w:left="720" w:hanging="360"/></w:pPr>'
                  '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r></w:p>' % (_rpr(26, True, DARK), esc(text)))

    def h2(self, text: str):
        """2단계 제목(2.1 등). 양식의 들여쓴 글 그대로, 굵게."""
        check(text)
        self._add('<w:p><w:pPr><w:keepNext w:val="1"/><w:spacing w:before="220" w:after="80" w:line="276" w:lineRule="auto"/>'
                  '<w:ind w:left="240" w:firstLine="0"/></w:pPr>'
                  '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r></w:p>' % (_rpr(22, True, INK), esc(text)))

    def h3(self, text: str):
        """팀(코어·이동·액티비티·요식업) 소제목 — 왼쪽 파란 띠."""
        check(text)
        self._add('<w:p><w:pPr><w:keepNext w:val="1"/><w:pBdr><w:left w:val="single" w:sz="24" w:space="6" w:color="%s"/></w:pBdr>'
                  '<w:spacing w:before="240" w:after="100" w:line="276" w:lineRule="auto"/><w:ind w:left="300" w:firstLine="0"/></w:pPr>'
                  '<w:r>%s<w:t xml:space="preserve">%s</w:t></w:r></w:p>' % (BLUE, _rpr(21, True, DARK), esc(text)))

    def line(self, label: str, value: str):
        """양식의 「라벨 : 값」 글머리 줄."""
        check(label + value)
        self._add('<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'
                  '<w:spacing w:after="60" w:line="276" w:lineRule="auto"/><w:ind w:left="720" w:hanging="360"/>'
                  '<w:rPr><w:color w:val="000000"/></w:rPr></w:pPr>'
                  '<w:r>%s<w:t xml:space="preserve">%s : </w:t></w:r>%s</w:p>'
                  % (_rpr(20, True), esc(label), _runs(value, 20)))

    def bullet(self, text: str):
        check(text)
        self._add('<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr>'
                  '<w:spacing w:after="50" w:line="276" w:lineRule="auto"/><w:ind w:left="720" w:hanging="360"/>'
                  '<w:rPr><w:color w:val="000000"/></w:rPr></w:pPr>%s</w:p>' % _runs(text, 20))

    def para(self, text: str, size=20, color=None, italic=False, after=100, indent=240):
        self._add('<w:p><w:pPr><w:spacing w:after="%d" w:line="288" w:lineRule="auto"/><w:ind w:left="%d"/><w:jc w:val="left"/></w:pPr>%s</w:p>'
                  % (after, indent, _runs(text, size, False, color, italic)))

    def note(self, text: str):
        """표·그림 아래 작은 출처·주의 글."""
        self.para(text, size=17, color=MUTED, after=120, indent=240)

    # -- 표 -----------------------------------------------------------
    def table(self, header, rows, widths, align=None, first_bold=False, size=17, zebra=True, keep=False):
        """폭 합계 9,345 dxa 로 맞춘 표. widths 는 비율(정수)."""
        total = sum(widths)
        w = [round(TABLE_W * x / total) for x in widths]
        w[-1] += TABLE_W - sum(w)
        align = align or ["left"] * len(header)
        grid = "".join('<w:gridCol w:w="%d"/>' % x for x in w)

        def cell(text, width, fill, bold, al, hdr=False):
            check(str(text), "표")
            shd = '<w:shd w:val="clear" w:color="auto" w:fill="%s"/>' % fill if fill else ""
            paras = []
            for chunk in str(text).split("\n"):
                paras.append('<w:p><w:pPr>%s<w:keepLines w:val="1"/><w:spacing w:before="0" w:after="0" w:line="264" w:lineRule="auto"/>'
                             '<w:jc w:val="%s"/></w:pPr>%s</w:p>' % ('<w:keepNext w:val="1"/>' if (keep or hdr) else "", "center" if hdr else al,
                                                                    _runs(chunk, size, bold or hdr, INK if not hdr else "000000")))
            return ('<w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/><w:tcBorders>'
                    '<w:top w:val="single" w:sz="4" w:space="0" w:color="C9D3DF"/><w:left w:val="single" w:sz="4" w:space="0" w:color="C9D3DF"/>'
                    '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="C9D3DF"/><w:right w:val="single" w:sz="4" w:space="0" w:color="C9D3DF"/>'
                    '</w:tcBorders>%s<w:tcMar><w:top w:w="70" w:type="dxa"/><w:left w:w="90" w:type="dxa"/>'
                    '<w:bottom w:w="70" w:type="dxa"/><w:right w:w="90" w:type="dxa"/></w:tcMar><w:vAlign w:val="%s"/></w:tcPr>%s</w:tc>'
                    % (width, shd, "center" if hdr else "top", "".join(paras)))

        trs = ['<w:tr><w:trPr><w:cantSplit w:val="1"/><w:tblHeader w:val="1"/></w:trPr>%s</w:tr>'
               % "".join(cell(h, w[i], HEAD_FILL, True, "center", hdr=True) for i, h in enumerate(header))]
        for r, row in enumerate(rows):
            fill = ZEBRA if zebra and r % 2 == 1 else None
            trs.append('<w:tr><w:trPr><w:cantSplit w:val="1"/></w:trPr>%s</w:tr>'
                       % "".join(cell(v, w[i], fill, first_bold and i == 0, align[i]) for i, v in enumerate(row)))
        self._add('<w:tbl><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:jc w:val="center"/><w:tblLayout w:type="fixed"/>'
                  '<w:tblLook w:val="0420"/></w:tblPr><w:tblGrid>%s</w:tblGrid>%s</w:tbl>' % (TABLE_W, grid, "".join(trs)))
        self.gap(100, keep_next=keep)

    def kpi(self, cards):
        """요약 카드 — [(큰 값, 이름, 설명)] 3~4개. 값은 파랑 큰 글씨."""
        n = len(cards)
        w = TABLE_W // n
        cells = []
        for value, label, note in cards:
            check(value + label + note, "카드")
            cells.append('<w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/><w:tcBorders><w:top w:val="single" w:sz="18" w:space="0" w:color="%s"/>'
                         '<w:left w:val="nil"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="C9D3DF"/><w:right w:val="single" w:sz="4" w:space="0" w:color="FFFFFF"/></w:tcBorders>'
                         '<w:shd w:val="clear" w:color="auto" w:fill="F5F8FC"/><w:tcMar><w:top w:w="110" w:type="dxa"/><w:left w:w="140" w:type="dxa"/>'
                         '<w:bottom w:w="110" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tcMar></w:tcPr>'
                         '<w:p><w:pPr><w:keepNext w:val="1"/><w:spacing w:after="20"/></w:pPr>%s</w:p>'
                         '<w:p><w:pPr><w:keepNext w:val="1"/><w:spacing w:after="30"/></w:pPr>%s</w:p>'
                         '<w:p><w:pPr><w:spacing w:after="0" w:line="252" w:lineRule="auto"/></w:pPr>%s</w:p></w:tc>'
                         % (w, BLUE, _runs(value, 40 if len(value) <= 6 else (30 if len(value) <= 9 else 24), True, BLUE), _runs(label, 18, True, INK), _runs(note, 15, False, MUTED)))
        grid = "".join('<w:gridCol w:w="%d"/>' % w for _ in cards)
        self._add('<w:tbl><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:jc w:val="center"/><w:tblLayout w:type="fixed"/>'
                  '<w:tblLook w:val="0000"/></w:tblPr><w:tblGrid>%s</w:tblGrid><w:tr><w:trPr><w:cantSplit w:val="1"/></w:trPr>%s</w:tr></w:tbl>'
                  % (w * n, grid, "".join(cells)))
        self.gap(120)

    def callout(self, title: str, lines, kind="info"):
        """한 칸 상자 — info(파랑) / warn(주황). 제목 한 줄 + 글머리 줄."""
        fill, bar = {"info": ("EEF5FD", BLUE), "warn": ("FDF3E8", "EB6834")}[kind]
        paras = ['<w:p><w:pPr><w:keepNext w:val="1"/><w:spacing w:after="40"/></w:pPr>%s</w:p>' % _runs(title, 19, True, DARK)]
        for line in lines:
            check(line, "상자")
            paras.append('<w:p><w:pPr><w:spacing w:after="30" w:line="264" w:lineRule="auto"/><w:ind w:left="200" w:hanging="200"/></w:pPr>%s%s</w:p>'
                         % (_runs("· ", 17, False, MUTED), _runs(line, 17, False, INK)))
        self._add('<w:tbl><w:tblPr><w:tblW w:w="%d" w:type="dxa"/><w:jc w:val="center"/><w:tblLayout w:type="fixed"/><w:tblLook w:val="0000"/></w:tblPr>'
                  '<w:tblGrid><w:gridCol w:w="%d"/></w:tblGrid><w:tr><w:trPr><w:cantSplit w:val="1"/></w:trPr><w:tc><w:tcPr><w:tcW w:w="%d" w:type="dxa"/>'
                  '<w:tcBorders><w:top w:val="nil"/><w:left w:val="single" w:sz="30" w:space="0" w:color="%s"/><w:bottom w:val="nil"/><w:right w:val="nil"/></w:tcBorders>'
                  '<w:shd w:val="clear" w:color="auto" w:fill="%s"/><w:tcMar><w:top w:w="100" w:type="dxa"/><w:left w:w="160" w:type="dxa"/>'
                  '<w:bottom w:w="80" w:type="dxa"/><w:right w:w="140" w:type="dxa"/></w:tcMar></w:tcPr>%s</w:tc></w:tr></w:tbl>'
                  % (TABLE_W, TABLE_W, TABLE_W, bar, fill, "".join(paras)))
        self.gap(120)

    # -- 그림 ---------------------------------------------------------
    def figure(self, png: Path, caption: str, width_cm=15.0, max_height_cm=None):
        from docxfill import png_size
        w_px, h_px = png_size(str(png))
        width = width_cm
        if max_height_cm and width * h_px / w_px > max_height_cm:
            width = max_height_cm * w_px / h_px
        p = self.doc.add_paragraph()
        self.sect.addprevious(p._p)
        p.alignment = 1
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = 60
        p.add_run().add_picture(str(png), width=Cm(width))
        self.figures += 1
        caption = "그림 %d. %s" % (self.figures, caption)
        check(caption, "캡션")
        self._add('<w:p><w:pPr><w:spacing w:after="160"/><w:jc w:val="center"/></w:pPr>%s</w:p>' % _runs(caption, 17, False, MUTED))

    # -- 저장 ---------------------------------------------------------
    def all_text(self):
        parts = [p.text for p in self.doc.paragraphs]
        for tb in self.doc.tables:
            for row in tb.rows:
                parts += [c.text for c in row.cells]
        return "\n".join(parts)

    def save(self, out: Path):
        check(self.all_text(), out.name)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))
        reopened = Document(str(out))                      # 저장 뒤 다시 열어 본다
        check("\n".join(p.text for p in reopened.paragraphs), out.name)
        return out


def _with_ns(xml: str) -> str:
    """루트 요소에 이름공간 선언을 붙여 lxml 이 읽게 한다."""
    return re.sub(r"^(<[\w:]+)", r"\1 " + nsdecls("w", "wp", "a", "pic", "r"), xml, count=1)
