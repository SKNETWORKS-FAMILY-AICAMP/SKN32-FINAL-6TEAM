# 파일: program/산출물양식/_build/build_아키텍처_v3.py
import os
import re
import sys
import time
import tempfile
import subprocess
from collections import Counter
from copy import deepcopy
from pathlib import Path

sys.dont_write_bytecode = True
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.text.paragraph import Paragraph
from docxfill import Doc, check, BANNED
from docx_tables import insert_table_after
from fit_images import fit

FORMS = HERE.parent
W2 = FORMS / "w2"

BANNED_MAP = {
    "\u2014": "-", "\u2013": "-", "\u2192": " > ", "\u21d2": " > ",
    "\u2605": "", "\u2606": "", "\u2713": "O", "\u2714": "O",
    "\u2715": "X", "\u2717": "X", "\u25b8": "-", "\u2026": "...",
    "\u21c4": "<>",
    "\u2460": "(1)", "\u2461": "(2)", "\u2462": "(3)",
    "\u2463": "(4)", "\u2464": "(5)", "\u2465": "(6)",
    "\u2212": "-", "\uff1d": "=",
}


def clean(s):
    s = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1(\2)", s)
    s = re.sub(r"\[([^\]]+)\]\(<?[^)]*>?\)", r"\1", s)
    s = s.replace("**", "").replace("`", "")
    for a, b in BANNED_MAP.items():
        s = s.replace(a, b)
    return re.sub(r"  +", " ", s).strip()


def text(s):
    # HTML 줄바꿈은 Word 줄바꿈으로 옮긴다.
    s = clean(re.sub(r"<br\s*/?>", "\n", s, flags=re.I))
    check(s)
    if re.search(
        r"[\u3001\u3002\u3040-\u30ff\u31f0-\u31ff"
        r"\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff"
        r"\uff66-\uff9f\U00020000-\U000323af]",
        s,
    ):
        raise ValueError("허용하지 않은 문자: " + s)
    return s


def norm(s):
    return " ".join(s.split())


def parse(md, stop):
    # 작업 기록부터 제외한다. 본문에는 정해진 치환만 적용한다.
    for title in stop:
        md = re.split(
            r"(?m)^\*\*" + re.escape(title) + r"\*\*\s*$",
            md,
            maxsplit=1,
        )[0]

    blocks = []
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line or re.fullmatch(r"[-*_]{3,}", line):
            continue

        if line.startswith("|"):
            rows = []
            while True:
                cells = re.split(r"(?<!\\)\|", line.strip("|"))
                separator = all(
                    re.fullmatch(r"\s*:?-+:?\s*", c) for c in cells
                )
                if not separator:
                    rows.append([
                        text(c.replace(r"\|", "|")) for c in cells
                    ])
                if i == len(lines) or not lines[i].lstrip().startswith("|"):
                    break
                line = lines[i].strip()
                i += 1

            if len(rows) < 2 or any(len(r) != len(rows[0]) for r in rows):
                raise ValueError("표 열 수 또는 본문 행 오류")
            blocks.append(("table", rows))
        elif re.match(r"^#{1,3} ", line):
            marks, value = line.split(" ", 1)
            blocks.append(("h" + str(len(marks)), text(value)))
        elif re.fullmatch(r"\*\*.+\*\*", line):
            blocks.append(("bold", text(line)))
        else:
            blocks.append(("para", text(line)))

    return blocks


def split_source(config):
    blocks = parse(
        config["src"].read_text(encoding="utf-8-sig"),
        config["stop"],
    )
    if [b[0] for b in blocks[:2]] != ["h1", "table"]:
        raise ValueError("표지 메타데이터 구조 변경")

    meta = blocks[1][1]
    sections = {}
    headings = {}
    current = None

    for kind, value in blocks[2:]:
        if kind in ("h2", "h3"):
            key = value.split()[0]
            if key in sections:
                raise ValueError("절 중복: " + key)
            current = key
            sections[key] = []
            headings[key] = value
        else:
            if current is None:
                raise ValueError("절 밖의 본문")
            sections[current].append((kind, value))

    if list(sections) != list(config["anchors"]):
        raise ValueError("정본 절 순서 또는 양식 대응 변경")

    return blocks[0][1], meta, sections, headings


def paragraph_texts(doc):
    # 병합 셀을 중복 방문하지 않고 줄바꿈도 보존한다.
    for p in doc.element.body.iter():
        if p.tag.endswith("}p"):
            yield Paragraph(p, doc._body).text


class ExactDoc(Doc):
    def normalize_template(self):
        # 질문과 머리행 원문은 일괄 정규화하지 않는다.
        pass


def prepare(config, d, image=True):
    title, meta, sections, headings = split_source(config)
    doc = Document(config["template"])
    tables = list(doc.tables)
    paragraphs = list(doc.paragraphs)

    expected = Counter()
    protected = []
    original_tables = [t._tbl for t in tables]
    pending = []
    serial = 0

    def expect(s):
        if norm(s):
            expected[norm(s)] += 1

    def protect_cell(cell):
        protected.append((cell._tc, cell.text))

    def set_cell(cell, value):
        cell.text = value

    def add_fragment(after, xml):
        wrapper = parse_xml(
            "<w:root " + nsdecls("w") + ">" + xml + "</w:root>"
        )
        for child in list(wrapper):
            after.addnext(child)
            after = child
        return after

    def put(after, value, bold=False):
        expect(value)
        items = [("b", value)] if bold else [value]
        return add_fragment(after, d.build(items))

    anchors = {}
    keep = set()
    for key, (kind, index) in config["anchors"].items():
        if kind == "t":
            table = tables[index]
            actual = table.cell(0, 0).text
            anchors[key] = table._tbl
            protect_cell(table.cell(0, 0))
        else:
            p = paragraphs[index]
            actual = p.text
            anchors[key] = p._p
            keep.add(p._p)
            protected.append((p._p, actual))

        if norm(actual) != norm(headings[key]):
            raise ValueError("양식 제목 불일치: " + key)

    # 표지의 기수·팀명은 사용자 지정값을 사용한다.
    cover = tables[0].cell(0, 0)
    cover.text = "SK 네트웍스 Family AI 32기 : 6팀 A-COPilot"
    cover.add_paragraph(title)
    expect(title)

    for row in tables[1].rows:
        protect_cell(row.cells[0])

    values = {r[0]: r[1] for r in meta[1:]}
    for key in ("제출 일자", "깃허브 경로", "작성 팀원"):
        if not values[key].startswith("[미확보]"):
            raise ValueError("표지 미확보 표시 변경: " + key)

    seen = set()
    for row in tables[1].rows:
        key = row.cells[0].text
        set_cell(row.cells[1], values[key])
        expect(key)
        expect(values[key])
        seen.add(key)

    # 양식에 없는 메타데이터는 기존 표에 행으로 추가한다.
    for key, value in meta[1:]:
        if key not in seen:
            cells = tables[1].add_row().cells
            set_cell(cells[0], key)
            set_cell(cells[1], value)
            expect(key)
            expect(value)

    security = config.get("security")
    if security:
        label = paragraphs[security["heading"]]
        keep.add(label._p)
        protected.append((label._p, label.text))
        for index in security["questions"]:
            keep.add(paragraphs[index]._p)

    # 예시 본문만 지운다. 표·절 제목·보안 질문은 유지한다.
    first_body = doc.element.body.index(next(iter(anchors.values())))
    for p in paragraphs:
        if (
            doc.element.body.index(p._p) >= first_body
            and p._p not in keep
            and p.text.strip()
        ):
            p._p.getparent().remove(p._p)

    for index in config["main_tables"].values():
        for cell in tables[index].rows[0].cells:
            protect_cell(cell)

    def queue_table(after, rows, native=None, preserve_labels=False):
        nonlocal serial
        serial += 1
        marker = "@@표%d@@" % serial
        after = add_fragment(after, d.build([marker]))
        for row in rows:
            for value in row:
                expect(value)
        pending.append((marker, rows, native, preserve_labels))
        return after

    image_table = tables[config["image_table"]]
    image_cell = image_table.cell(0, 0)
    image_cell.text = ""

    if image:
        fragment = parse_xml(
            "<w:root " + nsdecls("w") + ">"
            + d.image_xml(str(config["fig"]), 6.2)
            + "</w:root>"
        )
        image_cell._tc.replace(image_cell.paragraphs[0]._p, fragment[0])

    for key, content in sections.items():
        after = anchors[key]

        if key == config.get("image_before"):
            after.addnext(image_table._tbl)
            after = image_table._tbl

        native = config["main_tables"].get(key)
        native_used = False
        security_done = False

        for kind, value in content:
            if security and key == security["section"]:
                if (
                    kind == "bold"
                    and value == paragraphs[security["heading"]].text
                ):
                    after = paragraphs[security["heading"]]._p
                    continue

                if (
                    kind == "table"
                    and value[0] == ["질문 칸", "값"]
                    and not security_done
                ):
                    questions = security["questions"]
                    if len(value[1:]) != len(questions):
                        raise ValueError("보안 질문 수 변경")

                    for index, row in zip(questions, value[1:]):
                        p = paragraphs[index]
                        original = re.sub(
                            r"^\[ 확인완료 \]\s*", "", p.text
                        )
                        if original != row[0]:
                            raise ValueError("보안 질문 원문 불일치")

                        status = re.match(r"\[[^\]]+\]", row[1])
                        if not status:
                            raise ValueError("보안 판정 상태 없음")

                        # 질문 원문은 유지하고 기존 확인완료 표시만 바꾼다.
                        p.text = status.group() + " " + original
                        protected.append((p._p, p.text))

                    after = paragraphs[questions[-1]]._p
                    security_done = True

            if kind == "table":
                if (
                    key == config.get("image_in")
                    and value[0] == ["(이미지 삽입)"]
                ):
                    if len(value) != 2:
                        raise ValueError("그림 명세 구조 변경")

                    after.addnext(image_table._tbl)
                    after = image_table._tbl

                    # 그림은 원래 칸에, 정본의 명세는 덧붙인 행에 넣는다.
                    cell = image_table.add_row().cells[0]
                    set_cell(cell, value[1][0])
                    expect(value[1][0])

                elif native is not None and not native_used:
                    actual = [
                        norm(c.text) for c in tables[native].rows[0].cells
                    ]
                    if actual != [norm(c) for c in value[0]]:
                        raise ValueError("머리행 불일치: " + key)

                    after = queue_table(
                        after,
                        value,
                        tables[native],
                        native == config.get("label_table"),
                    )
                    native_used = True
                else:
                    after = queue_table(after, value)
            else:
                after = put(after, value, kind == "bold")

        if native is not None and not native_used:
            raise ValueError("양식 표에 대응할 정본 표 없음: " + key)
        if security and key == security["section"] and not security_done:
            raise ValueError("보안 판정 누락")

    # 모든 본문 표는 docx_tables를 거쳐 실제 Word 표로 만든다.
    # 기존 표와 대응하면 원래 표·머리행을 유지하여 내용을 옮긴다.
    from docx.table import Table, _Cell

    for marker, rows, native, preserve_labels in pending:
        # 정본에 없는 캡션 문장은 만들지 않는다.
        cap = insert_table_after(doc, marker, "", rows[0], rows[1:])
        donor = cap._p.getnext()
        marker_p = next(p for p in doc.paragraphs if p.text == marker)

        if native is not None:
            source_rows = list(Table(donor, doc._body).rows)[1:]
            existing = list(native.rows)[1:]

            if preserve_labels:
                pools = {}
                for row in existing:
                    key = row.cells[0].text
                    pools.setdefault(key, []).append(row)
                    protect_cell(row.cells[0])

                last = native.rows[0]._tr
                for values_row, source_row in zip(rows[1:], source_rows):
                    pool = pools.get(values_row[0], [])
                    if pool:
                        target = pool.pop(0)
                        for j, value in enumerate(values_row[1:], 1):
                            set_cell(target.cells[j], value)
                        last.addnext(target._tr)
                        last = target._tr
                    else:
                        added = deepcopy(source_row._tr)
                        last.addnext(added)
                        last = added

                if any(pools.values()):
                    raise ValueError("정본에 없는 양식 질문 행")
            else:
                if len(existing) > len(source_rows):
                    raise ValueError("양식 행을 삭제해야 하는 표")
                while len(native.rows) - 1 < len(source_rows):
                    native.add_row()
                for row, values_row in zip(
                    list(native.rows)[1:], rows[1:]
                ):
                    for cell, value in zip(row.cells, values_row):
                        set_cell(cell, value)

            donor.addprevious(native._tbl)
            donor.getparent().remove(donor)

        cap._p.getparent().remove(cap._p)
        marker_p._p.getparent().remove(marker_p._p)

    def validate(result):
        actual = Counter(
            norm(v) for v in paragraph_texts(result) if norm(v)
        )
        # 여러 문단·줄바꿈을 포함한 셀도 셀 전체로 대조한다.
        cell_counts = Counter(
            norm(_Cell(tc, result._body).text)
            for tc in result.element.body.iter()
            if tc.tag.endswith("}tc")
        )
        actual |= cell_counts

        missing = expected - actual
        if missing:
            raise ValueError(
                "정본 누락: " + repr(list(missing.items())[:5])
            )

        for value in paragraph_texts(result):
            text(value)
            if any(ch in value for ch in BANNED) or "@@표" in value:
                raise ValueError("금지 문자 또는 표 자리표시 잔존")
            if "[ 확인완료 ]" in value:
                raise ValueError("예시 보안 판정 잔존")

        if image and len(result.inline_shapes) != 1:
            raise ValueError("그림은 한 개여야 한다")

    # 질문·절 제목·표 머리행의 원문을 확인한다.
    for element, original in protected:
        if element.tag.endswith("}tc"):
            value = _Cell(element, doc._body).text
        else:
            value = Paragraph(element, doc._body).text
        if value != original:
            raise ValueError(
                "양식 질문·제목·머리행 변경: " + original
            )

    body_tables = [
        e for e in doc.element.body if e.tag.endswith("}tbl")
    ]
    retained = [e for e in body_tables if e in original_tables]
    if retained != original_tables:
        raise ValueError("원본 표 삭제 또는 순서 변경")

    validate(doc)
    return doc, validate, len(pending)


def build(config):
    # 임시 작업이나 출력 저장에 앞서 필수 그림을 검사한다.
    if not config["fig"].is_file():
        raise FileNotFoundError(config["fig"])

    with tempfile.TemporaryDirectory(
        prefix="_tmp_v3_", dir=HERE
    ) as tmp:
        d = ExactDoc(
            str(config["template"]),
            str(Path(tmp) / "package"),
        )
        doc, validate, count = prepare(config, d)

        from lxml import etree

        d.xml = etree.tostring(doc.element, encoding="unicode")
        stage = Path(tmp) / config["out"].name
        stage = Path(d.save(str(stage)))

        validate(Document(stage))
        fit(str(stage))
        subprocess.run(
            [
                sys.executable,
                str(HERE / "add_outline.py"),
                str(stage),
            ],
            check=True,
        )
        outline = stage.with_name(stage.stem + "_outline.docx")
        if outline.exists():
            os.replace(outline, stage)

        validate(Document(stage))

        # 후처리까지 마친 문서를 저장한다.
        # 처음부터 출력이 잠겨 있어도 미완성 문서를 남기지 않는다.
        target = config["out"]
        for attempt in range(5):
            try:
                os.replace(stage, target)
                break
            except PermissionError:
                time.sleep(0.4)
        else:
            target = target.with_name(target.stem + "_새버전.docx")
            os.replace(stage, target)
            print("원본이 잠겨 새 이름으로 저장한다:", target.name)

        print(
            "입력:", config["src"].name,
            "| 표", count,
            "| 그림 1 | 저장:", target.name,
        )


CONFIG = {
    "src": (
        W2 / "아키텍처" / "AI시스템아키텍처_코덱스_3차.md"
    ),
    "template": (
        FORMS
        / "[모델링 및 평가] AI 시스템 아키텍처 (멀티 에이전트 아키텍처)_양식.docx"
    ),
    "fig": (
        W2 / "표와시각자료" / "코덱스"
        / "ARCH_에이전트아키텍처_코덱스.png"
    ),
    "out": (
        FORMS / "[모델링 및 평가] AI 시스템 아키텍처_2차_여행.docx"
    ),
    "stop": [
        "Claude에게 넘기는 목록",
        "v11 담당에게 넘길 것",
        "반박 처리표",
    ],
    "anchors": {
        "1.": ("t", 2),
        "1.1": ("p", 6),
        "1.2": ("p", 8),
        "1.3": ("p", 15),
        "1.4": ("p", 17),
        "2.": ("t", 3),
        "2.1": ("p", 22),
        "2.2": ("p", 25),
        "2.3": ("p", 27),
        "2.4": ("p", 37),
        "3.": ("t", 6),
        "4.": ("t", 7),
        "4.1": ("p", 56),
        "4.2": ("p", 59),
        "4.3": ("p", 62),
        "4.4": ("p", 65),
        "5.": ("t", 8),
        "6.": ("t", 10),
    },
    "image_table": 4,
    "image_in": "2.1",
    "main_tables": {
        "2.2": 5,
        "5.": 9,
    },
}

if __name__ == "__main__":
    build(CONFIG)
