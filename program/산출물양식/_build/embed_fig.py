# -*- coding: utf-8 -*-
"""제출본 docx 의 특정 문단 뒤에 그림(+캡션)을 넣는다. 같은 그림은 두 번 넣지 않는다.
python embed_fig.py <docx 이름> <앵커 문단 시작 글자> <png> <캡션>"""
import os, sys, shutil
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
FORMS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
name, anchor, png, cap = sys.argv[1:5]
p = os.path.join(FORMS, "w2", "제출본", name)
d = Document(p)
if any(par.text.strip() == cap for par in d.paragraphs):
    print("이미 있음:", cap); sys.exit(0)
target = next((par for par in d.paragraphs if par.text.strip().startswith(anchor)), None)
if target is None:
    print("앵커 없음:", anchor); sys.exit(1)
pic = d.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
pic.add_run().add_picture(png, width=Inches(6.2))
c = d.add_paragraph(cap); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
target._p.addnext(c._p); target._p.addnext(pic._p)
d.save(p); shutil.copy(p, os.path.join(FORMS, name))
print("넣음:", cap, "| 앵커:", target.text[:30])
