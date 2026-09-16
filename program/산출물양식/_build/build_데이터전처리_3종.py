"""커머스 데이터 전처리 제출 산출물 3종 빌드.

근거:
- program/산출물양식/w3/데이터전처리/데이터전처리결과서_본문_2026-09-16.md
- program/산출물양식/w3/데이터전처리/학습결과서_본문_2026-09-16.md
- program/산출물양식/w3/데이터전처리/학습한모델_본문_2026-09-16.md

표지 셋·변경자는 팀장 지시로 비움.
표지의 제출 일자·깃허브 경로·작성 팀원과 모든 변경자 칸을 빈 문자열로 둔다.
기존 build_수집데이터보고서_v2.py의 setc·grow·fill 방식을 따른다.
표 밖에는 문단을 추가하지 않으며, 양식 원본은 덮어쓰지 않는다.

저장 위치: program/산출물양식/_build/
실행 환경: Python 3, python-docx
"""

import copy
import os
import re
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.shared import Pt


sys.stdout.reconfigure(encoding="utf-8")

HERE = Path(__file__).resolve().parent
FORMS = HERE.parent
OUT_DIR = FORMS / "w3" / "데이터전처리"

UNKNOWN = "모름 [미확보]"
COVER = "SK 네트웍스 Family AI 32기: 6팀"
DOMAIN_NOTE = "커머스(쇼핑몰) 시절 자료, 2026-09-08 여행 CS 로 전환 [실측]"

# 공개 금지 식별자는 스크립트에 원문 문자열로 기록하지 않는다.
PRIVATE_SERVER = "".join(chr(n) for n in (120, 54, 48, 48))

SOURCE_NAMES = (
    "데이터전처리결과서_본문_2026-09-16.md",
    "학습결과서_본문_2026-09-16.md",
    "학습한모델_본문_2026-09-16.md",
)

TEMPLATE_NAMES = (
    "[데이터 전처리] 데이터 전처리 결과서.docx",
    "[데이터 전처리] 머신러닝_딥러닝 학습 결과서.docx",
    "[데이터 전처리] 학습한 ML_DL 모델.docx",
)

OUTPUT_NAMES = (
    "[데이터 전처리] 데이터 전처리 결과서_3차_커머스.docx",
    "[데이터 전처리] 머신러닝_딥러닝 학습 결과서_3차_커머스.docx",
    "[데이터 전처리] 학습한 ML_DL 모델_3차_커머스.docx",
)

HISTORY_HEADER = ("변경일", "변경자", "변경내용", "영향 받는 항목", "비고")
EXPECTED_HEADERS = (
    {
        3: ("이슈 유형", "대상 필드", "판정 기준", "처리 방식", "처리량 / 비고"),
        4: ("처리 단계", "입력", "출력", "목적"),
        5: ("필드", "타입", "결측 처리", "중복 처리", "이상치 처리", "비고"),
        6: ("구분", "항목", "기준 / 결과", "설명 및 관리 방안"),
        7: HISTORY_HEADER,
    },
    {
        3: ("유형", "모델명", "선정 근거", "학습 데이터", "비고"),
        4: ("비교 지표", "산출식 / 측정 대상", "목표값", "비고"),
        5: ("파라미터", "값", "탐색 범위", "선정 근거", "적용 모델", "비고"),
        6: ("구분", "실험 #", "설정", "결과 및 비고"),
        7: HISTORY_HEADER,
    },
    {
        3: ("구성 요소", "유형", "기능", "의존성", "비고"),
        4: ("입·출력 필드", "타입", "shape / 길이", "설명"),
        5: ("항목", "값", "단위", "측정 조건", "서비스 타겟", "비고"),
        6: ("구분", "항목", "값 / 방법", "설명 및 관리 방안"),
        7: HISTORY_HEADER,
    },
)

FORBIDDEN = (
    "BERT-base-ko", "XGBoost", "Logistic", "TF-IDF", "8,000건",
    "2026.04.11", "2026.04.13", "79건", "416건", "124건",
    "MLflow", "joblib", "xgb_classifier", "sbert_embed",
    "random_state=42", PRIVATE_SERVER, "_____팀", "27기", "00기",
)


def nfc(text):
    return unicodedata.normalize("NFC", text)


def normalized_key(text):
    return re.sub(r"\s+", "", nfc(text))


def find_actual_name(directory, wanted):
    """NFC로 비교하되 os.listdir이 반환한 실제 파일명을 사용한다."""
    matches = [
        directory / actual
        for actual in os.listdir(directory)
        if nfc(actual) == nfc(wanted)
        and (directory / actual).is_file()
    ]
    if len(matches) != 1:
        raise RuntimeError(
            f"파일을 유일하게 찾을 수 없음: {wanted} (검색 결과 {len(matches)}개)"
        )
    return matches[0]


def public_text(text):
    text = str(text)
    text = text.replace(PRIVATE_SERVER + " GPU 서버", "GPU 서버")
    return text.replace(PRIVATE_SERVER, "GPU 서버")


def clean_md(text):
    """마크다운 장식만 제거하며 숫자와 등급 표시는 보존한다."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    text = text.replace(r"\|", "|")
    text = re.sub(r"<br\s*/?>", "\n", text, flags=re.IGNORECASE)
    return public_text(text.strip())


def parse_md_tables(text):
    tables = []
    lines = text.splitlines()
    i = 0

    def cells(line):
        return [
            clean_md(value)
            for value in re.split(r"(?<!\\)\|", line.strip()[1:-1])
        ]

    while i + 1 < len(lines):
        line = lines[i].strip()
        separator = lines[i + 1].strip()
        if not (
            line.startswith("|") and line.endswith("|")
            and separator.startswith("|") and separator.endswith("|")
        ):
            i += 1
            continue
        sep_cells = [
            value.strip() for value in separator[1:-1].split("|")
        ]
        if not sep_cells or not all(
            re.fullmatch(r":?-{3,}:?", value) for value in sep_cells
        ):
            i += 1
            continue

        header = cells(line)
        rows = []
        i += 2
        while i < len(lines):
            row_line = lines[i].strip()
            if not (row_line.startswith("|") and row_line.endswith("|")):
                break
            row = cells(row_line)
            if len(row) != len(header):
                raise ValueError(f"본문 표의 열 수 불일치: {header[0]}")
            rows.append(tuple(row))
            i += 1
        tables.append((tuple(header), rows))
    return tables


def source_table(tables, first, width, second=None):
    matches = [
        rows
        for header, rows in tables
        if len(header) == width
        and normalized_key(header[0]) == normalized_key(first)
        and (
            second is None
            or normalized_key(header[1]) == normalized_key(second)
        )
    ]
    if len(matches) != 1:
        raise ValueError(f"본문 표 식별 실패: {first}, {width}열")
    return copy.deepcopy(matches[0])


def setp(paragraph, text):
    """첫 run의 서식을 유지하고 나머지 내용 요소를 제거한다."""
    if paragraph.runs:
        run = paragraph.runs[0]
    else:
        run = paragraph.add_run("")
    for child in list(paragraph._p):
        if child is paragraph._p.pPr or child is run._r:
            continue
        paragraph._p.remove(child)
    run.text = public_text(text)
    return run


def setc(cell, text, bold=False):
    """기존 setc 방식: 첫 문단·첫 run만 남기고 셀을 채운다."""
    paragraph = cell.paragraphs[0]
    for child in list(cell._tc):
        if child is cell._tc.tcPr or child is paragraph._p:
            continue
        cell._tc.remove(child)
    run = setp(paragraph, text)
    run.bold = bold
    run.font.size = Pt(9)
    run.font.name = "맑은 고딕"


def grow(tbl, need):
    """머리행을 제외한 데이터 행을 need개까지 늘린다."""
    while len(tbl.rows) - 1 < need:
        tbl._tbl.append(copy.deepcopy(tbl.rows[-1]._tr))


def trim(tbl, need):
    """머리행을 제외한 데이터 행을 정확히 need개로 줄인다."""
    while len(tbl.rows) - 1 > need:
        row = tbl.rows[-1]._tr
        row.getparent().remove(row)


def fill(tbl, rows, start=1):
    """모든 칸을 채우고 남는 예시 행을 삭제한다."""
    rows = list(rows)
    if start != 1:
        raise ValueError("이 빌드의 본문 표는 머리행 다음 행부터 채운다.")
    width = len(tbl.columns)
    if not rows:
        rows = [tuple(UNKNOWN for _ in range(width))]
    for row in rows:
        if len(row) != width:
            raise ValueError(f"표 열 수 불일치: 기대 {width}, 실제 {len(row)}")
        for value in row:
            if len(public_text(value)) > 120:
                raise ValueError(
                    "120자를 넘는 셀은 자동 절단하지 않는다. 요약 또는 행 분리가 필요하다: "
                    + str(row[0])
                )

    grow(tbl, len(rows))
    trim(tbl, len(rows))

    for i, values in enumerate(rows, start=start):
        row = tbl.rows[i]
        row.height = None
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        # 복사된 데이터 행이 반복 머리행으로 동작하지 않도록 한다.
        for flag in list(row._tr.xpath("./w:trPr/w:tblHeader")):
            flag.getparent().remove(flag)
        for j, value in enumerate(values):
            setc(row.cells[j], value)

    header_properties = tbl.rows[0]._tr.get_or_add_trPr()
    if not header_properties.xpath("./w:tblHeader"):
        header_properties.append(OxmlElement("w:tblHeader"))


def unique_cells(table):
    seen = set()
    for row in table.rows:
        for cell in row.cells:
            if cell._tc not in seen:
                seen.add(cell._tc)
                yield cell


def validate_template(doc, index):
    expected_paragraphs = (24, 24, 25)[index]
    if len(doc.tables) != 8 or len(doc.paragraphs) != expected_paragraphs:
        raise ValueError(f"양식 구조가 변경됨: {TEMPLATE_NAMES[index]}")
    for table_index, expected in EXPECTED_HEADERS[index].items():
        actual = tuple(
            cell.text for cell in doc.tables[table_index].rows[0].cells
        )
        if tuple(map(normalized_key, actual)) != tuple(map(normalized_key, expected)):
            raise ValueError(f"T{table_index} 머리행 불일치: {TEMPLATE_NAMES[index]}")


def fill_cover(doc):
    cover_cells = list(unique_cells(doc.tables[0]))
    if len(cover_cells) != 1:
        raise ValueError("T0의 병합 셀 구조가 예상과 다르다.")
    setc(cover_cells[0], COVER)

    values = {
        "산출물 단계": "데이터 전처리",
        "제출 일자": "",
        "깃허브 경로": "",
        "작성 팀원": "",
    }
    found = set()
    for row in doc.tables[1].rows:
        key = normalized_key(row.cells[0].text)
        for label, value in values.items():
            if key == normalized_key(label):
                setc(row.cells[1], value)
                found.add(label)
    if found != set(values):
        raise ValueError("T1의 표지 항목을 모두 찾지 못했다.")


def preprocessing_data(tables):
    quality = source_table(tables, "이슈 유형", 5)
    pipeline = source_table(tables, "처리 단계", 4)
    fields = source_table(tables, "필드", 6)
    management = []

    # 파일 수와 데이터 레코드 수를 명확히 구분한다.
    for path, nature, raw, processed, scripts in source_table(
        tables, "데이터셋 경로(datasets/ 기준)", 5
    ):
        management.append((
            "파일 현황", path,
            f"raw {raw}; processed {processed}; scripts {scripts}",
            f"{nature}. 2026-09-16 [실측] 파일 수이며 레코드 수가 아님.",
        ))

    management.extend([
        ("원천", "K쇼핑 다운로드", "2026-08-20 [실측]; 83MB [실측]",
         "라벨링데이터만 다운로드하고 음성 원천은 제외."),
        ("원천", "확인 배치", "3개 [실측]",
         "K쇼핑은 라벨링데이터_231222_add에만 존재."),
    ])

    for category, intent, total, customer, sample in source_table(
        tables, "원천 카테고리", 5
    ):
        management.append((
            "K쇼핑 카테고리", category,
            f"전체 {total}; 고객 질의 {customer}",
            f"매핑 의도: {intent}; 표본: {sample}",
        ))

    management.extend([
        ("원천 범위", "카테고리", "7개 [실측]",
         "주문·결제·배송·반품·교환·AS·업무처리."),
        ("원천 범위", "AS·업무처리 제외", "15,635 + 227,108 = 242,743건 [실측]",
         "매핑된 의도가 없어 표본을 추출하지 않은 원천 규모."),
        ("학습 자료", "공개 QA", "5개 카테고리 × 300건 = 1,500건 [실측]",
         "주문·결제·배송·반품·교환에서 각 300건 [실측] 추출."),
        ("주문·배송", "정규화 산출물",
         "processed/orders.jsonl; processed/tracking.jsonl",
         "상품 행 / 주문·배송박스 단위. 원천 주문·상품·배송조회 행 수 모름 [미확보]."),
        ("코퍼스", "구성", "25문서 [실측]; 300청크 [실측]",
         "정책·FAQ 검색 자료. 학습 데이터 규모에 합산하지 않음."),
        ("코퍼스", "평균 길이", "문서 4,231자 [실측]; 청크 241자 [실측]",
         "청크 평균 길이는 공백 제외."),
    ])

    for scope, count in source_table(tables, "scope", 2):
        management.append(("코퍼스 구성", scope, count, "문서 수."))

    for group, size, limits in source_table(
        tables, "구분", 3, second="확인된 규모"
    ):
        management.append(("분할", group, size, limits))

    management.extend([
        ("분할 관리", "독립성", UNKNOWN,
         "미병합만으로 중복·오염 부재를 단정하지 않음. 상담 단위와 내용 중복 검증 필요."),
        ("주문·배송 활용", "정보 완전성", "확보 비율의 분자·분모 모름 [미확보]",
         "정규화 자료로 주문 상태·배송 이력 참조. 완전성을 수치로 보장하지 않음."),
        ("QA 활용", "세부 이슈 라벨", "키워드 규칙 기반",
         "기타 코드 편중을 매핑 실패로 단정하지 않으며 사람 검수 전 확정 정답으로 쓰지 않음."),
        ("후속 보정", "자료 성격", "judge 통과 LLM 응답 22건 [실측]",
         "실제 고객 응대 알파 로그가 아님."),
        ("후속 평가", "형식·품질 구분",
         "golden 파싱 70건/72건 [실측]; judge 통과 0건/72건 [실측]",
         "전처리 완료와 JSON 파싱만으로 모델 활용 적합성을 보장하지 않음."),
        ("입력 길이", "ContextPack",
         "6k~11k 토큰 [실측]; 기본 최대 길이 512토큰 [실측]",
         "입력 잘림 발생. 실제 토큰 길이를 측정하여 설정에 반영."),
        ("입력 길이", "check_lens.py", "현재 작업트리에 없음 [실측]",
         "커밋 33c9f8c^에서 확인. 긴 문맥은 입력·정답 보존 여부 점검."),
    ])

    for item, state, limits in source_table(tables, "관리 항목", 3):
        if item == "모델 산출물 확인":
            management.append((
                "재현성", item,
                "GPU 서버 파일·설정·SHA-256 확인 [실측 2026-09-17]",
                "학습한모델 본문의 후속 확인 반영. 로드 검증은 하지 않았음.",
            ))
        else:
            management.append(("재현성", item, state, limits))

    management.append((
        "후속 적용", "여행 CS", "여행 데이터로 다시 전처리·측정 필요",
        "원천·필드·매핑·품질·분할을 새로 검증. 커머스 실측값을 여행 결과로 전용하지 않음.",
    ))

    history = [
        (date, "", change, affected, note)
        for date, _author, change, affected, note
        in source_table(tables, "변경일", 5)
    ]

    paragraphs = {
        7: "전처리 목적 : 주문·배송, 상담 QA, 정책·FAQ를 정리하여 학습·검색·평가에 활용.",
        8: "전처리 범위 : " + DOMAIN_NOTE,
        9: "사용 도구 : Python, normalize.py, extract_and_map.py, check_corpus.py. 버전 모름 [미확보].",
        18: "분할 기준 : 분할 비율·시드·별도 검증셋 모름 [미확보]. 상담 단위·내용 중복 검증 필요.",
    }
    return paragraphs, {3: quality, 4: pipeline, 5: fields, 6: management, 7: history}


def training_data(tables):
    candidates = source_table(tables, "유형", 5)
    criteria = source_table(tables, "비교 지표", 4)
    parameters = source_table(tables, "파라미터", 6)

    # 본문에 실험 번호가 없으므로 번호를 만들어 붙이지 않는다.
    experiments = [
        ("실험 흐름", UNKNOWN,
         "준비 → 공개 학습 → 후속 학습 → 단독 추론 → 파싱·judge → 경로 비교",
         "출력 형식과 근거 정합성을 구분하여 채택 여부 판정."),
        ("학습·공개 단계", UNKNOWN, "K쇼핑 QA 1,500건 [실측]",
         "376/376 스텝 [실측]; 188스텝에서 재개."),
        ("공개 단계 손실", UNKNOWN, "후반부 189~376스텝 로그 [실측]",
         "loss 0.1987~0.2504, 평균 0.2219 (19회) [실측]."),
        ("공개 단계 요약", UNKNOWN, "train_loss 0.1104 [실측]",
         "최저 스텝 loss 0.1987 [실측]보다 낮음. 마지막 손실로 해석하지 않음."),
        ("공개 단계 해석", UNKNOWN, "재개 실행의 후반부 로그만 남음",
         "낮은 요약 손실의 원인은 [추정]. 시작부터의 손실 감소로 제시하지 않음."),
        ("학습·후속 단계", UNKNOWN, "golden judge 통과 응답 22건 [실측]",
         "12/12 스텝 [실측]; 요약 loss 0.5743 [실측]; 스텝 기록 1회 0.6091 [실측]."),
        ("후속 자료 출처", UNKNOWN, "2026-08-28_reeval_Proposed_v3.jsonl",
         "judge가 통과시킨 LLM 응답. 실제 고객 응대 알파 로그가 아님."),
        ("토큰 정확도 기록", UNKNOWN,
         "mean_token_accuracy: 공개 0.9548 [실측]; 후속 0.9101 [실측]",
         "정답·집계 토큰 수 모름 [미확보]. 비율로 해석하지 않는 학습 로그 기록값."),
        ("평가·golden", UNKNOWN,
         "72건 [실측]; Proposed 3회 [실측]; Proposed+FT 1회 [실측]",
         "Proposed: 72건 × 3회 = 216출력 [실측]. 반복 출력을 독립 고객 수로 세지 않음."),
        ("평가·holdout", UNKNOWN,
         "24건 [실측]; Proposed 3회 [실측]; Proposed+FT 1회 [실측]",
         "Proposed: 24건 × 3회 = 72출력 [실측]. 비교 결과로 프롬프트를 수정하지 않는 조건."),
        ("비교 구조", UNKNOWN,
         "Proposed: RAG+Context Broker; Proposed+FT: 파인튜닝 단독",
         "입력 근거·실행 구조가 다른 경로 비교. 가중치만 바꾼 통제 실험이 아님."),
        ("단독 추론", UNKNOWN, "시스템 지시문 + 고객 문의; 결정론적 생성",
         "RAG·Context Broker·ContextPack 미주입. 후속 단계 어댑터를 평가."),
    ]

    for dataset, metric, proposed, finetuned in source_table(
        tables, "평가 집합", 4
    ):
        experiments.append((
            f"{dataset} · {metric}", UNKNOWN,
            "Proposed: RAG+Context Broker\nProposed+FT: 파인튜닝 단독",
            f"Proposed: {proposed}\nProposed+FT: {finetuned}",
        ))

    experiments.extend([
        ("평가 범위", UNKNOWN, "단계별 독립 judge 비교", UNKNOWN),
        ("holdout 재실행", UNKNOWN, "rate limit으로 judge 호출 전 실패한 행 재실행",
         "최종 실측값 사용. 재실행 전후 통과율 동일."),
        ("검증 구성", UNKNOWN, "확인한 train.py",
         "학습용 검증 데이터·조기 종료·검증 점수 기반 최적 체크포인트 선택이 구성되어 있지 않음."),
        ("과적합 해석", UNKNOWN, "검증 손실 곡선·독립 토큰 평가 모름 [미확보]",
         "소량 후속 자료·golden 재사용으로 일반화 평가 제한. 근거 입력 부재도 실패 원인."),
        ("입력 길이 결함", UNKNOWN,
         "ContextPack 6k~11k 토큰 [실측]; 기본 최대 길이 512토큰 [실측]",
         "2026-08-31 [실측] 입력 잘림 확인. 실제 길이 측정과 길이 설정 전달 보완."),
        ("후속 RAG 통합", UNKNOWN, "evidence 대조·draft 복사 구분",
         "초안 복사와 judge 한계 확인. 직접 상태값 대조로 평가 보완. 최종 미채택."),
        ("방어지표", UNKNOWN, "fixture와 실제 LLM 출력을 구분",
         "holdout 방어지표는 judge 비교로 대체. 제안 라벨·fixture 문제로 측정 대상 미확보."),
        ("최종 판정", UNKNOWN, "DoD 문서상 부분 통과",
         "비교평가·미채택 결정 완료. 실제 알파 로그 학습과 holdout 방어지표 요건 미충족."),
        ("재검토 조건", UNKNOWN, "운영 자료·초안과 근거 불일치 사례·독립 평가 집합",
         "근거 상태값 직접 대조 필요. 동일 처방의 데이터 확대 중단."),
        ("모델 소재", UNKNOWN,
         r"GPU 서버 E:\dod28_ft; 11개 폴더, 9개에 어댑터 [실측 2026-09-17]",
         "ckpt_stage2 파일·SHA-256 확보. 로드 검증은 하지 않았음."),
        ("실패 실행", UNKNOWN, "ckpt_stage3_v4",
         "CUDA 메모리 부족으로 0/36스텝 종료 [실측 2026-09-17]. 폴더 비어 있음."),
        ("실패 실행", UNKNOWN, "ckpt_stage3_v5",
         "1/36스텝에서 CUDA 오류 종료 [실측 2026-09-17]. 폴더 비어 있음."),
        ("비교 제외", UNKNOWN, "2026-08-12 [실측]; provider=mock",
         "1.00 [실측: provider=mock 실행값, 분자·분모 모름 [미확보]]. 모델 성능이 아님."),
    ])

    history = [
        (UNKNOWN, "", "K쇼핑 QA 공개 학습", "응답 형식·도메인 어휘", "커머스 실험 기준 모델로 기록."),
        (UNKNOWN, "", "golden judge 통과 응답으로 후속 학습", "후속 학습 자료",
         "실제 알파 고객 로그 학습으로 표현하지 않음."),
        (UNKNOWN, "", "golden·holdout 단독 경로 비교", "근거 기반 응답 품질",
         "파인튜닝 단독 경로 미채택."),
        ("2026-08-31 [실측]", "", "ContextPack 길이 결함 확인·보완", "토큰 길이·학습 입력",
         "입력·정답 잘림 검증을 선행 조건으로 관리."),
        (UNKNOWN, "", "후속 RAG 통합에서 초안 복사 한계 확인", "근거 대조·평가 방식",
         "직접 상태값 대조로 보완. 최종 미채택 유지."),
        ("2026-09-08 [실측]", "", "커머스에서 여행 CS로 전환", "도메인·학습·평가 자료",
         "커머스 실험 결과를 여행 성능으로 전용하지 않음."),
        (UNKNOWN, "", "커밋 33c9f8c에서 학습 스크립트 79개 [실측] 삭제", "학습 코드",
         "구현 근거는 33c9f8c^에서 확인."),
        ("2026-09-16 [실측]", "", "저장소 모델 산출물 검색 0건 [실측]", "모델 인도 상태",
         "저장소 검색 결과와 GPU 서버 파일 존재를 구분."),
        ("2026-09-17 [실측]", "", "GPU 서버 모델 파일·SHA-256 확인", "ckpt_stage2·실행 재현",
         "11개 폴더 중 9개에 어댑터 [실측]. 로드 검증은 하지 않았음."),
        ("2026-09-17 [실측]", "", "3차 v4·v5 실패 실행 확인", "긴 ContextPack·GPU 메모리",
         "v4: 0/36스텝 메모리 부족; v5: 1/36스텝 CUDA 오류 [실측]."),
    ]

    paragraphs = {
        7: "목표 : 커머스 문의를 JSON으로 구조화하고 파인튜닝 경로의 RAG+Context Broker 대체 가능성을 검증.",
        8: "문제 정의 : " + DOMAIN_NOTE + ". 출력 형식과 정책·사실 근거를 함께 평가.",
        9: "사용 프레임워크 : PyTorch, Transformers, PEFT, TRL, Datasets. 버전 모름 [미확보].",
        18: "최종 선정 모델 : 파인튜닝 모델 미채택. Qwen/Qwen2.5-3B-Instruct + LoRA 실험 기록.",
    }
    return paragraphs, {3: candidates, 4: criteria, 5: parameters, 6: experiments, 7: history}


def trained_model_data(tables):
    components = source_table(tables, "구성 요소", 5)
    io_specs = source_table(tables, "입·출력 필드", 4)
    performance = source_table(tables, "항목", 6)
    management = [
        ("모델 개요", "기반 모델", "Qwen/Qwen2.5-3B-Instruct",
         "자기회귀 언어모델 기반 지도 파인튜닝. 순수 bf16 + LoRA [실측]."),
        ("모델 구조", "상세 규모", UNKNOWN,
         "정확한 파라미터 수·층 수·hidden dimension·학습 가능 파라미터 수 미확보."),
        ("적용 상태", "최종 판정", "파인튜닝 모델 미채택",
         "단독 경로와 후속 통합 실험 모두 최종 미채택."),
    ]

    for group, source, size, limits in source_table(tables, "구분", 4):
        management.append((
            "학습·평가 자료", group, f"{source}; {size}", limits,
        ))

    management.extend([
        ("데이터 관리", "전처리 표본과 학습 JSONL의 상세 대응", UNKNOWN,
         "전처리 표본 수를 임의 합산하여 학습 규모로 대체하지 않음."),
        ("데이터 관리", "datasets/", "Git 추적 대상에서 제외",
         "실제 고객 레코드·배송지 정보를 제출 문서에 수록하지 않음."),
    ])

    for library, purpose, version in source_table(tables, "라이브러리·구성", 3):
        management.append(("의존성", library, version, purpose))

    management.extend([
        ("실행 환경", "가상환경 경로", r"E:\dod28_ft\venv312",
         "폴더 이름으로 Python 실제 버전을 확정하지 않음."),
        ("모델 소재", "저장소 검색", "0건 [실측 2026-09-16]",
         "가중치·어댑터 설정·학습 설정 파일 검색 결과."),
        ("모델 소재", "GPU 서버", r"E:\dod28_ft [실측 2026-09-17]",
         "파일 존재·설정·SHA-256·기반 모델 스냅샷 확인. 로드 검증은 하지 않았음."),
        ("모델 소재", "체크포인트 폴더", "11개 [실측]; 9개에 어댑터 [실측]",
         "ckpt_stage1·ckpt_stage2·ckpt_stage3·ckpt_stage3_v2~v9."),
        ("실패 실행", "ckpt_stage3_v4", "CUDA 메모리 부족; 0/36스텝 [실측]",
         "학습 실패로 폴더 비어 있음."),
        ("실패 실행", "ckpt_stage3_v5", "CUDA 오류; 1/36스텝 [실측]",
         "학습 실패로 폴더 비어 있음."),
        ("평가 모델", "ckpt_stage2", "Proposed+FT 비교평가에 사용",
         "predict.py가 로딩하는 후속 단계 어댑터."),
    ])

    for item, value, status in source_table(tables, "모델 파일·구성 항목", 3):
        if item == "이전 단계 어댑터":
            # 긴 경로·해시·체크포인트를 나누되 숫자와 해시는 온전히 보존한다.
            prefix, tail = value.split(", SHA-256 ", 1)
            digest, checkpoints = tail.split(", 체크포인트 ", 1)
            management.extend([
                ("모델 파일", item, prefix, status),
                ("무결성", "이전 단계 SHA-256", digest, "확보 [실측]. 후속 단계 해시와 다름."),
                ("중간 저장", "이전 단계 체크포인트", checkpoints, "있음 [실측]."),
            ])
        else:
            management.append(("모델 파일", item, value, status))

    management.extend([
        ("저장", "종료 시 저장",
         "trainer.save_model(args.out); tokenizer.save_pretrained(args.out)",
         "모델과 토크나이저를 출력 경로에 저장."),
        ("저장", "중간 저장 설정", "40스텝 [실측]; 보관 제한 5개 [실측]",
         "train.py의 코드상 저장 방식."),
        ("학습 실행", "공개·후속 단계", "두 실행 모두 DONE_EXIT_0 [실측 2026-09-17]",
         "공개 단계 checkpoint-188 재개. 후속 단계는 공개 단계 어댑터 계승."),
        ("로딩", "기반 파일", "기반 모델·토크나이저의 로컬 경로 해석 또는 파일 확보",
         "33c9f8c^ 코드에 근거한 절차."),
        ("로딩", "기반 모델", "bf16로 CPU 로딩 후 CUDA 이동",
         "현재 어댑터 정상 로딩을 검증했다는 의미가 아님."),
        ("로딩", "어댑터 결합", "PeftModel.from_pretrained",
         "후속 단계 어댑터를 결합하고 평가 모드로 전환."),
        ("추론", "입력 구성", "시스템 지시문·고객 문의에 채팅 템플릿 적용",
         "비교평가 단독 경로에 RAG·ContextPack 미주입."),
        ("추론", "생성 방식", "샘플링을 사용하지 않는 결정론적 생성",
         "비교평가 출력 스키마에 유효한 정책 인용 전달 필드가 없음."),
        ("후처리", "JSON 추출·기록", "raw_output, prediction, parse_ok를 JSONL로 저장",
         "정규표현식과 json.loads 사용. 실패 시 prediction=null."),
        ("출력 검증", "스키마 한계", "엄격한 enum·업무 스키마 검증 없음",
         "JSON 파싱 성공을 정책 근거 정합성으로 간주하지 않음."),
        ("후속 통합", "근거 대조", "초안 복사·입력 잘림 문제 확인",
         "후속 실험 완료 후에도 미채택. 손실·파싱 개선만으로 채택하지 않음."),
        ("재검토 조건", "평가 범위",
         "ContextPack 전달·입력 잘림·정책 인용·초안 충돌·기권·실행 차단",
         "요건 충족 전 서비스 모델로 등록하지 않음."),
        ("인도", "파일 이전·무결성", "파일을 옮기고 SHA-256 대조",
         "실제 어댑터 로딩·추론 검증을 해야 실행 재현 완료로 기록."),
    ])

    history = [
        (UNKNOWN, "", "Qwen 기반 LoRA 공개 학습", "커머스 실험 모델",
         "커머스 학습 기록으로 보존."),
        (UNKNOWN, "", "golden judge 통과 응답으로 후속 학습", "공개 단계 어댑터",
         "실제 알파 고객 로그로 학습한 모델과 구분."),
        (UNKNOWN, "", "단독 추론의 형식·정책 근거 격차 확인", "파인튜닝 단독 경로",
         "제품 경로 미채택."),
        (UNKNOWN, "", "ContextPack 후속 실험·평가 보완", "입력 길이·초안 복사",
         "손실·파싱 개선만으로 채택하지 않음."),
        (UNKNOWN, "", "후속 통합까지 포함하여 최종 미채택", "모델 적용 전략",
         "동일 처방의 확대 중단, 실험 근거 보존."),
        ("2026-09-08 [실측]", "", "커머스에서 여행 CS로 전환", "도메인·모델 평가",
         "커머스 모델의 명칭·성능을 여행 모델로 승계하지 않음."),
        (UNKNOWN, "", "커밋 33c9f8c에서 학습 스크립트 79개 [실측] 삭제", "학습·추론 소스",
         "소스 확인 기준은 33c9f8c^."),
        ("2026-09-16 [실측]", "", "저장소 모델 산출물 검색 0건 [실측]", "모델 인도 상태",
         "GPU 서버의 파일 존재와 구분."),
        ("2026-09-17 [실측]", "", "GPU 서버 파일·설정·SHA-256·기반 스냅샷 확인", "ckpt_stage2",
         "로드 검증은 하지 않았음. 실행 재현 완료로 기록하지 않음."),
    ]

    paragraphs = {
        7: "모델 유형 : Qwen/Qwen2.5-3B-Instruct 기반 지도 파인튜닝. 순수 bf16 + LoRA [실측].",
        8: "담당 태스크 : 커머스 문의 구조화·고객 답변 생성. " + DOMAIN_NOTE,
        9: "프레임워크 / 버전 : PyTorch, Transformers, PEFT, TRL, Datasets / 모름 [미확보].",
        19: r"모델 파일 : GPU 서버 E:\dod28_ft\ckpt_stage2\adapter_model.safetensors [실측 2026-09-17]. 로드 미검증.",
    }
    return paragraphs, {3: components, 4: io_specs, 5: performance, 6: management, 7: history}


def collect_text(doc):
    """본문·중첩 표·머리글·바닥글의 모든 문단 및 셀 텍스트를 수집한다."""
    texts = []
    cell_texts = []

    def visit_table(table):
        for cell in unique_cells(table):
            cell_texts.append(cell.text)
            texts.extend(paragraph.text for paragraph in cell.paragraphs)
            for nested in cell.tables:
                visit_table(nested)

    def visit_container(container):
        texts.extend(paragraph.text for paragraph in container.paragraphs)
        for table in container.tables:
            visit_table(table)

    visit_container(doc)
    seen_parts = set()
    for section in doc.sections:
        for container in (
            section.header, section.first_page_header, section.even_page_header,
            section.footer, section.first_page_footer, section.even_page_footer,
        ):
            if container.part.partname not in seen_parts:
                seen_parts.add(container.part.partname)
                visit_container(container)
    return "\n".join(texts + cell_texts), cell_texts


def check_saved(path, expected, allowed_dates):
    """저장한 파일을 다시 열어 검사한다."""
    doc = Document(path)
    failures = []
    all_text, cells = collect_text(doc)

    print(f"\n검사: {path.name}")
    print(f"  표 개수: {len(doc.tables)}")
    print(f"  문단 개수: {len(doc.paragraphs)}")

    for word in FORBIDDEN:
        if word in all_text:
            label = "공개 금지 서버 식별자" if word == PRIVATE_SERVER else word
            failures.append(f"금지 부분 문자열 발견: {label}")

    if any(text == "000" for text in cells):
        failures.append("셀 텍스트가 정확히 000인 칸이 있음")

    if len(doc.tables) != 8:
        failures.append("표 개수가 원본과 다름")
    if len(doc.paragraphs) != expected["paragraph_count"]:
        failures.append("표 밖 문단 개수가 원본과 다름")

    if len(doc.tables) >= 2:
        blank_labels = {"제출 일자", "깃허브 경로", "작성 팀원"}
        found = set()
        for row in doc.tables[1].rows:
            key = normalized_key(row.cells[0].text)
            for label in blank_labels:
                if key == normalized_key(label):
                    found.add(label)
                    if row.cells[1].text != "":
                        failures.append(f"T1 {label} 칸이 빈 문자열이 아님")
            if key == normalized_key("산출물 단계"):
                if row.cells[1].text != "데이터 전처리":
                    failures.append("T1 산출물 단계 불일치")
        if found != blank_labels:
            failures.append("T1 필수 빈 칸의 라벨이 누락됨")

        cover = list(unique_cells(doc.tables[0]))
        if len(cover) != 1 or cover[0].text != COVER:
            failures.append("T0 표지 문구 불일치")

    if len(doc.tables) == 8:
        if doc.tables[2]._tbl.xml != expected["title_xml"]:
            failures.append("T2 제목 표가 변경됨")

        for table_index, rows in expected["rows"].items():
            table = doc.tables[table_index]
            if len(table.rows) != len(rows) + 1:
                failures.append(f"T{table_index} 행 개수 불일치")
                continue
            for i, values in enumerate(rows, start=1):
                actual = tuple(cell.text for cell in table.rows[i].cells)
                target = tuple(public_text(value) for value in values)
                if actual != target:
                    failures.append(f"T{table_index} 데이터 행 {i} 내용 불일치")
                if any(len(text) > 120 for text in actual):
                    failures.append(f"T{table_index} 데이터 행 {i}에 120자 초과 셀")

        for i, row in enumerate(doc.tables[7].rows[1:], start=1):
            if row.cells[1].text != "":
                failures.append(f"T7 데이터 행 {i} 변경자가 빈 문자열이 아님")
            for date in re.findall(r"\b20\d{2}-\d{2}-\d{2}\b", row.cells[0].text):
                if date not in allowed_dates:
                    failures.append(f"T7 본문에 없는 변경일: {date}")

    for index, value in expected["paragraphs"].items():
        if index >= len(doc.paragraphs) or doc.paragraphs[index].text != value:
            failures.append(f"P{index} 예시 문단 교체 불일치")

    if not any(DOMAIN_NOTE in p.text for p in doc.paragraphs[7:10]):
        failures.append("P7~P9에 커머스 자료·여행 전환 고지가 없음")

    if failures:
        for failure in failures:
            print("  [실패]", failure)
    else:
        print("  [통과] 금지 문구·예시 잔존·표지 빈 칸·변경자·행 수·문단 수 검사")
    return failures


def main():
    if not FORMS.is_dir() or not OUT_DIR.is_dir():
        raise FileNotFoundError("스크립트를 program/산출물양식/_build/에 저장해야 한다.")

    sources = [
        public_text(find_actual_name(OUT_DIR, name).read_text(encoding="utf-8-sig"))
        for name in SOURCE_NAMES
    ]
    parsed = [parse_md_tables(text) for text in sources]
    allowed_dates = set(re.findall(
        r"\b20\d{2}-\d{2}-\d{2}\b", "\n".join(sources)
    ))

    builders = (preprocessing_data, training_data, trained_model_data)
    prepared = []

    # 모든 양식과 내용 검증을 마친 뒤에 저장을 시작한다.
    for index, builder in enumerate(builders):
        source_path = find_actual_name(FORMS, TEMPLATE_NAMES[index])
        output_path = OUT_DIR / OUTPUT_NAMES[index]
        if source_path.resolve() == output_path.resolve():
            raise ValueError("양식 원본을 덮어쓸 수 없다.")

        doc = Document(source_path)
        validate_template(doc, index)
        title_xml = doc.tables[2]._tbl.xml
        original_paragraph_count = len(doc.paragraphs)

        paragraphs, table_rows = builder(parsed[index])
        fill_cover(doc)

        for paragraph_index, text in paragraphs.items():
            setp(doc.paragraphs[paragraph_index], text)

        for table_index, rows in table_rows.items():
            fill(doc.tables[table_index], rows)

        if len(doc.paragraphs) != original_paragraph_count:
            raise ValueError("표 밖 문단 개수가 변경되었다.")
        if doc.tables[2]._tbl.xml != title_xml:
            raise ValueError("T2 제목 표가 변경되었다.")

        expected = {
            "paragraph_count": original_paragraph_count,
            "title_xml": title_xml,
            "paragraphs": paragraphs,
            "rows": table_rows,
        }
        prepared.append((doc, output_path, expected))

    saved = []
    save_failures = []
    for doc, path, expected in prepared:
        try:
            doc.save(path)
            saved.append((path, expected))
            print("저장:", path)
        except (OSError, ValueError) as exc:
            save_failures.append(path.name)
            print(f"[저장 실패] {path.name}: {type(exc).__name__}")

    all_failures = list(save_failures)
    for path, expected in saved:
        try:
            failures = check_saved(path, expected, allowed_dates)
            all_failures.extend(f"{path.name}: {item}" for item in failures)
        except Exception as exc:
            all_failures.append(path.name)
            print(f"[검사 실패] {path.name}: {type(exc).__name__}")

    print()
    if all_failures or len(saved) != 3:
        print("자기 점검 실패. 위 결과를 확인하십시오.")
        sys.exit(1)
    print("자기 점검 통과: 저장한 세 파일을 다시 열어 확인했습니다.")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"[빌드 실패] {type(exc).__name__}: {public_text(str(exc))}")
        sys.exit(1)
