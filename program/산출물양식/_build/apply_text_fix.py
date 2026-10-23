# -*- coding: utf-8 -*-
"""제출본 docx·xlsx 의 문장을 JSON 교체 목록대로 바꾼다 (확인 필요 정리용).

python apply_text_fix.py <제출본 파일명> <json>
JSON: {"replacements":[{"old":"문단·셀 안에 있는 원문 조각","new":"바꿀 문장"}]}
문단(표 칸 포함)마다 old 가 들어 있으면 바꾼다. 어디에도 없으면 건너뛴다.
w2/제출본 과 원래 위치 둘 다 갱신한다.
"""
import json
import os
import re
import shutil
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
FORMS = os.path.dirname(HERE)
SUB = os.path.join(FORMS, "w2", "제출본")


def load(p):
    raw = open(p, encoding="utf-8").read().strip()
    raw = re.sub(r"^```[a-z]*\n|\n```$", "", raw)
    return json.loads(raw[raw.find("{"):raw.rfind("}") + 1])["replacements"]


def main():
    name, jp = sys.argv[1], sys.argv[2]
    f = os.path.join(SUB, name)
    reps = load(jp)
    hit = [0] * len(reps)
    if f.endswith(".docx"):
        from docx import Document
        d = Document(f)
        paras = list(d.paragraphs) + [p for tb in d.tables for r in tb.rows for c in r.cells for p in c.paragraphs]
        seen = set()
        for p in paras:
            if id(p._p) in seen:
                continue
            seen.add(id(p._p))
            t = p.text
            n = t
            for i, r in enumerate(reps):
                if r["old"] and r["old"] in n:
                    n = n.replace(r["old"], r["new"])
                    hit[i] += 1
            if n != t and p.runs:
                p.runs[0].text = n
                for x in p.runs[1:]:
                    x.text = ""
        d.save(f)
        allt = [p.text for p in d.paragraphs] + [c.text for tb in d.tables for r in tb.rows for c in r.cells]
    else:
        import openpyxl
        wb = openpyxl.load_workbook(f)
        ws = wb.worksheets[0]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    n = c.value
                    for i, r in enumerate(reps):
                        if r["old"] and r["old"] in n:
                            n = n.replace(r["old"], r["new"])
                            hit[i] += 1
                    c.value = n
        wb.save(f)
        allt = [str(c.value) for r in ws.iter_rows() for c in r if c.value]
    shutil.copy(f, os.path.join(FORMS, name))
    left = sum(t.count("확인 필요") for t in allt)
    print("교체 %d개 중 적용 %d | 남은 「확인 필요」 %d | 반영: w2/제출본 + 원래 위치" % (len(reps), sum(1 for h in hit if h), left))


if __name__ == "__main__":
    main()
