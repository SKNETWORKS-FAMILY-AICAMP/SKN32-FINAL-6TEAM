# -*- coding: utf-8 -*-
"""발표 자료 기계 검사 — 글자 넘침 · 작은 글자 · 금지어 · 발표자 노트 · 장 밖으로 나간 도형.

    python pptx_check.py <덱.pptx> [렌더 폴더]

★문체 지침(`program/산출물양식/_문체/글쓰기_지침.md` §4)의 기계 검사 몫이다. 이 검사를 통과해도
  읽어 보는 검수(제목만 이어 읽기 · 숫자 근거)는 따로 해야 한다.
★넘침은 PowerPoint 가 실제로 조판한 글 높이(TextFrame2.TextRange.BoundHeight)로 잰다 — 글자 수로 어림하지 않는다.
  PowerPoint 가 없는 곳(샌드박스 등)에서는 COM 검사를 건너뛰고 **건너뛰었다고 적는다**(통과로 치지 않는다).
결과: 화면에 요약, 덱 옆에 `<덱 이름>.check.json`. 실패가 하나라도 있으면 exit 1.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

sys.stdout.reconfigure(encoding="utf-8")

MIN_PT = 12.0        # 본문 최소 — 지침 §3-4
MIN_NOTE_PT = 10.0   # 출처·쪽번호·각주 최소
PT_PER_EMU = 1 / 12700

FORBIDDEN = [
    (r"사실표", "내부 작업 메모 이름"),
    (r"검수_클로드|지시문|_코덱스", "내부 작업 파일 이름"),
    (r"[A-Za-z]:\\|/Users/|\\Users\\|AppData", "로컬 경로"),
    (r"(?<![\d.])(?:\d{1,3}\.){3}\d{1,3}(?![\d.])", "IP 모양 숫자"),
    (r"webhooks/|discord(?:app)?\.com/api", "웹훅 주소"),
    (r"\w@\d", "계정@주소 모양"),
    (r"★|☆", "별표 기호"),
    (r"A-COP(?!ilot)", "옛 제품 이름(A-COP) — 제품은 triPilot, 팀은 A-COPilot"),
]
# 영문·숫자 바로 뒤 조사를 띄어 쓴 곳(「Case 를」) — 제출물은 붙인다(「Case를」).
JOSA_SPACE = re.compile(r"[A-Za-z0-9)\]] (?:을|를|이|가|은|는|에|에서|에게|의|로|으로|와|과|도|만)(?=[\s.,·)]|$)")


def runs_of(shape):
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            for r in para.runs:
                yield r
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for para in cell.text_frame.paragraphs:
                    for r in para.runs:
                        yield r


def static_checks(path: Path):
    prs = Presentation(str(path))
    W, H = prs.slide_width, prs.slide_height
    issues = []
    for n, slide in enumerate(prs.slides, start=1):
        texts = []
        for sh in slide.shapes:
            if sh.left is not None and (sh.left < -Emu(9525) or sh.top < -Emu(9525)
                                        or sh.left + sh.width > W + Emu(9525) or sh.top + sh.height > H + Emu(9525)):
                issues.append({"slide": n, "kind": "장 밖", "level": "fail",
                               "detail": f"{sh.name} 이 장 경계를 넘는다"})
            for r in runs_of(sh):
                t = r.text.strip()
                if not t:
                    continue
                texts.append(t)
                size = r.font.size.pt if r.font.size else None
                if t.isdigit() and len(t) <= 3:   # 쪽번호
                    continue
                if size is not None and size < MIN_NOTE_PT:
                    issues.append({"slide": n, "kind": "작은 글자", "level": "fail",
                                   "detail": f"{size}pt < {MIN_NOTE_PT}pt: {t[:30]}"})
                elif size is not None and size < MIN_PT and len(t) > 12:
                    issues.append({"slide": n, "kind": "작은 글자", "level": "warn",
                                   "detail": f"{size}pt 본문(12pt 미만): {t[:30]}"})
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        all_text = "\n".join(texts + [notes])
        for pat, why in FORBIDDEN:
            for m in re.finditer(pat, all_text):
                issues.append({"slide": n, "kind": "금지어", "level": "fail", "detail": f"{why}: 「{m.group(0)}」"})
        for m in JOSA_SPACE.finditer("\n".join(texts)):
            issues.append({"slide": n, "kind": "조사 띄어쓰기", "level": "warn", "detail": f"「{m.group(0)}」"})
        if n > 1 and len(notes) < 40:
            issues.append({"slide": n, "kind": "발표자 노트", "level": "fail",
                           "detail": f"노트가 없거나 너무 짧다({len(notes)}자)"})
    return len(prs.slides), issues


def com_checks(path: Path, render_dir: Path | None):
    try:
        import win32com.client  # noqa: PLC0415
        app = win32com.client.Dispatch("PowerPoint.Application")
    except Exception as e:  # noqa: BLE001 — PowerPoint 가 없는 환경
        return None, f"PowerPoint COM 을 못 열었다({type(e).__name__}) — 넘침 검사·렌더를 건너뛰었다"
    issues = []
    pres = app.Presentations.Open(str(path.resolve()), True, False, False)  # ReadOnly, Untitled, WithWindow
    try:
        sh_h = pres.PageSetup.SlideHeight
        for n in range(1, pres.Slides.Count + 1):
            slide = pres.Slides.Item(n)
            for i in range(1, slide.Shapes.Count + 1):
                sh = slide.Shapes.Item(i)
                if sh.HasTable:
                    bottom = sh.Top + sh.Height
                    if bottom > sh_h + 1:
                        issues.append({"slide": n, "kind": "표 넘침", "level": "fail",
                                       "detail": f"표 아래 끝 {bottom:.0f}pt > 장 높이 {sh_h:.0f}pt"})
                    continue
                if not sh.HasTextFrame or not sh.TextFrame2.HasText:
                    continue
                tf = sh.TextFrame2
                room = sh.Height - tf.MarginTop - tf.MarginBottom
                need = tf.TextRange.BoundHeight
                if need > room + 2:
                    issues.append({"slide": n, "kind": "글자 넘침", "level": "fail",
                                   "detail": f"{need:.0f}pt 필요 / {room:.0f}pt 자리: {tf.TextRange.Text[:30]!r}"})
            if render_dir:
                render_dir.mkdir(parents=True, exist_ok=True)
                slide.Export(str((render_dir / f"slide{n:02d}.png").resolve()), "PNG", 1600, 900)
    finally:
        pres.Close()
    return issues, None


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    path = Path(sys.argv[1])
    render_dir = Path(sys.argv[2]) if len(sys.argv) > 2 else None
    count, issues = static_checks(path)
    com, skipped = com_checks(path, render_dir)
    issues += com or []
    report = {"deck": path.name, "slides": count, "com_skipped": skipped, "issues": issues,
              "fail": sum(i["level"] == "fail" for i in issues), "warn": sum(i["level"] == "warn" for i in issues)}
    out = path.with_suffix(".check.json")
    try:
        out.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        out = None
    print(f"{path.name}: {count}장 · 실패 {report['fail']} · 경고 {report['warn']}"
          + (f" · ★{skipped}" if skipped else "") + (f" · 기록 {out.name}" if out else ""))
    for i in issues:
        print(f"  [{i['level']}] {i['slide']:>2}장 {i['kind']}: {i['detail']}")
    sys.exit(1 if report["fail"] or skipped else 0)


if __name__ == "__main__":
    main()
