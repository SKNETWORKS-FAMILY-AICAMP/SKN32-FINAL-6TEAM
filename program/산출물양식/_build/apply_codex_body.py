# -*- coding: utf-8 -*-
"""코덱스 본문 결과(빌드 스크립트 전문)를 적용하고 바로 검사한다.

python apply_codex_body.py <req|wbs|data|db> <코덱스 출력 파일>

1) 출력에서 스크립트와 변경요약을 가른다  2) build_<문서>_v3.py 로 저장(v2 는 안 건드린다)
3) 실행해 산출물을 만든다  4) 문서별 검사 — 하나라도 실패하면 제출 전 백업으로 되돌린다.
"""
import os
import re
import shutil
import subprocess
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
FORMS = os.path.dirname(HERE)
BACKUP = os.path.join(FORMS, "_backup", "2026-09-10_1611_제출전")

DOCS = {
    "req": ("build_요구사항정의서", "[기획] 요구사항 정의서_2차_여행.xlsx", None),
    "wbs": ("build_WBS", "[기획] WBS_2차_여행.xlsx", None),
    "data": ("build_수집데이터보고서", "[데이터 수집 및 저장] 수집 데이터 보고서_2차_여행.docx",
             "[데이터 수집 및 저장] 수집 데이터 보고서.docx"),
    "db": ("build_DB설계서", "[데이터 수집 및 저장] 데이터베이스저장소 설계 문서_2차_여행.docx",
           "[데이터 수집 및 저장] 데이터베이스_저장소 설계 문서.docx"),
}
BANNED = "—–→⇒★☆✓✔✕✗▸…⇄①②③④⑤⑥−"


def xlsx_text(path):
    import openpyxl
    ws = openpyxl.load_workbook(path).worksheets[0]
    return ws, [[c.value for c in r] for r in ws.iter_rows()]


def check_req(new, old):
    ws, rows = xlsx_text(new)
    _, orows = xlsx_text(old)
    ids = lambda rs: {m for r in rs for v in r if isinstance(v, str) for m in re.findall(r"REQ-\d{3}", v)}
    n, o = ids(rows), ids(orows)
    lost = sorted(o - n)
    merged_g = [str(m) for m in ws.merged_cells.ranges if m.min_col == 7 and m.max_col == 7 and m.min_row >= 5]
    print("  REQ 이전 %d → 지금 %d | 사라진 ID %s | G열 병합 %s" % (len(o), len(n), lost or 0, merged_g or 0))
    g = [str(ws.cell(r, 7).value or "") for r in range(5, ws.max_row + 1) if ws.cell(r, 4).value]
    bad_first = [x[:12] for x in g if x and not re.match(r"^\[(M|S|C|W|미확보)\]", x)]
    print("  G열 첫 토큰이 [M/S/C/W/미확보] 아닌 칸: %d %s" % (len(bad_first), bad_first[:3]))
    return not lost and not merged_g


def check_wbs(new, old):
    _, rows = xlsx_text(new)
    _, orows = xlsx_text(old)
    tid = lambda rs: {int(str(r[0]).strip()) for r in rs if r and r[0] is not None and str(r[0]).strip().isdigit()}
    n, o = tid(rows), tid(orows)
    lost = sorted(o - n)
    # 같은 번호의 세부 업무가 바뀌었나 (번호 고정 검사)
    om = {int(str(r[0]).strip()): str(r[2])[:25] for r in orows if r and r[0] is not None and str(r[0]).strip().isdigit()}
    nm = {int(str(r[0]).strip()): str(r[2])[:25] for r in rows if r and r[0] is not None and str(r[0]).strip().isdigit()}
    moved = [k for k in om if k in nm and om[k][:10] != nm[k][:10]]
    print("  Task 이전 %d → 지금 %d | 사라진 ID %s | 세부 업무 앞머리가 바뀐 ID %d개 %s"
          % (len(o), len(n), lost or 0, len(moved), moved[:8]))
    return not lost


def check_docx(new, form):
    from docx import Document
    S, O = Document(os.path.join(FORMS, form)), Document(new)
    t = "\n".join(p.text for p in O.paragraphs) + "".join(c.text for tb in O.tables for r in tb.rows for c in r.cells)
    ex = sorted({(ti, ri) for ti, tb in enumerate(O.tables) for ri, r in enumerate(tb.rows)
                 for c in r.cells if "(예" in c.text or "예:" in c.text})
    # 양식 표의 머리글·질문 칸(0열)이 그대로 있나 — 머리글이 같은 표끼리 짝지어 비교
    # 머리글은 0행 0열만 비교한다 — 「라벨 | 값」 한 줄 표는 0행 1열이 값 칸이라 바뀌는 게 정상이다
    heads = lambda tb: tb.rows[0].cells[0].text.strip()
    lost = []
    used = set()
    for ts in S.tables[2:]:
        hs = heads(ts)
        cand = [i for i, to in enumerate(O.tables) if i not in used and heads(to) == hs]
        if not cand:
            lost.append(("머리글 사라짐", hs[:20]))
            continue
        to = O.tables[cand[0]]
        used.add(cand[0])
        two_col = len(set(c.text for c in ts.rows[0].cells)) == 2
        if two_col:   # 「라벨 | 값」 표는 0열 라벨이 같아야 한다
            for ri, r in enumerate(ts.rows):
                a = r.cells[0].text.strip()
                b = to.rows[ri].cells[0].text.strip() if ri < len(to.rows) else ""
                if a and not a.startswith("기타") and a != b:
                    lost.append((a[:16], b[:16]))
    ban = [ch for ch in BANNED if ch in t]
    print("  예시 잔존 %s | 양식 칸 바뀜 %s | 표지 00기 %s | 금지문자 %s"
          % (ex or 0, lost or 0, "00기" in t, ban or 0))
    return not ex and not lost and "00기" not in t


def main():
    key, out_file = sys.argv[1], sys.argv[2]
    base, product, form = DOCS[key]
    raw = open(out_file, encoding="utf-8").read()
    raw = re.sub(r"^```[a-z]*\n|\n```\s*$", "", raw.strip())
    script, _, summary = raw.partition("=====변경요약=====")
    assert "import" in script and len(script) > 2000, "스크립트가 아니다"
    dst = os.path.join(HERE, base + "_v3.py")
    open(dst, "w", encoding="utf-8", newline="\n").write(script.rstrip() + "\n")
    open(os.path.join(HERE, base + "_v3_변경요약.md"), "w", encoding="utf-8").write(summary.strip() + "\n")
    print("저장:", os.path.basename(dst), len(script), "자")
    r = subprocess.run([sys.executable, dst], cwd=HERE, capture_output=True, text=True,
                       encoding="utf-8", errors="replace", env=dict(os.environ, PYTHONIOENCODING="utf-8"))
    print("  실행 종료코드", r.returncode, "|", (r.stdout + r.stderr).strip()[-400:])
    new = os.path.join(FORMS, product)
    old = os.path.join(BACKUP, product)
    ok = r.returncode == 0
    if ok:
        ok = {"req": lambda: check_req(new, old), "wbs": lambda: check_wbs(new, old),
              "data": lambda: check_docx(new, form), "db": lambda: check_docx(new, form)}[key]()
    if not ok:
        shutil.copy(old, new)
        print("  ✗ 검사 실패 — 제출 전 백업으로 되돌렸다:", product)
    else:
        print("  통과:", product)


if __name__ == "__main__":
    main()
