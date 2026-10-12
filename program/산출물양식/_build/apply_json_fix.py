# -*- coding: utf-8 -*-
"""에이전트·코덱스가 낸 JSON 수정 목록을 적용한다 (결정 14·15 반영용).

python apply_json_fix.py <data|db|wbs|plan|req> <json 파일 또는 -(표준입력)>
- data: build_수집데이터보고서_v2.py 교체 → 실행 → 양식 칸·예시 검사
- db  : build_DB설계서_v3.py 교체 → 실행 → 같은 검사
- wbs : xlsx 셀 교체(cells) + 새 행(new_rows) → Task ID 보존 검사
- plan·req: apply_d1415.py 로 넘긴다
정확히 1회가 아닌 교체는 건너뛴다. 실패하면 원래대로 되돌린다. 통과하면 w2/제출본 에 복사.
"""
import copy
import json
import os
import re
import shutil
import subprocess
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
FORMS = os.path.dirname(HERE)
SUBMIT = os.path.join(FORMS, "w2", "제출본")
sys.path.insert(0, HERE)
import apply_codex_body as A  # noqa: E402

CFG = {
    "data": ("build_수집데이터보고서_v2.py", "[데이터 수집 및 저장] 수집 데이터 보고서_2차_여행.docx",
             "[데이터 수집 및 저장] 수집 데이터 보고서.docx"),
    "db": ("build_DB설계서_v3.py", "[데이터 수집 및 저장] 데이터베이스저장소 설계 문서_2차_여행.docx",
           "[데이터 수집 및 저장] 데이터베이스_저장소 설계 문서.docx"),
}
WBS = "[기획] WBS_2차_여행.xlsx"


def load_json(src):
    raw = sys.stdin.read() if src == "-" else open(src, encoding="utf-8").read()
    raw = re.sub(r"^```[a-z]*\n|\n```$", "", raw.strip())
    return json.loads(raw[raw.find("{"):raw.rfind("}") + 1])


def publish(prod):
    os.makedirs(SUBMIT, exist_ok=True)
    shutil.copy(os.path.join(FORMS, prod), os.path.join(SUBMIT, prod))
    print("  통과 — w2/제출본 에 복사:", prod)


def fix_script(key, j):
    script, prod, form = CFG[key]
    sp = os.path.join(HERE, script)
    t = open(sp, encoding="utf-8").read()
    reps = j.get("replacements", [])
    ok_reps = [r for r in reps if t.count(r["old"]) == 1]
    print("교체", len(reps), "| 적용", len(ok_reps), "| 건너뜀", len(reps) - len(ok_reps))
    new = t
    for r in ok_reps:
        new = new.replace(r["old"], r["new"])
    bak, pbak = sp + ".before_fix", os.path.join(FORMS, prod) + ".before_fix"
    shutil.copy(sp, bak)
    shutil.copy(os.path.join(FORMS, prod), pbak)
    open(sp, "w", encoding="utf-8", newline="\n").write(new)
    r = subprocess.run([sys.executable, sp], cwd=HERE, capture_output=True, text=True, encoding="utf-8",
                       errors="replace", env=dict(os.environ, PYTHONIOENCODING="utf-8"))
    print("  실행", r.returncode, "|", (r.stdout + r.stderr).strip()[-250:])
    ok = r.returncode == 0 and A.check_docx(os.path.join(FORMS, prod), form)
    if not ok:
        shutil.copy(bak, sp)
        shutil.copy(pbak, os.path.join(FORMS, prod))
        print("  실패 — 되돌렸다")
        return
    for x in ok_reps:
        print("  -", x.get("why", "")[:80])
    publish(prod)


def fix_wbs(j):
    import openpyxl
    p = os.path.join(FORMS, WBS)
    bak = p + ".before_fix"
    shutil.copy(p, bak)
    wb = openpyxl.load_workbook(p)
    ws = wb.worksheets[0]
    col = {c: i + 1 for i, c in enumerate("ABCDEFGH")}
    done = 0
    for c in j.get("cells", []):
        cell = ws.cell(int(c["row"]), col[c.get("col", "C")])
        v = str(cell.value or "")
        if v.count(c["old"]) == 1:
            cell.value = v.replace(c["old"], c["new"])
            done += 1
    last = max(r for r in range(1, ws.max_row + 1)
               if ws.cell(r, 1).value is not None and str(ws.cell(r, 1).value).strip().isdigit())
    ids = {int(str(ws.cell(r, 1).value).strip()) for r in range(2, last + 1)
           if ws.cell(r, 1).value is not None and str(ws.cell(r, 1).value).strip().isdigit()}
    added = 0
    for nr in j.get("new_rows", []):
        if int(nr["id"]) in ids:
            continue
        last += 1
        if ws.cell(last, 2).value not in (None, ""):
            print("  새 행 자리가 비어 있지 않아 멈춘다:", last)
            break
        vals = [nr["id"], nr.get("B"), nr.get("C"), nr.get("D"), nr.get("E", "예정"), nr.get("F"), nr.get("G"), nr.get("H", "보통")]
        for ci in range(1, 9):
            s, d = ws.cell(last - 1, ci), ws.cell(last, ci)
            d.value = vals[ci - 1]
            if s.has_style:
                d.font, d.border, d.fill = copy.copy(s.font), copy.copy(s.border), copy.copy(s.fill)
                d.alignment, d.number_format = copy.copy(s.alignment), s.number_format
        added += 1
    wb.save(p)
    print("셀 교체", done, "/", len(j.get("cells", [])), "| 새 행", added)
    if not A.check_wbs(p, os.path.join(A.BACKUP, WBS)):
        shutil.copy(bak, p)
        print("  실패 — 되돌렸다")
        return
    publish(WBS)


def main():
    key, src = sys.argv[1], sys.argv[2]
    if key in ("plan", "req"):
        subprocess.run([sys.executable, os.path.join(HERE, "apply_d1415.py"), key, src], cwd=HERE)
        return
    j = load_json(src)
    fix_wbs(j) if key == "wbs" else fix_script(key, j)


if __name__ == "__main__":
    main()
