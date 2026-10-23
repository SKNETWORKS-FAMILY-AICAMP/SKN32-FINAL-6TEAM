OUT_DIR = "program/산출물양식/w2/표와시각자료/코덱스"

# 모든 사업 수치와 표시 설정의 단일 입력점.
# [결정]은 팀이 정한 가격·범위이며 관측값이 아니다.
# [v3 인용]은 과거 원본 재현이며 현재 판매 전망이 아니다.
P = {
    "환율": 1500,                       # [추정] v3 §7 예산 환율
    "VAT": 0.10,                        # [추정] v3의 손익 계산 가정
    "제휴수익": 0,                      # [추정] 계약 전 기준 손익에서 제외
    "개월_연": 12,
    "달러_가격단위": 1000,
    "표시_만원": 10000,
    "표시_억원": 100000000,
    "백분율": 100,

    # [외부] v3 §7 및 교차검토의 공개 가격 기록.
    # 무료 할당·대량 할인 제외. 요청 필드·교통 모드별 SKU 확인 필요.
    "API": {
        "상세": {
            "이름": "Place Details Enterprise",
            "USD": 20, "분모": 1000, "단위": "건",
        },
        "검색": {
            "이름": "Text Search Pro",
            "USD": 32, "분모": 1000, "단위": "건",
        },
        "행렬": {
            "이름": "Route Matrix Essentials",
            "USD": 5, "분모": 1000, "단위": "요소",
        },
        "리뷰": {
            "이름": "Details Enterprise + Atmosphere",
            "USD": 0.025, "분모": 1, "단위": "건",
        },
        "그라운딩": {
            "이름": "Search Grounding",
            "USD": 0.035, "분모": 1, "단위": "회",
        },
    },

    "상품": {
        "결정13": {
            "이름": "결정 13 예산",
            "가격": 49000,
            "가격등급": "결정",
            "가격상태": "VAT 포함 팀당 가격 결정; 판매 검증 전",
            "일": 7, "인원": 4, "도시": 1, "지역": "서울",
            "환불률": 0.03,
            "결제비율": 0.04,
            "결제설명": "비용과 수익구조 §4-A: VAT 포함 결제액의 4%",
            "기계_상한": 10000,
            "사용량": {}, "데이터_여유": 0, "LLM_정액": 0,
            "Case": None, "Case_단가": None,
            "지원_직접분": 15, "인계율": None, "인계_분": None,
            "지원_분단가": 500, "대응": 1000, "플랫폼": 0,
        },
        "기본": {
            "이름": "v3 원본 재현 · 현재 판매 전망 아님",
            "가격": 59000,
            "가격등급": "v3 인용",
            "가격상태": "v3 원본 재현 · 현재 판매 전망 아님",
            "일": 7, "인원": 4, "도시": 2,
            "지역": "지원 지역",
            "환불률": 0.03,              # [추정] v3·클로드 유지
            "결제비율": 0.034,           # [외부] 국내 일반카드 참고요율
            "결제설명": "국내 공개 참고요율 적용; 해외 계약요율 미확보",
            "사용량": {
                "상세": 120, "검색": 20, "행렬": 250,
                "리뷰": 10, "그라운딩": 10,
            },
            "데이터_여유": 665,          # [추정] v3의 재시도·가격 여유
            "LLM_정액": 1500,            # [추정] v3 예산
            "Case": None,
            "Case_단가": None,
            "지원_직접분": 6,
            "인계율": None,
            "인계_분": None,
            "지원_분단가": 500,
            "대응": 1000,
            "플랫폼": 0,                 # [추정] v3 미분리; 무료라는 뜻 아님
        },
        "단기": {
            "이름": "비교안 - 3일·2인",
            "가격": 39000,
            "가격등급": "추정",
            "가격상태": "비교안; 팀 결정 아님",
            "일": 3, "인원": 2, "도시": 1,
            "지역": "서울",
            "항목": 12,
            "환불률": 0.05,
            "결제비율": 0.045,
            "결제설명": "해외 결제 예산 가정; 계약요율 미확보",
            "사용량": {
                "상세": 96, "검색": 4, "행렬": 40,
                "리뷰": 0, "그라운딩": 0,
            },
            "데이터_여유": 0,
            "LLM_정액": None,
            "Case": 5,
            "Case_단가": 50,
            "지원_직접분": None,
            "인계율": 0.20,
            "인계_분": 10,
            "지원_분단가": 500,
            "대응": 1500,
            "플랫폼": 500,
        },
    },

    # [추정] 코덱스 본문: 중복 제거 적격 도달  /  연결  /  최초 결제.
    "인바운드_채널": [
        {"이름": "개인 에이전트 커뮤니티", "도달": 8000,
         "연결률": 0.50, "결제율": 0.10, "획득지출": 4000000},
        {"이름": "인바운드 콘텐츠·검색", "도달": 20000,
         "연결률": 0.20, "결제율": 0.10, "획득지출": 4800000},
        {"이름": "숙박·체험 소개", "도달": 10000,
         "연결률": 0.40, "결제율": 0.10, "획득지출": 5600000},
    ],
    "CAC_후보": [0, 8500, 12000, 15000, 20000, 25000, 35000, 40000],
    "판매량_후보": {
        "부진_기준대비": 0.50,
        "v3_참조": 8000,                 # 인바운드 기준 판매량으로 사용 금지
        "첫해_판매개월": 6,
    },
    "v3_CAC": 8500,
    "고정비": {
        "개발": 120000000,
        "사업운영": 30000000,
        "기본인프라": 18000000,
        "관리회계보안": 12000000,
    },
    "첫해_일회성": 20000000,
    "예비운전자금": 30000000,

    # 각 원문의 비용 악화 입력. 데이터는 사용량 견적이 아닌 정액 예산.
    # 두 상품에 동일한 충격을 가한 통제 실험이 아님을 T8에 명시한다.
    "비용악화": {
        "기본": {                       # [추정] v3 §7
            "환불률": 0.06, "결제비율": 0.034,
            "데이터": 12000, "LLM": 2500,
            "지원_분": 18, "지원_분단가": 500,
            "대응": 2000, "플랫폼": 0,
        },
        "단기": {                       # [추정] 코덱스 본문
            "환불률": 0.10, "결제비율": 0.055,
            "데이터": 6000, "LLM": 1000,
            "Case": 9, "인계율": 0.30, "인계_분": 10,
            "지원_분단가": 500,
            "대응": 3000, "플랫폼": 1000,
        },
    },

    "시장": {
        "연도": 2025,
        "시작월": 1,
        "종료월": 11,
        "방한방문": 17410000,             # [외부] 두 본문 인용 누적치
        "개별여행비율": 0.829,            # [외부] 기존 문서 인용
        "동행수_범위": [1, 4],            # [추정] 클로드의 상한 범위
        "서울방문률": None,              # [미확보]
        "에이전트사용률": None,           # [미확보]
        "SOM": None,                     # 산정하지 않음
    },
    "실증": {
        "현금상한": 5000000,
        "역할인원": 6,
        "주당시간": 20,
        "주수": 6,
        "완료여행_범위": [20, 30],
        "조건위반_목표": 0,
        "최종검토": "2026-10-26",
        "원가검토": "2026-09-24",
        "통합검토": "2026-10-12",
    },
    "기존측정": {
        "커머스_로그행": 216,
        "커머스_USD_Case": 0.0021651125,
        "이전_원_Case": 4.06,             # 역사적 원화 기록; 재환산하지 않음
        "여행차용_원_범위": [3, 37],       # [추정] 여행 실측 아님
    },
    "경쟁": {
        "TripIt_연_USD": 49,
        "토모_상담": 29000,
        "토모_패키지": 249000,
    },

    # 아래 수치는 출력 크기·표시 정밀도이며 사업 입력이 아니다.
    "표시": {
        "원_소수": 0,
        "경계_소수": 2,
        "비율_소수": 1,
        "시장_소수": 1,
        "DPI": 220,
        "폰트": 10,
        "문서폰트": 9,
        "문서_가로인치": 11.69,
        "문서_세로인치": 8.27,
        "문서_여백인치": 0.45,
        "F1_크기": (12, 7),
        "F1_표본": 1600,
        "F1_Y상한": 60000,
        "F1_X여유": 1.12,
        "F2_크기": (16, 7),
        "F3_크기": (11, 8),
    },
}

import sys
sys.dont_write_bytecode = True

from decimal import Decimal, ROUND_HALF_UP, ROUND_CEILING
from pathlib import Path

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import Polygon, Rectangle
from matplotlib.ticker import FuncFormatter

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


BUILD_REL = Path("program/산출물양식/_build")
SOURCES = {
    "코덱스": "program/산출물양식/w2/기획서/기획서_본문_코덱스.md",
    "클로드": "program/산출물양식/w2/기획서/기획서_본문_클로드.md",
    "교차검토": "program/산출물양식/w2/기획서/교차검토_클로드_1차.md",
    "v3": "team_branch/me/travel-agent-business-review_v3.md",
    "결정 11·13": "program/plan/A-COP_구현계획서_v11.md",
    "결정 13 예산": "program/plan/A-COP_비용과_수익구조.md",
    "확정 페르소나": "team_branch/jh/확정_시나리오.md",
}
COLORS = {"결정13": "#315E7D", "기본": "#315E7D", "단기": "#A46937"}
GRAY = "#9C9C9C"
LIGHT = "#EEEEEE"

ZERO = Decimal(0)
ONE = Decimal(1)


def D(value):
    return Decimal(str(value))


def total(values):
    return sum(values, ZERO)


def number(value, places=0):
    quantum = ONE.scaleb(-places)
    rounded = D(value).quantize(quantum, rounding=ROUND_HALF_UP)
    return f"{rounded:,.{places}f}"


def cell(value, unit="", grade="추정", places=0):
    if value is None:
        return "[미확보]"
    return f"{number(value, places)}{unit} [{grade}]"


def won(value, grade="추정", places=None):
    if places is None:
        places = P["표시"]["원_소수"]
    return cell(value, "원", grade, places)


def pct(value, grade="추정"):
    return cell(
        D(value) * D(P["백분율"]), "%",
        grade, P["표시"]["비율_소수"],
    )


def ceiling(value):
    return int(D(value).to_integral_value(rounding=ROUND_CEILING))


def label(key):
    product = P["상품"][key]
    return (
        f"{won(product['가격'], product['가격등급'])}"
        f" / {product['이름']}"
    )


def scope(key):
    product = P["상품"][key]
    grade = product["가격등급"]
    prefix = "최대 " if key == "결정13" else ""
    text = (
        f"{prefix}{cell(product['일'], '일', grade)} · "
        f"{cell(product['인원'], '인', grade)} · "
        f"{cell(product['도시'], '도시', grade)} · {product['지역']}"
    )
    if "항목" in product:
        text += f" · {cell(product['항목'], '항목')}"
    if key == "결정13":
        text += (
            "\n4인 팀 인당 " + won(D(product["가격"]) / D(product["인원"]), "결정")
            + "; 확정 페르소나: 대만인 친구 2명"
            + "\n2인 팀 인당 " + won(D(product["가격"]) / D(2), "결정")
        )
    elif key == "기본":
        text += "\n2도시는 결정 13의 서울 1도시와 다름"
    return text


def find_root():
    starts = [Path.cwd()]
    if "__file__" in globals():
        starts.append(Path(__file__).resolve().parent)
    checked = set()
    for start in starts:
        for candidate in (start, *start.parents):
            if candidate in checked:
                continue
            checked.add(candidate)
            if (candidate / BUILD_REL / "docx_tables.py").is_file():
                return candidate
    raise FileNotFoundError(
        "저장소 루트에서 실행하거나 스크립트를 저장소 내부에 두세요: "
        + str(BUILD_REL / "docx_tables.py")
    )


def api_unit(key):
    rate = P["API"][key]
    return D(rate["USD"]) / D(rate["분모"]) * D(P["환율"])


def api_recipe(key):
    product = P["상품"][key]
    if "기계_상한" in product:
        return (
            "외부 데이터·LLM 합산 상한 " + won(product["기계_상한"])
            + "; 세부 배분·실제 사용량 [미확보]"
        )
    pieces = []
    for sku, usage in product["사용량"].items():
        if usage:
            rate = P["API"][sku]
            pieces.append(
                f"{rate['이름']}: {cell(usage, rate['단위'])}"
                f" × {won(api_unit(sku))}/{rate['단위']}"
            )
    pieces.append(f"추가 여유 {won(product['데이터_여유'])}")
    return "\n".join(pieces)


def economy(key, stressed=False):
    product = P["상품"][key]
    refund = D(product["환불률"])
    pg_rate = D(product["결제비율"])
    data_parts = {
        sku: D(usage) * api_unit(sku)
        for sku, usage in product["사용량"].items()
    }
    data = total(data_parts.values()) + D(product["데이터_여유"])

    if product["LLM_정액"] is None:
        llm = D(product["Case"]) * D(product["Case_단가"])
    else:
        llm = D(product["LLM_정액"])

    if product["지원_직접분"] is None:
        minutes = (
            D(product["Case"]) * D(product["인계율"])
            * D(product["인계_분"])
        )
    else:
        minutes = D(product["지원_직접분"])

    support = minutes * D(product["지원_분단가"])
    response = D(product["대응"])
    platform = D(product["플랫폼"])

    if stressed:
        stress = P["비용악화"][key]
        refund = D(stress["환불률"])
        pg_rate = D(stress["결제비율"])
        # 악화 시나리오가 기준 비용보다 낮아지지 않게 한다.
        # 정액 예산을 넘는 환율·사용량 수정도 계산에 반영된다.
        data = max(data, D(stress["데이터"]))
        llm = max(llm, D(stress["LLM"]))
        if "지원_분" in stress:
            stress_minutes = D(stress["지원_분"])
        else:
            stress_minutes = (
                D(stress["Case"]) * D(stress["인계율"])
                * D(stress["인계_분"])
            )
        minutes = max(minutes, stress_minutes)
        support = max(support, minutes * D(stress["지원_분단가"]))
        response = max(response, D(stress["대응"]))
        platform = max(platform, D(stress["플랫폼"]))

    gross = D(product["가격"])
    before_refund = gross / (ONE + D(P["VAT"]))
    refund_loss = before_refund * refund
    net = before_refund - refund_loss + D(P["제휴수익"])
    costs = {
        "결제": gross * pg_rate,
        "데이터": data,
        "LLM": llm,
        "지원": support,
        "대응": response,
        "플랫폼": platform,
    }
    if "기계_상한" in product:
        costs = {
            "결제": gross * pg_rate,
            "기계(외부 데이터·LLM)": D(product["기계_상한"]),
            "지원": support,
            "대응": response,
            "플랫폼": platform,
        }
    variable = total(costs.values())
    return {
        "가격": gross,
        "VAT제외": before_refund,
        "환불차감": refund_loss,
        "환불률": refund,
        "결제비율": pg_rate,
        "순매출": net,
        "비용": costs,
        "데이터상세": data_parts,
        "지원분": minutes,
        "변동비": variable,
        "기여Ⅰ": net - variable,
    }


def context():
    sales = total(
        D(c["도달"]) * D(c["연결률"]) * D(c["결제율"])
        for c in P["인바운드_채널"]
    )
    spend = total(D(c["획득지출"]) for c in P["인바운드_채널"])
    if sales <= ZERO:
        raise ValueError("인바운드 기준 판매량은 양수여야 합니다.")
    fixed = total(D(v) for v in P["고정비"].values())
    if fixed <= ZERO:
        raise ValueError("손익분기 그림의 고정비는 양수여야 합니다.")
    return {
        "판매": sales,
        "획득지출": spend,
        "CAC": spend / sales,
        "고정비": fixed,
        "상품": {key: economy(key) for key in P["상품"]},
        "악화": {key: economy(key, True) for key in P["비용악화"]},
        "개별방문": (
            D(P["시장"]["방한방문"]) * D(P["시장"]["개별여행비율"])
        ),
    }


def break_even(cm2, fixed):
    if cm2 <= ZERO:
        return None, None
    annual = ceiling(fixed / cm2)
    monthly = ceiling(D(annual) / D(P["개월_연"]))
    return annual, monthly


def paired(ctx, getter):
    return "\n".join(
        f"{label(key)}: {getter(key, model)}"
        for key, model in ctx["상품"].items()
    )


def make_tables(ctx):
    specs = []
    notes = {}

    def add(tid, title, header, rows, note):
        anchor = f"{tid} {title}"
        # docx_tables.insert_tables의 실제 명세: 4요소 튜플.
        specs.append((anchor, f"{title} - 근거 등급 병기", header, rows))
        notes[anchor] = (
            note
            + "\n주 상품: 결정 13, 서울·최대 7일·4인, 팀당 49,000원 VAT 포함 [결정]. "
            "4인 인당 12,250원, 확정 페르소나 대만인 친구 2명은 인당 24,500원. "
            "v3 원본 재현 · 현재 판매 전망 아님: 2도시로 결정 13의 서울 1도시와 다름. "
            "비교안 - 3일·2인은 팀 결정 아님."
        )

    add(
        "T1", "PGBBB 요약",
        ["파트", "핵심 주장", "근거·상태", "다음 파트로", "상세 위치"],
        [
            ["Pain", "변경 발생 시 예약·이동·조건을 여행자가 재확인한다.",
             "인바운드 고객의 발생 빈도·손실 [미확보]",
             "재확인 부담과 일정 실패를 줄여야 한다.", "문제 정의"],
            ["Gap", "변경 감지·조건 보존·업체 인계를 연결한다.",
             "일정 생성·알림 상품은 존재 [외부]; 당사 우위 [미확보]",
             "같은 사건의 해결 결과로 비교한다.", "문제 정의 / 경쟁 비교"],
            ["Buyer", "여행 대표자를 우선 지불자 후보로 검증한다.",
             "실제 구매자·사업자 계약·인바운드 결제율 [미확보]",
             "가격·접점·채널별 구매를 검증한다.", "시장조사 및 BM 분석"],
            ["Build", "사실 조회와 코드 검증으로 일정 조정·인계·안내를 연결한다.",
             "공통 기반·감시 관측 표는 원문 확인 [실측]; "
             "Trip 저장·경로 조회·실행 게이트는 원문 기준 미구현",
             "호출량·지원 시간·실패 처리량을 계측한다.",
             "시스템 / 모델링 / 데이터"],
            ["Business", "제공 원가와 획득비를 이용권 순매출에서 회수한다.",
             paired(ctx, lambda k, m: "기여이익Ⅰ " + won(m["기여Ⅰ"])),
             "실제 CAC와 고정비 회수 가능한 판매량을 검증한다.",
             "결정 13 및 비교 시나리오 / CAC 손익분기 / 하방·자금"],
        ],
        "Buyer·Business의 인바운드 근거가 가장 약하다. "
        "경쟁사 기능의 부재나 당사의 성능 우위가 입증된 것은 아니다. "
        "근거: 두 본문 프로젝트 주제, 교차검토 §3.",
    )

    pilot = P["실증"]
    effort = (
        D(pilot["역할인원"]) * D(pilot["주당시간"]) * D(pilot["주수"])
    )
    pilot_range = " ~ ".join(
        cell(v, "건") for v in pilot["완료여행_범위"]
    )
    add(
        "T2", "승인 요청",
        ["승인 대상", "범위", "투입", "기대 결과", "다음 단계 조건", "미충족 시"],
        [
            ["결정 13 실증 준비", label("결정13") + "\n" + scope("결정13"),
             "가격·도시·인원은 결정; 원가·수용성은 검증 전",
             "결정 상품과 구현 범위 일치",
             "지원 상한과 결정 11의 언어 경로 구현·검증",
             "실증 범위 조정"],
            ["제한 실증 개발", "서울 일정 검증·변경 감지·조정·인계; 링크·알림은 고객 선택 언어",
             f"현금 상한 {won(pilot['현금상한'])}\n"
             f"역할 {cell(pilot['역할인원'], '명', '실측')}\n"
             f"공수 {cell(effort, '시간')}\n"
             f"최종 검토 {pilot['최종검토']} [실측]",
             "구현 증거·여행 원가·고객 연결 가능성",
             f"필수 조건 위반 {cell(pilot['조건위반_목표'], '건')} 목표; "
             "언어 경로·데이터 권한 확인",
             "범위 축소 또는 유료 실증 보류"],
            ["유료 현장 실증", "선택한 동일 상품의 결제·사용·환불 측정",
             f"완료 여행 {pilot_range}; 위 현금 한도 내 단계 집행",
             "구매와 사용 결과 연결; CAC·원가 측정",
             "기여이익Ⅱ 양수·고객 행동 개선·지원 가능 시간 확인",
             "가격·채널·상품 재설계"],
            ["상용 확장", "상시 운영·고객 획득 확대",
             "이번 승인 대상 제외; 추가 자금 [미확보]",
             "반복 가능한 판매 구조",
             "손익분기 획득량과 운영 능력 입증",
             "채용·광고 확대 보류"],
        ],
        "상품은 결정 13으로 확정했다. 결정 11: 링크·알림은 우리가 고객 선택 언어로 "
        "번역하며 기본은 한국어·영어; API는 한국어로 제공하고 에이전트가 번역한다. "
        "역할 인원은 문서상 계획이며 "
        "실제 실명·가용시간은 미확보다. 근거: 코덱스 프로젝트 주제·R&R, "
        "교차검토 §3. 날짜의 [실측]은 일정 기준선 기록을 뜻한다.",
    )

    market = P["시장"]
    low_group = min(market["동행수_범위"])
    high_group = max(market["동행수_범위"])
    low_trips = ctx["개별방문"] / D(high_group)
    high_trips = ctx["개별방문"] / D(low_group)
    market_money = "\n".join(
        f"{label(key)}: "
        f"{cell(low_trips * D(product['가격']) / D(P['표시_억원']), '억원', places=P['표시']['시장_소수'])}"
        f" ~ "
        f"{cell(high_trips * D(product['가격']) / D(P['표시_억원']), '억원', places=P['표시']['시장_소수'])}"
        for key, product in P["상품"].items()
    )
    period = (
        f"{market['연도']}년 {market['시작월']}~{market['종료월']}월 [외부]"
    )
    add(
        "T3", "시장 규모",
        ["TAM·SAM·SOM", "대상", "산식", "규모", "등급", "한계"],
        [
            ["기초: 방한 방문", period, "원문 누적치",
             cell(market["방한방문"], "명", "외부"), "[외부]",
             "연간치·고유 고객 수가 아니다."],
            ["기초: 개별 방문", "개별여행 비율을 적용한 조건부 방문량",
             f"{cell(market['방한방문'], '명', '외부')} × "
             f"{pct(market['개별여행비율'], '외부')}",
             cell(ctx["개별방문"], "명"), "[추정]",
             "두 외부 입력의 기간·모집단 정합 [미확보]"],
            ["TAM", "개별 방한객이 공동 여행마다 구매하는 조건부 상한",
             f"개별 방문량 ÷ 동행 {cell(low_group, '인')} ~ "
             f"{cell(high_group, '인')} × 상품 가격",
             f"{cell(low_trips, '여행')} ~ {cell(high_trips, '여행')}\n"
             + market_money,
             "[추정]",
             "평균 동행 수·구매율 [미확보]. 결제액 상한이며 순매출 아님. "
             "비교안의 인원 제한을 반영한 유효시장 아님."],
            ["SAM", "서울 방문·개인 에이전트·상품·언어·지원 조건 충족 고객",
             "TAM × 서울 방문률 × 에이전트 사용률 × 추가 적합 조건",
             "[미확보]", "[미확보]",
             "핵심 비율과 교차분포를 확보하기 전 수치화하지 않는다."],
            ["SOM", "실제 획득 가능한 유료 여행",
             "검증된 채널별 도달 × 연결률 × 결제율",
             "산정 안 함 [미확보]", "[미확보]",
             "1,200건은 미검증 획득 시나리오이며 판매 가정·SOM이 아니다."],
        ],
        "근거: 클로드 §3-2, 코덱스 시장 규모의 기간·모집단 주의사항, "
        "교차검토 §3. 요청된 누적 방문량으로 통일하며 연간화하지 않는다. "
        "동행 수 범위는 공통 상한 검토용 가정이다.",
    )

    rival = P["경쟁"]
    add(
        "T4", "경쟁 비교",
        ["대상", "유형", "해결하는 일", "확인된 사실", "차별화 가설", "남은 검증"],
        [
            ["Mindtrip", "직접 경쟁", "일정 생성·수정·개인화·여행 중 추천",
             "공식 안내와 사용자 관찰을 원문에서 인용 [외부]; "
             "일부 선제 기능은 관찰 범위에서 미확인",
             "외부 에이전트 맥락을 받아 지속 감시·재검증",
             "같은 서울 변경 과제 시험 [미확보]; 기능 부재로 단정하지 않음"],
            ["트리플·놀유니버스", "인접 여행 앱", "일정·추천·공동 편집·기록",
             "공식 기능 안내·개발 발표 [외부]; 발표와 출시 구분",
             "변경 후 필수 예약과 조건의 연결 복구",
             "인바운드 접근성·최신 제공 범위·동일 과제 비교 [미확보]"],
            ["TripIt·Wanderlog", "유료 관리", "예약 통합·일정 관리·변동 알림",
             f"TripIt Pro 연 {cell(rival['TripIt_연_USD'], '달러', '외부')}; "
             "Wanderlog 유료 상품 존재 [외부]",
             "서울 활동 변경에서 실행 가능한 재계획·인계까지 연결",
             "단기 이용권과 연간 상품 구분; 인바운드 사용률 [미확보]"],
            ["토모트래블·트립플랜", "인적 설계·지원", "맞춤 설계·상담·여행 지원",
             f"토모 상담 {won(rival['토모_상담'], '외부')}, "
             f"패키지 {won(rival['토모_패키지'], '외부')}; "
             "트립플랜은 별도 서비스 안내",
             "반복 확인을 자동화해 고객·운영자 부담 절감",
             "한국인 대상 인접 근거; 인바운드 결제율·지원 품질로 확대 불가"],
            ["OpenClaw·Hermes", "범용 대체재·사용 접점", "검색·도구 호출·자동화·메모리",
             "공식 도구·외부 연결 기능 [외부]",
             "여행 상태·근거·실패 처리를 직접 구성하는 부담 절감",
             "동일 모델·데이터 조건에서 여행 신뢰성·총비용 비교 [미확보]"],
            ["지도·공지·예약처 수작업", "현상 유지", "고객이 직접 확인·연락·변경",
             "제품 정의의 현재 행동 가설 [추정]; 실제 관찰 [미확보]",
             "작업 시간·반복 입력·예약 누락 감소",
             "짝지은 과제와 현장 기록; 자동화가 확인 부담을 늘리는지도 측정"],
        ],
        "외국인 대상 한국 여행 서비스의 경쟁 조사는 미수행 [미확보]. "
        "공개 기능·가격은 지정된 원문 기록의 재인용이며 실행 시 최신 조회를 하지 않는다. "
        "근거: 두 본문의 경쟁 비교, v3 §3·출처 목록, 교차검토 §4.",
    )

    # T5는 요청된 비교 열 순서, T6와 나머지 계열은 결정13 우선.
    cost_order = ("기본", "결정13", "단기")
    cost_rows = [
        ["상품·범위"] + [label(k) + "\n" + scope(k) for k in cost_order],
        ["가격 상태"] + [P["상품"][k]["가격상태"] for k in cost_order],
        ["순매출"] + [won(ctx["상품"][k]["순매출"], places=2) for k in cost_order],
        ["기계(외부 데이터·LLM)"] + [
            won(P["상품"][k]["기계_상한"]) + " 합산 상한"
            if "기계_상한" in P["상품"][k] else
            won(ctx["상품"][k]["비용"]["데이터"] + ctx["상품"][k]["비용"]["LLM"])
            for k in cost_order
        ],
        ["기계 산출 근거"] + [api_recipe(k) for k in cost_order],
        ["사람 지원"] + [
            cell(ctx["상품"][k]["지원분"], "분") + " × "
            + won(P["상품"][k]["지원_분단가"]) + "/분 = "
            + won(ctx["상품"][k]["비용"]["지원"]) for k in cost_order
        ],
        ["장애·보상"] + [won(ctx["상품"][k]["비용"]["대응"]) for k in cost_order],
        ["결제"] + [
            pct(P["상품"][k]["결제비율"]) + " × VAT 포함 결제액 = "
            + won(ctx["상품"][k]["비용"]["결제"], places=2) for k in cost_order
        ],
        ["환불률"] + [pct(P["상품"][k]["환불률"]) for k in cost_order],
        ["플랫폼 변동분"] + [
            won(ctx["상품"][k]["비용"]["플랫폼"]) + "; 미분리는 무료를 뜻하지 않음"
            for k in cost_order
        ],
        ["변동비 / CAC 제외"] + [
            won(ctx["상품"][k]["변동비"], places=2) for k in cost_order
        ],
        ["기여이익Ⅰ"] + [won(ctx["상품"][k]["기여Ⅰ"], places=2) for k in cost_order],
    ]
    add(
        "T5", "원가 예산 대조",
        ["항목", "v3 원본 재현", "결정 13 예산", "비교안"], cost_rows,
        "항목 열 외 세 상품 열은 지정 순서다. 근거: 비용과 수익구조 §4-A, v3 §7, "
        "코덱스 본문 비교안. 결정 13은 기계 상한 10,000원과 사람 지원 상한 "
        "15분 × 분당 500원 = 7,500원을 반영한 예산 [추정]이다. "
        "결제는 문서대로 순매출이 아닌 VAT 포함 결제액의 4% = 1,960원이다. "
        "순매출 = 49,000 / 1.1 × 0.97; "
        "기여이익Ⅰ = 순매출 - (10,000 + 7,500 + 1,000 + 49,000 × 0.04). "
        "기계 세부 배분·실제 지원 시간·해외 계약요율은 [미확보]. "
        "v3 기계비는 공개 단가 × 가정 사용량 + 여유 + LLM 정액 예산; "
        "비교안 LLM은 Case 수 × Case 예산이다. 커머스 실측은 여행 실측이 아니다.",
    )

    rows = [
        ["상품 범위"] + [scope(k) for k in P["상품"]],
        ["가격·상태"] + [
            label(k) + "\n" + P["상품"][k]["가격상태"] for k in P["상품"]
        ],
        ["VAT 제외 매출"] + [won(m["VAT제외"], places=2) for m in ctx["상품"].values()],
        ["환불률"] + [pct(m["환불률"]) for m in ctx["상품"].values()],
        ["환불 매출 차감"] + [won(m["환불차감"], places=2) for m in ctx["상품"].values()],
        ["순매출"] + [won(m["순매출"], places=2) for m in ctx["상품"].values()],
    ]
    cost_names = dict.fromkeys(
        cost for model in ctx["상품"].values() for cost in model["비용"]
    )
    for cost in cost_names:
        rows.append(
            [cost] + [
                won(m["비용"][cost], places=2) if cost in m["비용"]
                else "별도 분리 없음; 기계 합산 행 참조"
                if cost in ("데이터", "LLM")
                else "데이터·LLM 개별 행 참조"
                for m in ctx["상품"].values()
            ]
        )
    rows.extend([
        ["변동비 합계 / CAC 제외"]
        + [won(m["변동비"], places=2) for m in ctx["상품"].values()],
        ["기여이익Ⅰ"]
        + [won(m["기여Ⅰ"], places=2) for m in ctx["상품"].values()],
        ["순매출 대비 기여이익Ⅰ"]
        + [pct(m["기여Ⅰ"] / m["순매출"]) for m in ctx["상품"].values()],
    ])
    add(
        "T6", "결정 13 및 비교 시나리오",
        ["항목"] + [label(k) for k in P["상품"]], rows,
        "순매출 = 가격 / (1 + VAT) × (1 - 환불률) + 제휴수익. "
        "기여이익Ⅰ = 순매출 - 제공 변동비. 결제 수수료는 VAT 포함 가격에 적용한다. "
        "후속 계산은 반올림 전 값으로 수행한다. 결정 13의 기계 합산 예산은 "
        "데이터·LLM으로 재차 더하지 않는다. 서로 다른 범위·원가의 시나리오다. "
        "유료 감시 호출·번역을 포함한 기계 실제 사용량과 지원 시간은 [미확보].",
    )

    cac_rows = []
    candidates = sorted({D(v) for v in P["CAC_후보"]} | {ctx["CAC"]})
    for key, model in ctx["상품"].items():
        for cac in candidates:
            cm2 = model["기여Ⅰ"] - cac
            annual, monthly = break_even(cm2, ctx["고정비"])
            verdict = "고정비 회수 가능량 검증 필요"
            if cm2 < ZERO:
                verdict = "판매 증가 시 손실 확대"
            elif cm2 == ZERO:
                verdict = "판매해도 고정비 회수 불가"
            if cac == D(P["v3_CAC"]):
                verdict += "; v11 §11-A · 한국인 채널 값 - 인바운드 측정 아님"
            if cac == ctx["CAC"]:
                verdict += "; 코덱스 판 가정"
            cac_rows.append([
                label(key), won(cac), won(cm2),
                cell(annual, "건") if annual is not None else "성립 안 함 [추정]",
                cell(monthly, "건") if monthly is not None else "해당 없음 [추정]",
                verdict,
            ])
        cac_rows.append([
            label(key),
            "정확한 기여이익Ⅰ과 동일\n"
            + won(model["기여Ⅰ"], places=P["표시"]["경계_소수"])
            + " (반올림 표시)",
            won(ZERO),
            "성립 안 함 [추정]",
            "해당 없음 [추정]",
            "수직 점근선; 이보다 큰 CAC에서는 건별 손실",
        ])
    add(
        "T7", "CAC 함수 손익분기",
        ["상품", "CAC", "기여이익Ⅱ", "연 손익분기", "월 평균·올림", "판정"],
        cac_rows,
        f"고정비 {won(ctx['고정비'])}/년. 기여이익Ⅱ = 기여이익Ⅰ - CAC. "
        "기여이익Ⅱ가 양수일 때만 연 손익분기 = ceil(고정비 ÷ 기여이익Ⅱ). "
        f"월 표시는 ceil(연 손익분기 ÷ {cell(P['개월_연'], '개월', '외부')})이며 "
        "월별 계절성 예측이 아니다. CAC 후보는 민감도 입력 [추정]; "
        "실제 인바운드 CAC는 [미확보].",
    )

    down_sales = ctx["판매"] * D(P["판매량_후보"]["부진_기준대비"])
    first_sales = (
        ctx["판매"] * D(P["판매량_후보"]["첫해_판매개월"])
        / D(P["개월_연"])
    )
    if down_sales <= ZERO or first_sales <= ZERO:
        raise ValueError("판매 부진·첫해 판매량은 양수여야 합니다.")

    t8_rows, funding_notes = [], []
    for key, model in ctx["상품"].items():
        scenarios = [
            ("미검증 획득 시나리오", ctx["판매"], model, ctx["CAC"]),
            ("미검증 획득 시나리오 절반·획득지출 유지", down_sales, model,
             ctx["획득지출"] / down_sales),
            ("미검증 획득 시나리오·기간 제한", first_sales, model, ctx["CAC"]),
        ]
        if key in ctx["악화"]:
            scenarios[2:2] = [
                ("미검증 획득 시나리오·비용 악화", ctx["판매"],
                 ctx["악화"][key], ctx["CAC"]),
                ("미검증 획득 시나리오 절반·비용 악화", down_sales,
                 ctx["악화"][key], ctx["획득지출"] / down_sales),
            ]
        for name, sales, selected, cac in scenarios:
            cm2 = selected["기여Ⅰ"] - cac
            profit = sales * cm2 - ctx["고정비"]
            t8_rows.append([
                f"{label(key)}\n{name}\nCAC {won(cac)}",
                cell(sales, "건"),
                won(cm2),
                won(profit),
            ])
        first_profit = first_sales * (model["기여Ⅰ"] - ctx["CAC"]) - ctx["고정비"]
        funding = (
            max(ZERO, -first_profit) + D(P["첫해_일회성"])
            + D(P["예비운전자금"])
        )
        funding_notes.append(
            f"{label(key)}: 첫해 영업손실 충당 "
            f"{won(max(ZERO, -first_profit))} + 일회성 "
            f"{won(P['첫해_일회성'])} + 예비 운전자금 "
            f"{won(P['예비운전자금'])} = 필요자금 근사치 {won(funding)}."
        )
        if key not in ctx["악화"]:
            continue
        stress = ctx["악화"][key]
        funding_notes.append(
            "해당 상품 비용 악화 입력: "
            f"환불 {pct(stress['환불률'])}, 결제 {pct(stress['결제비율'])}, "
            + ", ".join(
                f"{name} {won(value)}" for name, value in stress["비용"].items()
                if name != "결제"
            )
            + "."
        )

    funnel_note = "\n".join(
        f"{channel['이름']}: "
        f"{cell(channel['도달'], '적격 도달')} × "
        f"{pct(channel['연결률'])} 연결 × "
        f"{pct(channel['결제율'])} 결제 = "
        f"{cell(D(channel['도달']) * D(channel['연결률']) * D(channel['결제율']), '건')}; "
        f"획득지출 {won(channel['획득지출'])}"
        for channel in P["인바운드_채널"]
    )
    add(
        "T8", "하방·자금",
        ["미검증 획득 시나리오", "획득 건수 [추정]", "기여이익Ⅱ", "조건부 연 손익"],
        t8_rows,
        f"연 손익 = 연 판매 × 기여이익Ⅱ - 연 고정비 {won(ctx['고정비'])}. "
        "이자·법인세 전 조건부 손익이다.\n"
        + funnel_note
        + f"\n미검증 획득 시나리오 {cell(ctx['판매'], '건/년')}; 판매 가정·SOM 아님. "
        f"획득지출 {won(ctx['획득지출'])}; CAC {won(ctx['CAC'])}는 코덱스 판 가정.\n"
        f"첫해 판매기간 {cell(P['판매량_후보']['첫해_판매개월'], '개월')}에도 "
        "연 고정비 전액을 반영한다.\n"
        + "\n".join(funding_notes)
        + "\n결정 13 비용 악화 예산은 [미확보]로 별도 계산하지 않는다. "
        "v3 원본 재현의 비용 악화 입력은 v3, 비교안은 코덱스 본문을 따른다. "
        "동일 충격의 통제 비교가 아니다. 데이터 악화 예산은 정액 가정이며 "
        "공개 API 계산액이 이를 넘으면 큰 값을 쓴다. "
        "여기서는 미검증 획득 시나리오를 유지해 획득비 영향을 분리했다. "
        "첫해 일회성 예산에 실증 현금 한도를 포함하므로 중복 가산하지 않는다. "
        "필요자금은 정산·세금 시차를 단순화한 근사치이며 월별 현금계획은 [미확보].",
    )

    add(
        "T9", "미결정 관리",
        ["ID", "무엇", "영향", "확보 방법", "담당", "기한"],
        [
            ["U1", "인바운드 CAC·채널·실제 지불자 [미확보]",
             "판매량·손익·청구 상대",
             "모집 / 연결 / 결제 / 환불 로그와 채널별 획득 지출",
             "[미배정] 사업 검증 책임자 지정 필요", "모집 설계 전·실증 종료"],
            ["U2", "SAM 비율·교차분포·평균 동행 수 [미확보]",
             "유효시장·구매 단위",
             "동일 기간 관광 통계와 대상 고객 설문",
             "[미배정]", "시장 규모 확정 전"],
            ["U3", "결정 11 언어 경로 확정; 번역 품질·이해도 [미확보]",
             "링크·알림은 우리가 고객 선택 언어로 번역; 기본 한국어·영어",
             "API는 한국어로 제공하고 에이전트가 번역; 두 경로 이해도 시험",
             "역할 B·C 제안; 실명 미확정", "모집 전"],
            ["U4", "결정 13 상품 확정; 가격 수용성·원가 [미확보]",
             "서울·최대 7일·4인·49,000원 VAT 포함 [결정]",
             "상한 준수·지원 시간·가격 수용성을 파일럿에서 검증",
             "팀·승인자 실명 [미확보]", "유료 실증 전"],
            ["U5", "해외 카드·정산 계약요율 [미확보]",
             "결제 원가·현금 시차",
             "결제수단별 견적·환불 수수료·정산 조건 확인",
             "[미배정]; 재무 역할 지정 필요", "결제수단 결정 시"],
            ["U6", "여행 인계율·처리 시간·LLM 사용량 [미확보]",
             "지원·모델 원가",
             "Case별 호출 로그·지원 작업 타이머·파일럿",
             "역할 D·E·평가 담당 제안; 실제 배정 재확인",
             "실증 종료"],
            ["U7", "API 사용량·필드·교통 모드·권한 [미확보]",
             "데이터비·일정 성립",
             "공개 단가에 SKU별 실제 과금 단위·로그 대조",
             "역할 B·D·E 제안",
             pilot["원가검토"] + " [추정]"],
            ["U8", "감시 소스의 최신성·비용 [미확보]",
             "선제 감시 제공 가능성",
             "지원 대상별 소스·갱신 시각·호출 정책 실증",
             "역할 E 제안", pilot["통합검토"] + " [추정]"],
            ["U9", "Trip 저장·경로 조회·실행 게이트 미완성",
             "조정·출발 안내·실행 안전",
             "마이그레이션·경합·조회·차단 시험",
             "역할 A·B 제안; WBS와 담당 불일치 확인",
             pilot["통합검토"] + " [추정]"],
            ["U10", "정책 수치·RAG 원문 전달 계약 불일치",
             "취소·환급 판정",
             "실행 재현 후 구조화 정책과 원문 버전 연결",
             "역할 D 제안", "판정 평가 전"],
            ["U11", "운영 담당·복구·가용성·가용 공수 [미확보]",
             "출시 가능성·상용 고정비",
             "실명 배정·작업 가능 시간 확인·운영표·복구 훈련",
             "프로젝트 책임자·역할 E 제안", "첫 검토·판매 전"],
            ["U12", "인바운드 직접 경쟁 조사 [미확보]",
             "구매 이유·차별화",
             "외국인 대상 한국 여행 서비스의 동일 과제 비교",
             "[미배정] 사업 검증 담당 필요", "유료 제안 전"],
        ],
        "역할 문자는 코덱스 판의 제안이며 확정 인사가 아니다. "
        "클로드 판의 사업 검증 담당 미배정과 담당 문서 불일치를 유지한다. "
        "기술 담당이 있다는 사실만으로 사업 검증 책임자가 배정된 것으로 보지 않는다. "
        "근거: 두 본문 R&R, 교차검토 §3·§4. ID·표 번호·문서 판명은 식별자다.",
    )
    return specs, notes


def save_docx(root, out, specs, notes):
    sys.path.insert(0, str(root / BUILD_REL))
    from docx_tables import insert_tables

    path = out / "기획서_표모음.docx"
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(P["표시"]["문서_가로인치"])
    section.page_height = Inches(P["표시"]["문서_세로인치"])
    margin = Inches(P["표시"]["문서_여백인치"])
    section.left_margin = section.right_margin = margin
    section.top_margin = section.bottom_margin = margin

    normal = doc.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal.font.size = Pt(P["표시"]["문서폰트"])
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"), "맑은 고딕"
    )
    for index, (anchor, caption, header, rows) in enumerate(specs):
        if index:
            doc.add_page_break()
        # 제목의 전체 텍스트와 insert_tables의 앵커가 정확히 일치한다.
        doc.add_heading(anchor, level=1)
        doc.add_paragraph(notes[anchor])
        if index == 0:
            doc.add_paragraph(
                "등급: [결정] 팀이 정한 가격·범위 / [v3 인용] 과거 원본 재현 / "
                "[실측] 관측·기록 확인(가격 결정 제외) / [외부] 공개 자료의 원문 인용 / "
                "[추정] 가정·계산·예산·목표 / [미확보] 근거 미확보. "
                "계산 결과는 원 단위 표시에서 반올림한다."
            )
        if any(len(row) != len(header) for row in rows):
            raise ValueError(f"표 열 수 불일치: {anchor}")

    doc.add_paragraph("입력 근거")
    for name, source in SOURCES.items():
        doc.add_paragraph(f"{name}: {source}")
    doc.save(path)

    saved, count = insert_tables(str(path), specs)
    if Path(saved).resolve() != path.resolve():
        raise RuntimeError(
            "지정 파일이 잠겨 도구가 대체 파일에 저장했습니다. "
            f"Word에서 파일을 닫고 다시 실행하세요: {saved}"
        )

    # 넓은 표와 여러 페이지에 걸치는 표의 가독성을 보정한다.
    doc = Document(path)
    if count != len(specs) or len(doc.tables) != len(specs):
        raise RuntimeError("표 삽입 개수가 명세와 다릅니다.")
    usable_width = (
        doc.sections[0].page_width
        - doc.sections[0].left_margin
        - doc.sections[0].right_margin
    )
    for table in doc.tables:
        table.autofit = False
        width = int(usable_width / len(table.columns))
        for column in table.columns:
            column.width = width
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in table.rows:
            for table_cell in row.cells:
                table_cell.width = width
                if "[미확보]" in table_cell.text or "[미배정]" in table_cell.text:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EEEEEE")
                    table_cell._tc.get_or_add_tcPr().append(shade)
    doc.save(path)
    return path


def setup_plots():
    font_manager.findfont("Malgun Gothic", fallback_to_default=False)
    plt.rcParams.update({
        "font.family": "Malgun Gothic",
        "axes.unicode_minus": False,
        "font.size": P["표시"]["폰트"],
        "axes.spines.top": False,
        "axes.spines.right": False,
        "savefig.facecolor": "white",
    })


def plot_cac(ctx, out):
    fig, ax = plt.subplots(figsize=P["표시"]["F1_크기"])
    maximum = max(
        max(float(v) for v in P["CAC_후보"]),
        max(float(m["기여Ⅰ"]) for m in ctx["상품"].values()),
    )
    xmax = maximum * P["표시"]["F1_X여유"]
    ymax = P["표시"]["F1_Y상한"]
    xgrid = np.linspace(0, xmax, P["표시"]["F1_표본"])

    for index, (key, model) in enumerate(ctx["상품"].items()):
        cm1 = float(model["기여Ⅰ"])
        color = COLORS[key]
        # CAC >= 기여이익Ⅰ 영역은 정의하지 않는다.
        x = xgrid[xgrid < cm1]
        if len(x):
            y = float(ctx["고정비"]) / (cm1 - x)
            ax.plot(x, y, color=color, lw=2, label=label(key),
                    linestyle=":" if key == "기본" else "-")
        if cm1 > 0:
            ax.axvline(cm1, color=color, linestyle="--", alpha=0.8)
            ax.text(
                cm1, ymax * (0.96 - index * 0.17),
                "  기여이익Ⅰ 경계\n  "
                + won(model["기여Ⅰ"], places=P["표시"]["경계_소수"]),
                color=color, va="top", ha="left",
            )

    for volume, title, style in [
        (ctx["판매"], "미검증 획득 시나리오", ":"),
        (D(P["판매량_후보"]["v3_참조"]), "v3 판매 참조·인바운드 근거 아님", "-."),
    ]:
        ax.axhline(float(volume), color=GRAY, linestyle=style, lw=1.2)
        ax.text(
            xmax * 0.98, float(volume), title + " " + cell(volume, "건"),
            ha="right", va="bottom", color="#555555",
            bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.8},
        )

    primary = ctx["상품"]["결정13"]
    for index, (cac, caption) in enumerate([
        (D(P["v3_CAC"]), "v11 §11-A · 한국인 채널 값 - 인바운드 측정 아님"),
        (ctx["CAC"], "코덱스 판 가정"),
    ]):
        cm2 = primary["기여Ⅰ"] - cac
        if cm2 > ZERO:
            y = float(ctx["고정비"] / cm2)
            ax.scatter(float(cac), y, color=COLORS["결정13"], zorder=5)
            ax.annotate(
                "CAC " + won(cac) + "\n" + caption,
                xy=(float(cac), y), xytext=(0.46, 0.52 - index * 0.17),
                textcoords="axes fraction",
                arrowprops={"arrowstyle": "-", "color": GRAY},
                fontsize=P["표시"]["폰트"] - 2,
            )
    ax.set_xlim(0, xmax)
    ax.set_ylim(0, ymax)
    ax.set_xlabel("CAC (원) [추정 후보] - 실제 인바운드 CAC [미확보]")
    ax.set_ylabel("연 손익분기 판매량 (건) [추정]")
    ax.set_title("F1  CAC-손익분기 곡선")
    ax.xaxis.set_major_formatter(FuncFormatter(lambda x, _: f"{x:,.0f}"))
    ax.yaxis.set_major_formatter(FuncFormatter(lambda y, _: f"{y:,.0f}"))
    ax.grid(axis="y", color=LIGHT)
    ax.legend(loc="upper left")
    fig.text(
        0.02, 0.02,
        f"고정비 {won(ctx['고정비'])}/년. 곡선 = 고정비 / (기여이익Ⅰ - CAC); 정수 올림은 T7 참조.\n"
        "CAC < 기여이익Ⅰ 이면 기여, = 이면 0, > 이면 팔수록 손해. 같거나 크면 손익분기 없음.\n"
        "주 상품: 결정 13, 서울·최대 7일·4인. v3 원본 재현의 2도시는 결정 13의 서울 1도시와 다름.\n"
        "1,200건은 미검증 획득 시나리오이며 판매 가정·SOM 아님. 상단 범위 밖 곡선은 잘라 표시한다.",
        fontsize=P["표시"]["폰트"] - 1,
    )
    fig.tight_layout(rect=(0, 0.18, 1, 1))
    path = out / "F1_CAC_손익분기.png"
    fig.savefig(path, dpi=P["표시"]["DPI"])
    plt.close(fig)
    return path


def plot_waterfall(ctx, key, out):
    model = ctx["상품"][key]
    entries = [("순매출", model["순매출"], "합계")]
    entries.extend(
        (name, -value, "차감") for name, value in model["비용"].items()
    )
    entries.extend([
        ("기여이익Ⅰ", model["기여Ⅰ"], "합계"),
        ("CAC\n코덱스 판 가정", -ctx["CAC"], "차감"),
        ("기여이익Ⅱ", model["기여Ⅰ"] - ctx["CAC"], "합계"),
    ])
    fig, ax = plt.subplots(figsize=P["표시"]["F2_크기"])
    current = ZERO
    levels = [ZERO]
    for index, (name, value, kind) in enumerate(entries):
        if kind == "합계":
            bottom = ZERO
            height = value
            current = value
            color = COLORS[key]
        else:
            bottom = current + value
            height = -value
            current += value
            color = GRAY
        ax.bar(
            index, float(height), bottom=float(bottom),
            color=color, width=0.65, edgecolor="white",
        )
        top = max(bottom, bottom + height)
        low = min(bottom, bottom + height)
        levels.extend([top, low])
        text_y = top if value >= ZERO or kind == "차감" else low
        ax.annotate(
            won(value), (index, float(text_y)),
            xytext=(0, 7 if text_y >= ZERO else -7),
            textcoords="offset points",
            ha="center", va="bottom" if text_y >= ZERO else "top",
            fontsize=P["표시"]["폰트"] - 1,
        )
        if index < len(entries) - 1:
            ax.plot(
                [index + 0.325, index + 0.675],
                [float(current), float(current)],
                color=GRAY, linestyle=":", lw=1,
            )

    span = float(max(levels) - min(levels)) or 1.0
    ax.set_ylim(float(min(levels)) - span * 0.12,
                float(max(levels)) + span * 0.18)
    ax.axhline(0, color="#555555", lw=0.8)
    ax.set_xticks(
        range(len(entries)),
        [entry[0].replace("기계(외부 데이터·LLM)", "기계\n외부 데이터·LLM")
         for entry in entries],
    )
    ax.set_ylabel("원 / 최초 결제 여행 [추정]")
    ax.yaxis.set_major_formatter(FuncFormatter(lambda y, _: f"{y:,.0f}"))
    ax.set_title("F2  기여이익 폭포 - " + label(key) + "\n" + scope(key))
    ax.grid(axis="y", color=LIGHT)
    ax.set_axisbelow(True)
    fig.text(
        0.02, 0.02,
        f"CAC {won(ctx['CAC'])}: 코덱스 판 가정. 실제 CAC·여행 원가는 [미확보].\n"
        "순매출은 VAT·환불 차감 후; 결제 수수료는 VAT 포함 가격에 적용한다. "
        "기여이익Ⅰ·Ⅱ는 합계로서 다시 더하지 않는다.\n"
        "CAC < 기여이익Ⅰ 이면 기여, = 이면 0, > 이면 팔수록 손해.\n"
        "v3 원본 재현 · 현재 판매 전망 아님: 2도시로 결정 13의 서울 1도시와 다름. 비교안은 팀 결정 아님.",
        fontsize=P["표시"]["폰트"] - 1,
    )
    fig.tight_layout(rect=(0, 0.18, 1, 0.94))
    suffix = {"결정13": "기본", "기본": "v3_원본_재현", "단기": "비교안"}[key]
    path = out / f"F2_기여이익_폭포_{suffix}.png"
    fig.savefig(path, dpi=P["표시"]["DPI"])
    plt.close(fig)
    return path


def plot_market(ctx, out):
    market = P["시장"]
    fig, ax = plt.subplots(figsize=P["표시"]["F3_크기"])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")

    half = 0.43
    ratio = float(D(market["개별여행비율"]))
    first_bottom = half * ratio

    # 위 두 단계의 가로폭만 방문량 비율을 표현한다.
    upper = Polygon(
        [(0.5 - half, 0.88), (0.5 + half, 0.88),
         (0.5 + first_bottom, 0.69), (0.5 - first_bottom, 0.69)],
        closed=True, facecolor=COLORS["결정13"], edgecolor="white",
    )
    ax.add_patch(upper)
    ax.text(
        0.5, 0.785,
        "방한 "
        + cell(D(market["방한방문"]) / D(P["표시_만원"]), "만 명", "외부")
        + f"\n{market['연도']}년 {market['시작월']}~{market['종료월']}월 누적 [외부]",
        ha="center", va="center", color="white",
    )

    ax.add_patch(Rectangle(
        (0.5 - first_bottom, 0.49), first_bottom * 2, 0.16,
        facecolor=COLORS["결정13"], alpha=0.72, edgecolor="white",
    ))
    ax.text(
        0.5, 0.57,
        "개별 "
        + cell(ctx["개별방문"] / D(P["표시_만원"]), "만 명")
        + "\n방한 방문량 × "
        + pct(market["개별여행비율"], "외부"),
        ha="center", va="center", color="white",
    )

    # 미확보 단계는 같은 폭의 독립 상자: 면적·폭으로 시장 크기를 암시하지 않는다.
    for y, text in [
        (0.27, "SAM [미확보]\n서울 방문률 · 에이전트 사용률 · 추가 적합 조건"),
        (0.05, "SOM 산정 안 함 [미확보]\n검증된 획득 채널·전환 실적 필요"),
    ]:
        ax.add_patch(Rectangle(
            (0.14, y), 0.72, 0.15,
            facecolor=LIGHT, edgecolor=GRAY, hatch="///",
        ))
        ax.text(0.5, y + 0.075, text, ha="center", va="center", color="#333333")

    for start, end in [(0.685, 0.655), (0.475, 0.425), (0.255, 0.205)]:
        ax.annotate(
            "", xy=(0.5, end), xytext=(0.5, start),
            arrowprops={"arrowstyle": "->", "color": GRAY},
        )

    ax.set_title("F3  시장 깔때기 - 조건부 상한과 미확보 단계", pad=15)
    fig.text(
        0.05, 0.025,
        "개별 방문량은 외부 입력의 곱으로 계산한 [추정]이며 만 명 단위로 반올림했다.\n"
        "기간·모집단 정합은 [미확보]. 연간치·고유 구매자·실제 유료 수요가 아니다.\n"
        "회색 빗금 상자의 폭·면적은 수량을 뜻하지 않는다. "
        "1,200건은 미검증 획득 시나리오이며 판매 가정·SOM 아님.\n"
        "주 상품: 결정 13, 서울·최대 7일·4인·49,000원 VAT 포함 [결정].\n"
        "v3 원본 재현 · 현재 판매 전망 아님: 2도시로 결정 13의 서울 1도시와 다름.",
        fontsize=P["표시"]["폰트"] - 1,
    )
    fig.tight_layout(rect=(0, 0.19, 1, 1))
    path = out / "F3_시장_깔때기.png"
    fig.savefig(path, dpi=P["표시"]["DPI"])
    plt.close(fig)
    return path


def validate(ctx):
    if D(P["환율"]) <= ZERO or D(P["VAT"]) < ZERO:
        raise ValueError("환율·VAT 입력을 확인하세요.")
    for channel in P["인바운드_채널"]:
        for rate in ("연결률", "결제율"):
            if not ZERO <= D(channel[rate]) <= ONE:
                raise ValueError(f"채널 비율 오류: {channel['이름']} / {rate}")
    for key, model in ctx["상품"].items():
        if not ZERO <= model["환불률"] < ONE:
            raise ValueError(f"환불률 오류: {key}")
        if model["순매출"] <= ZERO:
            raise ValueError(f"순매출은 양수여야 합니다: {key}")
        if any(value < ZERO for value in model["비용"].values()):
            raise ValueError(f"음수 비용 입력: {key}")
        if model["순매출"] - model["변동비"] != model["기여Ⅰ"]:
            raise AssertionError("기여이익 계산 불일치")
        cm2 = model["기여Ⅰ"] - ctx["CAC"]
        annual, _ = break_even(cm2, ctx["고정비"])
        if annual is not None:
            assert D(annual) * cm2 >= ctx["고정비"]
            assert D(annual - 1) * cm2 < ctx["고정비"]
        assert break_even(ZERO, ctx["고정비"]) == (None, None)
        assert break_even(-ONE, ctx["고정비"]) == (None, None)


def main():
    root = find_root()
    setup_plots()                       # 글꼴 문제는 파일 저장 전에 확인
    ctx = context()
    validate(ctx)
    specs, notes = make_tables(ctx)

    out = Path(OUT_DIR)
    if not out.is_absolute():
        out = root / out
    out.mkdir(parents=True, exist_ok=True)

    paths = [save_docx(root, out, specs, notes), plot_cac(ctx, out)]
    paths.extend(plot_waterfall(ctx, key, out) for key in ("결정13", "단기"))
    paths.append(plot_market(ctx, out))

    for key, model in ctx["상품"].items():
        print(
            label(key),
            "| 기계", won(
                model["비용"]["기계(외부 데이터·LLM)"] if "기계_상한" in P["상품"][key]
                else model["비용"]["데이터"] + model["비용"]["LLM"]
            ),
            "| 기여이익Ⅰ", won(model["기여Ⅰ"]),
        )
    print("미검증 획득 시나리오:", cell(ctx["판매"], "건/년"), "/ 판매 가정·SOM 아님")
    print("CAC 코덱스 판 가정:", won(ctx["CAC"]), "/ 실제 값 [미확보]")
    for path in paths:
        print(path.resolve())


if __name__ == "__main__":
    main()
