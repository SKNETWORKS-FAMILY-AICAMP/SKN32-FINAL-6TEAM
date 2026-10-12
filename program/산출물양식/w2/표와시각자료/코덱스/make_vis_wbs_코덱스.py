from pathlib import Path
from datetime import date, datetime, timedelta
from collections import OrderedDict
import math
import re
import unicodedata

import numpy as np
import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from matplotlib.colors import LinearSegmentedColormap
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


ROOT = Path.cwd()
SOURCE = ROOT / "program/산출물양식/w2/제출본/[기획] WBS_2차_여행.xlsx"
OUTPUT = ROOT / "program/산출물양식/w2/표와시각자료/코덱스"

YEAR = 2026
BLUE = "#2a78d6"
ORANGE = "#eb6834"
GRAY = "#737373"
DARK = "#333333"

plt.rcParams.update({
    "font.family": "Malgun Gothic",
    "axes.unicode_minus": False,
    "font.size": 7.5,
    "text.color": DARK,
    "axes.labelcolor": DARK,
    "xtick.color": GRAY,
    "ytick.color": DARK,
    "axes.edgecolor": "#cccccc",
    "savefig.dpi": 200,
    "figure.facecolor": "white",
})


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def label(value):
    return re.sub(r"\[[^\]]*\]", "", clean(value)).strip()


def wrap_label(value, limit=48):
    lines, current, width = [], "", 0
    for char in label(value):
        char_width = 2 if unicodedata.east_asian_width(char) in "WF" else 1
        if current and width + char_width > limit:
            lines.append(current.rstrip())
            current, width = "", 0
        current += char
        width += char_width
    if current.strip():
        lines.append(current.strip())
    return "\n".join(lines) or "미기재"


def parse_date(value, epoch):
    if isinstance(value, (datetime, date)):
        return date(YEAR, value.month, value.day)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        converted = from_excel(value, epoch)
        if isinstance(converted, (datetime, date)):
            return date(YEAR, converted.month, converted.day)
    text = clean(value)
    match = re.fullmatch(
        r"(?:(?:\d{4})[./-])?(\d{1,2})[./-](\d{1,2})(?:\s.*)?", text
    )
    if not match:
        match = re.fullmatch(r"(\d{1,2})월\s*(\d{1,2})일?", text)
    if not match:
        raise ValueError(f"날짜를 해석할 수 없습니다: {value!r}")
    return date(YEAR, int(match[1]), int(match[2]))


def split_owners(value):
    parts = re.split(r"[,，、/／;；&＆+＋·ㆍ\n\r]+", str(value or ""))
    names = list(dict.fromkeys(label(part) for part in parts if label(part)))
    return names or ["미지정"]


def read_tasks():
    workbook = load_workbook(SOURCE, data_only=True, read_only=False)
    try:
        candidates = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(min_row=1, max_row=30, max_col=7):
                values = [clean(cell.value).replace(" ", "") for cell in row]
                if (
                    values[0].upper() in ("ID", "TASKID")
                    and "주요업무" in values[1]
                    and "세부업무" in values[2]
                ):
                    candidates.append((sheet, row[0].row))
                    break

        if len(candidates) != 1:
            raise ValueError(
                f"A~G열의 지정된 머리글을 가진 시트가 하나여야 합니다. "
                f"발견한 시트 수: {len(candidates)}"
            )

        sheet, header_row = candidates[0]
        merged_values = {}
        for region in sheet.merged_cells.ranges:
            origin = sheet.cell(region.min_row, region.min_col).value
            for row_index in range(region.min_row, region.max_row + 1):
                for column_index in range(region.min_col, min(region.max_col, 7) + 1):
                    merged_values[(row_index, column_index)] = origin

        def value(row_index, column_index):
            return merged_values.get(
                (row_index, column_index),
                sheet.cell(row_index, column_index).value,
            )

        tasks = []
        current_major = ""
        for row_index in range(header_row + 1, sheet.max_row + 1):
            major = label(value(row_index, 2))
            if major:
                current_major = major

            # 병합된 ID의 하위 행을 별도 작업으로 중복 집계하지 않는다.
            raw_id = sheet.cell(row_index, 1).value
            task_id = clean(raw_id)
            if not task_id or task_id.upper() == "ID":
                continue
            if re.fullmatch(r"(합계|총계|소계|총\s*\d+\s*개)", task_id):
                continue
            if isinstance(raw_id, float) and raw_id.is_integer():
                task_id = str(int(raw_id))

            detail = label(value(row_index, 3))
            if not detail and not current_major:
                continue

            try:
                start = parse_date(value(row_index, 6), workbook.epoch)
                end = parse_date(value(row_index, 7), workbook.epoch)
            except (ValueError, TypeError, OverflowError) as error:
                raise ValueError(
                    f"{sheet.title} 시트 {row_index}행, ID {task_id}: {error}"
                ) from error

            if end < start:
                raise ValueError(f"ID {task_id}: 마감일이 시작일보다 빠릅니다.")

            tasks.append({
                "id": task_id,
                "major": current_major or "주요 업무 미지정",
                "detail": detail or current_major,
                "owners": split_owners(value(row_index, 4)),
                "start": start,
                "end": end,
            })

        if len(tasks) != 125:
            raise ValueError(
                f"작업 수가 요청한 125개와 다릅니다: {len(tasks)}개. "
                "원본의 ID, 병합 셀 및 작업 행 구성을 확인하세요."
            )
        return tasks
    finally:
        workbook.close()


def draw_gantt(tasks):
    groups = OrderedDict()
    for task in tasks:
        groups.setdefault(task["major"], []).append(task)

    milestones = [
        (date(YEAR, 9, 15), "중간발표 09-15", ORANGE),
        (date(YEAR, 10, 26), "최종발표 10-26", GRAY),
    ]
    first = min(min(task["start"] for task in tasks), milestones[0][0])
    last = max(max(task["end"] for task in tasks), milestones[1][0])
    left_date, right_date = first - timedelta(days=2), last + timedelta(days=3)

    rows = []
    position = 0.0
    for major, members in groups.items():
        start = min(task["start"] for task in members)
        end = max(task["end"] for task in members)
        title = wrap_label(f"{major} ({len(members)}개)", 46)
        height = 0.17 * len(title.splitlines()) + 0.24
        rows.append({
            "summary": True, "y": position + height / 2,
            "top": position, "height": height, "text": title,
            "start": start, "end": end,
        })
        position += height

        for task in members:
            title = wrap_label(f'{task["id"]}  {task["detail"]}', 48)
            height = max(0.23, 0.16 * len(title.splitlines()) + 0.08)
            rows.append({
                "summary": False, "y": position + height / 2,
                "top": position, "height": height, "text": title,
                "start": task["start"], "end": task["end"],
            })
            position += height
        position += 0.12

    # 가로 폭은 6.3인치로 고정하고, 모든 작업을 읽을 수 있도록 세로를 늘린다.
    bottom_inches, top_inches = 0.38, 1.48
    figure_height = position + bottom_inches + top_inches
    fig = plt.figure(figsize=(6.3, figure_height))
    axis_bottom = bottom_inches / figure_height
    axis_height = position / figure_height
    text_ax = fig.add_axes([0.035, axis_bottom, 0.425, axis_height])
    time_ax = fig.add_axes([0.485, axis_bottom, 0.495, axis_height])

    for ax in (text_ax, time_ax):
        ax.set_ylim(position, 0)
        ax.set_yticks([])
        for spine in ax.spines.values():
            spine.set_visible(False)

    text_ax.set_xlim(0, 1)
    text_ax.set_xticks([])
    time_ax.set_xlim(mdates.date2num(left_date), mdates.date2num(right_date))

    total_days = (right_date - left_date).days
    tick_interval = max(1, math.ceil(total_days / (7 * 16)))
    time_ax.xaxis.set_major_locator(
        mdates.WeekdayLocator(byweekday=mdates.MO, interval=tick_interval)
    )
    time_ax.xaxis.set_major_formatter(mdates.DateFormatter("%m/%d"))
    time_ax.tick_params(
        axis="x", top=True, labeltop=True, bottom=False,
        labelbottom=False, labelsize=6.5, rotation=90, length=2, pad=3,
    )
    time_ax.grid(axis="x", color="#e4e4e4", linewidth=0.4)
    time_ax.set_axisbelow(True)

    for index, row in enumerate(rows):
        y, height = row["y"], row["height"]
        if row["summary"]:
            for ax in (text_ax, time_ax):
                ax.axhspan(row["top"], row["top"] + height,
                           color="#edf3fa", zorder=0)
            duration = (row["end"] - row["start"]).days + 1
            subtitle = (
                f'{row["start"]:%m/%d} - {row["end"]:%m/%d} · {duration}일'
            )
            text_ax.text(
                0.01, y - 0.075, row["text"], va="center",
                fontsize=7.6, fontweight="bold", linespacing=1.15,
            )
            text_ax.text(
                0.01, row["top"] + height - 0.065, subtitle,
                va="center", fontsize=6.6, color=GRAY,
            )
        else:
            if index % 2 == 0:
                for ax in (text_ax, time_ax):
                    ax.axhspan(row["top"], row["top"] + height,
                               color="#fafafa", zorder=0)
            text_ax.text(
                0.025, y, row["text"], va="center",
                fontsize=7.0, linespacing=1.15,
            )

        start_number = mdates.date2num(row["start"])
        width = (row["end"] - row["start"]).days + 1
        time_ax.barh(
            y, width, left=start_number,
            height=0.14 if row["summary"] else 0.10,
            color=BLUE, alpha=1 if row["summary"] else 0.62,
            edgecolor="none", zorder=3,
        )

    for milestone, _, color in milestones:
        time_ax.axvline(
            mdates.date2num(milestone), color=color,
            linewidth=0.9, linestyle=(0, (3, 3)), zorder=4,
        )

    def header_y(inches_from_top):
        return 1 - inches_from_top / figure_height

    fig.text(0.035, header_y(0.23), "주요 업무별 작업 일정",
             fontsize=12, fontweight="bold", va="top")
    fig.text(0.035, header_y(0.53),
             f"2026년 · 작업 {len(tasks)}개 · 주요 업무 {len(groups)}개",
             fontsize=8, va="top")
    fig.text(0.035, header_y(0.76),
             "굵은 막대: 주요 업무 기간 · 가는 막대: 개별 작업",
             fontsize=7, color=GRAY, va="top")
    for x, (_, title, color) in zip((0.49, 0.75), milestones):
        fig.text(x, header_y(0.76), title,
                 fontsize=7, color=color, va="top")
    fig.text(0.035, 0.11 / figure_height,
             "기간은 시작일과 마감일을 포함하며, 주요 업무 기간은 소속 작업의 전체 범위입니다.",
             fontsize=6.5, color=GRAY, va="bottom")

    fig.savefig(OUTPUT / "WBS_간트_코덱스.png", dpi=200)
    plt.close(fig)


def draw_workload(tasks):
    first = min(task["start"] for task in tasks)
    last = max(task["end"] for task in tasks)
    first_monday = first - timedelta(days=first.weekday())
    last_monday = last - timedelta(days=last.weekday())
    weeks = [
        first_monday + timedelta(weeks=index)
        for index in range((last_monday - first_monday).days // 7 + 1)
    ]
    owners = sorted({owner for task in tasks for owner in task["owners"]})
    if "미지정" in owners:
        owners.remove("미지정")
        owners.append("미지정")

    owner_index = {owner: index for index, owner in enumerate(owners)}
    counts = np.zeros((len(owners), len(weeks)), dtype=int)
    for task in tasks:
        for column, monday in enumerate(weeks):
            sunday = monday + timedelta(days=6)
            if task["start"] <= sunday and task["end"] >= monday:
                for owner in task["owners"]:
                    counts[owner_index[owner], column] += 1

    cmap = LinearSegmentedColormap.from_list(
        "작업수", ["#f1f1f1", "#dbe9f9", BLUE]
    )
    maximum = max(1, int(counts.max()))
    owner_labels = [wrap_label(owner, 18) for owner in owners]
    row_height = max(
        0.38, 0.15 * max(len(text.splitlines()) for text in owner_labels) + 0.12
    )
    figure_height = max(3.5, 1.85 + row_height * len(owners))
    fig = plt.figure(figsize=(6.3, figure_height))
    ax = fig.add_axes([
        0.20, 0.85 / figure_height, 0.775,
        (figure_height - 1.90) / figure_height,
    ])
    plotted = ax.imshow(
        counts, cmap=cmap, vmin=0, vmax=maximum,
        interpolation="nearest", aspect="auto",
    )
    ax.set_yticks(np.arange(len(owners)), labels=owner_labels)
    ax.set_xticks(
        np.arange(len(weeks)),
        labels=[monday.strftime("%m/%d") for monday in weeks],
    )
    ax.tick_params(axis="y", labelsize=7, length=0, pad=5)
    ax.tick_params(axis="x", labelsize=6.5, length=0, pad=5, rotation=90)
    ax.set_xticks(np.arange(-0.5, len(weeks), 1), minor=True)
    ax.set_yticks(np.arange(-0.5, len(owners), 1), minor=True)
    ax.grid(which="minor", color="white", linewidth=0.8)
    ax.tick_params(which="minor", bottom=False, left=False)
    for spine in ax.spines.values():
        spine.set_visible(False)

    for row in range(len(owners)):
        for column in range(len(weeks)):
            number = counts[row, column]
            ax.text(
                column, row, str(number), ha="center", va="center",
                fontsize=6.5,
                color="white" if number >= maximum * 0.70 else DARK,
            )

    fig.text(0.035, 1 - 0.20 / figure_height,
             "담당자별 주간 작업 수", fontsize=12,
             fontweight="bold", va="top")
    fig.text(0.035, 1 - 0.49 / figure_height,
             "2026년 · 해당 주에 하루 이상 진행되는 작업을 1건으로 집계",
             fontsize=7.5, va="top")
    fig.text(0.035, 1 - 0.71 / figure_height,
             "공동 담당 작업은 각 담당자에게 1건씩 반영 · 주간 기준: 월요일 - 일요일",
             fontsize=7, color=GRAY, va="top")

    color_ax = fig.add_axes([0.57, 0.22 / figure_height, 0.39, 0.10 / figure_height])
    colorbar = fig.colorbar(plotted, cax=color_ax, orientation="horizontal")
    ticks = sorted(set([0, maximum // 2, maximum]))
    colorbar.set_ticks(ticks)
    colorbar.ax.tick_params(labelsize=6, length=2, pad=2)
    colorbar.outline.set_visible(False)
    fig.text(0.035, 0.39 / figure_height,
             "가로축 날짜: 각 주의 월요일",
             fontsize=6.5, color=GRAY, va="center")
    fig.text(0.57, 0.39 / figure_height,
             "작업 수 (건)", fontsize=6.5, color=GRAY, va="center")

    fig.savefig(OUTPUT / "WBS_담당별부하_코덱스.png", dpi=200)
    plt.close(fig)


if __name__ == "__main__":
    tasks = read_tasks()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    draw_gantt(tasks)
    draw_workload(tasks)
    print(f"작업 {len(tasks)}개 시각화 완료: {OUTPUT}")