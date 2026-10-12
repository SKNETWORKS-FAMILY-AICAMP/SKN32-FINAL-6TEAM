from pathlib import Path
from collections import Counter
import re
import textwrap
import unicodedata

import openpyxl
import numpy as np
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import ListedColormap
from matplotlib.patches import Patch
from matplotlib.ticker import MaxNLocator

plt.rcParams.update({
    "font.family": "Malgun Gothic",
    "axes.unicode_minus": False,
    "font.size": 9,
    "text.color": "#333333",
    "axes.labelcolor": "#333333",
    "xtick.color": "#555555",
    "ytick.color": "#333333",
    "figure.facecolor": "white",
    "axes.facecolor": "white",
    "savefig.facecolor": "white",
})

ROOT = Path.cwd()
SOURCE = ROOT / "program/산출물양식/w2/제출본/[기획] 요구사항 정의서_2차_여행.xlsx"
OUTPUT = ROOT / "program/산출물양식/w2/표와시각자료/코덱스"
BLUE = "#2a78d6"
ORANGE = "#eb6834"
GRAY = "#888888"
GRID = "#dddddd"
WIDTH = 6.3
DPI = 200

# 현재 원본 F열에는 DoD 표기가 없다.
# 명시된 번호만 연결하며, 요구사항 내용으로 연결을 추정하지 않는다.
DOD_MARKER = re.compile(r"\bDoD\b", re.IGNORECASE)
NUMBER = r"\d{1,3}(?!\d)"
ITEM = rf"{NUMBER}(?:\s*[-~–—∼]\s*(?:DoD\s*[-:#]?\s*)?{NUMBER})?"
DOD_REFERENCE = re.compile(
    rf"\bDoD\b\s*(?:번호\s*)?[:：#\-]?\s*[\[(]?\s*"
    rf"(?P<numbers>{ITEM}(?:\s*(?:[,，/·ㆍ;]|및|와|과)\s*"
    rf"(?:DoD\s*[-:#]?\s*)?{ITEM})*)",
    re.IGNORECASE,
)


def clean(value):
    return "" if value is None else str(value).strip()


def display_label(value):
    value = re.sub(r"\[[^\]]*\]", "", value)
    value = value.replace("Trip", "여행").replace("Team", "팀")
    return " ".join(value.split())


def parse_dod(value, requirement_id):
    value = unicodedata.normalize("NFKC", clean(value))
    references = list(DOD_REFERENCE.finditer(value))
    covered = [(match.start(), match.end()) for match in references]

    for marker in DOD_MARKER.finditer(value):
        if not any(start <= marker.start() < end for start, end in covered):
            print(f"확인 필요: {requirement_id}의 DoD 표기에 번호가 없습니다.")

    result = set()
    for match in references:
        sequence = re.sub(
            r"\bDoD\b\s*[-:#]?\s*", "", match["numbers"], flags=re.IGNORECASE
        )
        for item in re.finditer(
            r"(\d+)(?:\s*[-~–—∼]\s*(\d+))?", sequence
        ):
            start = int(item.group(1))
            end = int(item.group(2) or start)
            if not 1 <= start <= end <= 26:
                raise ValueError(
                    f"{requirement_id}: DoD 번호 범위를 확인하세요: {item.group(0)}"
                )
            result.update(range(start, end + 1))
    return result


def read_requirements():
    if not SOURCE.is_file():
        raise FileNotFoundError(f"워크스페이스 루트에서 실행하세요: {SOURCE}")

    workbook = openpyxl.load_workbook(SOURCE, data_only=True)
    try:
        candidates = []
        for sheet in workbook.worksheets:
            rows = [
                row for row in range(5, sheet.max_row + 1)
                if re.fullmatch(
                    r"REQ-\d+", clean(sheet.cell(row, 4).value), re.IGNORECASE
                )
            ]
            if rows:
                candidates.append((sheet, rows))

        if len(candidates) != 1:
            raise ValueError("요구사항 시트를 하나로 식별할 수 없습니다.")

        sheet, rows = candidates[0]

        # 병합 셀은 메모리에서만 해석한다. 원본 셀에는 값을 쓰지 않는다.
        merged_values = {}
        for merged in sheet.merged_cells.ranges:
            if merged.min_col > 3 or merged.max_row < 5:
                continue
            value = sheet.cell(merged.min_row, merged.min_col).value
            for row in range(max(5, merged.min_row), merged.max_row + 1):
                for column in range(merged.min_col, min(3, merged.max_col) + 1):
                    merged_values[row, column] = value

        records = []
        for row in rows:
            categories = [
                clean(merged_values.get((row, column), sheet.cell(row, column).value))
                for column in range(1, 4)
            ]
            requirement_id = clean(sheet.cell(row, 4).value).upper()
            if not all(categories):
                raise ValueError(f"{requirement_id}: A~C 분류에 빈 값이 있습니다.")
            records.append({
                "major": categories[0],
                "middle": categories[1],
                "minor": categories[2],
                "id": requirement_id,
                "requirement": clean(sheet.cell(row, 5).value),
                "dod": parse_dod(sheet.cell(row, 6).value, requirement_id),
            })

        ids = [record["id"] for record in records]
        if len(records) != 62 or len(set(ids)) != 62:
            raise ValueError(
                f"고유한 요구사항 62개가 필요합니다. 행 {len(records)}개, "
                f"고유 ID {len(set(ids))}개입니다."
            )
        return records
    finally:
        workbook.close()


def category_colors(records):
    majors = list(dict.fromkeys(record["major"] for record in records))
    palette = [BLUE, ORANGE]
    return {
        major: palette[index] if index < len(palette) else GRAY
        for index, major in enumerate(majors)
    }


def draw_counts(records, colors):
    major_counts = Counter(record["major"] for record in records)
    middle_counts = Counter(
        (record["major"], record["middle"]) for record in records
    )
    grouped_keys = [
        key for major in major_counts for key in middle_counts if key[0] == major
    ]
    major_labels = [display_label(major) for major in major_counts]
    middle_labels = [
        "\n".join(textwrap.wrap(
            f"{display_label(major)} / {display_label(middle)}",
            width=23,
            break_long_words=True,
            break_on_hyphens=False,
        ))
        for major, middle in grouped_keys
    ]

    top_height = max(1.3, len(major_counts) * 0.42)
    bottom_height = sum(max(1, label.count("\n") + 1) * 0.25 + 0.12
                        for label in middle_labels)
    height = top_height + bottom_height + 2.0

    figure = plt.figure(figsize=(WIDTH, height))
    grid = figure.add_gridspec(
        2, 1,
        left=0.45, right=0.94, top=1 - 0.85 / height, bottom=0.65 / height,
        height_ratios=[top_height, bottom_height],
        hspace=0.65,
    )

    panels = [
        (
            figure.add_subplot(grid[0]),
            major_labels,
            list(major_counts.values()),
            [colors[major] for major in major_counts],
            "대분류별 요구사항 수",
        ),
        (
            figure.add_subplot(grid[1]),
            middle_labels,
            [middle_counts[key] for key in grouped_keys],
            [colors[key[0]] for key in grouped_keys],
            "중분류별 요구사항 수",
        ),
    ]

    for axis, labels, values, bar_colors, title in panels:
        positions = np.arange(len(labels))
        axis.barh(positions, values, color=bar_colors, height=0.60)
        axis.set_yticks(positions, labels, fontsize=8.5)
        axis.invert_yaxis()
        axis.set_xlim(0, max(values) * 1.20)
        axis.xaxis.set_major_locator(MaxNLocator(integer=True, nbins=5))
        axis.set_xlabel("요구사항 수(개)", fontsize=8.5)
        axis.set_title(title, loc="left", fontsize=10, pad=10)
        axis.set_axisbelow(True)
        axis.grid(axis="x", color=GRID, linewidth=0.5)
        axis.tick_params(axis="y", length=0, pad=7)
        for spine in axis.spines.values():
            spine.set_visible(False)
        for position, value in zip(positions, values):
            axis.text(
                value + max(values) * 0.025, position, str(value),
                va="center", fontsize=9, fontweight="bold",
            )

    figure.suptitle(
        f"분류별 요구사항 현황 · 전체 {len(records)}개",
        fontsize=13, fontweight="bold", y=1 - 0.20 / height,
    )
    figure.text(
        0.04, 0.12 / height,
        "각 분류 수준의 합계는 62개이며, 중분류는 대분류 안에서 집계합니다.",
        fontsize=8, color="#666666",
    )
    figure.savefig(OUTPUT / "REQ_분류별_코덱스.png", dpi=DPI)
    plt.close(figure)


def draw_matrix(records):
    records = sorted(records, key=lambda record: int(record["id"].split("-")[1]))
    matrix = np.zeros((len(records), 26), dtype=np.uint8)
    for row, record in enumerate(records):
        for number in record["dod"]:
            matrix[row, number - 1] = 1

    connections = int(matrix.sum())
    linked_rows = int(np.count_nonzero(matrix.any(axis=1)))
    linked_columns = int(np.count_nonzero(matrix.any(axis=0)))

    height = 12.6
    figure, axis = plt.subplots(figsize=(WIDTH, height))
    figure.subplots_adjust(left=0.16, right=0.975, top=0.895, bottom=0.07)
    figure.suptitle(
        "요구사항과 완료 기준 연결",
        fontsize=13, fontweight="bold", y=0.985,
    )
    figure.text(
        0.04, 0.956,
        f"요구사항 62개 · 완료 기준 26개 · 명시된 연결 {connections}개",
        fontsize=9,
    )
    figure.text(
        0.04, 0.936,
        f"연결된 요구사항 {linked_rows}개 · 연결된 완료 기준 {linked_columns}개",
        fontsize=8.5, color="#666666",
    )

    axis.imshow(
        matrix,
        cmap=ListedColormap(["white", BLUE]),
        vmin=0, vmax=1,
        interpolation="nearest",
        aspect="auto",
        origin="upper",
    )
    axis.set_xticks(np.arange(26), [str(number) for number in range(1, 27)])
    axis.set_yticks(np.arange(len(records)), [record["id"] for record in records])
    axis.tick_params(
        axis="x", top=True, labeltop=True, bottom=True, labelbottom=True,
        length=0, pad=5, labelsize=7.5,
    )
    axis.tick_params(axis="y", length=0, pad=5, labelsize=7.5)
    axis.set_xlabel("완료 기준 번호", fontsize=9, labelpad=8)
    axis.xaxis.set_label_position("top")
    axis.set_ylabel("요구사항 식별자", fontsize=9, labelpad=7)
    axis.set_xticks(np.arange(-0.5, 26, 1), minor=True)
    axis.set_yticks(np.arange(-0.5, len(records), 1), minor=True)
    axis.grid(which="minor", color=GRID, linewidth=0.35)
    axis.tick_params(which="minor", length=0)
    for spine in axis.spines.values():
        spine.set_color("#aaaaaa")
        spine.set_linewidth(0.6)

    figure.legend(
        handles=[
            Patch(facecolor=BLUE, edgecolor="none", label="명시된 연결"),
            Patch(facecolor="white", edgecolor="#aaaaaa", label="명시된 연결 없음"),
        ],
        loc="lower left", bbox_to_anchor=(0.035, 0.020),
        frameon=False, ncol=2, fontsize=8,
    )
    note = (
        "원본 인수조건에 완료 기준 번호가 없어 모든 칸을 비워 두었습니다."
        if connections == 0 else
        "인수조건에 명시된 번호만 표시합니다. 빈칸은 연결이 명시되지 않았음을 뜻합니다."
    )
    figure.text(0.04, 0.012, note, fontsize=7.8, color="#666666")
    figure.savefig(OUTPUT / "REQ_DoD추적_코덱스.png", dpi=DPI)
    plt.close(figure)


def main():
    records = read_requirements()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    draw_counts(records, category_colors(records))
    draw_matrix(records)
    print(f"저장 완료: {OUTPUT / 'REQ_분류별_코덱스.png'}")
    print(f"저장 완료: {OUTPUT / 'REQ_DoD추적_코덱스.png'}")


if __name__ == "__main__":
    main()