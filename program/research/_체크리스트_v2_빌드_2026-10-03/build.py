# -*- coding: utf-8 -*-
"""v2 표를 합쳐 md · xlsx · json 을 만든다. 검토 에이전트의 정정은 v2_fixes.FIX 로 덮어쓴다(없으면 그대로).

    python build.py            # 결과: out/ 아래
"""
import collections
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from v2_data import ROWS, OUT, FACTS          # noqa: E402
from v2_meta import META, METRICS, VOCAB       # noqa: E402

try:
    from v2_fixes import FIX, SOURCES_NOTE, CROSS, FACTS2   # noqa: E402
    FACTS = FACTS2
except ImportError:
    FIX, SOURCES_NOTE, CROSS = {}, "", []

OUTDIR = HERE / "out"
OUTDIR.mkdir(exist_ok=True)
ITEMS = {i["id"]: i for i in json.loads((HERE / "items.json").read_text(encoding="utf-8"))}
CATS = {"P": "장소", "O": "열림", "R": "예약", "T": "시간·짜임새", "M": "길·교통", "W": "날씨", "F": "먹기", "C": "돈", "G": "일행", "L": "당일", "X": "범위 밖", "H": "신뢰"}
ORDER = list("POTHRMWFCGLX")
PREFIX = {"MIG/": "app/infrastructure/db/migrations/", "TO/": "travel_ops/", "INF/": "infra/travel/", "E/": "travel_ops/mobility/engine/", "G:": "config/guardrails.yaml:"}
SYM = {"Y": "●", "P": "◐", "N": "·", "C": "◇"}
PATHS = [("g", "생성"), ("r", "등록"), ("w", "감시"), ("a", "응답")]


def expand(text: str) -> str:
    for k, v in PREFIX.items():
        text = text.replace(k, v)
    return text


def cell(c: str):
    letter, _, note = c.partition("|")
    return letter, note


rows = []
for rid, g, r, w, a, ev, miss, note in ROWS:
    d = {"id": rid, "g": g, "r": r, "w": w, "a": a, "ev": ev, "miss": miss, "note": note}
    fx = dict(FIX.get(rid, {}))
    adds = {k: fx.pop(k) for k in ("note_add", "ev_add", "miss_add") if k in fx}
    d.update(fx)
    for k, tgt in (("note_add", "note"), ("ev_add", "ev"), ("miss_add", "miss")):
        if k in adds:
            d[tgt] = (d[tgt] + " · " if d[tgt] else "") + adds[k]
    rows.append(d)

by_id = {d["id"]: d for d in rows}


def v2_status(d):
    L = [cell(d[p])[0] for p in "grw"]
    if any(x in "YPC" for x in L):
        return "일부 경로·조건에서 판정"
    return "표시만(판정 없음)" if cell(d["a"])[0] == "P" else "어느 경로에도 없음"


for d in rows:
    d["status"] = v2_status(d)
    d["meta"] = META[d["id"]]
    d["v1"] = ITEMS[d["id"]]["status"]
    d["text"] = ITEMS[d["id"]]["text"]
    d["unit"] = ITEMS[d["id"]]["unit"]

cnt = {p: collections.Counter(cell(d[p])[0] for d in rows) for p, _ in PATHS}
st = collections.Counter(d["status"] for d in rows)
move = collections.Counter((d["v1"], d["status"]) for d in rows)
pri = collections.Counter(d["meta"][3] for d in rows)
act = collections.Counter(d["meta"][4] for d in rows)
N = len(rows)


def pct(a, b):
    return f"{a}/{b} = {a * 100 // b}%" if b else "분모 0"


# ───────────────────────── json ─────────────────────────
dump = {
    "counts": {p: dict(cnt[p]) for p, _ in PATHS}, "status": dict(st), "moves": {f"{a}→{b}": v for (a, b), v in move.items()},
    "priority": dict(pri), "action": dict(act), "n": N,
    "rows": [{"id": d["id"], "cat": CATS[d["id"][0]], "text": d["text"], "unit": d["unit"], "v1": d["v1"], "status": d["status"],
              **{p: cell(d[p]) for p, _ in PATHS}, "method": d["meta"][0], "criterion": d["meta"][1], "auto": d["meta"][2],
              "priority": d["meta"][3], "action": d["meta"][4], "ev": expand(d["ev"]), "miss": d["miss"], "note": d["note"]} for d in rows],
    "out": {k: {"status": v[0], "text": v[1], "item": ITEMS[k]["text"]} for k, v in OUT.items()},
    "facts": [(a, b, expand(c)) for a, b, c in FACTS], "metrics": METRICS, "vocab": VOCAB,
    "cross": CROSS, "sources": SOURCES_NOTE,
}
(OUTDIR / "v2_table.json").write_text(json.dumps(dump, ensure_ascii=False, indent=1), encoding="utf-8")

# ───────────────────────── md ─────────────────────────
L = []
w = L.append
n_none, n_part, n_show = st["어느 경로에도 없음"], st["일부 경로·조건에서 판정"], st["표시만(판정 없음)"]
v1_none = sum(1 for d in rows if d["v1"] == "아직 없음")
v1_none_found = sum(1 for d in rows if d["v1"] == "아직 없음" and d["status"] != "어느 경로에도 없음")
v1_done_over = sum(1 for d in rows if d["v1"] == "서버가 이미 판정")
w("# 여행자 확인 체크리스트 v2 — 경로별 판정 현황과 평가 지표 (2026-10-03)")
w("")
w("> v1([여행자_확인_체크리스트_조사_2026-10-02.md](여행자_확인_체크리스트_조사_2026-10-02.md), 72가지)은 그대로 둔다. 이 문서는 v1 의 「서버 현황」 주장을 **코드에서 반증**하고, 같은 확인이 **어느 길(생성·등록·감시·응답)에서 도는지**로 다시 쪼갠 2판이다.")
w("")
w("> 같은 내용의 필터·정렬용 엑셀: [여행자_확인_체크리스트_v2_2026-10-03.xlsx](여행자_확인_체크리스트_v2_2026-10-03.xlsx) · 웹 페이지(비공개 아티팩트): https://claude.ai/artifact/WNKrhzbt6Ga64rhaXWnRCe · 만든 과정·교차검증 원본: [_체크리스트_v2_빌드_2026-10-03/](_체크리스트_v2_빌드_2026-10-03/)")
w("")
w("## 한눈에")
w("")
w(f"- **결론**: 사람이 하는 확인 72가지 중 서버 일정 판정 범위 안은 **{N}가지**. 그중 **{n_part}가지({pct(n_part, N)})는 일부 길·조건에서 이미 판정**하고, **{n_show}가지는 보여 주기만** 하고, **{n_none}가지({pct(n_none, N)})는 어느 길에도 없다**. 모든 길에서 판정하는 항목은 하나도 없다.")
v1_judged = sum(1 for d in rows if d["v1"] == "아직 없음" and d["status"] == "일부 경로·조건에서 판정")
v1_shown = v1_none_found - v1_judged
w(f"- **v1 에서 달라진 것**: v1 은 「아직 없음」 {v1_none}가지라 했는데 그중 **{v1_judged}가지는 일부 길에서 이미 판정**하고 {v1_shown}가지는 보여 주기만 한다. 「서버가 이미 판정」 {v1_done_over}가지는 **{v1_done_over}가지 모두 일부 길·조건에서만**이라 과장이었다.")
w(f"- **가장 큰 구멍 셋**: ① 외부에서 받은 일정(등록 길)은 같은 곳 두 번·끼니 누락·요일 불일치를 못 본다 ② 웹 화면은 결제·예산·식이 조건을 서버에 보내지 않아(서버 입구는 받는다) 웹 경로에서는 그 검사가 꺼져 있다 ③ 감시 두 벌 중 시나리오 쪽에는 「바꾼 뒤 일정 전체 재판정」이 없다.")
w("- **정할 것**: 아래 「적용 순서」 1단계(우선순위 1 항목 규칙 추가)를 서버 세션에 넘길지, 그리고 결함일 비율 측정(§7)을 먼저 한 번 돌릴지.")
w("")
w("| 할 일 | 상태 | 진행 |")
w("|---|---|---|")
w(f"| v1 72가지의 서버 현황을 코드로 반증 | 끝 | {N}/{N} = 100% (범위 안 {N}가지, 범위 밖 {len(OUT)}가지 제외) |")
w("| 길(생성·등록·감시·응답)별 판정 표 | 끝 | 66/66 = 100% |")
w("| 항목별 확인 방법·통과 기준·우선순위 | 끝 | 66/66 = 100% |")
w("| 정량 평가 고리(지표 6개·판정 어휘 4개) 설계 | 끝 | 설계만 — 아직 측정 0회 |")
w("| 평가 고리 실측(결함일 비율) | 미착수 | 0/1 = 0% |")
w("| 서버에 규칙 추가 | 미착수 | 0/{} = 0% (우선순위 1 중 규칙 추가) |".format(sum(1 for d in rows if d["meta"][3] == 1 and d["meta"][4] in ("규칙 추가", "배선"))))
w("")
w("## 1. 무엇이고 왜 필요한가")
w("")
w("사람은 여행 일정표를 믿기 전에 **같은 질문을 속으로 반복**한다 — 그날 여나, 닫기 전에 닿나, 같은 곳을 두 번 가진 않나, 길이 이어지나. Validation(검증) 에이전트가 할 일은 이 질문을 **빠짐없이 하고, 결과를 숫자로 내는 것**이다. 그래서 질문 목록이 곧 **채점표(rubric)** 가 된다.")
w("")
w("멘토 피드백(`program/산출물양식/w5/발표자료/assets/mentor_feedback.md` 기능·검증 절)의 요구는 둘이다 — ① 사람이 체크하는 것을 전부 모아 체크리스트를 만들 것 ② **추천 품질을 숫자로** 보일 것. v1 은 ①을 했고, 이 2판은 ①을 코드로 확인해 고치고 ②의 틀(§7)을 더한다.")
w("")
w("외부 근거: AI 가 짠 여행 356건을 사람이 점검한 연구에서 **일정의 하루 중 43%에 결함이 있었다**(그때 닫힘 22.7% · 마감 뒤 16.6% · 같은 곳 두 번 5.1% · 지그재그 9.5%). 같은 잣대를 우리 일정에 대면 「결함일 비율」로 직접 비교된다(§7).")
w("")
w("## 2. 네 길 — 같은 확인도 어느 길에서 도는지가 다르다")
w("")
w("| 길 | 언제 | 코드 위치 | 판정기 |")
w("|---|---|---|---|")
w("| 생성 | 우리 생성기가 일정을 **짤 때** | `travel_ops/planner.py` | 짜면서 같은 판정기로 고침(최대 3회) |")
w("| 등록 | 고객·외부 에이전트가 만든 일정을 **받을 때**(REST·MCP·읽은 일정 확인) | `itinerary_checks.py` · `trip_api.py` · `intake/` | `check_itinerary` |")
w("| 감시 | 여행 **중** 점검·대체(3분 주기·03:00 새벽 확인·고객 신고) | `trip_watch*.py` · `itinerary_changes.py` · `replan.py` · `dawn_check.py` | 항목 하나 점검 + (Case 쪽만) 바꾼 뒤 전체 재판정 |")
w("| 응답 | 결과가 **고객 화면·API 응답**에 나가는가 | `trip_api.py` · `place_info.py` · 알림 문구 | 표시일 뿐 판정이 아님 |")
w("")
w("기호: ● 판정한다 · ◐ 일부 조건·일부 경우만 · ◇ 만드는 방식으로 보장(판정 코드는 아님) · `·` 안 한다.")
w("")
w(f"경로별 칸 개수(66항목): " + " · ".join(f"{name} ●{cnt[p].get('Y',0)} ◐{cnt[p].get('P',0)} ◇{cnt[p].get('C',0)} ·{cnt[p].get('N',0)}" for p, name in PATHS))
w("")
w("## 3. 경로별 판정표")
w("")
w("`다음 조치`: 규칙 추가(값은 이미 있는데 판정이 안 걸림) · 입력 연결(고객 입력이 서버에 안 닿음) · 배선(만들어 둔 코드를 아무도 안 부름) · 소스 필요(데이터 자체가 없음) · 고객 확인 · 유지. 우선순위 1~3.")
w("")
for cat in ORDER:
    sub = [d for d in rows if d["id"][0] == cat]
    if not sub:
        continue
    w(f"### {CATS[cat]}")
    w("")
    w("| ID | 확인 | 생성 | 등록 | 감시 | 응답 | 지금 | 우선 | 다음 조치 |")
    w("|---|---|:-:|:-:|:-:|:-:|---|:-:|---|")
    for d in sub:
        sy = [SYM[cell(d[p])[0]] for p, _ in PATHS]
        w(f"| {d['id']} | {d['text'].replace('|', '/')} | {sy[0]} | {sy[1]} | {sy[2]} | {sy[3]} | {d['status']} | {d['meta'][3]} | {d['meta'][4]} |")
    w("")
w("범위 밖(서버 일정 판정 대상이 아님): " + " · ".join(f"{k}({v[1]})" for k, v in OUT.items()))
w("")
w("## 4. 서버 공통 사실 — 표 밖의 구멍")
w("")
for i, (a, b, c) in enumerate(FACTS, 1):
    w(f"{i}. **{a}** — {b} (`{expand(c)}`)")
w("")
w("## 5. 값을 모를 때 — 경로마다 다르게 다룬다")
w("")
w("계획서 결정 15(「값을 모르는 상황을 만들지 않는다 — 대체 소스를 두고, 대체까지 안 되면 서버를 끈다」)와 실제 코드가 어긋나는 곳이 표에 반복해서 나온다.")
w("")
w("| 값 | 등록 | 생성 | 감시(대체 후보) |")
w("|---|---|---|---|")
w("| 영업시간 모름 | 통과(건너뜀) | 통과, 단 `coverage.hours_known n/m` 으로 공개 | 식당은 탈락 · 활동은 통과 |")
w("| 결제 수단 모름 | 통과 | 통과 | 식당은 탈락 · 활동은 안 봄 |")
w("| 가격 모름 | 0원으로 합산(통과 쪽) | 0원으로 합산, 분모는 `coverage` 에서만 | 활동은 경고 · 식당은 안 봄 |")
w("| 소스 호출 실패 | — | — | fatal(바꾸지 않음·알리지 않음 — 결정 15) |")
w("")
w("→ **같은 「모름」이 길마다 통과/탈락으로 갈린다.** 판정 결과 어휘에 `unknown` 을 두면(§7) 이 갈림이 숫자로 보인다.")
w("")
w("## 6. 우선순위 1 항목 — 길별 근거와 값이 없을 때")
w("")
for d in rows:
    if d["meta"][3] != 1:
        continue
    w(f"**{d['id']} {d['text']}** — {d['status']} · 다음 조치 {d['meta'][4]}")
    for p, name in PATHS:
        letter, note = cell(d[p])
        w(f"- {name} {SYM[letter]} {note or '없음'}")
    w(f"- 값이 없을 때: {d['miss']}")
    w(f"- 근거: `{expand(d['ev'])}`")
    w(f"- 확인 방법 → 통과 기준: {d['meta'][0]} → {d['meta'][1]}")
    if d["note"]:
        w(f"- 메모: {d['note']}")
    w("")
w("나머지 항목의 길별 설명·근거 `파일:줄`·값이 없을 때는 엑셀 파일의 「경로별 상세」 시트에 있다.")
w("")
w("## 7. 정량 평가 고리")
w("")
w("체크리스트를 **채점표**로 쓰면 일정 하나를 넣었을 때 항목마다 한 줄 결과가 나온다.")
w("")
w("결과 어휘(`itinerary-verifier` 방식): " + " · ".join(f"`{a}` {b}" for a, b in VOCAB))
w("")
w("★`unknown` 은 통과가 아니다. 값이 없어 못 본 것을 `pass` 로 세면 모름이 성공률을 올려 보이게 한다(전역 규칙 「조용한 스킵 금지」와 같은 이유).")
w("")
w("| 지표 | 계산 | 왜 |")
w("|---|---|---|")
for a, b, c in METRICS:
    w(f"| {a} | {b} | {c} |")
w("")
w("예시(계산 방식만 보인다 — 실측 값이 아니다): 7일 일정 7일 중 영업 마감 뒤 방문이 2일, 같은 곳 두 번이 1일(겹치는 날 없음)이면 결함일 비율 = 3/7 = 43%, 유형별로는 마감 뒤 2/7 · 중복 1/7.")
w("")
w("★이 문서의 표는 **코드를 읽은 결과**이지 **일정을 흘려 잰 결과**가 아니다. 기존 평가 러너(`eval/runners/travel_scenarios.py`)는 같은 판정기로 시나리오 10건(하루 1종·서울 1도시)에서 위반 0건을 확인하는 용도라, **결함일 비율은 아직 잰 적이 없다.**")
w("")
w("## 8. 적용 순서")
w("")
p1_rule = [d for d in rows if d["meta"][3] == 1 and d["meta"][4] in ("규칙 추가", "배선")]
w(f"1. **우선순위 1 규칙 추가**({len(p1_rule)}가지: " + ", ".join(d["id"] for d in p1_rule) + "): 값은 이미 있는데 판정이 안 걸린 길에 판정만 더한다. 대표: T8(등록 길 같은 곳 중복) · T7(등록 길 끼니 누락) · T10(요일 대조) · O1(공휴일 연결) · T9(되돌아감) · T5(시나리오 감시 재판정).")
w("2. **입력 연결**: 웹 화면이 결제·예산·식이를 서버에 보내게 하고, 설문 음식 코드를 식이 조건으로 바꾸는 코드(`conds_from_survey`)·이동 일행 조건(`party_of`)을 실제로 채운다.")
w("3. **배선**: 만들어 두고 아무도 안 부르는 `dining/sweep.py` · `dining/tick.py` · `watch.py` 를 감시 루프에 건다 — 또는 안 쓸 거면 문서에서 「있다」고 하지 않는다.")
w("4. **평가 고리 측정**: 체크리스트 id 별 결과를 세는 채점기를 `eval/` 에 붙여 결함일 비율·모름 비율·판정 범위를 1회 측정한다.")
w("5. **소스 필요 항목은 보류**: 예약·메뉴·접근성 같은 데이터가 아예 없는 19가지는 소스를 정할 때까지 「모름 표시」로 둔다.")
w("")
w("## 9. 교차검증 기록")
w("")
import re as _re


def plain(t):
    for a, c in (("입니다", "이다"), ("됩니다", "된다"), ("않습니다", "않는다"), ("습니다", "다"), ("못했습니다", "못했다"), ("봤습니다", "봤다"), ("뒀습니다", "뒀다")):
        t = t.replace(a, c)
    return t


w(plain(SOURCES_NOTE) if SOURCES_NOTE else "- 방법: ①하위 조사 3개 ②Codex ③검토 전용 Claude.")
w("")
for line in CROSS:
    w(f"- {plain(line)}")
w("")
w("## 10. 한계")
w("")
w("- **코드만 읽었다.** 라이브 DB 에 영업시간·결제·가격 칸이 실제로 몇 % 채워져 있는지, 운영에서 이동 계산기(`mobility_data_dir`)가 켜져 있는지, 구글·UTIC 키가 있는지는 보지 않았다. 꺼져 있으면 표의 P 중 일부는 N 이 된다.")
w("- 「N」은 한국어·영어 키워드 검색으로 안 나왔다는 뜻이다. 약어·JSON 키·런타임에 조립되는 문자열·다른 이름의 같은 개념은 놓칠 수 있다. SQL 안(요식 원장 함수)은 하위 조사들이 읽지 않았고 검토 에이전트가 일부 확인했다.")
w("- 「호출처 0건」은 `app/`·`scripts/` 이름 검색이다. 크론·동적 호출은 놓칠 수 있다.")
w("- 응답 칸은 서버가 내보내는 값 기준이다. 프런트가 실제로 그리는지는 보지 않았다.")
w("- 위 표는 변하는 코드의 2026-10-03 시점 스냅샷이다.")
(OUTDIR / "checklist_v2.md").write_text("\n".join(L) + "\n", encoding="utf-8")

# ───────────────────────── xlsx ─────────────────────────
from openpyxl import Workbook                 # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side   # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

wb = Workbook()
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD = PatternFill("solid", fgColor="1F3A5F")
FILL = {"Y": "C6E7C9", "P": "FFF2B3", "N": "E4E4E4", "C": "CFE2F6"}
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, names, widths):
    ws.append(names)
    for i, (n, wd) in enumerate(zip(names, widths), 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "C2"


# 1 요약
ws = wb.active
ws.title = "요약"
ws.column_dimensions["A"].width = 34
for col in "BCDEF":
    ws.column_dimensions[col].width = 16
ws["A1"] = "여행자 확인 체크리스트 v2 — 요약 (2026-10-03)"
ws["A1"].font = Font(bold=True, size=14)
ws.append([])
ws.append(["서버 일정 판정 범위 안 항목", N, "", "범위 밖(제외)", len(OUT)])
ws.append(["일부 길·조건에서 이미 판정", n_part, pct(n_part, N)])
ws.append(["보여 주기만(판정 없음)", n_show, pct(n_show, N)])
ws.append(["어느 길에도 없음", n_none, pct(n_none, N)])
ws.append([])
ws.append(["길별 칸 수", "● 판정", "◐ 일부", "◇ 구성 보장", "· 안 함"])
for c in ws[ws.max_row]:
    c.font = Font(bold=True)
for p, name in PATHS:
    ws.append([name, cnt[p].get("Y", 0), cnt[p].get("P", 0), cnt[p].get("C", 0), cnt[p].get("N", 0)])
ws.append([])
ws.append(["v1 상태 → v2 상태", "항목 수"])
for c in ws[ws.max_row]:
    c.font = Font(bold=True)
for (a, b), v in sorted(move.items()):
    ws.append([f"{a} → {b}", v])
ws.append([])
ws.append(["우선순위", "항목 수"])
for c in ws[ws.max_row]:
    c.font = Font(bold=True)
for k in (1, 2, 3):
    ws.append([f"{k}순위", pri.get(k, 0)])
ws.append([])
ws.append(["다음 조치", "항목 수"])
for c in ws[ws.max_row]:
    c.font = Font(bold=True)
for k, v in act.most_common():
    ws.append([k, v])
ws.append([])
ws.append(["서버 공통 사실(표 밖의 구멍)"])
ws[ws.max_row][0].font = Font(bold=True)
for a, b, c in FACTS:
    ws.append([a, b])
    ws.cell(row=ws.max_row, column=2).alignment = WRAP
    ws.merge_cells(start_row=ws.max_row, start_column=2, end_row=ws.max_row, end_column=6)
    ws.row_dimensions[ws.max_row].height = 60

# 2 체크리스트
ws = wb.create_sheet("체크리스트")
header(ws, ["ID", "분류", "확인 항목", "단위", "v1 상태", "v2 상태", "생성", "등록", "감시", "응답", "확인 방법", "통과 기준", "자동화", "우선", "다음 조치"],
       [6, 9, 44, 7, 14, 20, 6, 6, 6, 6, 40, 40, 8, 6, 12])
for d in rows:
    ws.append([d["id"], CATS[d["id"][0]], d["text"], d["unit"], d["v1"], d["status"]] + [SYM[cell(d[p])[0]] for p, _ in PATHS]
              + [d["meta"][0], d["meta"][1], d["meta"][2], d["meta"][3], d["meta"][4]])
    r = ws.max_row
    for j, (p, _) in enumerate(PATHS):
        c = ws.cell(row=r, column=7 + j)
        c.fill = PatternFill("solid", fgColor=FILL[cell(d[p])[0]])
        c.alignment = Alignment(horizontal="center", vertical="top")
    for col in range(1, 16):
        ws.cell(row=r, column=col).border = BORDER
        if col in (3, 11, 12):
            ws.cell(row=r, column=col).alignment = WRAP
for k, v in OUT.items():
    ws.append([k, CATS[k[0]], ITEMS[k]["text"], ITEMS[k]["unit"], ITEMS[k]["status"], v[0], "", "", "", "", v[1]])
ws.auto_filter.ref = f"A1:O{ws.max_row}"

# 3 경로별 상세
ws = wb.create_sheet("경로별 상세")
header(ws, ["ID", "확인 항목", "생성", "등록", "감시", "응답", "근거 파일:줄", "값이 없을 때", "메모"], [6, 34, 44, 44, 44, 40, 48, 40, 44])
for d in rows:
    ws.append([d["id"], d["text"]] + [f"{SYM[cell(d[p])[0]]} {cell(d[p])[1]}".strip() for p, _ in PATHS] + [expand(d["ev"]), d["miss"], d["note"]])
    r = ws.max_row
    for col in range(1, 10):
        ws.cell(row=r, column=col).alignment = WRAP
        ws.cell(row=r, column=col).border = BORDER
    for j, (p, _) in enumerate(PATHS):
        ws.cell(row=r, column=3 + j).fill = PatternFill("solid", fgColor=FILL[cell(d[p])[0]])
ws.auto_filter.ref = f"A1:I{ws.max_row}"

# 4 평가 지표
ws = wb.create_sheet("평가 지표")
header(ws, ["지표", "계산", "왜"], [32, 56, 70])
ws.freeze_panes = "A2"
for a, b, c in METRICS:
    ws.append([a, b, c])
    for col in range(1, 4):
        ws.cell(row=ws.max_row, column=col).alignment = WRAP
ws.append([])
ws.append(["판정 결과 어휘", "뜻"])
for c in ws[ws.max_row]:
    c.font = Font(bold=True)
for a, b in VOCAB:
    ws.append([a, b])

# 5 교차검증
ws = wb.create_sheet("교차검증 기록")
header(ws, ["항목", "내용"], [28, 120])
ws.freeze_panes = "A2"
for line in ([SOURCES_NOTE] if SOURCES_NOTE else []) + CROSS:
    ws.append(["", line])
    ws.cell(row=ws.max_row, column=2).alignment = WRAP

xlsx_path = OUTDIR / "checklist_v2.xlsx"
wb.save(xlsx_path)
print("ok", N, dict(st), {p: dict(cnt[p]) for p, _ in PATHS})
